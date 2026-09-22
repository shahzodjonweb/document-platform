"""Serialized OCR/conversion/editor acceptance tests including irreversible removal."""
import io
from pathlib import Path
import pytest
from PIL import Image,ImageDraw,ImageFont
from pypdf import PdfReader
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.utils import ImageReader
from reportlab.platypus import Table,TableStyle
from reportlab.lib import colors
from processors import execute,inspect_file,normalize_parameters,ProcessorError
from processors.sandbox import execute_sandbox
from processors.editor import _font

FONT=Path(__file__).resolve().parent.parent/'processors/assets/fonts/NotoSans-Regular.ttf'


def pdf(path,text='Account OLD VALUE',image=None,form=False,font='Helvetica'):
    c=Canvas(str(path),pagesize=(400,600),pageCompression=0)
    c.setFont(font,18);c.drawString(40,550,text)
    if image:c.drawImage(ImageReader(image),40,400,100,80)
    if form:c.acroForm.textfield(name='customer',x=40,y=430,width=260,height=45,fontSize=12)
    c.showPage();c.save();return path


def scan(path,text):
    image=Image.new('RGB',(1500,320),'white')
    ImageDraw.Draw(image).text((50,100),text,font=ImageFont.truetype(str(FONT),58),fill='black')
    image.save(path);return path


def run(feature,source,commands,out,other=None,**kwargs):
    return execute_sandbox(feature,[source]+([other] if other else []),{'commands':commands,**kwargs},out)


def output(result):return Path(result['artifacts'][0]['path'])


@pytest.mark.parametrize(('language','text','expected'),[('eng','HELLO DOCUMENT 2026','HELLO'),('rus','ПРИВЕТ ДОКУМЕНТ 2026','ПРИВЕТ'),('uzb','SALOM HUJJAT 2026','HUJJAT')])
def test_printed_ocr_in_three_languages(tmp_path,language,text,expected):
    source=scan(tmp_path/'scan.png',text)
    result=execute_sandbox('ocr.extract_text',[source],{'language':language},tmp_path/'out')
    assert expected in output(result).read_text()
    assert result['actual_page_units']==1
    assert 'printed_text_ocr_review_required' in result['metadata']['warnings']


def test_searchable_scanned_pdf_has_real_text_layer(tmp_path):
    source=scan(tmp_path/'scan.png','HELLO DOCUMENT 2026')
    result=execute_sandbox('ocr.searchable_pdf',[source],{'language':'eng'},tmp_path/'out')
    page=PdfReader(output(result)).pages[0]
    assert 'HELLO' in page.extract_text() and len(page.images)>0


def test_pdf_to_editable_word_preserves_cyrillic(tmp_path):
    from docx import Document
    source=pdf(tmp_path/'source.pdf','Привет документ',font=_font())
    result=execute_sandbox('convert.pdf_to_docx',[source],{'mode':'text'},tmp_path/'out')
    doc=Document(output(result))
    assert any('Привет документ' in p.text for p in doc.paragraphs)
    assert not list(doc.inline_shapes),'Text conversion cannot substitute page screenshots'
    assert 'editable_text_reflowed_layout_not_preserved' in result['metadata']['warnings']


def test_pdf_tables_to_xlsx_neutralizes_formulas(tmp_path):
    from openpyxl import load_workbook
    source=tmp_path/'table.pdf';canvas=Canvas(str(source),pagesize=(400,600))
    table=Table([['Item','Value'],['Alpha','=2+2'],['Beta','20']],colWidths=[150,150],rowHeights=30)
    table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),1,colors.black)]));table.wrapOn(canvas,300,200);table.drawOn(canvas,40,350);canvas.save()
    result=execute_sandbox('convert.pdf_to_xlsx',[source],{},tmp_path/'out')
    book=load_workbook(output(result),data_only=False);sheet=book.worksheets[0]
    assert sheet['B2'].value=='=2+2' and sheet['B2'].data_type=='s'
    assert sheet['A2'].value=='Alpha'


def test_scan_word_requires_explicit_ocr_mode(tmp_path):
    image=scan(tmp_path/'scan.png','HELLO DOCUMENT')
    c=Canvas(str(tmp_path/'scan.pdf'),pagesize=(600,200));c.drawImage(str(image),0,0,600,128);c.save()
    with pytest.raises(ProcessorError) as error:
        execute_sandbox('convert.pdf_to_docx',[tmp_path/'scan.pdf'],{'mode':'text'},tmp_path/'out')
    assert error.value.code=='scan_requires_ocr'


def test_multilingual_text_highlight_and_note_are_serialized(tmp_path):
    source=pdf(tmp_path/'source.pdf')
    commands=[{'type':'add_text','page':1,'x':40,'y':100,'text':'Salom Привет','font_size':18},
              {'type':'highlight','page':1,'x':38,'y':35,'width':180,'height':24},
              {'type':'note','page':1,'x':300,'y':60,'text':'Примечание'}]
    result=run('editor.visual',source,commands,tmp_path/'out')
    reader=PdfReader(output(result));page=reader.pages[0]
    assert 'Salom Привет' in page.extract_text()
    annotations=[a.get_object() for a in page['/Annots']]
    assert {a['/Subtype'] for a in annotations}=={'/Highlight','/Text'}
    assert any(a.get('/Contents')=='Примечание' for a in annotations)


def test_existing_text_changes_content_not_overlay(tmp_path):
    source=pdf(tmp_path/'source.pdf')
    result=run('editor.existing_text',source,[{'type':'replace_text','page':1,'find':'OLD','replacement':'NEW'}],tmp_path/'out')
    reader=PdfReader(output(result));page=reader.pages[0]
    assert 'Account NEW VALUE' in page.extract_text() and 'OLD' not in page.extract_text()
    assert b'OLD' not in page.get_contents().get_data()
    assert result['metadata']['text_replacements']==1


def test_complex_text_encoding_fails_explicitly(tmp_path):
    source=pdf(tmp_path/'source.pdf','Привет',font=_font())
    with pytest.raises(ProcessorError) as error:
        run('editor.existing_text',source,[{'type':'replace_text','page':1,'find':'Привет','replacement':'Пока'}],tmp_path/'out')
    assert error.value.code=='unsupported_text_encoding'


def test_interactive_unicode_form_has_value_and_embedded_appearance(tmp_path):
    source=pdf(tmp_path/'source.pdf',form=True)
    assert inspect_file(source)['form_fields'][0]['name']=='customer'
    result=run('editor.fill_forms',source,[{'type':'fill_form','field':'customer','value':'Павел Salom'}],tmp_path/'out')
    reader=PdfReader(output(result))
    assert reader.get_fields()['customer']['/V']=='Павел Salom'
    widget=reader.pages[0]['/Annots'][0].get_object()
    appearance=widget['/AP']['/N'].get_object()
    assert appearance.get_data() and appearance['/Resources']['/Font']
    import pypdfium2 as pdfium
    with pdfium.PdfDocument(output(result)) as document:
        document.init_forms()
        page=document[0];bitmap=page.render(scale=2);image=bitmap.to_pil().convert('RGB')
        crop=image.crop((90,260,590,320))
        assert sum(1 for r,g,b in crop.get_flattened_data() if r<150 and g<150 and b<150)>80
        bitmap.close();page.close()


def test_replace_and_remove_real_image_objects(tmp_path):
    red=tmp_path/'red.png';blue=tmp_path/'blue.png'
    Image.new('RGB',(100,80),'red').save(red);Image.new('RGB',(100,80),'blue').save(blue)
    source=pdf(tmp_path/'source.pdf',image=red)
    replaced=run('editor.replace_images',source,[{'type':'replace_image','page':1,'image_index':0,'input_index':1}],tmp_path/'replaced',other=blue)
    image=PdfReader(output(replaced)).pages[0].images[0].image.convert('RGB')
    assert image.getpixel((50,40))[2]>240
    removed=run('editor.replace_images',source,[{'type':'remove_image','page':1,'image_index':0}],tmp_path/'removed')
    assert len(PdfReader(output(removed)).pages[0].images)==0


def test_signature_and_draw_are_real_output_content(tmp_path):
    signature=tmp_path/'signature.png';Image.new('RGBA',(100,30),(0,0,0,180)).save(signature)
    source=pdf(tmp_path/'source.pdf')
    commands=[{'type':'signature','page':1,'x':40,'y':100,'width':140,'height':45,'input_index':1},
              {'type':'draw','page':1,'points':[[40,200],[100,240],[150,200]],'color':'#235CC5'}]
    result=run('editor.visual',source,commands,tmp_path/'out',other=signature)
    page=PdfReader(output(result)).pages[0]
    assert len(page.images)==1 and b' m' in page.get_contents().get_data()


def test_redaction_removes_recoverable_text_and_raster_pixels(tmp_path):
    source=tmp_path/'source.pdf';canvas=Canvas(str(source),pagesize=(400,600),pageCompression=0)
    canvas.setFont('Helvetica',18);canvas.drawString(40,550,'PUBLIC TOP');canvas.drawString(40,470,'SECRET123');canvas.save()
    command={'type':'redact','page':1,'x':25,'y':105,'width':250,'height':60}
    result=run('editor.redact',source,[command],tmp_path/'out',accept_rasterization=True)
    reader=PdfReader(output(result));page=reader.pages[0]
    assert page.extract_text()=='' and not reader.get_fields() and not page.get('/Annots')
    assert b'SECRET123' not in output(result).read_bytes()
    raster=page.images[0].image.convert('RGB')
    assert raster.getpixel((150,260))==(0,0,0)
    recovered=execute_sandbox('ocr.extract_text',[output(result)],{'language':'eng'},tmp_path/'recovery')
    text=output(recovered).read_text()
    assert 'PUBLIC' in text and 'SECRET123' not in text
    assert result['metadata']['redaction_mode']=='rasterized_destructive'


def test_redaction_requires_explicit_rasterization_acknowledgement():
    with pytest.raises(ProcessorError) as error:
        normalize_parameters('editor.redact',{'commands':[{'type':'redact','page':1,'x':1,'y':1,'width':10,'height':10}]})
    assert error.value.code=='rasterization_consent_required'


@pytest.mark.parametrize('command',[{'type':'add_text','page':1,'x':1,'y':1,'text':'x','file_path':'/etc/passwd'}, {'type':'insert_image','page':1,'x':1,'y':1,'width':10,'height':10,'input_index':-1}, {'type':'replace_text','page':1,'find':'','replacement':'x'}])
def test_unsafe_editor_commands_rejected(command):
    with pytest.raises(ProcessorError):normalize_parameters('editor.visual',{'commands':[command]})


def test_parent_form_actions_are_rejected(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import NameObject,DictionaryObject,TextStringObject
    source=pdf(tmp_path/'source.pdf',form=True)
    writer=PdfWriter(clone_from=source)
    field=writer.root_object['/AcroForm']['/Fields'][0].get_object()
    field[NameObject('/AA')]=DictionaryObject({NameObject('/K'):DictionaryObject({NameObject('/S'):NameObject('/JavaScript'),NameObject('/JS'):TextStringObject('unsafe()')})})
    writer.write(tmp_path/'active.pdf')
    with pytest.raises(ProcessorError) as error:inspect_file(tmp_path/'active.pdf')
    assert error.value.code=='active_content_unsupported'


def test_javascript_uri_and_action_chains_rejected(tmp_path):
    from pypdf import PdfWriter
    from pypdf.annotations import Link
    from pypdf.generic import NameObject,DictionaryObject
    source=pdf(tmp_path/'source.pdf')
    for name,url,chain in [('uri','javascript:alert(1)',False),('chain','https://example.com',True)]:
        writer=PdfWriter(clone_from=source)
        annotation=Link(rect=(20,20,80,40),url=url)
        if chain:annotation['/A'][NameObject('/Next')]=DictionaryObject({NameObject('/S'):NameObject('/Launch')})
        writer.add_annotation(0,annotation);target=tmp_path/f'{name}.pdf';writer.write(target)
        with pytest.raises(ProcessorError) as error:inspect_file(target)
        assert error.value.code=='active_content_unsupported'


def test_displayed_page_sizes_follow_rotation(tmp_path):
    from pypdf import PdfWriter
    writer=PdfWriter();writer.add_blank_page(width=200,height=300).rotate(90);source=tmp_path/'rotated.pdf';writer.write(source)
    assert inspect_file(source)['page_sizes']==[{'width':300.0,'height':200.0}]
