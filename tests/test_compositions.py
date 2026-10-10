"""Every composition keeps text readable, unobstructed and in its box, in every theme.

A composition arranges covers, dividers and closing slides and frames content
slides (apps/studio/compositions). check.problems reads the rendered file back
and reports text below its contrast, text on a picture or across a line,
overlapping or off-slide text, and cover or divider text that does not fit.
"""
import io

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from apps.studio.compositions import COMPOSITIONS, Canvas, compose
from apps.studio.compositions import sample
from apps.studio.compositions.check import problems
from apps.studio.slides import palette

compose('classic')  # registers every composition
NAMES = sorted(COMPOSITIONS)
# Colours that stress contrast: a dark accent, a mid-tone one and a light one.
ACCENTS = [('#1F3A68', '#E0A458'), ('#D94A26', '#2F80ED'), ('#F4D35E', '#2A9D8F')]


@pytest.mark.parametrize('theme', ['light', 'dark', 'bold'])
@pytest.mark.parametrize('name', NAMES)
def test_text_stays_readable_and_in_place(name, theme, tmp_path):
    found = []
    for accent, secondary in ACCENTS:
        for with_photos in (False, True):
            style = sample._style(name, accent, secondary, theme)
            path = tmp_path / 'deck.pptx'
            result = sample.render(style, path, with_photos)
            found += [f'{accent} photos={with_photos}: {problem}'
                      for problem in problems(path, result['metadata']['layouts'], accent)]
    assert not found, '\n'.join(found)


@pytest.mark.parametrize('name', NAMES)
def test_long_cyrillic_titles_fit(name, tmp_path):
    path = tmp_path / 'deck.pptx'
    result = sample.render(sample._style(name, '#5B2A86', '#E0A458', 'light'), path, locale='ru')
    assert not problems(path, result['metadata']['layouts']), problems(path, result['metadata']['layouts'])


def test_the_checker_finds_what_it_is_for():
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    canvas = Canvas(slide, palette('#1F3A68'), {}, 'cover')
    canvas.background('surface')
    canvas.accent((0, 0, 0.2, 7.5))
    canvas.rect((1, 1, 4, 2), 'accent', alpha=0.3)
    canvas.line((1, 5), (9, 5), 'ink')
    buffer = io.BytesIO()
    from PIL import Image
    Image.new('RGB', (10, 10), (200, 0, 0)).save(buffer, 'PNG')
    slide.shapes.add_picture(buffer, Inches(10), Inches(1), Inches(2), Inches(2))

    def text(box, value, colour):
        frame = slide.shapes.add_textbox(*(Inches(v) for v in box)).text_frame
        run = frame.paragraphs[0].add_run()
        run.text, run.font.size, run.font.color.rgb = value, Pt(14), RGBColor.from_string(colour)

    text((1.2, 1.2, 3, 1), 'white on a pale tint', 'FFFFFF')
    text((1, 4.8, 6, 0.5), 'a line through me', '111111')
    text((10.2, 1.2, 1.5, 1), 'on a picture', 'FFFFFF')
    text((1.5, 1.5, 2, 1), 'over other text', '111111')
    found = '\n'.join(problems(deck, accent='#1F3A68'))
    for expected in ('white on a pale tint', 'has a line across it', 'sits on a picture', 'overlaps'):
        assert expected in found, found
