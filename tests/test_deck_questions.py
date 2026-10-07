"""A deck's questions share slides at the end and count toward the slides asked for.

"10 ta slayd … oxirida 5 ta savol" made fifteen slides, the last five each
holding one short question. Now it makes ten: nine written slides and one
slide of five questions, with the answers in the speaker notes.
"""
import pytest
from pptx import Presentation

from apps.studio.domain import SLIDES, create_draft, generation_quote, unpack
from apps.studio.slides import render_pptx
from tests.test_platform import account
from tests.test_studio import paid


def question(n, options=True):
    return {'id': f'q{n}', 'stem': f'Which of these best describes stage {n} of the water cycle in a river basin?',
            'options': ['Evaporation from the surface', 'Condensation into clouds', 'Precipitation as rain',
                        'Collection in lakes'] if options else [],
            'answer': 'Evaporation from the surface', 'explanation': '', 'topic': 'Water', 'marks': 1}


def deck(sections, questions):
    return {'title': 'Suv aylanishi', 'citations': [], 'questions': questions, 'sections': [
        {'id': f's{i + 1}', 'heading': f'Heading {i + 1}', 'body': 'First point.\nSecond point.', 'notes': ''}
        for i in range(sections)]}


@pytest.mark.parametrize('count,slides', [(1, 1), (5, 1), (6, 2), (10, 2)])
def test_questions_are_grouped_five_to_a_slide(tmp_path, count, slides):
    result = render_pptx(deck(4, [question(n) for n in range(1, count + 1)]), tmp_path / 'deck.pptx', 'uz')
    assert result['page_count'] == 4 + slides
    shown = Presentation(str(tmp_path / 'deck.pptx'))
    last = shown.slides[-1]
    text = '\n'.join(shape.text_frame.text for shape in last.shapes if shape.has_text_frame)
    assert 'Savollar' in text and 'A) Evaporation from the surface' in text
    assert f'{count}.' in text
    assert 'Javoblar' in last.notes_slide.notes_text_frame.text
    assert result['metadata']['layouts'][-slides:] == ['questions'] * slides


@pytest.mark.django_db
@pytest.mark.parametrize('brief,total,written,question_slides', [
    ('Ochiq darsga 10 ta slayd kerak — Orol dengizi. Oxirida 5 ta savol.', 10, 9, 1),
    ('12 ta slayd, oxirida 10 ta savol', 12, 10, 2),
    ('Suv aylanishi haqida slayd, 5 ta savol bilan', 5, 4, 1),
    ('8 ta slayd: suv aylanishi', 8, 8, 0),
])
def test_question_slides_count_toward_the_slides_asked_for(settings, brief, total, written, question_slides):
    settings.DEBUG = True
    customer = paid(settings, account())
    draft = create_draft(customer, {'feature_id': SLIDES, 'prompt': brief})
    data = unpack(draft.encrypted_data)
    assert data['options']['length'] == total
    assert data['options']['question_slides'] == question_slides
    assert len(data['content']['sections']) == written
    # The price covers every slide, question slides included.
    quote = generation_quote(customer, draft.id, draft.version)
    assert quote.policy['generation_bounds']['output_pages'] == total + 1


@pytest.mark.django_db
def test_the_language_asked_for_is_the_language_written(settings):
    import json
    from apps.studio import provider
    from tests.test_provider_contract import CONFIG
    settings.DEBUG = True
    customer = paid(settings, account())
    english = unpack(create_draft(customer, {'feature_id': SLIDES, 'prompt': 'Best friend mavzusida 5 ta slayd, ingliz tilida'}).encrypted_data)
    assert english['output_locale'] == 'en'
    customer.locale = 'uz'
    customer.save(update_fields=['locale'])
    uzbek = unpack(create_draft(customer, {'feature_id': SLIDES, 'prompt': 'НЕФТ ВА ГАЗ ҚУДУҚЛАРИ, 6 ta slayd'}).encrypted_data)
    assert uzbek['output_locale'] == 'uz' and uzbek['options']['uz_script'] == 'cyrl'
    sent = provider.request_body(CONFIG, uzbek, 'ai.pptx')
    assert json.loads(sent['input'])['output_script'] == 'cyrillic'
    assert 'Never mix the two alphabets' in sent['instructions']


def test_question_labels_follow_the_alphabet(tmp_path):
    result = render_pptx(deck(1, [question(1)]), tmp_path / 'c.pptx', 'uz', style={'uz_script': 'cyrl'})
    shown = Presentation(str(tmp_path / 'c.pptx'))
    text = '\n'.join(shape.text_frame.text for shape in shown.slides[-1].shapes if shape.has_text_frame)
    assert 'Саволлар' in text and result['page_count'] == 2
