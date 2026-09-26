"""Independent studio confirmations, provider boundary, ownership and role tests."""
import copy
import io
import json
import uuid
from datetime import timedelta
import pytest
from django.core import signing
from django.test import Client
from django.utils import timezone
from pypdf import PdfReader
from pptx import Presentation
from apps.core.errors import DomainError
from apps.core.models import Job,UsageLedger
from apps.core.services import submit_job,execute_job,storage_path
from apps.studio import workflows,provider
from apps.studio.domain import create_draft,update_draft,generation_quote,unpack,draft_data
from apps.studio.models import SavedDefinition,WorkflowRun,GenerationDraft,EducationProject,PracticeAttempt,ProviderUsage
from apps.studio.rendering import render_pptx
from operations.integrations import save_config
from tests.test_platform import account,upload,login_client
from tests.test_studio import paid,content

pytestmark=pytest.mark.django_db

@pytest.fixture(autouse=True)
def local(settings):
    settings.DEBUG=True;settings.ENABLE_BETA_TOOLS=True;settings.COMMERCE_SANDBOX_ENABLED=True;settings.LOCAL_SYNC_JOBS=True


def workflow(a,steps=None):
    return SavedDefinition.objects.create(account=a,kind='workflow',name='Safe workflow',definition={'steps':steps or [{'feature_id':'pdf.rotate','parameters':{'angle':90}},{'feature_id':'pdf.extract_pages','parameters':{'pages':'1'}}]})


def confirmation(a,row,asset):
    q=workflows.quote(a,row.id,[str(asset.id)])
    return {'version':q['version'],'input_ids':q['input_ids'],'confirmed':True,'confirmation_token':q['confirmation_token']}


def test_workflow_requires_real_confirmation_and_bound_inputs(settings):
    a=paid(settings,account());row=workflow(a);first=upload(a);second=upload(a,name='other.pdf');client=login_client(a)
    body={'version':row.version,'input_ids':[str(first.id)],'confirmed':True}
    url=f'/api/v1/workflows/{row.id}/run'
    assert client.post(url,body,content_type='application/json',HTTP_IDEMPOTENCY_KEY='missing-confirmation').status_code==409
    body=confirmation(a,row,first);body['input_ids']=[str(second.id)]
    assert client.post(url,body,content_type='application/json',HTTP_IDEMPOTENCY_KEY='changed-confirmation').status_code==409
    assert not WorkflowRun.objects.exists() and not Job.objects.exists()


def test_workflow_quote_rejects_invalid_parameters_and_terminal_output_chain(settings):
    a=paid(settings,account());asset=upload(a)
    for steps in ([{'feature_id':'pdf.rotate','parameters':{'angle':13}}],[{'feature_id':'pdf.to_images'},{'feature_id':'pdf.rotate'}]):
        row=workflow(a,steps)
        with pytest.raises(DomainError):workflows.quote(a,row.id,[str(asset.id)])
    assert not Job.objects.exists()


def test_workflow_once_even_after_definition_changes(settings):
    a=paid(settings,account());asset=upload(a);row=workflow(a);body=confirmation(a,row,asset)
    run,token=workflows.start(a,row.id,body,'once-workflow-key');run=workflows.execute(run,token)
    assert run.status=='succeeded' and len(run.job_ids)==2
    row.version+=1;row.save(update_fields=['version'])
    same,token=workflows.start(a,row.id,body,'once-workflow-key')
    assert same.id==run.id and token is None and Job.objects.count()==2
    assert UsageLedger.objects.filter(kind='consume',meter='file_tasks').count()==2
    with pytest.raises(DomainError,match='idempotency_conflict'):
        workflows.start(a,row.id,{**body,'version':9},'once-workflow-key')


def test_workflow_crash_replays_existing_settled_child(settings,monkeypatch):
    a=paid(settings,account());asset=upload(a);row=workflow(a);body=confirmation(a,row,asset)
    run,token=workflows.start(a,row.id,body,'crash-workflow-key');real=workflows.execute_job
    def crashed(identifier):
        real(identifier)
        raise RuntimeError('synthetic process failure')
    monkeypatch.setattr(workflows,'execute_job',crashed)
    with pytest.raises(RuntimeError):workflows.execute(run,token)
    monkeypatch.setattr(workflows,'execute_job',real)
    run,token=workflows.start(a,row.id,body,'crash-workflow-key');run=workflows.execute(run,token)
    assert run.status=='succeeded' and Job.objects.count()==2
    assert UsageLedger.objects.filter(kind='consume',meter='file_tasks').count()==2


def test_workflow_expired_tampered_and_foreign_confirmation(settings,monkeypatch):
    a=paid(settings,account());other=paid(settings,account(43));asset=upload(a);row=workflow(a);body=confirmation(a,row,asset)
    with pytest.raises(DomainError):workflows.start(other,row.id,body,'foreign-workflow-key')
    with pytest.raises(DomainError,match='confirmation_required'):workflows.start(a,row.id,{**body,'confirmation_token':body['confirmation_token']+'x'},'tampered-workflow-key')
    def expired(*args,**kwargs):raise signing.SignatureExpired()
    monkeypatch.setattr(workflows.signing,'loads',expired)
    with pytest.raises(DomainError,match='quote_expired'):workflows.start(a,row.id,body,'expired-workflow-key')
    assert not Job.objects.exists()


def live_configuration():
    save_config('ai',{'mode':'openai','model':'mock-model','api_key':'sk-local-test-not-a-live-key'})


def raw_content():
    return {'title':'Verified answer','answer_supported':True,'citations':[],'sections':[{'id':'s1','heading':'Answer','body':'A reviewed response.','notes':''}],'questions':[]}


def request_data():
    return {'output_locale':'ru','prompt':'Explain this','source_text':'Короткий текст.','excerpts':[],'content':raw_content(),'options':{'question_count':0}}


def test_provider_exact_schema_token_budget_and_no_redirect(monkeypatch):
    payload={'status':'completed','output':[{'type':'message','content':[{'type':'output_text','text':json.dumps(raw_content())}]}],'usage':{'input_tokens':1200,'output_tokens':200}}
    captured={}
    class Response:
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def read(self,limit):captured['read_limit']=limit;return json.dumps(payload).encode()
    class Opener:
        def open(self,request,timeout):captured.update(request=request,timeout=timeout);return Response()
    monkeypatch.setattr(provider.urllib.request,'build_opener',lambda handler:Opener())
    cfg={'model':'mock-model','api_key':'never-log-this'}
    value,usage=provider.generate(cfg,request_data(),'ai.pdf_topic',uuid.uuid4(),token_limit=24000)
    request=json.loads(captured['request'].data)
    assert request['store'] is False and request['text']['format']['strict'] is True and 'tools' not in request
    assert captured['request'].full_url=='https://api.openai.com/v1/responses' and captured['read_limit']==2_000_001
    assert usage['output_tokens']==200 and value['title']=='Verified answer'
    payload['output'][0]['content'][0]['text']=json.dumps({**raw_content(),'unexpected':'private'})
    with pytest.raises(DomainError,match='provider_failed'):provider.generate(cfg,request_data(),'ai.pdf_topic','test',token_limit=24000)
    with pytest.raises(ValueError):provider._NoRedirect().redirect_request(None,None,302,None,None,'https://evil.invalid')
    oversized=request_data();oversized['source_text']='Ж'*3000
    with pytest.raises(DomainError,match='generation_limit'):provider.enforce_input_budget(cfg,oversized,'ai.pdf_topic',4000)


def test_provider_question_overrun_releases_credits_and_hides_outputs(settings,monkeypatch):
    a=paid(settings,account());live_configuration();draft=create_draft(a,{'feature_id':'ai.pdf_topic','source_text':'Hello','options':{'length':1,'question_count':0}})
    quote=generation_quote(a,draft.id,draft.version)
    raw=raw_content();raw['questions']=[{'id':'q1','stem':'2+2?','options':['4'],'answer':'4','explanation':'Pairs','topic':'math','marks':1}]
    monkeypatch.setattr(provider,'generate',lambda *args,**kwargs:(raw,{'input_tokens':100,'output_tokens':100}))
    job,_=submit_job(a,quote.id,'provider-overrun');job=execute_job(job.id)
    assert job.status=='failed' and job.error_code=='generation_limit' and not job.artifacts.exists()
    assert not UsageLedger.objects.filter(job=job,kind='consume').exists()


def test_provider_model_pin_rejects_changed_config(settings):
    a=paid(settings,account());live_configuration();draft=create_draft(a,{'feature_id':'ai.pdf_topic','source_text':'Hello','options':{'length':1}});quote=generation_quote(a,draft.id,1)
    save_config('ai',{'mode':'openai','model':'different-model','api_key':''})
    with pytest.raises(DomainError,match='provider_changed'):submit_job(a,quote.id,'changed-provider-model')


def test_patch_authoring_fields_and_foreign_source_isolation(settings):
    a=account();other=account(43);foreign=upload(other);draft=create_draft(a,{'feature_id':'ai.pptx','source_text':'Original'})
    updated=update_draft(a,draft.id,{'version':1,'title':'Новый заголовок','source_text':'Новый текст','prompt':'Summarize','output_locale':'ru','output_format':'pptx','options':{'length':1,'template_id':'classroom'}})
    data=draft_data(updated)
    assert data['title']=='Новый заголовок' and data['content']['sections'][0]['body']=='Новый текст'
    assert data['output_format']=='pptx' and data['options']['template_style']['accent']=='#3b6588'
    assert data['options']['length']==1, 'local authoring follows the material given'
    with pytest.raises(DomainError):update_draft(a,draft.id,{'version':2,'source_ids':[str(foreign.id)]})
    assert GenerationDraft.objects.count()==1


def test_learner_pptx_never_contains_private_speaker_notes(tmp_path):
    material=content();material['sections'][0]['notes']='SECRET-TEACHER-KEY-123'
    material['sections'][0]['id']='s1'
    path=tmp_path/'learner.pptx';render_pptx(material,path,role='learner_material',style={'accent':'#3b6588'})
    deck=Presentation(path)
    assert all('SECRET-TEACHER' not in slide.notes_slide.notes_text_frame.text for slide in deck.slides)
    assert any(shape.has_text_frame and 'Учитель' in shape.text for slide in deck.slides for shape in slide.shapes)


def test_editor_rejects_unsafe_and_unentitled_draft_commands():
    a=account();asset=upload(a);client=login_client(a)
    created=client.post('/api/v1/editor/documents',{'file_id':str(asset.id)},content_type='application/json').json();url='/api/v1/editor/documents/'+created['id']
    for command in ({'type':'add_text','page':1,'x':1,'y':1,'text':'Hello','path':'/etc/passwd'},{'type':'redact','page':1,'x':1,'y':1,'width':10,'height':10}):
        response=client.patch(url,{'version':1,'commands':[command]},content_type='application/json')
        assert response.status_code in (400,403)
    assert client.get(url).json()['commands']==[]


def test_expired_projects_are_inaccessible_and_cleaned():
    from apps.studio.domain import pack
    from apps.core.services import cleanup_expired
    a=account();project=EducationProject.objects.create(account=a,title='Expired',encrypted_content=pack(content()),expires_at=timezone.now()-timedelta(seconds=1))
    PracticeAttempt.objects.create(project=project,idempotency_key='past-practice',request_hash='a'*64,result={'weak_topics':['addition']})
    client=login_client(a);assert client.get(f'/api/v1/education/projects/{project.id}').status_code==404
    cleanup_expired();assert not EducationProject.objects.exists() and not PracticeAttempt.objects.exists()


def source_document(a):
    from django.core.files.uploadedfile import SimpleUploadedFile
    from reportlab.pdfgen.canvas import Canvas
    from apps.core.services import upload_file
    stream=io.BytesIO();canvas=Canvas(stream);canvas.drawString(45,700,'The river flows north.');canvas.save()
    return upload_file(a,SimpleUploadedFile('source.pdf',stream.getvalue(),'application/pdf'))


@pytest.mark.parametrize('case',['supported','forged'])
def test_a_document_built_from_a_source_must_quote_it_exactly(settings,monkeypatch,case):
    a=paid(settings,account());asset=source_document(a);live_configuration()
    draft=create_draft(a,{'feature_id':'ai.pdf_topic','prompt':'Summarise the river report.','source_ids':[str(asset.id)]})
    raw=raw_content();raw['sections'][0]['body']='The river flows north.'
    raw['citations']=[{'asset_id':str(asset.id),'page':1,'quote':'The river flows north.'}]
    if case=='forged':raw['citations'][0]['quote']='A sentence absent from the source'
    monkeypatch.setattr(provider,'generate',lambda *args,**kwargs:(raw,{'input_tokens':100,'output_tokens':100}))
    q=generation_quote(a,draft.id,1);job,_=submit_job(a,q.id,'source-cite-'+case);job=execute_job(job.id)
    if case=='forged':
        assert job.status=='failed' and job.error_code=='invalid_source_citation'
        assert not UsageLedger.objects.filter(job=job,kind='consume').exists()
    else:
        assert job.status=='succeeded',job.error_code
        text=''.join(p.extract_text() for p in PdfReader(storage_path(job.artifacts.get().file.object_key)).pages)
        assert 'north' in text and 'p. 1' in text
        assert job.settled_meters['ai_credits']<=q.meters['ai_credits']


def test_provider_completion_does_not_overwrite_concurrent_draft_edit(settings,monkeypatch):
    a=paid(settings,account());live_configuration();draft=create_draft(a,{'feature_id':'ai.pdf_topic','source_text':'Original','options':{'length':1}})
    quote=generation_quote(a,draft.id,1)
    def while_editing(*args,**kwargs):
        update_draft(a,draft.id,{'version':1,'title':'Customer new title','source_text':'Customer new body'})
        return raw_content(),{'input_tokens':100,'output_tokens':100}
    monkeypatch.setattr(provider,'generate',while_editing)
    job,_=submit_job(a,quote.id,'concurrent-draft-edit');job=execute_job(job.id)
    assert job.status=='succeeded',job.error_code
    draft.refresh_from_db();data=unpack(draft.encrypted_data)
    assert draft.version==2 and data['title']=='Customer new title' and data['source_text']=='Customer new body'


def test_used_workflow_definition_can_be_deleted_without_erasing_history(settings):
    a=paid(settings,account());asset=upload(a);row=workflow(a);body=confirmation(a,row,asset)
    run,token=workflows.start(a,row.id,body,'workflow-delete-history');run=workflows.execute(run,token)
    response=login_client(a).delete(f'/api/v1/workflows/{row.id}')
    assert response.status_code==200,response.content
    run.refresh_from_db();assert run.definition_id is None and len(run.job_ids)==2 and run.snapshot['workflow_id']==str(row.id)
