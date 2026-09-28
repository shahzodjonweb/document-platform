"""The twenty slide layouts: every one drawn, readable, and safe to re-upload.

Each layout is rendered in every theme against accents chosen to break a
palette, and every run of text is checked against whatever it actually sits on
— its own panel, the shape beneath it, or the slide. A layout the model chose
but whose content cannot fill it must fall back, never fail.
"""
import io
import random
import zipfile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Emu

from apps.studio import layouts
from apps.studio.slide_layouts import DRAW, PHOTO_ZONES
from apps.studio.slides import contrast_ratio, render_pptx

THEMES = ['light', 'dark', 'bold']
ACCENTS = ['#FFFFE0', '#808080', '#255e49', '#6E1230']


def item(label='', text='', value=''):
    return {'label': label, 'text': text, 'value': value}


def section(n, layout, heading, body='', items=(), columns=(), query=''):
    return {'id': f's{n}', 'heading': heading, 'body': body, 'notes': 'Presenter note.',
            'layout': layout, 'items': list(items), 'columns': list(columns), 'image_query': query}


EVERY = [
    section(1, 'cover', 'Q3 Results', 'What changed, and what we do next.', query='office team'),
    section(2, 'agenda', 'What we will cover', items=[item('Revenue'), item('Customers', 'Who stayed'),
                                                       item('Product'), item('Next quarter')]),
    section(3, 'section', 'Part one', 'How the quarter went'),
    section(4, 'bullets', 'Revenue grew', 'Revenue up 18 percent\nCosts flat since April\nMargin widened'),
    section(5, 'two_column', 'Two things changed', items=[item('Enterprise', 'Renewals held.'),
                                                          item('Self-serve', 'Conversion improved.')]),
    section(6, 'statement', 'The headline', 'Net retention reached a new high.'),
    section(7, 'quote', 'What customers said', 'The first tool our whole team uses every day.',
            items=[item('Dana Reyes', 'Head of Operations')]),
    section(8, 'big_number', 'One number', items=[item('Net revenue retention', 'Up from 104%', '118%')]),
    section(9, 'stats', 'At a glance', items=[item('Revenue', 'yoy', '+18%'), item('Customers', '', '1,240'),
                                              item('Churn', '', '1.9%')]),
    section(10, 'timeline', 'How we got here', items=[item('Pilot', '', 'Jan'), item('Launch', '', 'Mar'),
                                                      item('Series A', '', 'Jun')]),
    section(11, 'process', 'Onboarding', items=[item('Sign up', 'Two fields'), item('Import', 'Files'),
                                                item('Go live', 'Day one')]),
    section(12, 'comparison', 'Build or buy', columns=['Build', 'Buy'],
            items=[item('Cost', 'High', 'Monthly'), item('Time', 'Months', 'Weeks')]),
    section(13, 'pros_cons', 'Expand to Europe', items=[item('Large market', '', 'pro'),
                                                        item('Regulatory cost', '', 'con')]),
    section(14, 'cards', 'What it does', items=[item('Merge', 'Combine files', 'Core'), item('Compress', 'Shrink'),
                                                item('Sign', 'Signatures')]),
    section(15, 'matrix', 'Where we stand', columns=['Internal', 'External'],
            items=[item('Strengths', 'Fast'), item('Weaknesses', 'Thin'), item('Opportunities', 'Schools'),
                   item('Threats', 'Rivals')]),
    section(16, 'chart', 'Revenue by region', columns=['USD thousands'],
            items=[item('North America', '', '420'), item('Europe', '', '310'), item('Asia', '', '185')]),
    section(17, 'table', 'Plans compared', columns=['Plan', 'Pages', 'Price'],
            items=[item('Free', '7', '$0'), item('Plus', '15', '$9')]),
    section(18, 'image_split', 'Our new office', 'Opened in August\nRoom for forty', query='modern office'),
    section(19, 'image_full', 'A quarter of building', 'The team at the offsite.', query='team meeting'),
    section(20, 'closing', 'Thank you', 'Questions welcome'),
]


def photo(width=1600, height=1000):
    buffer = io.BytesIO()
    Image.new('RGB', (width, height), (120, 150, 170)).save(buffer, 'JPEG', quality=80)
    return {'jpeg': buffer.getvalue(), 'width': width, 'height': height, 'provider': 'test', 'id': 1}


PHOTOS = {'s1': photo(), 's18': photo(1000, 1500), 's19': photo()}


def deck(sections, title='Deck'):
    return {'title': title, 'questions': [], 'citations': [], 'sections': sections}


def render(tmp_path, sections, name='deck', **kwargs):
    path = tmp_path / f'{name}.pptx'
    result = render_pptx(deck(sections), path, **kwargs)
    return path, result, Presentation(str(path))


# ---------------------------------------------------------------- catalogue


def test_there_are_twenty_layouts_and_each_one_can_be_drawn():
    assert len(layouts.LAYOUT_IDS) == 20
    assert set(layouts.LAYOUT_IDS) <= set(DRAW), set(layouts.LAYOUT_IDS) - set(DRAW)
    assert set(layouts.PHOTO_LAYOUTS) | {'cover'} == set(PHOTO_ZONES)


@pytest.mark.parametrize('kind', layouts.LAYOUT_IDS)
def test_every_fallback_chain_ends_somewhere_that_always_accepts(kind):
    seen, current = [], kind
    while current not in seen:
        seen.append(current)
        current = layouts.CATALOGUE[current].fallback
    assert 'section' in seen, f'{kind} falls back through {seen} and never reaches one that always accepts'
    assert layouts.accepts('section', {}, [], 3, 5, False)


def test_every_layout_is_described_to_the_model():
    text = layouts.guide(photos=2)
    for kind in layouts.LAYOUT_IDS:
        assert f'- {kind}:' in text
    assert 'image_split' not in layouts.guide(photos=0).split('\n', 1)[1].split('Vary')[0]


# ---------------------------------------------------------------- drawing


@pytest.mark.parametrize('theme', THEMES)
@pytest.mark.parametrize('accent', ACCENTS)
def test_every_layout_draws_in_every_theme(tmp_path, theme, accent):
    path, result, presentation = render(tmp_path, EVERY, style={'accent': accent, 'deck_theme': theme},
                                        photos=PHOTOS)
    assert result['page_count'] == len(EVERY)
    assert result['metadata']['layouts'] == [s['layout'] for s in EVERY], 'each chose and each held'
    for slide in presentation.slides:
        assert str(slide.shapes[0].fill.fore_color.rgb) == accent.lstrip('#').upper()
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                assert shape.text_frame.word_wrap is True
                assert shape.text_frame.auto_size == MSO_AUTO_SIZE.NONE


def _fill(shape):
    try:
        return str(shape.fill.fore_color.rgb) if shape.fill.type == MSO_FILL_TYPE.SOLID else None
    except Exception:
        return None


def _contains(shape, x, y):
    return shape.left <= x <= shape.left + shape.width and shape.top <= y <= shape.top + shape.height


def _required(run):
    size = run.font.size.pt if run.font.size else 12
    return 3.0 if size >= 18 or (run.font.bold and size >= 14) else 4.5


@pytest.mark.parametrize('theme', THEMES)
@pytest.mark.parametrize('accent', ACCENTS)
def test_every_run_of_text_reads_against_what_it_sits_on(tmp_path, theme, accent):
    """The contrast walker: each run against its own fill, the shape beneath it, or the slide."""
    _, _, presentation = render(tmp_path, EVERY, style={'accent': accent, 'deck_theme': theme}, photos=PHOTOS)
    failures = []
    for number, slide in enumerate(presentation.slides, 1):
        ground = str(slide.background.fill.fore_color.rgb)
        shapes = list(slide.shapes)
        for position, shape in enumerate(shapes):
            if getattr(shape, 'has_table', False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        behind = str(cell.fill.fore_color.rgb)
                        for paragraph in cell.text_frame.paragraphs:
                            for run in paragraph.runs:
                                ratio = contrast_ratio(str(run.font.color.rgb), behind)
                                if ratio < _required(run):
                                    failures.append((number, 'table', run.text, round(ratio, 2)))
                continue
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            x, y = shape.left + shape.width // 2, shape.top + shape.height // 2
            behind = _fill(shape) or next((_fill(under) for under in reversed(shapes[:position])
                                           if _fill(under) and _contains(under, x, y)), None) or ground
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if not run.text.strip():
                        continue
                    ratio = contrast_ratio(str(run.font.color.rgb), behind)
                    if ratio < _required(run):
                        failures.append((number, run.text[:30], round(ratio, 2)))
    assert not failures, failures


def test_no_text_ever_sits_on_a_photo(tmp_path):
    _, _, presentation = render(tmp_path, EVERY, photos=PHOTOS)
    for number, slide in enumerate(presentation.slides, 1):
        pictures = [shape for shape in slide.shapes if shape.shape_type == 13 and shape.width > Emu(914400 * 2)]
        for picture in pictures:
            for shape in slide.shapes:
                if not (shape.has_text_frame and shape.text_frame.text.strip()):
                    continue
                overlap = (shape.left < picture.left + picture.width and picture.left < shape.left + shape.width
                           and shape.top < picture.top + picture.height and picture.top < shape.top + shape.height)
                assert not overlap, f'slide {number}: {shape.text_frame.text[:30]!r} is on the photo'


def test_a_photo_is_cropped_to_its_zone_not_stretched(tmp_path):
    _, _, presentation = render(tmp_path, EVERY, photos=PHOTOS)
    portrait = [s for s in presentation.slides[17].shapes if s.shape_type == 13][0]
    assert portrait.crop_top > 0 and portrait.crop_bottom > portrait.crop_top, 'a tall photo is cropped top and bottom'
    assert portrait.crop_left == 0 and portrait.crop_right == 0


def test_a_deck_with_every_layout_and_photos_can_be_uploaded_again(tmp_path):
    """Charts are shapes, so no workbook is embedded for our own upload check to refuse."""
    from apps.core.services import engine_inspect
    path, result, _ = render(tmp_path, EVERY, photos=PHOTOS)
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        assert not [n for n in names if '/embeddings/' in n or n.endswith('.bin')], names
        media = [n for n in names if n.startswith('ppt/media/')]
        assert media and all(n.endswith('.jpg') for n in media), media
        for name in names:
            if name.endswith('.rels'):
                assert b'TargetMode="External"' not in archive.read(name)
    checked = engine_inspect(path)
    assert checked['kind'] == 'pptx' and checked['page_count'] == result['page_count']
    assert path.stat().st_size < 10 * 1024 * 1024


def test_the_licence_record_of_each_photo_travels_with_the_deck(tmp_path):
    _, result, _ = render(tmp_path, EVERY, photos=PHOTOS)
    slides = [entry['slide'] for entry in result['metadata']['photos']]
    assert slides == [1, 18, 19]
    assert all('jpeg' not in entry for entry in result['metadata']['photos'])


# ---------------------------------------------------------------- fallback


@pytest.mark.parametrize('chosen,content,expected', [
    ('stats', section(2, 'stats', 'x', items=[item('Revenue', '', 'grew'), item('Cost', '', 'fell')]), 'bullets'),
    ('chart', section(2, 'chart', 'x', columns=['Region', 'Change'],
                      items=[item('A', '', '-4'), item('B', '', '3')]), 'table'),
    ('chart', section(2, 'chart', 'x', items=[item('A', '', '-4'), item('B', '', '3')]), 'bullets'),
    ('chart', section(2, 'chart', 'x', items=[item('A', '', '4')]), 'bullets'),
    ('table', section(2, 'table', 'x', items=[item('a', 'b'), item('c', 'd')]), 'bullets'),
    ('quote', section(2, 'quote', 'x', 'Someone said it.'), 'statement'),
    ('cover', section(2, 'cover', 'x', 'line one\nline two'), 'section'),
    ('image_split', section(2, 'image_split', 'x', 'a\nb', query='x'), 'bullets'),
    ('image_full', section(2, 'image_full', 'x', 'One caption.', query='x'), 'statement'),
    ('matrix', section(2, 'matrix', 'x', items=[item('a', 'b'), item('c', 'd'), item('e', 'f')]), 'cards'),
    ('big_number', section(2, 'big_number', 'x', items=[item('Retention', '', 'high')]), 'statement'),
])
def test_a_layout_the_content_cannot_fill_falls_back(tmp_path, chosen, content, expected):
    sections = [section(1, 'cover', 'Deck', 'Sub'), content]
    _, result, _ = render(tmp_path, sections)
    assert result['page_count'] == 2
    assert result['metadata']['layouts'][1] == expected


def test_a_deck_that_never_chose_is_drawn_exactly_as_before(tmp_path):
    legacy = [{'id': f's{i}', 'heading': f'H{i}', 'body': body, 'notes': ''}
              for i, body in enumerate(['Sub', 'a\nb\nc', '', 'One line.', 'a\nb\nc\nd\ne'])]
    _, result, _ = render(tmp_path, legacy)
    assert result['metadata']['layouts'] == ['cover', 'bullets', 'section', 'statement', 'two_column']


def test_two_hundred_malformed_slides_never_break_the_count_or_the_file(tmp_path):
    """Whatever a model returns, the deck has the slides that were quoted."""
    from apps.core.services import engine_inspect
    chooser = random.Random(20260928)
    junk = ['', '0', '-3', 'lots', '1,5', '12%', 'x' * 200, '∞', '3.5k']
    sections = []
    for n in range(200):
        count = chooser.randint(0, 9)
        sections.append(section(
            n + 1, chooser.choice(layouts.LAYOUT_IDS), chooser.choice(['', 'Heading', 'H ' * 40]),
            chooser.choice(['', 'one', 'a\nb\nc', 'word ' * 90]),
            items=[item(chooser.choice(junk), chooser.choice(junk), chooser.choice(junk)) for _ in range(count)],
            columns=[chooser.choice(junk) for _ in range(chooser.randint(0, 4))],
            query=chooser.choice(['', 'office'])))
    for start in range(0, 200, 40):
        chunk = sections[start:start + 40]
        path, result, _ = render(tmp_path, chunk, name=f'fuzz{start}',
                                 photos={s['id']: photo() for s in chunk[::7]})
        assert result['page_count'] == len(chunk)
        assert engine_inspect(path)['page_count'] == len(chunk)


# ---------------------------------------------------------------- values


@pytest.mark.parametrize('value,number', [
    ('1,200', 1200), ('3.5k', 3500), ('$4.2m', 4.2e6), ('18%', 18), ('1,5', 1.5),
    ('12 000', 12000), ('2,5 млн', 2.5e6), ('lots', None), ('', None), ('-4', -4),
])
def test_figures_are_read_the_way_people_write_them(value, number):
    assert layouts.parse_number(value) == number


@pytest.mark.parametrize('raw,clean', [
    ('Modern Office', 'modern office'),
    ('офис в Ташкенте', ''),
    ('team <script> meeting!!', 'team script meeting'),
    ('one two three four five six seven', 'one two three four five'),
])
def test_a_photo_search_is_only_ever_a_few_plain_english_words(raw, clean):
    assert layouts.clean_query(raw) == clean


# ---------------------------------------------------------------- pipeline


@pytest.fixture
def live_customer(settings):
    from tests.test_page_fill import live
    return live(settings)


def _answer(sections):
    return ({'title': 'Deck', 'answer_supported': True, 'citations': [], 'questions': [],
             'sections': sections}, {'input_tokens': 10, 'output_tokens': 10})


@pytest.mark.django_db
def test_a_layout_the_model_chose_survives_to_the_deck(live_customer, monkeypatch):
    from apps.core.services import execute_job, storage_path, submit_job
    from apps.studio import provider
    from apps.studio.domain import SLIDES, create_draft, draft_data, generation_quote
    chosen = [section(1, 'cover', 'Deck', 'Sub'), EVERY[8], EVERY[9], EVERY[15]]
    for index, entry in enumerate(chosen):
        entry = dict(entry)
        chosen[index] = {**entry, 'id': f's{index + 1}'}
    monkeypatch.setattr(provider, 'generate', lambda *args, **kwargs: _answer(chosen))
    draft = create_draft(live_customer, {'feature_id': SLIDES, 'prompt': 'A 4 slide deck on results.'})
    quote = generation_quote(live_customer, draft.id, draft.version)
    job, _ = submit_job(live_customer, quote.id, f'layouts-{draft.id}')
    job = execute_job(job.id)

    assert job.status == 'succeeded', job.error_code
    stored = draft_data(type(draft).objects.get(pk=draft.pk))['content']['sections']
    assert [s['layout'] for s in stored] == ['cover', 'stats', 'timeline', 'chart']
    presentation = Presentation(str(storage_path(job.artifacts.get().file.object_key)))
    assert len(presentation.slides) == 4
    assert any('118%' in s.text_frame.text or '+18%' in s.text_frame.text
               for s in presentation.slides[1].shapes if s.has_text_frame)


@pytest.mark.django_db
def test_editing_a_deck_keeps_each_slide_layout(live_customer):
    from apps.studio.domain import SLIDES, create_draft, draft_data, update_draft
    draft = create_draft(live_customer, {'feature_id': SLIDES, 'prompt': 'A 3 slide deck.'})
    content = draft_data(draft)['content']
    content['sections'][1].update(layout='stats', items=[item('Revenue', '', '+18%'), item('Churn', '', '2%')])
    updated = update_draft(live_customer, draft.id, {'version': draft.version, 'content': content})
    kept = draft_data(updated)['content']['sections'][1]
    assert kept['layout'] == 'stats' and kept['items'][0]['value'] == '+18%'
    # The outline form of an edit keeps it too.
    outline = [{'id': s['id'], 'title': s['heading'], 'body': s['body'], 'layout': 'cards' if i == 2 else s['layout'],
                'items': s['items']} for i, s in enumerate(draft_data(updated)['content']['sections'])]
    again = update_draft(live_customer, draft.id, {'version': updated.version, 'outline': outline})
    assert draft_data(again)['content']['sections'][2]['layout'] == 'cards'


@pytest.mark.django_db
def test_a_layout_no_model_could_choose_is_refused_from_a_client(live_customer):
    from apps.core.errors import DomainError
    from apps.studio.domain import SLIDES, create_draft, draft_data, update_draft
    draft = create_draft(live_customer, {'feature_id': SLIDES, 'prompt': 'A 3 slide deck.'})
    content = draft_data(draft)['content']
    content['sections'][1]['layout'] = 'hexagon'
    with pytest.raises(DomainError, match='invalid_parameters'):
        update_draft(live_customer, draft.id, {'version': draft.version, 'content': content})


@pytest.mark.django_db
def test_a_document_never_grows_slide_fields(live_customer):
    from apps.studio.domain import DOCUMENT, create_draft, draft_data
    draft = create_draft(live_customer, {'feature_id': DOCUMENT, 'prompt': 'A 2 page guide.'})
    for entry in draft_data(draft)['content']['sections']:
        assert set(entry) == {'id', 'heading', 'body', 'notes'}


@pytest.mark.django_db
def test_a_revision_sends_each_slide_back_with_its_layout(live_customer, monkeypatch):
    import json as _json
    from apps.studio import provider
    from apps.studio.domain import SLIDES, create_draft, draft_data, pack, unpack
    draft = create_draft(live_customer, {'feature_id': SLIDES, 'prompt': 'A 3 slide deck.'})
    stored = unpack(draft.encrypted_data)
    stored['content']['sections'][1].update(layout='stats', body='x',
                                            items=[item('Revenue', '', '+18%'), item('Churn', '', '2%')])
    draft.encrypted_data = pack(stored)
    draft.save(update_fields=['encrypted_data'])
    revised = create_draft(live_customer, {'prompt': 'Make slide two punchier.',
                                           'options': {'revise_draft_id': str(draft.id)}})
    sent = _json.loads(provider.request_body({'model': 'm', 'api_key': 'k'}, unpack(revised.encrypted_data),
                                             SLIDES)['input'])
    assert sent['outline']['sections'][1]['layout'] == 'stats'
    assert sent['outline']['sections'][1]['items'][0]['value'] == '+18%'
    assert draft_data(revised)['output_format'] == 'pptx'
