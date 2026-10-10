"""Classic designs carry decorative art and photo covers, and text is still never set on a picture.

Designs on other compositions place their own art; tests/test_compositions.py and
tests/test_deck_designs.py check those.
"""
import io

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from apps.studio.deck_designs import DESIGN_IDS, DESIGNS
from apps.studio.slides import COVER_PANEL_ZONE, palette, render_pptx
from tests.test_deck_designs import deck


def full_deck():
    content = deck()
    content['sections'].append({'id': 's4', 'heading': 'Thank you', 'body': 'Questions welcome', 'notes': '',
                                'layout': 'closing', 'items': [], 'columns': [], 'image_query': ''})
    return content


def photo():
    from PIL import Image
    buffer = io.BytesIO()
    Image.new('RGB', (1200, 800), (40, 90, 140)).save(buffer, 'JPEG')
    return {'jpeg': buffer.getvalue(), 'width': 1200, 'height': 800}


def box(shape):
    return shape.left, shape.top, shape.left + shape.width, shape.top + shape.height


def overlaps(a, b):
    return a[0] < b[2] and b[0] < a[2] and a[1] < b[3] and b[1] < a[3]


def pictures(slide):
    return [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]


def written(slide):
    return [shape for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text.strip()]


CLASSIC = [key for key in DESIGN_IDS if DESIGNS[key]['composition'] == 'classic']


@pytest.mark.parametrize('design', CLASSIC)
def test_art_never_sits_under_text(design, tmp_path):
    path = tmp_path / 'deck.pptx'
    render_pptx(full_deck(), path, style={'deck_design': design})
    presentation = Presentation(str(path))
    art = DESIGNS[design]['art']
    for index, slide in enumerate(presentation.slides):
        assert str(slide.shapes[0].fill.fore_color.rgb) == DESIGNS[design]['accent'].lstrip('#').upper()
        for picture in pictures(slide):
            for text in written(slide):
                assert not overlaps(box(picture), box(text)), (design, index, text.text_frame.text)
    drawn = [len(pictures(slide)) for slide in presentation.slides]
    if not art:
        assert drawn == [0, 0, 0, 0]
    else:
        # Cover, divider and closing carry the pattern; a content slide only with `edge`.
        assert drawn == [1, 1 if art['edge'] else 0, 1, 1], (design, drawn)


def test_a_design_with_art_draws_it_in_place_of_the_cover_disc(tmp_path):
    from pptx.enum.shapes import MSO_SHAPE
    plain, patterned = tmp_path / 'plain.pptx', tmp_path / 'patterned.pptx'
    render_pptx(full_deck(), plain, style={'deck_design': 'forest'})
    render_pptx(full_deck(), patterned, style={'deck_design': 'corporate'})

    def discs(path):
        cover = Presentation(str(path)).slides[0]
        return [shape for shape in cover.shapes if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
                and shape.auto_shape_type == MSO_SHAPE.OVAL]
    assert len(discs(plain)) == 1 and not discs(patterned)


def test_the_same_pattern_is_stored_once(tmp_path):
    path = tmp_path / 'deck.pptx'
    content = full_deck()
    content['sections'].insert(3, {**content['sections'][2], 'id': 's9'})
    render_pptx(content, path, style={'deck_design': 'summit'})
    import zipfile
    with zipfile.ZipFile(path) as archive:
        media = [name for name in archive.namelist() if name.startswith('ppt/media/')]
    # Cover and closing share one picture, the two dividers another.
    assert len(media) == 2, media


@pytest.fixture
def classic_photo_design():
    """A classic design with a full-bleed photo cover, as the old photo category had."""
    from apps.studio.deck_designs import _art, _design
    DESIGNS['_classic_photo'] = {**_design('#1F4E79', '#F4A261', 'light', 'Georgia', 'Calibri', 'navy; travel',
                                    cover_photo=True, art=_art('waves', 'side', 0.25)), 'category': 'photo'}
    yield '_classic_photo'
    DESIGNS.pop('_classic_photo')


def test_a_photo_design_puts_the_picture_across_the_cover_and_the_title_on_a_panel(tmp_path, classic_photo_design):
    path = tmp_path / 'deck.pptx'
    result = render_pptx(full_deck(), path, style={'deck_design': classic_photo_design}, photos={'s1': photo()})
    assert result['metadata']['photos'][0]['slide'] == 1
    cover = Presentation(str(path)).slides[0]
    (picture,) = pictures(cover)
    assert picture.left == Inches(0.18) and picture.width == Inches(13.153)
    panel = Inches(COVER_PANEL_ZONE[0]), Inches(COVER_PANEL_ZONE[1]), \
        Inches(COVER_PANEL_ZONE[0] + COVER_PANEL_ZONE[2]), Inches(COVER_PANEL_ZONE[1] + COVER_PANEL_ZONE[3])
    shapes = list(cover.shapes)
    filled = [shape for shape in shapes if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
              and abs(shape.left - panel[0]) < 2 and abs(shape.top - panel[1]) < 2]
    assert filled and shapes.index(filled[0]) > shapes.index(picture), 'the panel is drawn over the photo'
    roles = palette(DESIGNS[classic_photo_design]['accent'], 'light')
    assert str(filled[0].fill.fore_color.rgb) == roles['cover_fill']
    for text in written(cover):
        if text.top < Inches(6.55):
            left, top, right, bottom = box(text)
            assert panel[0] <= left and right <= panel[2] and panel[1] <= top and bottom <= panel[3], text.text_frame.text


def test_without_a_picture_a_photo_design_has_its_ordinary_cover(tmp_path, classic_photo_design):
    path = tmp_path / 'deck.pptx'
    render_pptx(full_deck(), path, style={'deck_design': classic_photo_design})
    cover = Presentation(str(path)).slides[0]
    (art,) = pictures(cover)
    assert art.left > Inches(10), 'only the side pattern'


def test_a_dark_design_can_bring_its_own_ground(tmp_path):
    path = tmp_path / 'deck.pptx'
    render_pptx(full_deck(), path, style={'deck_design': 'chalkboard'})
    slide = Presentation(str(path)).slides[2]
    assert str(slide.background.fill.fore_color.rgb) == '1F3B2D'
