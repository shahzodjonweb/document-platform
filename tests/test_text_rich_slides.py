"""Slides carry real explanation, and a full slide still fits.

Customers found decks too thin to present from: five fragments of fourteen
words. A slide now holds four to six sentence-long bullets, and the item
layouts carry a sentence or two each. These tests fill every layout to the
length its guide allows, in all three languages, and require that nothing is
cut and nothing is set below a readable size.
"""
import pytest
from pptx import Presentation
from pptx.util import Pt

from apps.studio import layouts, pages
from apps.studio.slides import FLOOR_BODY_SIZE, MAX_BULLETS, bullets_of, render_pptx

SENTENCES = {
    'en': ('Regional teams cut delivery times by a third after moving stock closer to their largest '
           'customers last spring and hiring local drivers'),
    'ru': ('Региональные команды сократили сроки доставки на треть, перенеся склады ближе к крупнейшим '
           'клиентам прошлой весной и наняв местных водителей'),
    'uz': ('Hududiy jamoalar o‘tgan bahorda omborlarni eng yirik mijozlarga yaqinlashtirib, yetkazib '
           'berish muddatini uchdan bir qismga qisqartirdi va mahalliy haydovchilarni yolladi'),
}


def words(locale, count, offset=0):
    pool = (SENTENCES[locale] + ' ') * 6
    parts = pool.split()
    return ' '.join(parts[offset:offset + count])


def item(label, text, value=''):
    return {'label': label, 'text': text, 'value': value}


def full_deck(locale):
    """Every text-carrying layout at the most its guide asks the model for."""
    w = lambda count, offset=0: words(locale, count, offset)

    def slide(number, layout, heading, body='', items=None, columns=None):
        return {'id': f's{number}', 'heading': heading, 'body': body, 'notes': '', 'layout': layout,
                'items': items or [], 'columns': columns or [], 'image_query': ''}
    return {'title': w(5), 'questions': [], 'citations': [], 'sections': [
        slide(1, 'cover', w(7), w(14)),
        slide(2, 'bullets', w(8), '\n'.join(w(20, n * 3) for n in range(6))),
        slide(3, 'cards', w(5), items=[item(w(3, n), w(25, n)) for n in range(3)]),
        slide(4, 'cards', w(5), items=[item(w(3, n), w(25, n), 'new') for n in range(6)]),
        slide(5, 'process', w(6), items=[item(w(3, n), w(18, n)) for n in range(5)]),
        slide(6, 'stats', w(6), items=[item(w(3, n), w(15, n), value) for n, value in
                                       enumerate(('-33%', '118%', '4 200', '4.8'))]),
        slide(7, 'two_column', w(6), items=[item(w(3), w(50)), item(w(3, 4), w(50, 4))]),
        slide(8, 'matrix', w(6), items=[item(w(2, n), w(22, n)) for n in range(4)]),
        slide(9, 'timeline', w(6), items=[item(w(3, n), w(12, n), f'Q{n + 1}') for n in range(6)]),
        slide(10, 'agenda', w(6), items=[item(w(4, n), w(15, n)) for n in range(6)]),
        slide(11, 'big_number', w(6), items=[item(w(4), w(35), '118%')]),
        slide(12, 'closing', w(4), w(10)),
    ]}


@pytest.mark.parametrize('locale', ['en', 'ru', 'uz'])
def test_a_slide_filled_to_its_guide_keeps_every_word_at_a_readable_size(locale, tmp_path):
    path = tmp_path / f'{locale}.pptx'
    result = render_pptx(full_deck(locale), path, locale=locale, style={'deck_design': 'midnight'})
    assert not result['shortened'], 'text the guide allows must not be cut'
    assert result['metadata']['layouts'][1:11] == ['bullets', 'cards', 'cards', 'process', 'stats',
                                                   'two_column', 'matrix', 'timeline', 'agenda', 'big_number']
    deck = Presentation(str(path))
    for number, slide in enumerate(deck.slides, 1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            assert '…' not in shape.text_frame.text, (locale, number)
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if len(run.text) > 40:
                        assert run.font.size >= Pt(12), (locale, number, run.font.size.pt)


def test_a_list_slide_is_six_sentences_not_five_fragments():
    sentence = words('en', 20)
    kept, shortened = bullets_of('\n'.join([sentence] * MAX_BULLETS))
    assert len(kept) == 6 and not shortened and all(line == sentence for line in kept)
    kept, shortened = bullets_of('\n'.join([sentence] * 7))
    assert len(kept) == 6 and shortened


def test_a_list_that_still_runs_long_shrinks_before_it_is_cut(tmp_path):
    long = '\n'.join(words('ru', 20, n) + ' ' + words('ru', 18, n + 3) for n in range(6))
    deck = {'title': 'T', 'questions': [], 'citations': [], 'sections': [
        {'id': 's1', 'heading': 'Heading', 'body': long, 'notes': '', 'layout': 'bullets',
         'items': [], 'columns': [], 'image_query': ''}]}
    result = render_pptx(deck, tmp_path / 'long.pptx', locale='ru')
    sizes = {run.font.size.pt for shape in Presentation(str(tmp_path / 'long.pptx')).slides[0].shapes
             if shape.has_text_frame for paragraph in shape.text_frame.paragraphs
             for run in paragraph.runs if len(run.text) > 60}
    assert sizes and min(sizes) >= FLOOR_BODY_SIZE
    assert result['shortened'], 'what still cannot fit is cut, and the customer is told'


def test_bullets_use_the_leading_the_fit_was_measured_with(tmp_path):
    deck = {'title': 'T', 'questions': [], 'citations': [], 'sections': [
        {'id': 's1', 'heading': 'Heading', 'body': '\n'.join([words('en', 18)] * 5), 'notes': '',
         'layout': 'bullets', 'items': [], 'columns': [], 'image_query': ''}]}
    render_pptx(deck, tmp_path / 'leading.pptx')
    for shape in Presentation(str(tmp_path / 'leading.pptx')).slides[0].shapes:
        if shape.has_text_frame and len(shape.text_frame.paragraphs) == 5:
            for paragraph in shape.text_frame.paragraphs:
                run = paragraph.runs[0]
                assert abs(paragraph.line_spacing.pt - run.font.size.pt * 1.22) < 0.05
            break
    else:
        raise AssertionError('no bullet block found')


def test_the_model_is_asked_for_substance():
    from apps.studio.provider import writing_guidance
    guidance = writing_guidance({}, 5, 'pptx', 0, 5)
    assert '4 to 6 bullets' in guidance and 'full sentence of 12 to 20 words' in guidance
    assert f'more than {round(pages.target_words("pptx") * 1.3)} words' in guidance
    reference = layouts.reference()
    assert 'at most twice a deck' in reference and 'bullets on at most a third' not in reference
