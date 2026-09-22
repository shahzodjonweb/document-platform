"""Real branding artifacts and revision isolation, including unselected content."""
import copy
import io
import pytest
from PIL import Image
from pypdf import PdfReader
from pptx import Presentation
from django.core.files.uploadedfile import SimpleUploadedFile
from apps.core.errors import DomainError
from apps.core.services import upload_file,storage_path,submit_job,execute_job
from apps.studio.branding import prepare_branding,generation_inputs,render_style
from apps.studio.domain import create_draft,update_draft,draft_data,generation_quote
from apps.studio.revisions import prepare_revision,preserve_unselected,provider_content
from apps.studio.rendering import render_pdf,render_pptx
from tests.test_platform import account
from tests.test_studio import paid

pytestmark=pytest.mark.django_db

@pytest.fixture(autouse=True)
def local(settings):
    settings.DEBUG=True;settings.COMMERCE_SANDBOX_ENABLED=True;settings.ENABLE_BETA_TOOLS=True


def logo(a):
    picture=Image.new('RGB',(100,40),'#285cff');raw=io.BytesIO();picture.save(raw,format='PNG')
    return upload_file(a,SimpleUploadedFile('school-logo.png',raw.getvalue(),'image/png'))


def branding(a,asset):
    return prepare_branding(a,'ai.pdf_text',{'template_style':{'accent':'#255e49'},'branding':{'name':'Школа · Maktab','accent':'#285cff','logo_asset_id':str(asset.id)}})


def test_branding_real_logo_and_unicode_name_in_pdf_pptx(settings,tmp_path):
    a=paid(settings,account());asset=logo(a);options=branding(a,asset)
    payload={'source_ids':[],'options':options};style=render_style(a,payload)
    assert generation_inputs(payload)==[str(asset.id)]
    content={'title':'Learning','sections':[{'id':'s1','heading':'Topic','body':'A branded learning document.','notes':'Speaker notes'}],'questions':[]}
    pdf=tmp_path/'brand.pdf';slides=tmp_path/'brand.pptx'
    render_pdf(content,pdf,style=style);render_pptx(content,slides,style=style)
    reader=PdfReader(pdf);assert 'Школа' in reader.pages[0].extract_text() and len(reader.pages[0].images)==1
    deck=Presentation(slides)
    assert any(s.shape_type==13 for s in deck.slides[0].shapes)
    assert any('Школа' in s.text for s in deck.slides[0].shapes if s.has_text_frame)
    assert str(deck.slides[0].shapes[0].fill.fore_color.rgb)=='285CFF'


def test_branding_owner_plan_and_bytes_are_verified(settings):
    a=paid(settings,account());other=paid(settings,account(43));asset=logo(a)
    with pytest.raises(DomainError,match='file_unavailable'):branding(other,asset)
    with pytest.raises(DomainError,match='feature_not_in_plan'):branding(account(44),asset)
    options=branding(a,asset);storage_path(asset.object_key).write_bytes(b'changed-private-content')
    with pytest.raises(DomainError,match='file_changed'):render_style(a,{'options':options})
    with pytest.raises(DomainError,match='invalid_parameters'):prepare_branding(a,'ai.pdf_text',{'template_style':{'accent':'#000000'},'branding':{'logo_url':'https://example.invalid/private'}})


def source(a,fmt='pptx'):
    d=create_draft(a,{'feature_id':'ai.pptx' if fmt=='pptx' else 'ai.pdf_text','source_text':'Original first\n\nUntouched second','output_format':fmt,'options':{'length':2}})
    return d


def test_revision_forks_owner_version_and_exact_selection(settings):
    a=paid(settings,account());other=paid(settings,account(43));original=source(a)
    payload={'options':{'source_draft_id':str(original.id),'selected_section_ids':['s1']}}
    revised=prepare_revision(a,'ai.regenerate_slide',payload)
    assert len(provider_content(revised)['sections'])==1
    candidate=copy.deepcopy(revised['content']);candidate['sections'][0]['body']='Revised first';candidate['sections'][1]['body']='Do not replace this'
    merged=preserve_unselected(revised,candidate)
    assert merged['sections'][0]['body']=='Revised first' and merged['sections'][1]['body']=='Untouched second'
    assert draft_data(original)['content']['sections'][0]['body']=='Original first'
    with pytest.raises(DomainError,match='not_found'):prepare_revision(other,'ai.regenerate_slide',payload)
    with pytest.raises(DomainError,match='version_conflict'):prepare_revision(a,'ai.regenerate_slide',{'options':{**payload['options'],'source_draft_version':9}})
    with pytest.raises(DomainError,match='invalid_parameters'):prepare_revision(a,'ai.regenerate_slide',{'options':{**payload['options'],'selected_section_ids':['unknown']}})
    with pytest.raises(DomainError,match='invalid_parameters'):preserve_unselected(revised,candidate,provider=True)


def test_revision_actual_export_preserves_unselected_and_original(settings):
    a=paid(settings,account());original=source(a)
    d=create_draft(a,{'feature_id':'ai.regenerate_slide','prompt':'Simplify','options':{'source_draft_id':str(original.id),'selected_section_ids':['s1'],'length':2}})
    value=draft_data(d)['content'];value['sections'][0]['body']='Revised first';value['sections'][1]['body']='Malicious unrelated replacement'
    d=update_draft(a,d.id,{'version':1,'content':value});q=generation_quote(a,d.id,d.version)
    job,_=submit_job(a,q.id,'revision-export');job=execute_job(job.id);assert job.status=='succeeded',job.error_code
    deck=Presentation(storage_path(job.artifacts.get().file.object_key))
    texts='\n'.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
    assert 'Revised first' in texts and 'Untouched second' in texts and 'Malicious' not in texts
    assert draft_data(original)['content']['sections'][0]['body']=='Original first'


def test_branding_quote_binds_logo_and_failed_asset_releases(settings):
    a=paid(settings,account());asset=logo(a)
    d=create_draft(a,{'feature_id':'ai.pdf_text','source_text':'Branded content','options':{'length':1,'branding':{'name':'School','accent':'#285cff','logo_asset_id':str(asset.id)}}})
    q=generation_quote(a,d.id,d.version)
    assert q.input_ids==[str(asset.id)] and q.input_fingerprints[0]['sha256']==asset.sha256
    storage_path(asset.object_key).write_bytes(b'corruption')
    job,_=submit_job(a,q.id,'branding-corrupt');job=execute_job(job.id)
    assert job.status=='failed' and not job.artifacts.exists()


def test_revision_rejects_role_sensitive_and_non_document_sources(settings):
    from apps.studio.domain import pack,unpack
    a=paid(settings,account());original=source(a)
    payload={'options':{'source_draft_id':str(original.id),'selected_section_ids':['s1']}}
    for feature_id in ('teacher.lesson_plan','school.project_slides'):
        original.feature_id=feature_id;original.save(update_fields=['feature_id'])
        with pytest.raises(DomainError,match='invalid_parameters'):prepare_revision(a,'ai.rewrite',payload)
    original.feature_id='ai.images'
    data=unpack(original.encrypted_data);data['output_format']='png';original.encrypted_data=pack(data)
    original.save(update_fields=['feature_id','encrypted_data'])
    with pytest.raises(DomainError,match='invalid_parameters'):prepare_revision(a,'ai.rewrite',payload)


def test_revision_version_must_be_exact_integer(settings):
    a=paid(settings,account());original=source(a)
    for version in (True,'1',1.0):
        with pytest.raises(DomainError,match='version_conflict'):
            prepare_revision(a,'ai.rewrite',{'options':{'source_draft_id':str(original.id),'source_draft_version':version,'selected_section_ids':['s1']}})
