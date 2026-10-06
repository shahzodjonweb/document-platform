"""Dense card and argument slides retain their text inside the visible page."""
import os
import subprocess
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from apps.studio.layouts import clean_slide_fields
from apps.studio.slides import render_pptx


CARD_TITLE = 'Manage permissions across all your business teams'
CARD_BODY = ('Configure access policies and manage permissions consistently across departments '
             'while maintaining clear ownership boundaries.')
ARGUMENT = ('Centralized administration improves operational efficiency while providing visibility '
            'across every regional business unit.')
TRANSLATIONS = [
    ('en', CARD_TITLE, CARD_BODY, ARGUMENT),
    ('uz', 'Barcha bo‘limlar uchun kirish huquqlarini birgalikda boshqaring',
     'Kirish qoidalarini sozlang va bo‘limlarda ruxsatlarni izchil boshqaring hamda har bir mas’uliyatni aniq belgilang.',
     'Markazlashgan boshqaruv har bir hududiy bo‘limda ish samaradorligini oshirib barcha jamoalarning faoliyatini ko‘rsatadi.'),
    ('ru', 'Управляйте разрешениями во всех командах вашей компании',
     'Настраивайте правила доступа и согласованно управляйте разрешениями во всех подразделениях сохраняя понятную ответственность команды.',
     'Централизованное управление повышает эффективность работы обеспечивая прозрачность для каждого регионального подразделения нашей компании.'),
]


def section(kind, items):
    return {'id': kind, 'heading': kind.replace('_', ' ').title(), 'body': '', 'notes': '',
            **clean_slide_fields({'layout': kind, 'items': items, 'columns': []})}


def cards(count=6, tag='Enterprise', label=CARD_TITLE, text=CARD_BODY):
    return section('cards', [{'label': label, 'text': text, 'value': tag} for _ in range(count)])


def arguments(count=5, label=ARGUMENT):
    return section('pros_cons', [{'label': label, 'value': 'pro'} for _ in range(count)]
                   + [{'label': 'Additional operating costs', 'value': 'con'}])


def render(tmp_path, sections, locale='en'):
    path = tmp_path / 'fit.pptx'
    content = {'title': 'Renderer review', 'questions': [], 'citations': [], 'sections': [
        {'id': 'cover', 'heading': 'Renderer review', 'body': '', 'notes': '', 'layout': 'cover'},
        *sections,
    ]}
    result = render_pptx(content, path, locale=locale, style={'brand_name': 'Brand footer'})
    presentation = Presentation(str(path))
    assert result['page_count'] == len(presentation.slides) == len(sections) + 1
    return path, result, presentation


def body_frames(slide):
    return [shape for shape in slide.shapes if shape.has_text_frame and shape.text
            and Inches(2.5) <= shape.top < Inches(6.6)]


@pytest.mark.parametrize('count', [3, 4, 5, 6])
@pytest.mark.parametrize('tag', ['', 'Enterprise'])
def test_cards_preserve_realistic_limit_length_fields(tmp_path, count, tag):
    _, result, presentation = render(tmp_path, [cards(count, tag)])
    assert result['metadata']['layouts'] == ['cover', 'cards']
    assert result['shortened'] is False
    frames = body_frames(presentation.slides[1])
    assert len(frames) == count
    for shape in frames:
        assert CARD_TITLE in shape.text and CARD_BODY in shape.text
        assert (tag.upper() in shape.text) if tag else True
        assert shape.top + shape.height <= Inches(6.6)
        for paragraph in shape.text_frame.paragraphs:
            size = paragraph.runs[0].font.size.pt
            assert size >= (10 if paragraph.text == tag.upper() else 12)
            assert paragraph.line_spacing.pt >= size


@pytest.mark.parametrize('count', [5, 7])
def test_pros_cons_preserve_all_accepted_arguments(tmp_path, count):
    # Seven arguments on one side plus one on the other fills the eight-item contract.
    label = ARGUMENT if count == 5 else 'Clear ownership for every regional team'
    _, result, presentation = render(tmp_path, [arguments(count, label)])
    assert result['shortened'] is False
    frames = body_frames(presentation.slides[1])
    assert len(frames) == 2
    assert frames[0].text.count(label) == count
    assert 'Additional operating costs' in frames[1].text
    sizes = []
    for shape in frames:
        assert shape.top + shape.height <= Inches(6.4)
        for paragraph in shape.text_frame.paragraphs[1:]:
            sizes.append(paragraph.runs[0].font.size.pt)
    assert min(sizes) >= 13
    assert len(set(sizes)) == 1


@pytest.mark.parametrize('kind', ['cards', 'pros_cons'])
def test_unusually_long_words_are_safely_elided_and_reported(tmp_path, kind):
    # Word limits alone cannot protect a box from long identifiers.
    long_words = ' '.join(['VeryLongUnbrokenBusinessIdentifier'] * 7)
    content = cards(label=long_words, text=long_words) if kind == 'cards' else arguments(7, long_words)
    _, result, presentation = render(tmp_path, [content])
    assert result['shortened'] is True
    frames = body_frames(presentation.slides[1])
    assert any('…' in shape.text for shape in frames)
    assert len([shape for shape in frames if len(shape.text_frame.paragraphs) > 1]) == (6 if kind == 'cards' else 2)
    assert all(paragraph.text for shape in frames for paragraph in shape.text_frame.paragraphs)


@pytest.mark.parametrize('locale, title, body, argument', TRANSLATIONS)
def test_dense_translations_remain_complete(tmp_path, locale, title, body, argument):
    _, result, presentation = render(tmp_path, [cards(label=title, text=body), arguments(label=argument)], locale)
    assert result['shortened'] is False
    frames = body_frames(presentation.slides[1])
    assert sum(title in shape.text and body in shape.text for shape in frames) == 6
    assert sum(shape.text.count(argument) for shape in body_frames(presentation.slides[2])) == 5


@pytest.mark.parametrize('locale, title, body, argument', TRANSLATIONS)
@pytest.mark.parametrize('font_mode', ['office', 'dejavu'])
def test_native_render_keeps_dense_cards_and_arguments_inside_their_panels(
        tmp_path, monkeypatch, locale, title, body, argument, font_mode):
    """Optional integration test using a supplied headless LibreOffice binary.

    Set PPTX_TEST_SOFFICE to the bundled runtime's absolute soffice path. This
    checks native PDF glyph positions rather than repeating our fit estimate.
    """
    soffice = os.environ.get('PPTX_TEST_SOFFICE')
    if not soffice:
        pytest.skip('Set PPTX_TEST_SOFFICE for native slide-fit verification')
    import pdfplumber
    if font_mode == 'dejavu':
        # Linux substitutes these wider faces for the Office fonts. Explicitly
        # select them so the same regression runs on macOS as well as in CI.
        from apps.studio import slides, deck_designs
        monkeypatch.setattr(deck_designs, 'DEFAULT_FONTS', ('DejaVu Serif', 'DejaVu Sans'))
        monkeypatch.setattr(slides, 'BODY_FONT', 'DejaVu Sans')
        monkeypatch.setattr(slides, 'HEADING_FONT', 'DejaVu Serif')

    path, result, presentation = render(tmp_path, [cards(label=title, text=body), arguments(label=argument)], locale)
    assert result['shortened'] is False
    profile = (tmp_path / 'lo-profile').as_uri()
    subprocess.run([str(Path(soffice)), f'-env:UserInstallation={profile}', '--headless',
                    '--convert-to', 'pdf', '--outdir', str(tmp_path), str(path)],
                   check=True, capture_output=True, timeout=60)
    with pdfplumber.open(path.with_suffix('.pdf')) as document:
        assert len(document.pages) == 3
        if font_mode == 'dejavu':
            assert any('DejaVuSans' in char['fontname'] for char in document.pages[1].chars)
        for index, page in enumerate(document.pages[1:], 1):
            assert all(0 <= char['top'] < char['bottom'] <= page.height for char in page.chars)
            frames = body_frames(presentation.slides[index])
            words = [word for word in page.extract_words() if 178 <= word['top'] < 485]
            assert words
            for word in words:
                assert any(shape.left.pt - 1 <= word['x0'] and word['x1'] <= (shape.left + shape.width) / 12700 + 1
                           and shape.top.pt - 1 <= word['top']
                           and word['bottom'] <= (shape.top + shape.height) / 12700 + 3
                           for shape in frames), word
        # Some fallback-font PDFs have incomplete Unicode CMaps. Check visible
        # glyph coverage as well as bounds; exact authored text is pinned above.
        assert len(document.pages[1].chars) >= 0.9 * (len(title) + len(body)) * 6
        assert len(document.pages[2].chars) >= 0.9 * len(argument) * 5
