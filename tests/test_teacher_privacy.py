"""Teacher-only content never becomes a public artifact, including old links."""
import hashlib
import pytest
from django.utils import timezone
from datetime import timedelta
from django.test import Client
from pypdf import PdfReader
from pptx import Presentation
from apps.core.services import submit_job,execute_job,storage_path
from apps.studio.domain import create_draft,generation_quote
from apps.studio.models import ShareGrant
from tests.test_platform import account,login_client
from tests.test_studio import paid

pytestmark=pytest.mark.django_db


@pytest.fixture(autouse=True)
def local(settings):
    settings.DEBUG=True;settings.COMMERCE_SANDBOX_ENABLED=True


def generated(owner,feature,output_format='pdf'):
    questions=[{'id':'q1','stem':'Review question','options':['A','B'],'answer':'PRIVATE-ANSWER-428','explanation':'PRIVATE-REASON-916','marks':1}]
    draft=create_draft(owner,{'feature_id':feature,'source_text':'Teacher review notes about learner progress.','output_format':output_format,'options':{'length':1,'question_count':1,'questions':questions}})
    quote=generation_quote(owner,draft.id,draft.version)
    job,_=submit_job(owner,quote.id,'privacy-'+feature+'-'+output_format)
    done=execute_job(job.id);assert done.status=='succeeded',done.error_code
    return done


@pytest.mark.parametrize('feature',['teacher.feedback','teacher.suggest_marks','teacher.rubric','teacher.lesson_plan','teacher.syllabus'])
def test_teacher_only_exports_are_private_even_with_legacy_public_roles(settings,feature):
    owner=paid(settings,account());job=generated(owner,feature);client=login_client(owner)
    main=job.artifacts.exclude(role='teacher_key').get()
    assert main.role=='teacher_feedback'
    response=client.post('/api/v1/shares',{'artifact_id':str(main.id)},content_type='application/json')
    assert response.status_code==403 and response.json()['error']['code']=='share_role_forbidden'
    # Older workers used learner_material for every teacher feature. Existing
    # share links must fail too, without relying on a one-time data migration.
    main.role='learner_material';main.save(update_fields=['role'])
    token='legacy-private-feedback-link'
    ShareGrant.objects.create(account=owner,artifact=main,token_hash=hashlib.sha256(token.encode()).hexdigest(),expires_at=timezone.now()+timedelta(hours=1))
    assert client.post('/api/v1/shares',{'artifact_id':str(main.id)},content_type='application/json').status_code==403
    assert Client().get('/api/v1/shared/'+token).status_code==404


@pytest.mark.parametrize('output_format',['pdf','pptx'])
def test_answer_key_is_one_private_artifact_containing_actual_answers(settings,output_format):
    owner=paid(settings,account());job=generated(owner,'teacher.answer_key',output_format)
    assert job.artifacts.count()==1
    artifact=job.artifacts.get();assert artifact.role=='teacher_key'
    path=storage_path(artifact.file.object_key)
    if output_format=='pdf':text='\n'.join(page.extract_text() for page in PdfReader(path).pages)
    else:text='\n'.join(shape.text for slide in Presentation(path).slides for shape in slide.shapes if shape.has_text_frame)
    assert 'PRIVATE-ANSWER-428' in text and 'PRIVATE-REASON-916' in text
    assert login_client(owner).post('/api/v1/shares',{'artifact_id':str(artifact.id)},content_type='application/json').status_code==403


def test_learner_worksheet_remains_shareable_with_private_separate_key(settings):
    owner=paid(settings,account());job=generated(owner,'teacher.worksheet')
    learner=job.artifacts.get(role='learner_material');key=job.artifacts.get(role='teacher_key')
    learner_text='\n'.join(page.extract_text() for page in PdfReader(storage_path(learner.file.object_key)).pages)
    assert 'PRIVATE-ANSWER-428' not in learner_text and 'PRIVATE-REASON-916' not in learner_text
    client=login_client(owner)
    shared=client.post('/api/v1/shares',{'artifact_id':str(learner.id)},content_type='application/json')
    assert shared.status_code==200,shared.content
    assert Client().get(shared.json()['url']).status_code==200
    assert client.post('/api/v1/shares',{'artifact_id':str(key.id)},content_type='application/json').status_code==403
