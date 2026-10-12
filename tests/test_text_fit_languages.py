"""Text stays inside its box in every language, at every length a slide accepts.

Russian and Uzbek run a third longer than English, in longer words, and a box
sized for English text overflows in both unless it is measured. These tests
render decks that use every layout, at the lengths the writing guidance asks
the model for and at the hard limits the platform lets through
(`layouts.LIMITS`, the 160-character heading, 40-word bullets), in English,
Russian and Uzbek, on classic and composition designs, and read every text
frame and table cell back.

A box is measured the way the fitter measures it (`slide_layouts._text_lines`:
wrapped lines times the leading, plus paragraph spacing). Exact leading is
taken as written; relative leading is taken at 1.2 times the point size, which
is what a viewer's font metrics give one line. A box may come up 3 pt short,
the allowance the fitter itself keeps for ascenders and descenders. Text over
the limits may be cut, but only visibly, and the deck must say it was.

Documents: a callout of sixteen long Uzbek points breaks across pages instead
of failing the render, and long-word Russian and Uzbek prose written to the
page target keeps the page count that was asked for.
"""
import time
from types import SimpleNamespace

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu, Inches, Pt
from pypdf import PdfReader

from apps.studio import pages
from apps.studio.deck_designs import DESIGN_IDS, look
from apps.studio.doc_designs import clean_fields
from apps.studio.layouts import LIMITS, clean_slide_fields
from apps.studio.rendering import render_pdf
from apps.studio.slide_layouts import Ctx, Z, _fit_paragraphs, _paragraph_height, _text_lines, fit_headline
from apps.studio.slides import palette, render_pptx
from tests.test_text_rich_slides import full_deck, words as moderate_words

# ---------------------------------------------------------------- measuring

EMU_PER_PT = 12700
VIEWER_LEADING = 1.2
ALLOWANCE = 3
# A run that names no size is drawn at the template's default.
TEMPLATE_SIZE = 18
# The quote layout's 110 pt opening mark sits in the top of its own line.
DECORATIVE = {'“'}
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'


def _pieces(paragraph):
    """(text, size, bold) for each run and field of a paragraph, as a viewer draws it."""
    for element in paragraph._p:
        if element.tag == A + 'br':
            yield '\n', 0, False
        elif element.tag in (A + 'r', A + 'fld'):
            props = element.find(A + 'rPr')
            size = int(props.get('sz')) / 100 if props is not None and props.get('sz') else TEMPLATE_SIZE
            yield element.findtext(A + 't') or '', size, props is not None and props.get('b') in ('1', 'true')


def needed(paragraph, width):
    """Points of height one paragraph takes in a box `width` EMU wide."""
    pieces = list(_pieces(paragraph))
    drawn = [piece for piece in pieces if piece[0].strip()]
    if not drawn:
        return 0.0
    size = max(piece[1] for piece in drawn)
    props = paragraph._p.pPr
    inset = int(props.get('marL', 0)) + int(props.get('marR', 0)) if props is not None else 0
    lines = _text_lines(''.join(piece[0] for piece in pieces), size, max(1, width - inset),
                        bold=any(piece[2] for piece in drawn))
    spacing = paragraph.line_spacing
    leading = spacing.pt if hasattr(spacing, 'pt') else (spacing or 1.0) * size * VIEWER_LEADING
    before = paragraph.space_before.pt if paragraph.space_before is not None else 0
    after = paragraph.space_after.pt if paragraph.space_after is not None else 0
    return lines * leading + before + after


def _frame_problem(frame, width, height, where):
    text = ' '.join(''.join(piece[0] for piece in _pieces(paragraph)) for paragraph in frame.paragraphs).strip()
    if not text or text in DECORATIVE:
        return None
    need, box = sum(needed(paragraph, width) for paragraph in frame.paragraphs), height / EMU_PER_PT
    if need > box + ALLOWANCE:
        return f'{where}: needs {need:.1f} pt in a {box:.1f} pt box: {text[:60]!r}'
    return None


def _shapes(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shapes(shape.shapes)
        else:
            yield shape


def overflowing(deck, layouts=None):
    """Every text frame and table cell whose text needs more height than its box has."""
    presentation = deck if hasattr(deck, 'slides') else Presentation(str(deck))
    found = []
    for number, slide in enumerate(presentation.slides, 1):
        where = f'slide {number} ({layouts[number - 1] if layouts else "?"})'
        for shape in _shapes(slide.shapes):
            if shape.has_text_frame:
                frame = shape.text_frame
                found.append(_frame_problem(frame, shape.width - frame.margin_left - frame.margin_right,
                                            shape.height - frame.margin_top - frame.margin_bottom, where))
            if getattr(shape, 'has_table', False) and shape.has_table:
                grid = shape.table
                for r, row in enumerate(grid.rows):
                    for c, column in enumerate(grid.columns):
                        cell = grid.cell(r, c)
                        found.append(_frame_problem(
                            cell.text_frame, column.width - cell.margin_left - cell.margin_right,
                            row.height - cell.margin_top - cell.margin_bottom, f'{where} cell {r},{c}'))
    return [problem for problem in found if problem]


def drawn_text(path):
    texts = []
    for slide in Presentation(str(path)).slides:
        for shape in _shapes(slide.shapes):
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if getattr(shape, 'has_table', False) and shape.has_table:
                texts.extend(cell.text_frame.text for row in shape.table.rows for cell in row.cells)
    return '\n'.join(texts)


# ---------------------------------------------------------------- content

# Formal Russian and Uzbek run eight to nine characters a word; Uzbek carries
# o‘ g‘ (U+2018) and the tutuq belgisi ʼ (U+02BC).
PROSE = {
    'en': ('Regional distribution teams reduced average delivery times by roughly a third after relocating '
           'warehouse inventory closer to their largest institutional customers and recruiting experienced '
           'local drivers, which strengthened the reliability of the whole supply chain considerably'),
    'ru': ('Региональные распределительные подразделения сократили среднее время доставки приблизительно '
           'на треть, переместив складские запасы ближе к крупнейшим институциональным клиентам и привлекая '
           'опытных местных водителей, что существенно укрепило надёжность всей логистической цепочки'),
    'uz': ('Hududiy taqsimot bo‘linmalari omborlardagi zaxiralarni eng yirik muassasaviy mijozlarga '
           'yaqinlashtirib, yetkazib berishning o‘rtacha muddatini taxminan uchdan bir qismga qisqartirdi va '
           'tajribali mahalliy haydovchilarni ishga yollab, taʼminot zanjirining barqarorligini g‘ayrat bilan '
           'mustahkamladi'),
}
LOCALES = list(PROSE)
FIGURES = {'en': ('-33%', '118%', '4,200', '4.8'), 'ru': ('-33%', '118%', '4 200', '4,8'),
           'uz': ('-33%', '118%', '4 200', '4,8')}
# Designs that draw headings in the weight they are measured in: two classic
# looks and two compositions (arched and mosaic covers, dividers and closings).
REGULAR_DESIGNS = ('scholar', 'boardroom', 'counsel', 'samarkand')
# Designs whose heading face is set bold after the text was fitted.
BOLD_DESIGNS = ('midnight', 'broadsheet', 'sterling')
EVERY_LAYOUT = ['cover', 'agenda', 'section', 'bullets', 'two_column', 'statement', 'quote', 'big_number',
                'stats', 'timeline', 'process', 'comparison', 'pros_cons', 'cards', 'matrix', 'chart', 'table',
                'closing', 'questions']


def words(locale, count, offset=0):
    return ' '.join(((PROSE[locale] + ' ') * 30).split()[offset:offset + count])


def up_to(text, characters):
    """Whole words of `text`, no longer than `characters`."""
    while len(text) > characters:
        text = text.rsplit(' ', 1)[0]
    return text


def over(locale, offset=0):
    """Text longer than any field accepts: what reaches the renderer is the limit itself."""
    return ' '.join(((PROSE[locale] + ' ') * 30).split()[offset:])[:1000]


def item(label='', text='', value=''):
    return {'label': label, 'text': text, 'value': value}


def slide(number, layout, heading, body='', items=(), columns=()):
    """A section as validate_content leaves it: heading capped, slide fields bounded."""
    return {'id': f's{number}', 'heading': heading[:160], 'body': body, 'notes': '',
            **clean_slide_fields({'layout': layout, 'items': list(items), 'columns': list(columns)})}


def deck(sections, questions=(), title='Deck'):
    return {'title': title, 'citations': [], 'questions': list(questions), 'sections': sections}


def guide_deck(locale):
    """Every layout filled to what the writing guidance and the layout guide ask for."""
    w = lambda count, offset=0: words(locale, count, offset)
    figures = FIGURES[locale]
    return deck([
        slide(1, 'cover', w(8), w(14)),
        slide(2, 'agenda', w(8), items=[item(w(4, n), w(15, n)) for n in range(6)]),
        slide(3, 'section', w(8), w(12)),
        slide(4, 'bullets', w(8), '\n'.join(w(20, n * 3) for n in range(6))),
        slide(5, 'two_column', w(8), items=[item(w(3), w(50)), item(w(3, 4), w(50, 4))]),
        # At most 20 words, and no more than the 120 characters a statement accepts.
        slide(6, 'statement', w(8), up_to(w(20), 120)),
        slide(7, 'quote', w(8), w(45), [item(w(4), w(12))]),
        slide(8, 'big_number', w(8), items=[item(w(6), w(35), '118%')]),
        slide(9, 'stats', w(8), items=[item(w(3, n), w(15, n), figures[n]) for n in range(4)]),
        slide(10, 'timeline', w(8), items=[item(w(3, n), w(12, n), f'Q{n + 1} 2025') for n in range(6)]),
        slide(11, 'process', w(8), items=[item(w(3, n), w(18, n)) for n in range(5)]),
        slide(12, 'comparison', w(8), items=[item(w(3, n), w(10, n), w(10, n + 5)) for n in range(4)],
              columns=[w(2), w(2, 3)]),
        slide(13, 'pros_cons', w(8), items=[item(w(12, n), '', 'pro') for n in range(3)]
              + [item(w(12, n + 7), '', 'con') for n in range(3)]),
        slide(14, 'cards', w(8), items=[item(w(3, n), w(25, n), w(1, n)) for n in range(6)]),
        slide(15, 'matrix', w(8), items=[item(w(2, n), w(22, n)) for n in range(4)], columns=[w(3), w(3, 2)]),
        slide(16, 'chart', w(8), items=[item(w(3, n), '', str(10 + n * 7)) for n in range(8)], columns=[w(1)]),
        slide(17, 'table', w(8), items=[item(w(3, n), w(8, n), w(2, n)) for n in range(5)],
              columns=[w(2), w(2, 2), w(2, 4)]),
        slide(18, 'closing', w(4), w(10)),
    ], questions=[{'stem': w(15, n), 'options': [w(3, n + k) for k in range(4)], 'answer': 'A', 'marks': 1}
                  for n in range(5)], title=w(8))


def at_limits(layout, locale):
    """One slide of `layout` with every field at the most the platform lets through."""
    w = lambda count, offset=0: words(locale, count, offset)
    full = lambda offset=0: over(locale, offset)
    heading = over(locale)
    figures = FIGURES[locale]
    build = {
        'cover': lambda: slide(1, 'cover', heading, full()),
        'agenda': lambda: slide(2, 'agenda', heading, items=[item(full(n), full(n)) for n in range(6)]),
        'section': lambda: slide(2, 'section', heading, full()),
        # bullets_of passes six bullets of 40 words; the fitter has to hold them.
        'bullets': lambda: slide(2, 'bullets', heading, '\n'.join(w(40, n * 3) for n in range(6))),
        'two_column': lambda: slide(2, 'two_column', heading, items=[item(full(), full()), item(full(4), full(4))]),
        # The most a statement (120 characters), quote (45 words) and big number (12 characters) accept.
        'statement': lambda: slide(2, 'statement', heading, up_to(full(), 120)),
        'quote': lambda: slide(2, 'quote', heading, w(45), [item(full(), full())]),
        'big_number': lambda: slide(2, 'big_number', heading, items=[item(full(), full(), '118 000 000')]),
        'stats': lambda: slide(2, 'stats', heading,
                               items=[item(full(n), full(n), f'{figures[n]} {full(n)}') for n in range(4)]),
        'timeline': lambda: slide(2, 'timeline', heading,
                                  items=[item(full(n), full(n), f'20{n}0 {full(n)}') for n in range(6)]),
        'process': lambda: slide(2, 'process', heading, items=[item(full(n), full(n)) for n in range(5)]),
        'comparison': lambda: slide(2, 'comparison', heading,
                                    items=[item(full(n), full(n), full(n + 5)) for n in range(8)],
                                    columns=[full(), full(3)]),
        'pros_cons': lambda: slide(2, 'pros_cons', heading,
                                   items=[item(full(n), full(n), 'pro') for n in range(4)]
                                   + [item(full(n + 7), full(n), 'con') for n in range(4)],
                                   columns=[full(), full(2)]),
        'cards': lambda: slide(2, 'cards', heading, items=[item(full(n), full(n), full(n)) for n in range(6)]),
        'matrix': lambda: slide(2, 'matrix', heading, items=[item(full(n), full(n)) for n in range(4)],
                                columns=[full(), full(2)]),
        'chart': lambda: slide(2, 'chart', heading,
                               items=[item(full(n), full(n), f'{10 + n * 7} {full(n)}') for n in range(8)],
                               columns=[full()]),
        'table': lambda: slide(2, 'table', heading, items=[item(full(n), full(n), full(n)) for n in range(6)],
                               columns=[full(), full(2), full(4)]),
        'closing': lambda: slide(2, 'closing', heading, full()),
    }
    return build[layout]()


def short_cover(locale):
    return slide(1, 'cover', words(locale, 4), words(locale, 6))


def render(tmp_path, content, locale, design, name='deck'):
    path = tmp_path / f'{name}-{locale}-{design}.pptx'
    result = render_pptx(content, path, locale=locale, style={'deck_design': design})
    return path, result


# ---------------------------------------------------------------- the measure


def test_the_measure_finds_what_it_is_for():
    """A measure that never complains proves nothing: it must flag a box that is too small."""
    presentation = Presentation()
    page = presentation.slides.add_slide(presentation.slide_layouts[6])

    def box(text, width, height, size=14, shapes=None):
        shapes = shapes if shapes is not None else page.shapes
        frame = shapes.add_textbox(Inches(1), Inches(1), Inches(width), Inches(height)).text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        run = frame.paragraphs[0].add_run()
        run.text, run.font.size = text, Pt(size)
        return frame

    box(words('ru', 60), 3, 0.5)                   # many lines in one line's height
    box(words('en', 4), 6, 0.5)                    # one line with room to spare
    box('“', 1.0, 1.4, size=110)                   # the quote mark is decoration
    box(words('en', 60), 3, 0.5, shapes=page.shapes.add_group_shape().shapes)  # inside a group
    grid = page.shapes.add_table(1, 1, Inches(1), Inches(4), Inches(2), Inches(0.4)).table
    grid.cell(0, 0).text = words('uz', 40)
    found = overflowing(presentation)
    assert len(found) == 3, found
    assert 'Региональные' in found[0] and 'Regional' in found[1] and 'cell 0,0' in found[2]


# ---------------------------------------------------------------- at the guide's lengths


@pytest.mark.parametrize('design', REGULAR_DESIGNS)
@pytest.mark.parametrize('locale', LOCALES)
def test_every_layout_at_its_guide_length_stays_inside_its_boxes(tmp_path, locale, design):
    path, result = render(tmp_path, guide_deck(locale), locale, design)
    assert result['metadata']['layouts'] == EVERY_LAYOUT, 'each slide is drawn with the layout it chose'
    assert not overflowing(path, result['metadata']['layouts']), '\n'.join(
        overflowing(path, result['metadata']['layouts']))
    if '…' in drawn_text(path):
        assert result['shortened'], 'text that was visibly cut must be reported'


@pytest.mark.parametrize('locale', LOCALES)
def test_every_other_layout_keeps_guide_length_text_whole(tmp_path, locale):
    """The layouts tests/test_text_rich_slides.py does not fill are held to the same promise.

    With the same sentences, nothing a slide's guide allows is cut, in any language.
    """
    w = lambda count, offset=0: moderate_words(locale, count, offset)
    rich = full_deck(locale)['sections']
    extra = [
        slide(13, 'section', w(6), w(12)),
        slide(14, 'statement', w(6), up_to(w(20), 120)),
        slide(15, 'quote', w(6), w(45), [item(w(3), w(8))]),
        slide(16, 'comparison', w(6), items=[item(w(3, n), w(10, n), w(10, n + 5)) for n in range(4)],
              columns=[w(2), w(2, 3)]),
        slide(17, 'pros_cons', w(6), items=[item(w(12, n), '', 'pro') for n in range(3)]
              + [item(w(12, n + 7), '', 'con') for n in range(3)]),
        slide(18, 'chart', w(6), items=[item(w(3, n), '', str(10 + n * 7)) for n in range(8)], columns=[w(1)]),
        slide(19, 'table', w(6), items=[item(w(3, n), w(8, n), w(2, n)) for n in range(5)],
              columns=[w(2), w(2, 2), w(2, 4)]),
    ]
    questions = [{'stem': w(15, n), 'options': [w(3, n + k) for k in range(4)], 'answer': 'A', 'marks': 1}
                 for n in range(5)]
    path, result = render(tmp_path, deck(rich[:-1] + extra + rich[-1:], questions), locale, 'scholar')
    layouts = result['metadata']['layouts']
    assert layouts[11:] == ['section', 'statement', 'quote', 'comparison', 'pros_cons', 'chart', 'table',
                            'closing', 'questions']
    assert not result['shortened'], 'text the guide allows must not be cut'
    assert '…' not in drawn_text(path)
    assert not overflowing(path, layouts), '\n'.join(overflowing(path, layouts))


@pytest.mark.xfail(strict=True, reason=(
    'Left as it is: matrix (slide_layouts.matrix) gives axis names 0.35 in of the cells\' height, so four cells '
    'of 22 words — what the matrix guide allows, with axes it calls optional — are cut in Russian and Uzbek'))
@pytest.mark.parametrize('locale', ['ru', 'uz'])
def test_a_matrix_with_axis_names_keeps_guide_length_cells_whole(tmp_path, locale):
    w = lambda count, offset=0: moderate_words(locale, count, offset)
    matrix = slide(2, 'matrix', w(6), items=[item(w(2, n), w(22, n)) for n in range(4)],
                   columns=[w(2), w(2, 3)])
    path, result = render(tmp_path, deck([short_cover(locale), matrix]), locale, 'scholar')
    assert result['metadata']['layouts'] == ['cover', 'matrix']
    assert not result['shortened'], 'text the guide allows must not be cut'


def test_designs_with_bold_headings_keep_guide_length_text_inside_its_boxes(tmp_path):
    """slides.apply_fonts sets the heading face bold after drawing, so it is fitted bold from the start.

    Fitted at regular weight, cover, closing and divider titles, quotes and
    big-number labels wrapped to extra lines bold and ran out of their boxes.
    """
    found = []
    for design in BOLD_DESIGNS:
        assert look({'deck_design': design})['heading_bold'], design
        for locale in LOCALES:
            path, result = render(tmp_path, guide_deck(locale), locale, design)
            found += [f'{design} {locale} {problem}'
                      for problem in overflowing(path, result['metadata']['layouts'])]
    assert not found, '\n'.join(found)


def _bold_design_of_each_composition():
    chosen = {}
    for design in sorted(DESIGN_IDS):
        design_look = look({'deck_design': design})
        if design_look['heading_bold'] and design_look['composition'] not in ('classic', 'editorial', 'sidebar'):
            chosen.setdefault(design_look['composition'], design)
    return sorted(chosen.items())


@pytest.mark.parametrize('composition, design', _bold_design_of_each_composition())
def test_every_composition_keeps_bold_headings_inside_their_boxes(tmp_path, composition, design):
    """The compositions the test above does not draw, in Russian and Uzbek: their titles are their own boxes."""
    for locale in ('ru', 'uz'):
        path, result = render(tmp_path, guide_deck(locale), locale, design)
        assert not overflowing(path, result['metadata']['layouts']), '\n'.join(
            overflowing(path, result['metadata']['layouts']))


# ---------------------------------------------------------------- at the hard limits

ACROSS_LIMITS = ('scholar', 'counsel')


@pytest.mark.parametrize('layout', EVERY_LAYOUT[:-1])
def test_text_at_the_input_limits_is_cut_visibly_and_never_runs_out_of_its_box(tmp_path, layout):
    """Four of these once ran out of their boxes here.

    Six 40-word bullets, measured by total width rather than word wrap, settled
    a line taller than the body. Stats and chart values with words after the
    figure were fitted to the width only. A 14-word Russian closing title
    stopped shrinking at 18 pt in the arch composition's title box, uncut.
    """
    found = []
    for locale in LOCALES:
        if layout == 'cover':
            content = deck([at_limits('cover', locale), slide(2, 'closing', words(locale, 4))],
                           title=over(locale)[:160])
        else:
            content = deck([short_cover(locale), at_limits(layout, locale)])
        for design in ACROSS_LIMITS:
            path, result = render(tmp_path, content, locale, design)
            assert result['metadata']['layouts'] == ['cover', 'closing' if layout == 'cover' else layout], \
                (locale, design)
            found += [f'{design} {locale} {problem}'
                      for problem in overflowing(path, result['metadata']['layouts'])]
            text = drawn_text(path)
            if layout == 'statement':
                # A statement accepts no more than it can show: drawn whole, nothing to report.
                assert '…' not in text and not result['shortened'], (locale, design)
            else:
                assert result['shortened'], f'{locale} {design}: the cut must be reported'
                assert '…' in text, f'{locale} {design}: and visible on the slide'
    assert not found, '\n'.join(found)


@pytest.mark.parametrize('locale', LOCALES)
def test_questions_far_longer_than_a_slide_are_cut_visibly_and_stay_inside(tmp_path, locale):
    questions = [{'stem': over(locale, n)[:400], 'options': [over(locale, n + k)[:120] for k in range(4)],
                  'answer': 'A', 'marks': 1} for n in range(5)]
    path, result = render(tmp_path, deck([short_cover(locale)], questions), locale, 'scholar')
    assert result['metadata']['layouts'] == ['cover', 'questions']
    assert not overflowing(path, result['metadata']['layouts']), '\n'.join(
        overflowing(path, result['metadata']['layouts']))
    assert result['shortened'] and '…' in drawn_text(path)


def test_a_question_at_the_validation_limits_still_fits_its_slide(tmp_path):
    """validate_content keeps a 2000-character stem and eight 500-character options."""
    stem = (over('uz') + ' ' + over('uz', 40))[:2000]
    question = {'stem': stem, 'options': [over('uz', k)[:500] for k in range(8)], 'answer': 'A', 'marks': 1}
    path, result = render(tmp_path, deck([short_cover('uz')], [question]), 'uz', 'counsel')
    assert result['metadata']['layouts'] == ['cover', 'questions']
    assert not overflowing(path, result['metadata']['layouts']), '\n'.join(
        overflowing(path, result['metadata']['layouts']))
    assert result['shortened'] and '…' in drawn_text(path)


def test_five_questions_at_the_validation_limits_render_in_seconds(tmp_path):
    """Cutting a word and measuring everything again took about 15 s for this one slide.

    The premium plan allows 150 questions: thirty such slides. The time is
    this process's CPU time, so a busy machine does not fail it.
    """
    questions = [{'stem': (over('ru', n) + ' ' + over('ru', n + 40))[:2000],
                  'options': [over('ru', n + k)[:500] for k in range(8)], 'answer': 'A', 'marks': 1}
                 for n in range(5)]
    start = time.process_time()
    path, result = render(tmp_path, deck([short_cover('ru')], questions), 'ru', 'scholar')
    assert time.process_time() - start < 3
    assert result['metadata']['layouts'] == ['cover', 'questions'] and result['shortened']
    assert not overflowing(path, result['metadata']['layouts'])


def _cut_word_by_word(zone, paragraphs):
    """How a block used to be cut: the tallest paragraph's last word, then everything measured again."""
    fitted = [dict(paragraph) for paragraph in paragraphs]
    available = Emu(zone.height).pt - 3
    while sum(_paragraph_height(paragraph, zone) for paragraph in fitted) > available:
        candidates = [paragraph for paragraph in fitted if len(paragraph['text'].rstrip('…')) > 1]
        if not candidates:
            break
        paragraph = max(candidates, key=lambda value: _paragraph_height(value, zone))
        text = paragraph['text'].rstrip('…')
        parts = text.split()
        paragraph['text'] = (' '.join(parts[:-1]) if len(parts) > 1 else text[:-1]).rstrip(',;:.') + '…'
    return [paragraph['text'] for paragraph in fitted]


@pytest.mark.parametrize('locale', LOCALES)
def test_the_fast_cut_leaves_what_a_word_by_word_cut_left(locale):
    """Searching for where to stop must not cut more than cutting one word at a time did."""
    question = [dict(text=over(locale)[:600], size=12, floor=12, bold=True, after=2),
                dict(text='  ·  '.join(over(locale, k)[:200] for k in range(4)), size=11, floor=11)]
    equal = [dict(text=words(locale, 30, n % 2), size=14, floor=14) for n in range(3)]  # ties go to the first
    identifier = [dict(text='VeryLongUnbrokenBusinessIdentifier' * 3, size=14, floor=14, bold=True)]
    label = [dict(text=words(locale, 10), size=28, floor=28, after=8),
             dict(text=words(locale, 40, 3), size=16, floor=16, spacing=1.22)]
    for zone, paragraphs in ((Z(0, 0, 10.9, 1.2), question), (Z(0, 0, 3.2, 2.0), equal),
                             (Z(0, 0, 1.5, 0.5), identifier), (Z(0, 0, 5.8, 1.6), label)):
        ctx = SimpleNamespace(shortened=False, heading_bold=False)
        assert [paragraph['text'] for paragraph in _fit_paragraphs(ctx, zone, paragraphs)] == \
            _cut_word_by_word(zone, paragraphs)
        assert ctx.shortened


def test_a_headline_still_too_tall_at_its_floor_is_cut_visibly_and_reported():
    """A title used to stop at 18 pt, whole, and run out of a narrow composition box."""
    presentation = Presentation()
    page = presentation.slides.add_slide(presentation.slide_layouts[6])

    def context(bold=False):
        return Ctx(page, palette('255E49'), {}, 0, 2, 18, 'ru', heading_bold=bold)

    ctx, zone = context(), Z(1, 1, 3.4, 1.0)
    paragraph = fit_headline(ctx, zone, over('ru')[:160], size=44, lines=3)
    size = paragraph.runs[0].font.size.pt
    assert size == 18 and paragraph.text.endswith('…') and ctx.shortened
    assert _text_lines(paragraph.text, size, zone.width) * size * 1.08 <= 72
    # Bold is measured as drawn, and a title that fits keeps its size and every word.
    ctx = context(bold=True)
    paragraph = fit_headline(ctx, Z(1, 1, 8, 2.1), 'Quarterly review', size=44, lines=3)
    assert (paragraph.text, paragraph.runs[0].font.size.pt, paragraph.runs[0].font.bold) == \
        ('Quarterly review', 44, True)
    assert not ctx.shortened


@pytest.mark.parametrize('design', ['scholar', 'midnight'])
def test_text_that_fits_keeps_the_sizes_it_always_had(tmp_path, design):
    """Measuring by height, word wrap and weight changes nothing for text that already fit."""
    content = deck([
        slide(1, 'cover', 'Quarterly review', 'Results for the board'),
        slide(2, 'section', 'Results', 'What changed'),
        slide(3, 'statement', 'One idea', 'Delivery is a third faster'),
        slide(4, 'quote', 'Voices', 'It simply arrives on time now', [item('Ana Lee', 'Buyer')]),
        slide(5, 'big_number', 'Growth', items=[item('Faster delivery', 'Across every region', '33%')]),
        slide(6, 'stats', 'Numbers', items=[item('Faster', '', '-33%'), item('Growth', '', '118%')]),
        slide(7, 'chart', 'By region', items=[item('North', '', '42'), item('South', '', '17')]),
        slide(8, 'bullets', 'Next', 'Hire drivers\nMove stock'),
        slide(9, 'closing', 'Thank you', 'Questions welcome'),
    ])
    path, result = render(tmp_path, content, 'en', design)
    assert result['metadata']['layouts'] == ['cover', 'section', 'statement', 'quote', 'big_number', 'stats',
                                             'chart', 'bullets', 'closing']
    assert not result['shortened']
    sizes = {}
    for page in Presentation(str(path)).slides:
        for shape in _shapes(page.shapes):
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    if paragraph.runs and paragraph.text not in sizes:
                        sizes[paragraph.text] = paragraph.runs[0].font.size.pt
    assert {text: sizes[text] for text in (
        'Quarterly review', 'Results', 'Delivery is a third faster', 'It simply arrives on time now', '33%',
        'Faster delivery', '-33%', '118%', '42', 'Thank you', 'Hire drivers')} == {
        'Quarterly review': 44, 'Results': 34, 'Delivery is a third faster': 40,
        'It simply arrives on time now': 30, '33%': 110, 'Faster delivery': 28, '-33%': 52, '118%': 52, '42': 14,
        'Thank you': 44, 'Hire drivers': 20}
    numbers = Presentation(str(path)).slides[1]._element.iter(A + 'fld')
    assert [field.find(A + 'rPr').get('sz') for field in numbers] == ['1000']


def test_the_limits_these_decks_are_built_at_are_the_platforms():
    """If a limit moves, the decks above must be rebuilt at the new one."""
    assert LIMITS == {'label': 160, 'text': 600, 'value': 160}
    for locale in LOCALES:
        assert len(over(locale)) > max(LIMITS.values()) and len(at_limits('cards', locale)['heading']) == 160
        cards = at_limits('cards', locale)['items']
        assert [len(cards[0][key]) for key in ('label', 'text', 'value')] == [160, 600, 160]
        assert all(len(line.split()) == 40 for line in at_limits('bullets', locale)['body'].split('\n'))
    assert all(design in DESIGN_IDS for design in REGULAR_DESIGNS + BOLD_DESIGNS)
    assert not any(look({'deck_design': design})['heading_bold'] for design in REGULAR_DESIGNS)
    assert look({'deck_design': 'counsel'})['composition'] == 'arch'
    assert look({'deck_design': 'samarkand'})['composition'] == 'mosaic'


@pytest.mark.xfail(strict=True, reason=(
    'Left as it is: a cover, divider or closing slide draws only the first line of its body (render_pptx keeps '
    'lines[:1] for a cover; slide_layouts.cover, section_ and closing use lines[0]) and render_pptx counts only '
    'bullets_of\'s cut as shortening, so every further line is dropped with no ellipsis and shortened=False'))
@pytest.mark.parametrize('layout', ['cover', 'section', 'closing'])
def test_a_body_line_a_one_line_slide_cannot_show_is_reported(tmp_path, layout):
    body = 'Hududiy bo‘linmalar muddatni qisqartirdi\nMahalliy haydovchilar taʼminotni mustahkamladi'
    sections = [slide(1, 'cover', 'Taʼminot', body if layout == 'cover' else 'Yo‘l xaritasi'),
                slide(2, 'bullets' if layout == 'cover' else layout, 'Keyingi qadam',
                      'Bir\nIkki' if layout == 'cover' else body)]
    path, result = render(tmp_path, deck(sections), 'uz', 'scholar')
    assert 'Mahalliy haydovchilar' in drawn_text(path) or result['shortened'], \
        'a line that is not drawn has to be reported as cut'


def test_the_slide_number_is_drawn_at_the_size_of_the_footer_it_sits_in(tmp_path):
    """The field carries its own size: its paragraph has no run, and the template's 18 pt overflowed the box."""
    path, _ = render(tmp_path, deck([short_cover('en'), slide(2, 'bullets', 'Heading', 'One\nTwo')]), 'en',
                     'scholar')
    field = Presentation(str(path)).slides[1]._element.find('.//' + A + 'fld')
    assert field is not None
    assert field.find(A + 'rPr').get('sz') == '1000'


# ---------------------------------------------------------------- documents

DOCUMENT_PROSE = {
    'ru': ['Региональные распределительные подразделения сократили среднее время доставки приблизительно на треть.',
           'Складские запасы переместились ближе к крупнейшим институциональным клиентам.',
           'Опытные местные водители существенно укрепили надёжность всей логистической цепочки.'],
    'uz': ['Hududiy taqsimot bo‘linmalari yetkazib berishning o‘rtacha muddatini taxminan uchdan bir qismga '
           'qisqartirdi.',
           'Omborlardagi zaxiralar eng yirik muassasaviy mijozlarga yaqinlashtirildi.',
           'Tajribali mahalliy haydovchilar taʼminot zanjirining barqarorligini g‘ayrat bilan mustahkamladi.'],
}


def prose(locale, characters):
    """Paragraphs of five sentences, one per line, as the model writes a section."""
    sentences, paragraphs, total = DOCUMENT_PROSE[locale], [], 0
    while total < characters:
        paragraph = ' '.join(sentences[(len(paragraphs) * 5 + n) % len(sentences)] for n in range(5))
        paragraphs.append(paragraph)
        total += len(paragraph) + 1
    return '\n'.join(paragraphs)[:characters].rsplit(' ', 1)[0]


def test_a_callout_of_sixteen_long_uzbek_points_breaks_across_pages(tmp_path):
    """A box longer than a page used to be one table cell, and failed the whole job with a LayoutError."""
    points = [item(f'{n + 1}-band. ' + prose('uz', 140), prose('uz', 900)[:600]) for n in range(16)]
    section = {'id': 's1', 'heading': 'Asosiy xulosalar', 'body': prose('uz', 600), 'notes': '',
               **clean_fields({'layout': 'callout', 'items': points})}
    assert len(section['items']) == 16 and all(len(point['text']) == 600 for point in section['items'])
    result = render_pdf(deck([section], title='Taʼminot zanjiri'), tmp_path / 'callout.pdf', locale='uz')
    assert result['page_count'] > 1
    text = ''.join(page.extract_text() for page in PdfReader(tmp_path / 'callout.pdf').pages)
    assert all(f'{n}-band' in text for n in range(1, 17)), 'every point is printed, none dropped'


@pytest.mark.parametrize('sections', [1, 3, 5])
@pytest.mark.parametrize('locale', ['ru', 'uz'])
def test_long_word_prose_at_the_page_target_keeps_its_page_count(tmp_path, locale, sections):
    body = prose(locale, pages.target_chars('pdf'))
    content = deck([{'id': f's{n}', 'heading': f'{n + 1}', 'body': body, 'notes': ''} for n in range(sections)])
    result = render_pdf(content, tmp_path / f'{locale}-{sections}.pdf', locale=locale)
    assert result['page_count'] == sections, f'{locale} prose at the target must not spill'
    assert not result['shortened'], 'prose written to the target is never cut'
