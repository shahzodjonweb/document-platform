"""Decks wear one of over a hundred designs, chosen by the model or the customer; what the customer named still wins."""
import copy
import json

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from apps.core.services import execute_job, submit_job
from apps.studio import deck_designs, provider
from apps.studio.deck_designs import (CATEGORIES, DESIGN_IDS, DESIGNS, FONTS, PLACES, catalogue, look,
                                      reference, remember, wanted)
from apps.studio.domain import create_draft, generation_quote, unpack
from apps.studio.slides import _luminance, contrast_ratio, palette, render_pptx
from operations.integrations import save_config
from tests.test_platform import account
from tests.test_provider_contract import CONFIG, draft
from tests.test_provider_optimization import transport
from tests.test_studio import paid


def deck():
    return {'title': 'Tides', 'questions': [], 'citations': [], 'sections': [
        {'id': 's1', 'heading': 'Tides', 'body': 'Why the sea moves', 'notes': '', 'layout': 'cover',
         'items': [], 'columns': [], 'image_query': ''},
        {'id': 's2', 'heading': 'The moon pulls the water', 'body': 'Gravity\nTwo bulges\nTwice a day',
         'notes': '', 'layout': 'bullets', 'items': [], 'columns': [], 'image_query': ''},
        {'id': 's3', 'heading': 'Spring and neap', 'body': 'Sun and moon together', 'notes': '',
         'layout': 'section', 'items': [], 'columns': [], 'image_query': ''},
    ]}


def faces(presentation):
    found = set()
    for slide in presentation.slides:
        found |= {latin.get('typeface') for latin in slide._element.iter(qn('a:latin'))}
    return found


def theme_fonts(presentation):
    from apps.studio.slides import _theme_part
    from lxml import etree
    theme = etree.fromstring(_theme_part(presentation).blob)
    return [theme.find('.//' + qn(tag)).find(qn('a:latin')).get('typeface')
            for tag in ('a:majorFont', 'a:minorFont')]


@pytest.mark.parametrize('design', DESIGN_IDS)
def test_every_design_renders_its_own_colour_and_fonts(design, tmp_path):
    path = tmp_path / 'deck.pptx'
    result = render_pptx(deck(), path, style={'accent': '#255e49', 'deck_design': design})
    chosen = DESIGNS[design]
    presentation = Presentation(str(path))
    assert result['metadata']['design'] == design
    # The accent element is still shapes[0] and still the exact accent.
    for slide in presentation.slides:
        assert str(slide.shapes[0].fill.fore_color.rgb) == chosen['accent'].lstrip('#').upper()
    assert faces(presentation) <= {chosen['heading_font'], chosen['body_font']}
    assert chosen['body_font'] in faces(presentation)
    assert theme_fonts(presentation) == [chosen['heading_font'], chosen['body_font']]
    if chosen['composition'] == 'classic':
        ground = str(presentation.slides[1].background.fill.fore_color.rgb)
        assert (_luminance(ground) < 0.1) == (chosen['theme'] == 'dark')
    else:
        # A composition may set content on a panel over a coloured ground; the
        # page the text sits on is still the surface, dark on a dark design.
        roles = palette(chosen['accent'], chosen['theme'], chosen['secondary'], chosen['paper'])
        assert (_luminance(roles['surface']) < 0.1) == (chosen['theme'] == 'dark')


@pytest.mark.parametrize('design', DESIGN_IDS)
def test_every_design_keeps_text_readable(design):
    shown = look({'deck_design': design})
    roles = palette(shown['accent'], shown['theme'], shown['secondary'], shown['paper'], shown['accent_headings'])
    pairs = [('ink', 'surface', 7.0), ('muted', 'surface', 4.5), ('accent_text', 'surface', 4.5),
             ('heading', 'surface', 4.5), ('cover_ink', 'cover_fill', 4.5), ('band_ink', 'band', 4.5),
             ('card_ink', 'card', 7.0), ('soft_ink', 'accent_soft', 7.0)]
    for text, ground, target in pairs:
        assert contrast_ratio(roles[text], roles[ground]) >= target - 0.05, (design, text, ground)


# The ten designs decks wore before categories, exactly as they were.
CLASSIC = {'forest': ('#1E6B45', 'Cambria', 'Corbel'), 'boardroom': ('#1F3A68', 'Georgia', 'Calibri'),
           'midnight': ('#4F8CFF', 'Trebuchet MS', 'Calibri'), 'scholar': ('#7A1F35', 'Constantia', 'Cambria'),
           'spotlight': ('#D94A26', 'Arial', 'Arial'), 'lagoon': ('#0F766E', 'Candara', 'Candara'),
           'royal': ('#5B2A86', 'Georgia', 'Corbel'), 'classroom': ('#D9661F', 'Trebuchet MS', 'Trebuchet MS'),
           'graphite': ('#2B2F33', 'Arial', 'Arial'), 'gala': ('#D4A017', 'Georgia', 'Calibri')}


def test_over_a_hundred_designs_in_categories_with_the_classic_ten_unchanged():
    from pathlib import Path
    import re
    assert len(DESIGNS) >= 100 and len(CATEGORIES) >= 10
    assert list(CATEGORIES)[0] == 'classic'
    assert {key: (d['accent'], d['heading_font'], d['body_font']) for key, d in CATEGORIES['classic'][0].items()} == CLASSIC
    assert all(d['art'] is None and not d['cover_photo'] for d in CATEGORIES['classic'][0].values())
    signatures = {(d['accent'], d['theme'], d['paper'], d['heading_font'], d['body_font'], str(d['art']))
                  for d in DESIGNS.values()}
    assert len(signatures) == len(DESIGNS), 'no two designs are the same look'
    masks = Path(deck_designs.__file__).parent / 'assets' / 'deck'
    for category, (designs, names) in CATEGORIES.items():
        assert len(designs) >= 8 and len(names) == 3 and all(names), category
        # Within a category every design is its own colour.
        assert len({d['accent'].upper() for d in designs.values()}) == len(designs), category
    for key, design in DESIGNS.items():
        assert re.fullmatch(r'[a-z_]+', key) and design['category'] in CATEGORIES
        assert {design['heading_font'], design['body_font']} <= FONTS, key
        assert design['theme'] in ('light', 'dark', 'bold') and len(design['fits']) <= 120, key
        if design['art']:
            assert design['art']['place'] in PLACES and 0 < design['art']['strength'] <= 0.6, key
            assert (masks / f"{design['art']['pattern']}.png").is_file(), key
    assert {d['theme'] for d in DESIGNS.values()} == {'light', 'dark', 'bold'}
    # Most designs are arranged by a composition of their own, not only recoloured.
    from apps.studio.compositions import COMPOSITIONS, compose
    compose('classic')
    assert all(d['composition'] in COMPOSITIONS for d in DESIGNS.values())
    assert sum(d['composition'] == 'classic' for d in DESIGNS.values()) <= 22
    assert {d['composition'] for d in DESIGNS.values()} == set(COMPOSITIONS)
    assert all(DESIGNS[key]['cover_photo'] for key in CATEGORIES['photo'][0])


def test_the_list_the_model_reads_is_grouped_and_stays_small():
    full, plain = reference(), reference(photos=False)
    # It is in every deck call's instructions: keep it to a couple of thousand tokens.
    assert len(full) <= 9000
    assert 'Business:\n- corporate:' in full and 'Classic:\n- forest:' in full
    assert 'postcard:' in full and 'postcard:' not in plain and 'image_query' not in plain
    data = draft(4, 'pptx')
    assert 'postcard:' not in provider.request_body(CONFIG, data, 'ai.pptx')['instructions']


def test_the_picker_gets_every_design_with_the_colours_it_renders():
    groups = catalogue('uz')
    assert [group['id'] for group in groups] == list(CATEGORIES)
    assert groups[0]['name'] == 'Klassik'
    entries = [entry for group in groups for entry in group['designs']]
    assert [entry['id'] for entry in entries] == DESIGN_IDS
    midnight = next(entry for entry in entries if entry['id'] == 'midnight')
    assert midnight['colours']['accent'] == '#4F8CFF' and midnight['theme'] == 'dark'
    assert next(entry for entry in entries if entry['id'] == 'polar_night')['name'] == 'Polar Night'
    assert all(entry['preview'].startswith(f"/api/v1/studio/deck-designs/{entry['id']}.png?locale=uz&v=")
               for entry in entries)


def test_what_the_customer_named_wins_over_the_design(tmp_path):
    # A brand colour, or one named in the brief, replaces the design's accent.
    shown = look({'deck_design': 'spotlight', 'accent': '#1F4E9C', 'accent_fixed': True})
    assert shown['accent'] == '#1F4E9C' and shown['secondary'] is None and shown['heading_font'] == 'Arial'
    # The template's default green is not a choice anyone made.
    assert look({'deck_design': 'spotlight', 'accent': '#255e49'})['accent'] == '#D94A26'
    # "A dark deck" keeps the design's fonts but not its light paper.
    dark = look({'deck_design': 'scholar', 'deck_theme': 'dark'})
    assert dark['theme'] == 'dark' and dark['paper'] is None and dark['heading_font'] == 'Constantia'
    path = tmp_path / 'deck.pptx'
    render_pptx(deck(), path, style={'deck_design': 'lagoon', 'accent': '#9B1B1B', 'accent_fixed': True})
    assert str(Presentation(str(path)).slides[0].shapes[0].fill.fore_color.rgb) == '9B1B1B'


def test_a_deck_without_a_design_looks_exactly_as_before(tmp_path):
    path = tmp_path / 'deck.pptx'
    result = render_pptx(deck(), path, style={'accent': '#255e49'})
    presentation = Presentation(str(path))
    assert result['metadata']['design'] == ''
    assert str(presentation.slides[0].shapes[0].fill.fore_color.rgb) == '255E49'
    assert faces(presentation) <= {'Cambria', 'Corbel'}


def test_the_model_is_asked_once_per_deck():
    data = draft(20, 'pptx')
    first = provider.request_body(CONFIG, data, 'ai.pptx', (0, 8))
    later = provider.request_body(CONFIG, data, 'ai.pptx', (8, 16))
    assert first['text']['format']['schema']['properties']['design']['enum'] == DESIGN_IDS
    assert 'design' not in later['text']['format']['schema']['properties']
    # The list is in every call's instructions, so long decks keep one cached prefix.
    assert 'boardroom:' in first['instructions'] and first['instructions'] == later['instructions']
    chosen = copy.deepcopy(data)
    chosen['options']['template_style'] = {'deck_design': 'gala'}
    assert 'design' not in provider.request_body(CONFIG, chosen, 'ai.pptx')['text']['format']['schema']['properties']
    revising = copy.deepcopy(data)
    revising['revision'] = {'request': 'Fix a typo.', 'original_prompt': 'Tides.'}
    assert not wanted(revising)
    assert 'design' not in provider.request_body(CONFIG, draft(3, 'pdf'), 'ai.pdf_topic')['text']['format']['schema']['properties']


def test_the_choice_is_kept_and_only_replaced_when_asked():
    data = {'options': {'template_style': {'accent': '#255e49'}}}
    remember(data, {'design': 'not-a-design'})
    assert 'deck_design' not in data['options']['template_style']
    remember(data, {'design': 'royal'})
    remember(data, {'design': 'gala'})
    assert data['options']['template_style']['deck_design'] == 'royal'
    remember(data, {'design': 'gala'}, replace=True)
    assert data['options']['template_style']['deck_design'] == 'gala'


@pytest.mark.django_db
def test_a_brief_fixes_only_what_it_names(settings):
    settings.DEBUG = True
    customer = paid(settings, account())
    plain = unpack(create_draft(customer, {'feature_id': 'ai.pptx', 'prompt': 'Five slides on tides'}).encrypted_data)
    assert 'deck_theme' not in plain['options']['template_style']
    assert not plain['options']['template_style'].get('accent_fixed')
    named = unpack(create_draft(customer, {'feature_id': 'ai.pptx',
                                           'prompt': 'A dark deck in navy about tides'}).encrypted_data)
    style = named['options']['template_style']
    assert style['deck_theme'] == 'dark' and style['accent'] == '#16305C' and style['accent_fixed']


@pytest.mark.django_db
def test_a_generated_deck_wears_the_design_the_model_chose(settings, monkeypatch, tmp_path):
    settings.DEBUG = True
    customer = paid(settings, account())
    save_config('ai', {'mode': 'openai', 'model': CONFIG['model'], 'api_key': 'offline-only'})
    authored = create_draft(customer, {'feature_id': 'ai.pptx', 'prompt': 'Three slides about tides',
                                       'options': {'length': 3}})

    def answer(body, number):
        sent = json.loads(body['input'])
        reply = {'title': 'Tides', 'answer_supported': True,
                 'sections': [{'id': s['id'], 'heading': f'Tides {n}', 'body': 'Moon\nSun\nWater',
                               'notes': 'Say it.', 'layout': 'cover' if n == 0 else 'bullets', 'items': [],
                               'columns': [], 'image_query': ''}
                              for n, s in enumerate(sent['outline']['sections'])]}
        if 'design' in body['text']['format']['schema']['properties']:
            reply['design'] = 'lagoon'
        for field in ('citations', 'questions'):
            if field in body['text']['format']['schema']['properties']:
                reply[field] = []
        return reply

    transport(monkeypatch, answer)
    quote = generation_quote(customer, authored.id, authored.version)
    job, _ = submit_job(customer, quote.id, 'design-offline')
    job = execute_job(job.id)
    assert job.status == 'succeeded', job.error_code
    authored.refresh_from_db()
    assert unpack(authored.encrypted_data)['options']['template_style']['deck_design'] == 'lagoon'
    from apps.core import storage
    path = storage.local(job.artifacts.get().file.object_key)
    presentation = Presentation(str(path))
    assert str(presentation.slides[0].shapes[0].fill.fore_color.rgb) == '0F766E'
    assert 'Candara' in faces(presentation)


@pytest.mark.django_db
def test_a_design_picked_on_the_review_screen_beats_the_brief(settings, tmp_path):
    from apps.core.errors import DomainError
    settings.DEBUG = True
    customer = paid(settings, account())
    brief = 'A dark deck in green, 12 slides about tides'
    picked = unpack(create_draft(customer, {'feature_id': 'ai.pptx', 'prompt': brief,
                                            'options': {'deck_design': 'royal', 'pages': 7}}).encrypted_data)
    style = picked['options']['template_style']
    assert style['deck_design'] == 'royal' and 'deck_theme' not in style and not style.get('accent_fixed')
    assert picked['options']['length'] == 7 and picked['options']['deck_design'] == 'royal'
    # The model is not asked for a design the customer already chose.
    assert not wanted(picked)
    path = tmp_path / 'deck.pptx'
    render_pptx(deck(), path, style=style)
    assert str(Presentation(str(path)).slides[0].shapes[0].fill.fore_color.rgb) == '5B2A86'
    # "Auto" is the brief and the model, as before.
    auto = unpack(create_draft(customer, {'feature_id': 'ai.pptx', 'prompt': brief,
                                          'options': {'deck_design': 'auto'}}).encrypted_data)
    assert auto['options']['template_style']['deck_theme'] == 'dark'
    assert 'deck_design' not in auto['options']['template_style'] and wanted(auto)
    for wrong in ('nope', '', 7):
        with pytest.raises(DomainError, match='invalid_parameters'):
            create_draft(customer, {'feature_id': 'ai.pptx', 'prompt': brief, 'options': {'deck_design': wrong}})
    with pytest.raises(DomainError, match='invalid_parameters'):
        create_draft(customer, {'prompt': 'A guide to tides', 'options': {'deck_design': 'royal'}})



@pytest.mark.parametrize('design', DESIGN_IDS)
def test_every_design_keeps_its_text_readable_and_in_place(design, tmp_path):
    """The whole sample deck, with and without photos, through the composition checker."""
    from apps.studio.compositions import sample
    from apps.studio.compositions.check import problems
    found = []
    for with_photos in (False, True):
        path = tmp_path / 'deck.pptx'
        result = sample.render({'deck_design': design}, path, with_photos)
        found += problems(path, result['metadata']['layouts'], DESIGNS[design]['accent'])
    assert not found, '\n'.join(found)


def test_retired_designs_wear_a_design_that_exists():
    from apps.studio.deck_designs import RETIRED, current
    assert all(old not in DESIGNS for old in RETIRED)
    assert all(new in DESIGNS for new in RETIRED.values())
    for old, new in RETIRED.items():
        assert look({'deck_design': old})['id'] == new and current(old) == new
