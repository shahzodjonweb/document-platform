"""Figures and comparison rows survive layout validation and rendering."""
import pytest
from pptx import Presentation
from pptx.util import Inches

from apps.studio.layouts import clean_slide_fields, parse_number
from apps.studio.slides import render_pptx


def item(label='', text='', value=''):
    return {'label': label, 'text': text, 'value': value}


def render(tmp_path, layout, items, columns):
    section = {'id': 's2', 'heading': 'Figures', 'body': '', 'notes': '',
               **clean_slide_fields({'layout': layout, 'items': items, 'columns': columns})}
    content = {'title': 'Data', 'questions': [], 'citations': [], 'sections': [
        {'id': 's1', 'heading': 'Data', 'body': '', 'notes': '', 'layout': 'cover'}, section,
    ]}
    path = tmp_path / 'data.pptx'
    result = render_pptx(content, path)
    presentation = Presentation(str(path))
    assert result['page_count'] == len(presentation.slides) == 2
    return result, presentation.slides[1]


@pytest.mark.parametrize('value, expected', [
    ('5 ming', 5_000), ('5ming', 5_000), ('5 MING', 5_000),
    ('3 mlrd', 3_000_000_000), ('3mlrd', 3_000_000_000), ('3 MLRD', 3_000_000_000),
    ('2,5 mln', 2_500_000), ('2,5 МЛН', 2_500_000), ('1,2 млрд', 1_200_000_000),
    ('30 minutes', 30), ('30minutes', 30), ('30 mins', 30), ('30 km', 30),
    ('2 kilograms', 2), ('5 minglab', 5), ('4 billionaires', 4), ('2 millionaires', 2),
    ('5 thousand', 5_000), ('5 thousands', 5_000), ('2 million', 2_000_000), ('2 millions', 2_000_000),
    ('4 billion', 4_000_000_000), ('4 billions', 4_000_000_000), ('3 milliard', 3_000_000_000),
    ('1 тысяча', 1_000), ('3 тысячи', 3_000), ('5 тысяч', 5_000),
    ('1 миллион', 1_000_000), ('3 миллиона', 3_000_000), ('5 миллионов', 5_000_000),
    ('1 миллиард', 1_000_000_000), ('3 миллиарда', 3_000_000_000), ('5 миллиардов', 5_000_000_000),
    ('$4.2M', 4_200_000), ('3.5k', 3_500), ('12 000', 12_000), ('18%', 18),
])
def test_number_suffixes_are_complete_units(value, expected):
    assert parse_number(value) == expected


@pytest.mark.parametrize('first, second', [
    ('5 ming', '10k'), ('5 thousand', '10k'), ('2 million', '4m'), ('2 millions', '4mln'),
    ('4 billion', '8b'), ('3 milliard', '6mlrd'), ('5 тысяч', '10k'),
    ('2 миллиона', '4m'), ('4 миллиарда', '8b'), ('30 minutes', '60'),
])
def test_chart_bar_proportions_use_the_whole_unit(tmp_path, first, second):
    result, slide = render(tmp_path, 'chart', [item('First', value=first), item('Second', value=second)],
                           ['People'])
    bars = [shape for shape in slide.shapes
            if not shape.has_text_frame or not shape.text_frame.text]
    bars = [shape for shape in bars if shape.top >= Inches(2.85) and shape.height > Inches(0.3)]
    assert len(bars) == 2
    assert bars[0].width / bars[1].width == pytest.approx(0.5, abs=1e-6)
    assert result['metadata']['layouts'][1] == 'chart'
    assert result['shortened'] is False


@pytest.mark.parametrize('count', [2, 6, 8])
def test_chart_table_fallback_preserves_every_category_and_value(tmp_path, count):
    items = [item(f'Region {index}', value=str(index - 1)) for index in range(count)]
    result, slide = render(tmp_path, 'chart', items, ['Region', 'Change'])
    tables = [shape.table for shape in slide.shapes if shape.has_table]
    assert result['metadata']['layouts'][1] == 'table'
    assert len(tables) == 1
    assert [[cell.text for cell in row.cells] for row in tables[0].rows] == [
        ['Region', 'Change'], *[[entry['label'], entry['value']] for entry in items],
    ]
    assert result['shortened'] is False


@pytest.mark.parametrize('columns, expected', [
    (['Region', 'Change'], [['North', '-4 — First quarter'], ['South', '3 — Second quarter']]),
    (['Region', 'Context', 'Change'], [['North', 'First quarter', '-4'], ['South', 'Second quarter', '3']]),
])
def test_chart_table_fallback_keeps_optional_context(tmp_path, columns, expected):
    result, slide = render(tmp_path, 'chart', [item('North', 'First quarter', '-4'),
                                             item('South', 'Second quarter', '3')], columns)
    grid = next(shape.table for shape in slide.shapes if shape.has_table)
    assert [[cell.text for cell in row.cells] for row in list(grid.rows)[1:]] == expected
    assert result['shortened'] is False


@pytest.mark.parametrize('count', [6, 8])
def test_comparison_keeps_every_accepted_row(tmp_path, count):
    items = [item(f'Criterion {index}', f'First {index}', f'Second {index}') for index in range(count)]
    result, slide = render(tmp_path, 'comparison', items, ['Option A', 'Option B'])
    text = {shape.text_frame.text for shape in slide.shapes if shape.has_text_frame}
    assert result['metadata']['layouts'][1] == 'comparison'
    assert all(value in text for entry in items for value in entry.values())
    assert result['shortened'] is False


def test_comparison_keeps_populated_rows_without_a_label(tmp_path):
    items = [item('Cost', 'High', 'Low'), item('Time', 'Months', 'Weeks'), item('', 'Optional A', 'Optional B')]
    result, slide = render(tmp_path, 'comparison', items, ['Option A', 'Option B'])
    text = {shape.text_frame.text for shape in slide.shapes if shape.has_text_frame}
    assert {'Optional A', 'Optional B'} <= text
    assert result['shortened'] is False


def test_dense_comparison_keeps_all_rows_and_reports_shortened_cells(tmp_path):
    items = [item(f'Criterion {index}', 'Longword ' * 14, f'Value {index}') for index in range(8)]
    result, slide = render(tmp_path, 'comparison', items, ['Option A', 'Option B'])
    text = {shape.text_frame.text for shape in slide.shapes if shape.has_text_frame}
    assert all(entry['label'] in text and entry['value'] in text for entry in items)
    assert any(value.endswith('…') for value in text)
    assert result['shortened'] is True
