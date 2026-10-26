"""The names a customer dictates for a deck's title slide reach its cover, as written.

Students ask for their name, group and teacher on the first slide in many
ways ("Tayyorladi: …", "… deb yoz", "Birinchi betiga:", "сделала: …"). Left
to the model they were dropped, merged or never reached the cover, above all
in a change request. apps/studio/title_lines.py reads them without the model,
and only where the request plainly dictates them; these requests are shaped
like real ones, with invented names.
"""
import pytest
from pptx import Presentation

from apps.studio import title_lines
from apps.studio.title_lines import dictated

pytestmark = pytest.mark.django_db

ASKED = {
    'a block under "Ustki qismiga"': (
        "Pedagogika fani uchun Mustaqil ish\n\nUstki qismiga\n\nMustaqil ish\n\nTayyorladi: Karimova Dilnoza\n"
        "Qabul qildi: Rashidova Malika\n\n10ta slayd bolsin eng oxirida Xulosa\nRangi moviy bolsin",
        ['Mustaqil ish', 'Tayyorladi: Karimova Dilnoza', 'Qabul qildi: Rashidova Malika']),
    'a Russian line placed at the top': (
        "Нужно 10 слайдов для самостоятельной работы - историческая грамматика - в вузе. В начале напиши тему и "
        "вверх 205- группа Каримова Дилноза. Делай красиво и подробно",
        ['205- группа Каримова Дилноза']),
    'a change dictating the first page line by line': (
        'Birinchi betiga: "Hamshiralik ishi" yoʻnalishi\n202-guruh talabasi\nKarimova Dilnozaning\n'
        'Jamoada hamshiralik ishi fanidan tayyorlagan\nMustaqil taʼlim\nMavzu: Onalar va bolalar salomatligi\n\n'
        'Oʻqituvchi: Rashidova Malika\n\n3 ta reja birinchi betdan keyin',
        ['"Hamshiralik ishi" yoʻnalishi', '202-guruh talabasi', 'Karimova Dilnozaning',
         'Jamoada hamshiralik ishi fanidan tayyorlagan', 'Mustaqil taʼlim', 'Mavzu: Onalar va bolalar salomatligi',
         'Oʻqituvchi: Rashidova Malika']),
    'a Russian change with its label': (
        "напиши в начале слайда сделала: Каримова Дилноза 205-группа", ['сделала: Каримова Дилноза 205-группа']),
    'two names after the request to write them': (
        "Oila qadriyatlari\nShu mavzuda 8 ta slayd tayyorla\nIsm familiyamni ham yozib qo‘y slaydni 1 sahifasiga\n"
        "Karimova Dilnoza va Rashidova Malika deb",
        ['Karimova Dilnoza va Rashidova Malika']),
    'a name and class dictated with deb': (
        "Jurnalist mavzusida 8 ta slayd tayyorla Karimova Dilnoza 11 B sinf deb boshidagi kirish qismiga yozib qo‘y "
        "kim qilganligini ham bilishsin havorangda bo‘lsin",
        ['Karimova Dilnoza 11 B sinf']),
    'the name before tayyorladi': (
        "Adabiyot metodikasi mavzusiga 12 sahifa slayd qil , Karimova Dilnoza tayyorladi ,ochiq rangdan foydalan",
        ['Karimova Dilnoza tayyorladi']),
    'a name beside its group': (
        "Tarix fanidan 15ta slayd kerak-Karimova Dilnoza 026-136-guruh. Ilk oʻrta asr davlatlari",
        ['Karimova Dilnoza 026-136-guruh']),
    'a change adding a first page with the full name': (
        "Boshiga bitta sahifa qoʻsh va Karimova Dilnoza Aziz qizi 026-136-guruh deb yoz keyn Eftaliylar haqida ham "
        "yoz rasmlar qoʻsh",
        ['Karimova Dilnoza Aziz qizi 026-136-guruh']),
    'a long line dictated for the first page': (
        'Shunga boshidagi bir betga 133 IDUM 7 D sinf oʻquvchisi Rashidov Azizbekning "Ustoz otangdek ulugʻ" '
        'maqoli boʻyicha tayyorlagan taqdimoti deb yoz va ustoz oʻquvchilarning rasmini toʻrtta qoʻy',
        ['133 IDUM 7 D sinf oʻquvchisi Rashidov Azizbekning "Ustoz otangdek ulugʻ" maqoli boʻyicha tayyorlagan '
         'taqdimoti']),
    'labels in one sentence': (
        "Pedagogika fanidan 10 ta slayd. Tayyorladi: Karimova Dilnoza. Qabul qildi: Rashidova Malika. Rasmlar bilan",
        ['Tayyorladi: Karimova Dilnoza', 'Qabul qildi: Rashidova Malika']),
    'Russian labels on lines of their own': (
        "Презентация на тему экология, 10 слайдов.\nВыполнила: Каримова Дилноза\nПроверила: Рашидова Малика\n"
        "Группа: 305",
        ['Выполнила: Каримова Дилноза', 'Проверила: Рашидова Малика', 'Группа: 305']),
    'a title page listed inline': (
        "Boshiga titul varag'ini qo'sh: MUSTAQIL ISH, Tayyorladi: Karimova Dilnoza. Har bir sahifada rasm bo'lsin",
        ['MUSTAQIL ISH', 'Tayyorladi: Karimova Dilnoza']),
    'lines separated by slashes': (
        "Ustki qismiga MUSTAQIL ISH / TAYYORLADI: KARIMOVA DILNOZA / QABUL QILDI: RASHIDOVA MALIKA",
        ['MUSTAQIL ISH', 'TAYYORLADI: KARIMOVA DILNOZA', 'QABUL QILDI: RASHIDOVA MALIKA']),
    'initials': ("Tayyorladi: D. Karimova\nQabul qildi: M. Rashidova",
                 ['Tayyorladi: D. Karimova', 'Qabul qildi: M. Rashidova']),
    'a group on the next line': ("Tayyorladi: Karimova Dilnoza Aziz qizi\n021-104-guruh\n\n10 ta slayd kerak",
                                 ['Tayyorladi: Karimova Dilnoza Aziz qizi', '021-104-guruh']),
    'a name asked for in English': ("8 slides on volcanoes for school. My name: Dilnoza Karimova", ['Dilnoza Karimova']),
    'a title and a group inside the value': (
        "Bajardi: 301-guruh talabasi Yo'ldosheva Sevara\nTekshirdi: dots. Qodirov Anvar",
        ["Bajardi: 301-guruh talabasi Yo'ldosheva Sevara", 'Tekshirdi: dots. Qodirov Anvar']),
    'labels with dashes in a sentence': (
        "Tayyorladi - Sobirov Jasur, Qabul qildi - Nazarova Feruza. Raqamli iqtisodiyot 10 slayd",
        ['Tayyorladi - Sobirov Jasur', 'Qabul qildi - Nazarova Feruza']),
    'initials after the surname': (
        "Raqamli iqtisodiyot 10 slayd. Tayyorladi: Sobirov J.A. Qabul qildi: Nazarova F.",
        ['Tayyorladi: Sobirov J.A.', 'Qabul qildi: Nazarova F.']),
    'a block of fields for the title page': (
        "Titulga:\nFakultet: Iqtisodiyot\nYo'nalish: Buxgalteriya hisobi\nGuruh: 21-05\nTalaba: Sobirov Jasur",
        ['Fakultet: Iqtisodiyot', "Yo'nalish: Buxgalteriya hisobi", 'Guruh: 21-05', 'Talaba: Sobirov Jasur']),
    'roles beside who did the work': ("Tayyorladi: Sobirov Jasur\nO'qituvchi: Nazarova Feruza",
                                      ['Tayyorladi: Sobirov Jasur', "O'qituvchi: Nazarova Feruza"]),
    'a given name ending like a place word': ("Tayyorladi: Malika Ergasheva Anvar qizi",
                                              ['Tayyorladi: Malika Ergasheva Anvar qizi']),
    'a second byline after a comma': ('Работу выполнила Каримова Дилноза Азизовна, проверила Назарова Феруза',
                                      ['Работу выполнила Каримова Дилноза Азизовна', 'проверила Назарова Феруза']),
    'Uzbek Cyrillic dictated with deb': (
        'Собирова Мадина 8-Б синф ўқувчиси деб ёзиб қўй', ['Собирова Мадина 8-Б синф ўқувчиси']),
    'a Russian label with what the student studies': (
        'Презентация на тему экология, 10 слайдов. Выполнил: Собиров Жасур Анварович, студент 2 курса',
        ['Выполнил: Собиров Жасур Анварович, студент 2 курса']),
    'an English byline': ('10 slides about climate change. Prepared by John Miller, Group 12',
                          ['Prepared by John Miller, Group 12']),
    'a place word before the name': ("Muqovaga Karimova Dilnoza 302-guruh deb yozib qo'y", ['Karimova Dilnoza 302-guruh']),
    'a greeting before the name': ("Salom Ergasheva Malika 301-21 guruh. Tarix fanidan 10 slayd", ['Ergasheva Malika 301-21 guruh']),
    'a label of two words': ("Tayyorladi: Ergasheva Malika\nSinf rahbari: Xolmatova Sevara",
                             ['Tayyorladi: Ergasheva Malika', 'Sinf rahbari: Xolmatova Sevara']),
    'a block row that goes on to ask for more': (
        "Birinchi betiga:\nTayyorladi: Ergasheva Malika\nQabul qildi: Xolmatova Sevara. Rasmlar ko'p bo'lsin",
        ['Tayyorladi: Ergasheva Malika', 'Qabul qildi: Xolmatova Sevara']),
    'a group after gr.': ("Выполнил: Петров Игорь, гр. 301-21", ['Выполнил: Петров Игорь, гр. 301-21']),
    'a slide count beside the dictation': (
        "Inflyatsiya haqida 12 slayd kerak birinchi slaydga Ergasheva Malika 11-B sinf o'quvchisi deb yozib qo'y",
        ["Ergasheva Malika 11-B sinf o'quvchisi"]),
    'a hyphenated surname': ("Tayyorladi: Rasulova-Karimova Shahnoza", ['Tayyorladi: Rasulova-Karimova Shahnoza']),
    'a name for the cover in a change': ("Ismim Sobirova Mohinur bo'lsin", ['Sobirova Mohinur']),
    # A third review's: a sentence about other slides before the byline, given names ending like place
    # words, fillers, labels joined by a comma, initials with no colon, orders without one.
    'other slides spoken of first': ("Inflyatsiya haqida taqdimot, har bir slaydda rasm bo'lsin. Tayyorladi: Karimova Dilnoza",
                                     ['Tayyorladi: Karimova Dilnoza']),
    'a count of slides before the byline': ('10-slaydli taqdimot. Tayyorladi: Karimova Dilnoza', ['Tayyorladi: Karimova Dilnoza']),
    'a given name before a patronymic': ('Tayyorladi: Saida Karimovna Rashidova', ['Tayyorladi: Saida Karimovna Rashidova']),
    'a filler before a label': ('Tarix 10 ta slayd. Tayyorladi: Karimova Dilnoza\nVa Qabul qildi: Rashidova Malika',
                                ['Tayyorladi: Karimova Dilnoza', 'Qabul qildi: Rashidova Malika']),
    'labels joined by a comma': ("Tayyorladi: Ergasheva Malika, Fan o'qituvchisi: Xolmatova Sevara",
                                 ['Tayyorladi: Ergasheva Malika', "Fan o'qituvchisi: Xolmatova Sevara"]),
    'initials after a label with no colon': ('Выполнила Петрова О.С.', ['Выполнила Петрова О.С.']),
    'a group code with letters': ('Выполнила: Петрова Ольга, студентка группы ЭК-21', ['Выполнила: Петрова Ольга, студентка группы ЭК-21']),
    'two given names': ('Tayyorladi: Muhammad Ali Valiyev', ['Tayyorladi: Muhammad Ali Valiyev']),
    'an order with no colon': ('Please write my name Aziz Karimov on the first slide', ['Aziz Karimov']),
    # A fourth review's: capitals, list markers, a byline after where it goes, two bylines with no comma.
    'a request in capitals': ('BIRINCHI SLAYDGA SOBIROVA MADINA 302-GURUH DEB YOZ', ['SOBIROVA MADINA 302-GURUH']),
    'a given name ending like a surname': ("Iqtisodiyot fanidan 10 ta slayd. Titul Sobirova Madina 302-guruh deb yozib qo'y",
                                           ['Sobirova Madina 302-guruh']),
    'a numbered block': ('Birinchi slaydga:\n1. Sobirova Madina\n2. 302-guruh', ['Sobirova Madina', '302-guruh']),
    'bylines after where they go': ('Birinchi slaydga Tayyorladi Sobirova Madina, qabul qildi Nazarova Feruza',
                                    ['Tayyorladi Sobirova Madina', 'qabul qildi Nazarova Feruza']),
    'two bylines with no comma': ('Tayyorladi Sobirova Madina qabul qildi Nazarova Feruza',
                                  ['Tayyorladi Sobirova Madina', 'qabul qildi Nazarova Feruza']),
    'the first slide by its number': ('1-slaydga:\nTayyorladi: Sobirova Madina\nQabul qildi: Nazarova Feruza',
                                      ['Tayyorladi: Sobirova Madina', 'Qabul qildi: Nazarova Feruza']),
    'a byline after the request': ('Презентация 10 слайдов про Пушкина. Выполнила Каримова Дилноза', ['Выполнила Каримова Дилноза']),
    'only the byline beside a fact list': (
        'Tayyorladi: Karimova Dilnoza. Temuriylar davlati haqida 10 slayd. Asoschisi: Amir Temur. Ustozi: Shamsuddin Kulol',
        ['Tayyorladi: Karimova Dilnoza']),
    'co-authors a group row closes': ('Tayyorladi: Sobirov Jasur\nAbdullayeva Zarina\nNazarov Otabek\n302-guruh',
                                      ['Tayyorladi: Sobirov Jasur', 'Abdullayeva Zarina', 'Nazarov Otabek', '302-guruh']),
    'the name after where it goes': ('Put my name on the first slide: Emily Clark, Grade 9', ['Emily Clark, Grade 9']),
    'a block that ends where the request goes on': (
        "Inflyatsiya.\nBirinchi betiga:\nTayyorladi: Ergasheva Malika\nQabul qildi: Xolmatova Sevara\nIngliz tilida",
        ['Tayyorladi: Ergasheva Malika', 'Qabul qildi: Xolmatova Sevara']),
}

# Roles with nothing saying "title page" are left to the model, as before: a cast, a fact list or a lesson
# plan reads the same ("Student: Harry Potter", "Rahbari: Amir Temur"), and a misread line stays on a cover.
LEFT_TO_THE_MODEL = [
    "Topic : Concepts and Conceptual Analysis, Student: Karimova Dilnoza , slaydda mavzular oson tushuntirilsin",
    "Talaba: Sobirov Jasur\nO'qituvchi: Nazarova Feruza",
]

NOT_ASKED = [
    "Matematika fani uchun slayd kerak 10 sinf akademik litseydagilar uchun. Mavzusi haqiqiy sonlar",
    "Talabalar va o‘qituvchilar uchun 10 ta slayd kerak — Pedagogika fani",
    "O'qituvchi nutq madaniyati fanidan slayd tayyorlab ber. 10 sahifalik bolsin",
    "Alisher Navoiy hayoti va ijodi haqida 10 ta slayd",
    "Lermontov va Pushkin ijodi haqida 10 ta slayd, 9-sinf uchun",
    "Anton Chekhov hikoyalari, 11-sinf o'quvchilari uchun 8 slayd",
    "Oʻqituvchi vatan taraqqiyoti ning bunyodkori 10 ta slayd, yorqin rang Karimova Dilnoza uchun",
    "Xulosa deb yoz va rasmlar qo'sh",
    "Amir Temur davlati 12 ta slayd 7-sinf uchun",
    "Kirish qismiga mavzuning dolzarbligi deb yoz",
    "Kitob nomi: The Great Gatsby. Muallif: F. Scott Fitzgerald. Janr: roman",
    "Company name: Example Retail. 10 slides about our quarterly results",
    "Har bitta slaydga yaxshilab yoz orasiga mavzuga oid rasmlardan ham qo‘sh",
    "Mirzo Ulug'bek rasadxonasi deb nomlangan slayd qo'sh",
    'Birinchi slaydga "Erkin Vohidov ijodi" deb yoz',
    "10 ta slayd tayyorla. Mixail Lermontov o'z she'rlarida vatanni sevaman deb yozgan, shu haqida ham yoz",
    "1-slayd: Mavzu va Aleksandr Pushkin portreti\n2-slayd: Hayoti\n3-slayd: Ijodi\n4-slayd: Xulosa",
    "O'qituvchi: Mavzuni tushuntiradi\nO'quvchilar: Daftarga yozadi\nShu dars uchun 8 ta slayd",
    "201-guruh talabalari uchun Abdulla Oripov she'riyati haqida slayd",
    "Mavzu: Alisher Navoiy\nSinf: 7\n10 ta slayd kerak",
    # A second review's: another person's name, a fact list or a cast, a lesson plan, a byline inside a topic.
    "Yozuvchi ismi: Abdulla Qodiriy. Asari: O'tkan kunlar. 10 ta slayd",
    "Olimning ism familiyasi: Ivan Pavlov. Kashfiyotlari haqida slayd tayyorla",
    "Writer name: Charles Dickens. 10 slides about his life",
    "Davlat: Temuriylar. Poytaxti: Samarqand. Rahbari: Amir Temur. Yillar: 1370-1405",
    "Harry Potter characters: Student: Harry Potter; Teacher: Severus Snape; School: Hogwarts",
    "Студент: Родион Раскольников, герой романа Достоевского",
    "Ход урока:\nУчитель: приветствие учеников\nУченики: отвечают на вопросы",
    "Lesson plan for grade 5.\nTeacher: greets the students\nStudent: works in pairs",
    "The theory of evolution, presented by Charles Darwin in 1859. 10 slides",
    "Product: Model S. Made by: Tesla Motors",
    # A heading, a caption, a table, another slide.
    "Taqdimot mavzusini 'Mixail Lermontov' deb yoz",
    "Сарлавҳани «Жалолиддин Мангуберди» деб ёзинг",
    "Rasm ostiga Ivan Pavlov deb yoz",
    "Jadvalga Dmitriy Mendeleyev deb kirit",
    "Oxirgi slaydga yoz: Tayyorladi: Karimova Dilnoza",
    "На последнем слайде: Выполнила: Каримова Дилноза",
    "Учинчи слайдга Каримова Дилноза деб ёз",
    # An example, a topic beside an audience, the word "fish".
    "Ingliz tili: tanishuv. Misol yoz: My name is Anna Brown",
    "Amir Temur tilidan yoz: Mening ismim Amir Temur, men Kesh shahrida tug'ilganman",
    "302-guruh Zahiriddin Muhammad Bobur hayoti va ijodi, 10 ta slayd",
    "Киевская Русь, группа 205, 10 слайдов",
    "Animals of the Amazon. Fish: Red Piranha. Mammals: Pink Dolphin",
    "Har bir slayd boshida savol bo'lsin. Kimlar uchun: 302-guruh",
    # A third review's: an appositive in a topic, someone else's name, a lesson phrase, a cast or fact list
    # under a header, an assignment list, another slide, an example dialogue, a content correction.
    'Космический полет, выполнила Валентина Терешкова',
    "My hero's name: Amir Temur. 10 slides",
    'Выполнила: перевод текстов',
    'Mavzu: SSSR rahbarlari\nRahbar: Iosif Stalin\nRahbar: Nikita Xrushchev',
    'Report\nHarry Potter characters: Student: Harry Potter; Teacher: Severus Snape; School: Hogwarts',
    'Topshiriqlar taqsimoti:\n301-guruh Abdulla Oripov\n302-guruh Erkin Vohidov\nHar biriga 10 slayd',
    'Oxirgi slaydga: Karimova Dilnoza 302-guruh',
    'Ismimni yoz: Karimova Dilnoza, oxirgi slaydda',
    'Title slide: Greetings and introductions. Then dialogues.\nHello! My name is Anna Smith. I am from London.',
    "Muallifni Ivan Turgenev o'rniga Ivan Goncharov deb yoz",
    "Titulga mavzu sifatida 'Ivan Pavlov' deb yoz",
    'Hello, my name is Olga Petrova. I need 10 slides on the Romanovs; the title slide should have no names',
    # A fourth review's: an example byline, a credit or quiz answer, a title page said not to be, topic names
    # in a block, other pupils, a name and group as a fact, a quote corrected.
    'Как оформить титульный лист? Например, Выполнил: Петров Иван, Проверила: Сидорова Анна',
    'Космический полет. Выполнила Валентина Терешкова.',
    'Savol: raketani kim tayyorladi? Javob: Sergey Korolyov tayyorladi.',
    'Fizika: Atom bombasi loyihasi. Igor Kurchatov bajardi.',
    'No title slide. Harry Potter characters: Student: Harry Potter; Teacher: Severus Snape; School: Hogwarts',
    "Birinchi slaydga: Alisher Navoiy\n10 ta slayd bo'lsin",
    "Olimpiada g'olibi Karimova Dilnoza 9-A sinf o'quvchisi deb yoz",
    'Победитель олимпиады: Иванова Анна 205-группа',
    'Sobirov Jasur 302-guruh. Bu talaba haqida 10 slayd tayyorla',
    'Topshiriqlar taqsimoti:\n301-guruh Abdulla Oripov\n302-guruh Erkin Vohidov\nHar biriga 10 slayd',
    'In the quote change the name to Albert Einstein',
    # A fifth review's: an organisation credited, the thing done named before who did it, a line quoted to
    # come off, a byline far from a fact list, the title page as a rank or a book's cover, a template, a
    # thank-you slide, a character introducing themselves, portraits on the first slide.
    'Add a sources slide. Prepared by: World Health Organization',
    'Условные рефлексы. Опыт выполнил: Иван Павлов. 10 слайдов',
    'Периодический закон. Сделал: Дмитрий Менделеев. 8 слайдов',
    'Убери «Проверила: Рашидова Малика» с титульного слайда',
    "Hasanboy Do'smatov hayoti. Olimpiada chempioni titulini qo'lga kiritgan. Ustozi: Tulqin Kilichev. 10 ta slayd",
    'Kitob muqovasi qizil rangda edi. Muallif: Abdulla Qodiriy. Talaba: Otabek Yusufbekov',
    'Шаблон титульника:\nВыполнил: Иванов Иван\nПроверил: Петров Пётр',
    'Rahmat slaydiga yoz: Tayyorladi: Karimova Dilnoza',
    'На первом слайде Наполеон представляется: меня зовут Наполеон Бонапарт',
    'На первом слайде портреты:\nИван Павлов\nДмитрий Менделеев',
]


@pytest.mark.parametrize('case', list(ASKED))
def test_the_lines_a_customer_dictates_are_read_whole_and_in_order(case):
    text, lines = ASKED[case]
    assert dictated(text) == lines


@pytest.mark.parametrize('text', NOT_ASKED + LEFT_TO_THE_MODEL)
def test_a_name_in_the_topic_or_an_audience_is_not_a_title_line(text):
    assert dictated(text) == []


@pytest.mark.parametrize('text', [
    '1' + ' ' * 20000, '1-' + ' ' * 20000, 'вверху' + ' ' * 20000, 'Ab' + ' ' * 20000 + 'Cdova', 'deb yoz ' * 2500,
    'Ali Valiyeva 11-sinf ' * 1000, 'Birinchi betiga: ' * 1000, 'Karimova Dilnoza 026-136-guruh deb yoz. ' * 500,
    'Tayyorladi: Ab Cdova, ' * 1000, 'Ustki qismiga ' * 1400, 'MUSTAQIL ISH\nTayyorladi: Ab Cdova\n' * 600,
    'Ab Cdova tayyorladi ' * 1000, 'Ab, ' * 5000, "'" * 20000, '\n' * 20000, 'Ab Cdova deb yoz ' * 1200,
    'ab cdova deb yoz ' * 1200, '"Ab Cdova" deb yoz ' * 1000, 'Tayyorladi qismiga Ab Cdova deb yoz ' * 500,
    'Tayyorladi:\n' * 2000, 'ismim Ab Cdova ' * 1300, 'Birinchi betiga:\nAb Cdova\n\n' * 900, 'Ab Cdova\n' * 2000,
    ('yoz ' + 'ismim dr, ' * 2000)[:20000], ('напиши ' + 'моё имя Аб Цдова, ' * 1112)[:20000],
    'Tayyorladi: ' + '.' * 200 + 'xyz Abdu Karimova', 'Birinchi betiga:\n' + '.' * 500 + 'xyz Karimova Dilnoza',
    '. ' + ' ' * 19990 + '1', ('Birinchi betiga:\nTayyorladi: Ab Cdova' + ' ' * 2350 + 'x\n\n') * 8,
    (('Birinchi betiga: ' + 'dr. ' * 600 + 'slayd\n') * 10)[:20000], ('Birinchi betiga:\n' + 'A. ' * 800 + '\n') * 8,
])
def test_a_long_request_is_read_in_linear_time(text):
    import time
    started = time.perf_counter()
    dictated(text)
    assert time.perf_counter() - started < 0.5


# ---------------------------------------------------------------- the cover


def deck(cover_body='Tarixiy sharh', slides=3):
    sections = [{'id': 's1', 'heading': 'Ilk oʻrta asr davlatlari', 'body': cover_body, 'layout': 'cover'}]
    sections += [{'id': f's{n}', 'heading': f'Slide {n}', 'body': 'Text.', 'layout': 'bullets'}
                 for n in range(2, slides + 1)]
    return {'title': 'Deck', 'sections': sections, 'questions': [], 'citations': []}


def created(prompt, locale='uz'):
    return {'output_format': 'pptx', 'output_locale': locale, 'prompt': prompt, 'options': {}}


def changed(request, original='', locale='uz', **revision):
    return {'output_format': 'pptx', 'output_locale': locale, 'prompt': request, 'options': {},
            'revision': {'request': request, 'original_prompt': original, **revision}}


def cover(content):
    return content['sections'][0]['body']


def test_a_new_deck_carries_the_lines_its_description_dictates():
    data = created("Tarix fanidan 10 ta slayd. Tayyorladi: Karimova Dilnoza. Qabul qildi: Rashidova Malika")
    # The model merged the names into one line: they are written back as the customer wrote them.
    content = title_lines.apply(data, deck('Karimova Dilnoza · Rashidova Malika'))
    assert cover(content) == 'Tayyorladi: Karimova Dilnoza\nQabul qildi: Rashidova Malika'
    assert data['options']['title_lines'] == ['Tayyorladi: Karimova Dilnoza', 'Qabul qildi: Rashidova Malika']


def test_a_change_adds_its_lines_and_takes_nothing_off_the_cover():
    data = changed("напиши в начале слайда сделала: Каримова Дилноза 205-группа", locale='ru')
    content = title_lines.apply(data, deck('Историческая грамматика'))
    assert cover(content) == 'Историческая грамматика\nсделала: Каримова Дилноза 205-группа'


def test_a_change_that_dictates_nothing_leaves_the_cover_to_the_model():
    # The model is given the deck as it stands; nothing is put back or taken off for it.
    for request in ("Rasmlarni ko'paytir", "Ismimni olib tashla", "Taqdimotni ingliz tiliga o'gir"):
        content = deck('Tarixiy sharh\nPrepared by: Karimova Dilnoza')
        assert title_lines.apply(changed(request), content) == content, request


def test_a_line_is_written_in_the_alphabet_of_the_deck():
    from apps.studio.scripts import violations
    data = changed("напиши в начале слайда сделала: Каримова Дилноза 205-группа", locale='uz')
    content = title_lines.apply(data, deck('Tarixiy grammatika'))
    assert cover(content).split('\n')[1].startswith('sdelala: Karimova Dilnoza 205')
    assert not violations(content, 'uz')


def test_a_change_dictating_a_line_replaces_the_models_line_for_it():
    # A fuller name for the same person; a new role beside the others; a byline dictated again is that person.
    data = changed("Boshiga bitta sahifa qoʻsh va Karimova Dilnoza Aziz qizi 026-136-guruh deb yoz")
    content = title_lines.apply(data, deck('Tarixiy sharh\nKarimova Dilnoza 026-136-guruh'))
    assert cover(content) == 'Tarixiy sharh\nKarimova Dilnoza Aziz qizi 026-136-guruh'
    data = changed("Qabul qildi: Rashidova Malika")
    assert cover(title_lines.apply(data, deck('Tarixiy sharh\nTayyorladi: Karimova Dilnoza'))) == (
        'Tarixiy sharh\nTayyorladi: Karimova Dilnoza\nQabul qildi: Rashidova Malika')
    data = changed("Tayyorladi: Yusupov Bobur")
    assert cover(title_lines.apply(data, deck('Tarixiy sharh\nTayyorladi: Karimova Dilnoza'))) == (
        'Tarixiy sharh\nTayyorladi: Yusupov Bobur')
    # Another author added, not the first replaced; a teacher who shares the student's surname is another person.
    data = changed("Muqovaga yana bir ism qo'sh: Tayyorladi: Yusupova Lola")
    assert cover(title_lines.apply(data, deck('Tarixiy sharh\nTayyorladi: Karimova Dilnoza'))) == (
        'Tarixiy sharh\nTayyorladi: Karimova Dilnoza\nTayyorladi: Yusupova Lola')
    data = changed("Birinchi slaydga Qabul qildi: Karimova Nodira deb yoz")
    assert cover(title_lines.apply(data, deck('Tarixiy sharh\nTayyorladi: Karimova Dilnoza'))) == (
        'Tarixiy sharh\nTayyorladi: Karimova Dilnoza\nQabul qildi: Karimova Nodira')


def test_a_correction_is_made_in_the_line_it_corrects():
    data = changed("Sobirova Madina o'rniga Sobirova Mohinur deb yoz")
    assert cover(title_lines.apply(data, deck('Tarix\nTayyorladi: Sobirova Madina'))) == 'Tarix\nTayyorladi: Sobirova Mohinur'
    data = changed('Karimova Dilnoza emas, Rashidova Dilnoza deb yoz')
    assert cover(title_lines.apply(data, deck('Tarix\nTayyorladi: Karimova Dilnoza'))) == 'Tarix\nTayyorladi: Rashidova Dilnoza'
    # Already made by the model: nothing doubled.
    data = changed('Ismimni Sobirova Malika deb yoz')
    body = 'Tarix\nTayyorladi: Sobirova Malika\nQabul qildi: Nazarova Feruza'
    assert cover(title_lines.apply(data, deck(body))) == body


def test_the_models_subtitle_stays_whatever_it_shares_with_a_line():
    # A year, a given name, a role word: none makes the subtitle the customer's line.
    data = created("Birinchi betiga:\nMustaqil ish\nTayyorladi: Karimova Dilnoza\nToshkent 2025")
    assert cover(title_lines.apply(data, deck('Tarixiy sharh, 2025\nMustaqil ish'))) == (
        'Tarixiy sharh, 2025\nMustaqil ish\nTayyorladi: Karimova Dilnoza\nToshkent 2025')
    data = created("Tarix 10 slayd. Tayyorladi: Yusupov Bobur")
    assert cover(title_lines.apply(data, deck('Bobur hayoti va ijodi'))) == 'Bobur hayoti va ijodi\nTayyorladi: Yusupov Bobur'
    data = created("Tarix 10 slayd. Tayyorladi: Karimova Dilnoza")
    assert cover(title_lines.apply(data, deck("O'quvchilar uchun tarix"))) == (
        "O'quvchilar uchun tarix\nTayyorladi: Karimova Dilnoza")


def test_the_customers_lines_never_take_each_others_place():
    data = created("Birinchi betiga:\nTayyorladi: Ergasheva Malika\nTayyorladi: Yusupova Lola\nQabul qildi: Xolmatova Sevara")
    assert cover(title_lines.apply(data, deck('Pul muomalasi'))) == (
        'Pul muomalasi\nTayyorladi: Ergasheva Malika\nTayyorladi: Yusupova Lola\nQabul qildi: Xolmatova Sevara')
    data = created("Birinchi betiga:\nKarimova Dilnoza 202-guruh\nRashidova Malika 202-guruh")
    assert cover(title_lines.apply(data, deck('Pul muomalasi'))) == (
        'Pul muomalasi\nKarimova Dilnoza 202-guruh\nRashidova Malika 202-guruh')
    data = changed("Tayyorladi: Karimova Dilnoza Aziz qizi")
    assert cover(title_lines.apply(data, deck('Tarix\nTayyorladi: Karimova Dilnoza\n202-guruh talabasi\nQabul qildi: Rashidova Malika'))) == (
        'Tarix\nTayyorladi: Karimova Dilnoza Aziz qizi\n202-guruh talabasi\nQabul qildi: Rashidova Malika')


def test_the_models_own_version_of_a_line_is_never_doubled():
    # A fuller line, its own label, its own way of saying who, a label over its name, a placeholder.
    data = created('Tarix 10 slayd. Tayyorladi: Sobirov Jasur')
    assert cover(title_lines.apply(data, deck('Tarix\nTayyorladi: Sobirov Jasur, MT-21 guruh'))) == (
        'Tarix\nTayyorladi: Sobirov Jasur, MT-21 guruh')
    data = created('Презентация 10 слайдов. Выполнила: Каримова Дилноза', locale='ru')
    assert cover(title_lines.apply(data, deck('Экология\nАвтор: Каримова Д.'))) == 'Экология\nВыполнила: Каримова Дилноза'
    data = created("Biologiya 9-sinf uchun 10 ta slayd. Tayyorladi: Karimova Dilnoza")
    assert cover(title_lines.apply(data, deck("Hujayra tuzilishi\n9-sinf o'quvchisi Karimova Dilnoza"))) == (
        'Hujayra tuzilishi\nTayyorladi: Karimova Dilnoza')
    data = created("Tarix 10 slayd. Tayyorladi: Karimova Dilnoza. Qabul qildi: Rashidova Malika")
    assert cover(title_lines.apply(data, deck('Tarix\nTayyorladi:\nKarimova Dilnoza\nQabul qildi:\nRashidova Malika'))) == (
        'Tarix\nTayyorladi: Karimova Dilnoza\nQabul qildi: Rashidova Malika')
    data = created("Tarix 10 slayd. Tayyorladi: Karimova Dilnoza")
    assert cover(title_lines.apply(data, deck('Tarix\nTayyorladi: ______'))) == 'Tarix\nTayyorladi: Karimova Dilnoza'
    # Two people under one label become the customer's two lines.
    data = created('Tarix 10 slayd. Tayyorladi: Karimova Dilnoza. Qabul qildi: Rashidova Malika')
    assert cover(title_lines.apply(data, deck('Tarix\nTayyorladi: Karimova Dilnoza va Rashidova Malika'))) == (
        'Tarix\nTayyorladi: Karimova Dilnoza\nQabul qildi: Rashidova Malika')


def test_a_dictated_title_page_keeps_its_order_around_the_models_lines():
    prompt = ('Tarix 10 slayd.\nBirinchi betiga:\nToshkent davlat iqtisodiyot universiteti\nMustaqil ish\n'
              'Tayyorladi: Sobirova Madina\nQabul qildi: Nazarova Feruza')
    for body in ('Tarix\nTayyorladi: Sobirova Madina\nQabul qildi: Nazarova Feruza', 'Tarix\nSobirova Madina · Nazarova Feruza'):
        assert cover(title_lines.apply(created(prompt), deck(body))) == (
            'Tarix\nToshkent davlat iqtisodiyot universiteti\nMustaqil ish\nTayyorladi: Sobirova Madina\n'
            'Qabul qildi: Nazarova Feruza')


def test_the_outline_as_the_customer_left_it_says_which_lines_are_theirs():
    prompt = "Tarix fanidan 4 ta slayd. Tayyorladi: Karimova Dilnoza."
    stored = {'title_lines': ['Tayyorladi: Karimova Dilnoza']}
    data = {**created(prompt), 'options': dict(stored), 'content': deck('Tarix\nTayyorladi: Karimova Dilnoza Azizovna')}
    assert cover(title_lines.apply(data, deck('Tarix'))) == 'Tarix\nTayyorladi: Karimova Dilnoza Azizovna'
    data = {**created(prompt), 'options': dict(stored), 'content': deck('Tarix')}
    assert cover(title_lines.apply(data, deck('Tarix'))) == 'Tarix' and data['options']['title_lines'] == []


def test_a_line_carries_nothing_a_slide_cannot_print():
    data = created("Tarix fanidan 4 ta slayd. Tayyorladi: Karimova Dilnoza\ufffe\x00\nIlk davlatlar")
    assert cover(title_lines.apply(data, deck('Tarix'))) == 'Tarix\nTayyorladi: Karimova Dilnoza'


def test_lines_placed_twice_stay_as_they_are():
    data = created("Tarix 10 slayd. Tayyorladi: Karimova Dilnoza. Qabul qildi: Rashidova Malika")
    once = title_lines.apply(data, deck('Tarixiy sharh\nKarimova Dilnoza'))
    assert title_lines.apply(data, once) == once  # the outline, then the job


def test_a_long_cover_gives_up_the_models_lines_first():
    data = created("Birinchi betiga:\nMustaqil ish\nTayyorladi: Karimova Dilnoza\nQabul qildi: Rashidova Malika")
    content = title_lines.apply(data, deck('\n'.join(f'Satr {n}' for n in range(1, 8))))
    lines = cover(content).split('\n')
    assert len(lines) == title_lines.TITLE_LINES and lines[-3:] == [
        'Mustaqil ish', 'Tayyorladi: Karimova Dilnoza', 'Qabul qildi: Rashidova Malika']


def test_a_cover_outside_the_selection_or_a_document_is_never_touched():
    data = changed("Tayyorladi: Yusupov Bobur", selected_section_ids=['s2'])
    assert cover(title_lines.apply(data, deck('Tarixiy sharh'))) == 'Tarixiy sharh'
    document = {**created("Tayyorladi: Karimova Dilnoza"), 'output_format': 'pdf'}
    assert cover(title_lines.apply(document, deck('Tarixiy sharh'))) == 'Tarixiy sharh'
    one = created("Tayyorladi: Karimova Dilnoza")
    assert cover(title_lines.apply(one, deck('Tarixiy sharh', slides=1))) == 'Tarixiy sharh'


def test_a_title_line_is_never_taken_from_a_client():
    from apps.studio.domain import SLIDES, create_draft, unpack
    from tests.test_platform import account
    from django.conf import settings
    settings.DEBUG = True
    draft = create_draft(account(), {'feature_id': SLIDES, 'prompt': 'A 4 slide deck on volcanoes.',
                                     'options': {'title_lines': ['Anything a client sends']}})
    assert 'title_lines' not in unpack(draft.encrypted_data)['options']


# ---------------------------------------------------------------- the job


@pytest.fixture
def model(settings, monkeypatch):
    """A model that writes every deck's cover as one subtitle line: no names, whatever was asked."""
    from apps.studio import provider
    from tests.test_page_fill import live
    customer = live(settings)

    def generate(config, data, *args, **kwargs):
        count = len(data['content']['sections'])
        sections = [{'id': section['id'], 'heading': 'Ilk oʻrta asr davlatlari' if not n else f'Davr {n}',
                     'body': 'Tarixiy sharh' if not n else 'Davlatlar tarixi qisqacha bayon etiladi.',
                     'notes': '', 'layout': 'cover' if not n else 'bullets', 'items': [], 'columns': [],
                     'image_query': ''}
                    for n, section in enumerate(data['content']['sections'][:count])]
        return ({'title': 'Ilk oʻrta asr davlatlari', 'answer_supported': True, 'citations': [], 'questions': [],
                 'sections': sections}, {'input_tokens': 1, 'output_tokens': 1})
    monkeypatch.setattr(provider, 'generate', generate)
    return customer


def _run(customer, prompt, revising=None):
    from apps.core.services import execute_job, submit_job
    from apps.core.services import storage_path
    from apps.studio.domain import SLIDES, create_draft, generation_quote
    payload = {'feature_id': SLIDES, 'prompt': prompt}
    if revising:
        payload['options'] = {'revise_draft_id': str(revising.id)}
    draft = create_draft(customer, payload)
    quote = generation_quote(customer, draft.id, draft.version)
    job, _ = submit_job(customer, quote.id, f'title-{draft.id}')
    job = execute_job(job.id)
    assert job.status == 'succeeded', job.error_code
    cover = Presentation(str(storage_path(job.artifacts.get().file.object_key))).slides[0]
    texts = [paragraph.text for shape in cover.shapes if shape.has_text_frame
             for paragraph in shape.text_frame.paragraphs if paragraph.text.strip()]
    draft.refresh_from_db()
    return draft, texts


def test_names_reach_the_cover_of_a_deck_and_of_a_change_that_dictates_them(model):
    # The model never writes the names; the deck and the change that dictates a fuller one carry them anyway.
    first, cover = _run(model, "Tarix fanidan 4 ta slayd kerak-Karimova Dilnoza 026-136-guruh. Ilk oʻrta asr davlatlari")
    assert 'Karimova Dilnoza 026-136-guruh' in cover
    second, cover = _run(model, "Boshiga bitta sahifa qoʻsh va Karimova Dilnoza Aziz qizi 026-136-guruh deb yoz",
                         revising=first)
    assert 'Karimova Dilnoza Aziz qizi 026-136-guruh' in cover and 'Karimova Dilnoza 026-136-guruh' not in cover


def test_a_slide_asked_for_at_the_start_with_names_is_the_cover():
    from apps.studio.edits import read
    sections = deck(slides=5)['sections']
    for request in ("Boshiga bitta slayd qo'sh: Karimova Dilnoza 026-136-guruh",
                    "Add a slide at the beginning: prepared by Sobirov Jasur"):
        edit = read(request, fmt='pptx', sections=sections)
        assert edit.delta == 0 and edit.parts == ['title_page'], request


def test_a_null_or_a_reworded_description_never_keeps_title_lines():
    from apps.studio.domain import SLIDES, create_draft, pack, unpack, update_draft
    from tests.test_platform import account
    from django.conf import settings
    settings.DEBUG = True
    customer = account()
    draft = create_draft(customer, {'feature_id': SLIDES, 'prompt': 'A 4 slide deck on volcanoes.',
                                    'options': {'title_lines': None}})
    assert 'title_lines' not in unpack(draft.encrypted_data)['options']
    data = unpack(draft.encrypted_data)
    data['options']['title_lines'] = ['Tayyorladi: Karimova Dilnoza']
    draft.encrypted_data = pack(data)
    draft.save()
    saved = update_draft(customer, draft.id, {'version': draft.version, 'prompt': 'A 4 slide deck on earthquakes.'})
    assert 'title_lines' not in unpack(saved.encrypted_data)['options']


def test_long_title_lines_never_make_a_draft_too_large_to_save():
    from apps.studio.domain import SLIDES, create_draft, pack, unpack, update_draft
    from tests.test_platform import account
    from django.conf import settings
    settings.DEBUG = True
    customer = account()
    draft = create_draft(customer, {'feature_id': SLIDES, 'prompt': 'A 4 slide deck on volcanoes.'})
    # As a job leaves them: eight long Cyrillic lines, far past 4000 characters as JSON.
    lines = ['Выполнила: ' + 'Каримова Дилноза Азизовна, студентка ' * 4] * title_lines.TITLE_LINES
    data = unpack(draft.encrypted_data)
    data['options']['title_lines'] = lines
    draft.encrypted_data = pack(data)
    draft.save()
    # The web editor saving the deck sends its options back, the stored lines with them.
    saved = update_draft(customer, draft.id, {'version': draft.version, 'title': 'Volcanoes', 'options': data['options']})
    assert unpack(saved.encrypted_data)['options']['title_lines'] == lines
