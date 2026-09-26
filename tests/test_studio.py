import io
import json
import uuid
import pytest
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.test import Client
from django.utils import timezone
from pypdf import PdfReader
from pptx import Presentation
from apps.core.services import submit_job,execute_job,storage_path,create_quote
from apps.core.models import UsageLedger
from apps.core.errors import DomainError
from apps.studio.domain import create_draft,update_draft,generation_quote,draft_data,unpack
from apps.studio.models import EducationProject,GenerationDraft,ShareGrant
from apps.commerce.services import create_invoice,sandbox_pay
from operations.integrations import save_config,telegram_config
from operations.models import IntegrationConfig,AuditLog
from tests.test_platform import account,upload,login_client
from operations.auth import begin_session,COOKIE
from django.http import HttpResponse
pytestmark=pytest.mark.django_db

@pytest.fixture(autouse=True)
def studio_local(settings):
    settings.DEBUG=True
    settings.COMMERCE_SANDBOX_ENABLED=True


def paid(settings,a):
    settings.COMMERCE_SANDBOX_ENABLED=True
    inv,_=create_invoice(a,'premium','studio-premium-'+str(a.id));sandbox_pay(a,inv.id);a.refresh_from_db();return a

def content():
    return {'title':'O‘zbekcha — Русский','sections':[{'heading':'Bo‘lim / Раздел','body':'O‘qituvchi bilim beradi. Учитель помогает учиться.'}],'questions':[{'id':'q1','stem':'2 + 2?','options':['3','4'],'answer':'4','explanation':'Two pairs make four.','topic':'addition','marks':2}]}

def test_generation_pdf_pptx_real_outputs_and_no_fake_ai_charge(settings):
    a=account();cfg={'title':'O‘zbekcha — Русский','source_text':'O‘qituvchi bilim beradi. Учитель помогает учиться.'}
    for feature,fmt in (('ai.pdf_topic','pdf'),('ai.pptx','pptx')):
        d=create_draft(a,{**cfg,'feature_id':feature,'output_format':fmt});q=generation_quote(a,d.id,d.version)
        job,_=submit_job(a,q.id,'studio-'+fmt);job=execute_job(job.id)
        assert job.status=='succeeded',job.error_code
        assert job.warnings==['local_fixture_not_ai_generated']
        assert not UsageLedger.objects.filter(job=job,kind='consume').exists()
        asset=job.artifacts.get().file;path=storage_path(asset.object_key)
        if fmt=='pdf':assert 'Учитель' in ''.join(p.extract_text() for p in PdfReader(path).pages)
        else:
            pres=Presentation(path);assert any('Учитель' in s.text for sl in pres.slides for s in sl.shapes if s.has_text_frame)

def test_version_owner_and_quote_integrity():
    a=account();other=account(43);d=create_draft(a,{'feature_id':'ai.pdf_topic','source_text':'Hello'})
    assert b'Hello' not in bytes(d.encrypted_data)
    q=generation_quote(a,d.id,1)
    d=update_draft(a,d.id,{'version':1,'content':content()})
    with pytest.raises(DomainError,match='version_conflict'):submit_job(a,q.id,'stale-studio-quote')
    with pytest.raises(DomainError,match='version_conflict'):update_draft(a,d.id,{'version':1,'content':content()})
    assert login_client(other).get(f'/api/v1/generation/drafts/{d.id}').status_code==404

def test_editor_command_cannot_bypass_plan():
    a=account();asset=upload(a)
    with pytest.raises(DomainError,match='feature_not_in_plan'):
        create_quote(a,'editor.visual',[str(asset.id)],{'commands':[{'type':'redact','page':1,'x':0,'y':0,'width':10,'height':10}],'accept_rasterization':True})

def test_credentials_are_encrypted_and_never_rendered_or_audited(settings):
    token='123456789:abcdefghijklmnopqrstuvwxyz_1234567890'
    save_config('telegram',{'token':token,'username':'pdfmaster_bot','webapp_url':'http://localhost:3000/en/app'})
    assert telegram_config()['token']==token
    raw=bytes(IntegrationConfig.objects.get(pk='telegram').encrypted_secrets)
    assert token.encode() not in raw
    u=get_user_model().objects.create(username='testadmin',is_staff=True);u.groups.add(Group.objects.get_or_create(name='Administrator')[0])
    c=Client();resp=begin_session(HttpResponse(),u);c.cookies[COOKIE]=resp.cookies[COOKIE].value
    r=c.get('/ops/integrations');assert r.status_code==200;assert token.encode() not in r.content
    r=c.post('/ops/integrations',{'integration':'telegram','action':'save','reason':'Rotate bot configuration','username':'pdfmaster_bot','webapp_url':'http://localhost:3000/en/app','token':token});assert r.status_code==302
    assert token not in json.dumps(list(AuditLog.objects.values('reason','before','after')))
    assert login_client(account()).post('/ops/integrations',{}).status_code==302
