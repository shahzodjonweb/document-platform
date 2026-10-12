"""A change request is delivered as it was asked: every dictated title-page line, every part of the request.

Title pages: a student's title page is dictated line by line ("MUSTAQIL ISH",
"TAYYORLADI: …", "QABUL QILDI: …", the group). A cover or closing slide used
to draw only the first line of its body; every line is drawn now, in the
classic layout and in each composition's, fitted to its box, while a one-line
subtitle is drawn exactly as before. A PDF keeps such lines whole when it
shortens a long document to its page count.

Change requests: the model splits a change into its instructions first
(`edit_plan`, with the sections each goes in). An instruction whose sections
all came back exactly as they were was not applied, and goes into the one
quality repair — only those sections, with the instruction's own words. A new
look, a kind of picture, a file format, a test, "keep the rest" or spelling
leaves the sections as they were and is never sent there. Pictures a change
asked for and the result lacks go into the same repair, when the plan can add
any; pictures it asked to take away are taken away without a model call, from
the slides the plan or the request names and never from the whole deck.

The names below are made up.
"""
import json

import pytest
from pptx import Presentation

from apps.studio import quality, rendering
from apps.studio.compositions.check import problems
from apps.studio.deck_designs import look
from apps.studio.domain import create_draft, generation_quote, unpack
from apps.studio.slides import ZONES, render_pptx
from apps.core.services import execute_job, submit_job
from tests.test_generation_quality import customer  # noqa: F401  (the fixture)
from tests.test_generation_quality import ORIGINAL, by_id, generated_deck, model, reply, sent, slide, stages

# ---------------------------------------------------------------- title pages

LINES = {
    'en': ['INDEPENDENT WORK', 'Prepared by: Karimova Nodira Rustam qizi', 'Accepted by: Rahimova Zarina',
           'Group 021-104', 'Tashkent 2026'],
    'uz': ['MUSTAQIL ISH', 'TAYYORLADI: KARIMOVA NODIRA RUSTAM QIZI', 'QABUL QILDI: RAHIMOVA ZARINA',
           '021-104-guruh talabasi', 'Toshkent — 2026'],
}
# One design of the classic arrangement and three compositions whose subtitle
# box was too small for five lines: their title moves up to make room.
DESIGNS = ['atlas', 'groove', 'accolade', 'doodle']


def section(sid, heading, body, layout):
    return {'id': sid, 'heading': heading, 'body': body, 'notes': '', 'layout': layout, 'items': [], 'columns': [],
            'image_query': ''}


def title_deck(cover_body, closing_body='Thank you for listening'):
    return {'title': 'Pedagogika fanining rivojlanishi', 'questions': [], 'citations': [], 'sections': [
        section('s1', 'O‘zbekistonda pedagogika fanining rivojlanishi', cover_body, 'cover'),
        section('s2', 'Asosiy bosqichlar', 'Birinchi bosqich milliy maktablar tarmog‘i kengaygan davr.\n'
                'Ikkinchi bosqich pedagogika fakultetlari ochilgan davr.', 'bullets'),
        section('s3', 'E’tiboringiz uchun rahmat', closing_body, 'closing'),
    ]}


def paragraphs(slide):
    return [paragraph.text for shape in slide.shapes if shape.has_text_frame
            for paragraph in shape.text_frame.paragraphs if paragraph.text.strip()]


def box_of(slide, text):
    for shape in slide.shapes:
        if shape.has_text_frame and text in shape.text_frame.text:
            return shape
    raise AssertionError(f'{text!r} is not on the slide')


@pytest.mark.parametrize('design', DESIGNS)
@pytest.mark.parametrize('locale', ['en', 'uz'])
def test_every_dictated_title_page_line_is_drawn_on_the_cover_and_the_closing_slide(tmp_path, locale, design):
    body = '\n'.join(LINES[locale])
    path = tmp_path / 'deck.pptx'
    result = render_pptx(title_deck(body, body), path, locale=locale, style={'deck_design': design})
    deck = Presentation(str(path))
    assert result['metadata']['layouts'] == ['cover', 'bullets', 'closing']
    for slide in (deck.slides[0], deck.slides[2]):
        drawn = paragraphs(slide)
        for line in LINES[locale]:
            assert line in drawn, f'{design}: {line!r} is not drawn whole as its own line: {drawn}'
        sizes = {run.font.size.pt for shape in slide.shapes if shape.has_text_frame
                 for paragraph in shape.text_frame.paragraphs for run in paragraph.runs
                 if LINES[locale][1] in paragraph.text}
        assert sizes and min(sizes) >= 12, sizes
    assert not result['shortened']
    assert not problems(path, result['metadata']['layouts']), problems(path, result['metadata']['layouts'])


@pytest.mark.parametrize('design', DESIGNS)
def test_a_one_line_subtitle_is_drawn_where_and_as_it_always_was(tmp_path, design):
    path = tmp_path / 'deck.pptx'
    render_pptx(title_deck('Pedagogika tarixi bo‘yicha taqdimot'), path, locale='uz', style={'deck_design': design})
    cover = Presentation(str(path)).slides[0]
    shape = box_of(cover, 'Pedagogika tarixi bo‘yicha taqdimot')
    assert len(shape.text_frame.paragraphs) == 1
    if look({'deck_design': design})['composition'] == 'classic':
        zone = ZONES['cover_subtitle']
        assert (shape.left, shape.top, shape.width, shape.height) == (zone.left, zone.top, zone.width, zone.height)
        assert shape.text_frame.paragraphs[0].runs[0].font.size.pt == 18
    else:
        from apps.studio.compositions import Canvas, compose
        from apps.studio.slides import palette
        chosen = look({'deck_design': design})
        roles = palette(chosen['accent'], chosen['theme'], chosen['secondary'], chosen.get('paper'))
        scratch = Presentation()
        plan = compose(chosen['composition']).ground(
            Canvas(scratch.slides.add_slide(scratch.slide_layouts[6]), roles, chosen, 'cover'), 'cover')
        assert round(shape.top / 914400, 2) == round(plan.subtitle.box[1], 2)
        assert round(shape.height / 914400, 2) == round(plan.subtitle.box[3], 2)


def test_past_the_line_limit_the_rest_is_left_out_and_reported(tmp_path):
    from apps.studio.slide_layouts import TITLE_LINES
    extra = ['Ilmiy rahbar: Aliyev Sardor', 'Kafedra: Pedagogika', 'Fakultet: Tarix', 'Kurs: 2-kurs', 'Semestr: 1']
    lines = (LINES['uz'] + extra)[:TITLE_LINES + 1]
    path = tmp_path / 'deck.pptx'
    result = render_pptx(title_deck('\n'.join(lines)), path, locale='uz', style={'deck_design': 'atlas'})
    drawn = paragraphs(Presentation(str(path)).slides[0])
    assert all(line in drawn for line in lines[:TITLE_LINES]) and lines[TITLE_LINES] not in drawn
    assert result['shortened']


def test_a_list_marker_is_taken_off_a_title_page_line_but_an_initial_is_not(tmp_path):
    path = tmp_path / 'deck.pptx'
    render_pptx(title_deck('• Tayyorladi: N. Karimova\nA. Rahimova'), path, locale='uz',
                style={'deck_design': 'atlas'})
    drawn = paragraphs(Presentation(str(path)).slides[0])
    assert 'Tayyorladi: N. Karimova' in drawn and 'A. Rahimova' in drawn


def test_a_document_shortened_to_its_page_count_keeps_a_title_page_whole():
    title = {'id': 's1', 'heading': 'MUSTAQIL ISH', 'body': 'Tayyorladi: Karimova Dilnoza Rustam qizi\n021-104-guruh'}
    prose = {'id': 's2', 'heading': 'Kirish', 'body': 'Birinchi gap uzun. Ikkinchi gap ham bor. Uchinchi gap tugaydi.'}
    trimmed = rendering._trimmed([title, prose], 0.4)
    assert trimmed[0]['body'] == title['body']
    assert trimmed[1]['body'].endswith('…')


# ---------------------------------------------------------------- the edit plan, unit by unit

MOON = slide('The Moon pulls the oceans', 'Tides rise and fall twice a day because the Moon pulls the oceans.')
PILOTS = slide('Harbour pilots and the tide', 'Pilots time each arrival to the rising tide in narrow channels.')


def sections(**written):
    return {'title': 'Ocean tides', 'questions': [], 'citations': [],
            'sections': [{'id': sid, 'items': [], 'columns': [], 'image_query': '', **fields}
                         for sid, fields in written.items()]}


def test_a_planned_part_whose_sections_came_back_unchanged_is_unapplied():
    before = sections(s1=MOON, s2=PILOTS, s3=MOON | {'heading': 'Tide tables'})
    after = sections(s1=MOON, s2=PILOTS | {'body': 'A new paragraph about pilots.'},
                     s3=MOON | {'heading': 'Tide tables'})
    plan = [{'instruction': 'Rewrite the pilots slide', 'sections': ['s2']},
            {'instruction': 'Add a worked example to slides 1 and 3', 'sections': ['s1', 's3']},
            {'instruction': 'Make it friendlier', 'sections': []},
            {'instruction': 'Add a picture to slide 1', 'sections': ['s1', 's2']},
            'not an item', {'instruction': '', 'sections': ['s1']}]
    report = quality.Report()
    quality.check_plan(report, after, before, plan)
    # A part counts as done when any of its sections changed.
    assert report.unapplied == [(['s1', 's3'], 'Add a worked example to slides 1 and 3')]
    assert report.to_repair() == ['s1', 's3']
    assert report.summary() == {'unplaced': 1, 'unapplied': 2}
    text = quality.request_text(report, 'en', 'pptx', ['s1', 's3'])
    assert ('s1, s3: this part of the change request is not applied yet: "Add a worked example to slides 1 and 3". '
            'Apply it fully now.') in text
    assert 'pilots' not in text and 'friendlier' not in text


def test_a_change_to_selected_pages_is_held_only_to_those_pages():
    before = sections(s1=MOON, s2=PILOTS)
    plan = [{'instruction': 'Shorten both slides', 'sections': ['s1', 's2']},
            {'instruction': 'Retitle the first slide', 'sections': ['s1']}]
    report = quality.Report()
    quality.check_plan(report, before, before, plan, selected=['s2'])
    assert report.unapplied == [(['s2'], 'Shorten both slides')]
    assert report.to_repair() == ['s2']


# Parts of real requests that leave a section as it was, or that a repair of
# section text cannot do, with the change as edits.read reads it.
NOT_HELD = [
    ('Varoqlarni hammasini bir hil qil', {'redesign': True}),
    ('Make all slides look the same', {'redesign': True}),
    ('Dizaynni chiroyli qil', {'redesign': True}),
    ('Make the design nicer', {'redesign': True}),
    ('Orasidagi daftar varogʻiga oxshagan sahifalarni boshqasiga almashtir', {'redesign': True}),
    ('Yosh bolacha rasm qoʻyma', {'photo_style': 'photo'}),
    ('Realga oʻxshasin', {'photo_style': 'photo'}),
    ('Rasmlar kitoblarda beriladigan rasmga o‘xshashi kerak', {'photo_style': 'illustration'}),
    ('pdf shaklida yubor', {'format': 'pdf'}),
    ('Taqdimot shaklida qilib bering', {'format': 'pptx'}),
    ('Test tuz', {'questions': 'add'}),
    ('Тест туз', {'questions': 'add'}),
    ('Vaziyatli testlarni oʻchir', {'questions': 'remove'}),
    ('Hammasi avvalgidek bo‘lsin', {}),
    ('Keep everything else as it was', {}),
    ('hamma yozuvlar lotin alifbosida xatosiz yozilsin', {}),
    ('Krilcha harflarda yozilgan joylarni to‘g‘rila', {}),
    ('Imlosiga etibor ber', {}),
    ('Rus alifbosini ishlatma', {}),
    ('fix typos on slide 2 and slide 6', {}),
]


@pytest.mark.parametrize('instruction, edit', NOT_HELD)
def test_a_part_no_section_text_shows_is_not_held_to_the_plan(instruction, edit):
    before = deck_of(6)
    # Another slide changed, so the result is more than a new look.
    after = deck_of(6, s6=slide('Tides, part 6', 'A new point about the tides of the sea.'))
    report = quality.Report()
    quality.check_plan(report, after, before, [{'instruction': instruction, 'sections': ['s2', 's3']}], edit=edit)
    assert not report.unapplied and not report.issues


@pytest.mark.parametrize('instruction, edit', [
    ('Akbarali obrazini to‘liq yorit', {'redesign': True, 'photo_style': 'photo'}),
    ('Registon dizayni haqida yoz', {}),
    ('Lotin Amerikasi haqida yoz', {}),
    ('Make the design nicer', {}),
    ('Xulosani alohida slayd qilib yoz', {'format': 'pdf'}),
    ('Jinoyatga oid rasmlar ham qoʻsh', {'pictures': 'add'}),
    ('Copy the title exactly as it is: MUSTAQIL ISH', {}),
])
def test_a_part_the_text_shows_is_held_beside_a_new_look(instruction, edit):
    before = deck_of(6)
    after = deck_of(6, s6=slide('Tides, part 6', 'A new point about the tides of the sea.'))
    report = quality.Report()
    quality.check_plan(report, after, before, [{'instruction': instruction, 'sections': ['s2', 's3']}], edit=edit)
    assert report.unapplied == [(['s2', 's3'], instruction)]


def test_a_new_look_that_left_every_section_as_it_was_is_the_change_done():
    before = deck_of(6)
    plan = [{'instruction': 'Give slides 2 and 3 one feel', 'sections': ['s2', 's3']}]
    for edit in ({'redesign': True}, {'photo_style': 'photo'}, {'photo_style': 'illustration', 'pictures': None}):
        report = quality.Report()
        quality.check_plan(report, deck_of(6), before, plan, edit=edit)
        assert not report.unapplied and not report.issues, edit
    # A change asking for no new look that changed nothing was not applied.
    report = quality.Report()
    quality.check_plan(report, deck_of(6), before, plan, edit={'format': 'pdf'})
    assert report.unapplied == [(['s2', 's3'], 'Give slides 2 and 3 one feel')]


def test_a_part_listing_every_section_is_held_only_when_it_asks_for_more_information():
    before = deck_of(6)
    # The cover, a closing or a divider may be left out of "every section".
    for listed in (['s1', 's2', 's3', 's4', 's5', 's6'], ['s2', 's3', 's4', 's5', 's6']):
        report = quality.Report()
        quality.check_plan(report, before, before, [{'instruction': 'Friendlier for children', 'sections': listed}])
        assert not report.unapplied and not report.issues
    report = quality.Report()
    quality.check_plan(report, before, before, [{'instruction': 'Ko‘lproq malumot', 'sections': listed}],
                       edit={'expand': True})
    assert report.unapplied == [(listed, 'Ko‘lproq malumot')]


def picture_data(asked, fmt='pptx', **options):
    """A change asking for pictures (`asked`), on a plan that adds up to ten unless `image_cap` says otherwise.

    `images_wanted` is the pictures the document should have once changed, a
    deck's cover aside: those it had and those the change adds (domain.py).
    """
    return {'output_locale': 'en', 'output_format': fmt, 'options': {'image_cap': 10, **options},
            'revision': {'request': 'Add pictures', 'edit': {'pictures': asked}}}


def deck_of(count, **changes):
    written = {'s1': slide('Ocean tides', 'A short guide', layout='cover')}
    for number in range(2, count + 1):
        written[f's{number}'] = slide(f'Tides, part {number}', f'Point {number} about the tides of the sea.')
    written.update(changes)
    return sections(**written)


def test_pictures_asked_for_and_missing_go_to_the_repair_on_the_slides_that_suit_them():
    content = deck_of(7, s7=slide('Thank you', 'Questions?', layout='closing'),
                      s4=slide('Tide ranges', '', layout='stats') | {'items': [{'label': 'a', 'value': '1'}]})
    report = quality.Report()
    quality.check_pictures(report, content, picture_data('add', images_wanted=2))
    # Text slides first, spread through the deck; never the cover or the closing.
    assert report.issues['pictures_missing'] == ['s2', 's5']
    text = quality.request_text(report, 'en', 'pptx', report.to_repair())
    line = next(line for line in text.split('\n') if line.startswith('s2, s5:'))
    assert 'asks for pictures' in line and 'image_split' in line and 'image_query' in line
    # The picture the change was to add is the change done.
    pictured = deck_of(4, s3=slide('Tide pools', 'Rock pools at low tide.', layout='image_split')
                       | {'image_query': 'rock pool'})
    data = picture_data('add', images_wanted=1)
    data['revision']['before'] = deck_of(4)
    report = quality.Report()
    quality.check_pictures(report, pictured, data)
    assert not report.issues
    # One of the two it was to add is not.
    report = quality.Report()
    quality.check_pictures(report, pictured, {**data, 'options': {'images_wanted': 2, 'image_cap': 10}})
    assert report.issues['pictures_missing'] == ['s2']


def test_pictures_a_deck_already_had_do_not_answer_a_request_to_add_more():
    # "Jinoyatga oid rasmlar ham qoʻsh" on a deck with a cover photo and one
    # picture slide, which the model returned as it was: two to add, three in all.
    before = deck_of(5, s1=slide('Crime and law', 'A short guide', layout='cover') | {'image_query': 'courtroom'},
                     s3=slide('Evidence', 'Fingerprints and footprints.', layout='image_split')
                     | {'image_query': 'fingerprint'})
    data = picture_data('add', images_wanted=3)
    data['revision'].update(request='Jinoyatga oid rasmlar ham qoʻsh', before=before)
    report = quality.Report()
    quality.check_pictures(report, before, data)
    assert report.issues['pictures_missing'] == ['s2', 's4']
    # A new picture on the cover is not one the request asked for either.
    covered = deck_of(5, s1=before['sections'][0] | {'image_query': 'handcuffs'}, s3=before['sections'][2])
    report = quality.Report()
    quality.check_pictures(report, covered, data)
    assert report.issues['pictures_missing'] == ['s2', 's4']
    # One picture more than it had is one of the two: the other is still asked for.
    added = deck_of(5, s1=before['sections'][0], s3=before['sections'][2],
                    s5=slide('Police work', 'Patrols and arrests.', layout='image_split') | {'image_query': 'police car'})
    report = quality.Report()
    quality.check_pictures(report, added, data)
    assert report.issues['pictures_missing'] == ['s2']
    # Asked for one more, it is the change done.
    report = quality.Report()
    quality.check_pictures(report, added, {**data, 'options': {'images_wanted': 2, 'image_cap': 10}})
    assert not report.issues
    # And no more is asked for than the plan has room for: cover and slide 3 use two of three.
    report = quality.Report()
    quality.check_pictures(report, before, {**data, 'options': {'images_wanted': 3, 'image_cap': 3}})
    assert report.issues['pictures_missing'] == ['s2']


def test_pictures_added_to_a_deck_that_had_some_are_counted_with_them():
    """"Yana 2 ta rasm qo'shing" on twelve slides with four pictures (lens3/add_two_more): six in all.

    A model told "use 2 photo layouts" took two away and added none, and the
    repair could only bring the deck back to four.
    """
    pictures = {sid: slide(f'Coins {sid}', 'Old coins in a museum.', layout='image_split') | {'image_query': 'coins'}
                for sid in ('s3', 's6', 's9', 's11')}
    before = deck_of(12, s12=slide('Thank you', 'Questions?', layout='closing'), **pictures)
    data = picture_data('add', images_wanted=6, image_cap=12)
    data['revision'].update(request="Yana 2 ta rasm qo'shing", before=before)
    report = quality.Report()
    quality.check_pictures(report, before, data)
    assert len(report.issues['pictures_missing']) == 2
    # Two taken away and none added is four short.
    stripped = deck_of(12, s12=before['sections'][-1], s3=pictures['s3'], s6=pictures['s6'])
    report = quality.Report()
    quality.check_pictures(report, stripped, data)
    assert len(report.issues['pictures_missing']) == 4
    assert not {'s1', 's12'} & set(report.issues['pictures_missing'])
    # Six, whichever slides hold them, is the change done.
    done = deck_of(12, s12=before['sections'][-1], **pictures, s2=pictures['s3'] | {'heading': 'Coins s2'},
                   s4=pictures['s3'] | {'heading': 'Coins s4'})
    report = quality.Report()
    quality.check_pictures(report, done, data)
    assert not report.issues


def test_a_title_page_a_change_adds_to_a_document_gets_no_picture():
    titled = sections(s1=slide('', 'MUSTAQIL ISH\nTAYYORLADI: KARIMOVA NODIRA', layout=None),
                      s2=slide('Tides', 'Why the sea rises.', layout=None), s3=slide('Pilots', 'Pilots and tides.'))
    data = picture_data('every', 'pdf', images_wanted=2)
    data['revision']['edit']['parts'] = ['title_page']
    report = quality.Report()
    quality.check_pictures(report, titled, data)
    assert report.issues['pictures_missing'] == ['s2', 's3']


def test_no_pictures_are_asked_for_when_the_plan_can_add_none():
    # "Har bir slaydda rasm bo'lsin" on a free plan or while photos are off:
    # the model was told to use none, and a repair would only shorten each
    # slide for a picture that is never drawn.
    content = deck_of(6)
    for asked in ('add', 'every'):
        for fmt in ('pptx', 'pdf'):
            report = quality.Report()
            quality.check_pictures(report, content, picture_data(asked, fmt, images_wanted=6, image_cap=0))
            assert not report.issues, (asked, fmt)
    report = quality.Report()
    quality.check_pictures(report, content, {**picture_data('every'), 'options': {'images_wanted': 6}})
    assert not report.issues
    # Nor is the model held to a picture part of its plan it was told not to follow.
    plan = [{'instruction': 'Har bir slaydda rasm bo‘lsin', 'sections': ['s2', 's3']}]
    report = quality.Report()
    quality.check_plan(report, content, content, plan, edit={'pictures': 'every'}, pictures=False)
    assert not report.unapplied and not report.issues
    quality.check_plan(report, content, content, plan, edit={'pictures': 'every'})
    assert report.unapplied == [(['s2', 's3'], 'Har bir slaydda rasm bo‘lsin')]


def test_every_slide_asked_for_a_picture_gets_one_up_to_the_plans_cap():
    content = deck_of(6, s3=slide('Tide pools', 'Rock pools at low tide.', layout='image_split')
                      | {'image_query': 'rock pool'})
    report = quality.Report()
    quality.check_pictures(report, content, picture_data('every', images_wanted=5))
    assert report.issues['pictures_missing'] == ['s2', 's4', 's5', 's6']
    report = quality.Report()
    quality.check_pictures(report, content, picture_data('every', images_wanted=5, image_cap=3))
    assert report.issues['pictures_missing'] == ['s2', 's5']
    # A document's pages, its first included, take the image layout.
    pages = sections(s1=slide('Tides', 'Why the sea rises.', layout=None), s2=slide('Pilots', 'Pilots and tides.'))
    report = quality.Report()
    quality.check_pictures(report, pages, picture_data('every', 'pdf', images_wanted=2))
    assert report.issues['pictures_missing'] == ['s1', 's2']
    assert 'image layout' in quality.request_text(report, 'en', 'pdf', ['s1', 's2'])


class Job:
    id = 'job-offline'


def pictured_deck():
    """A cover with a photo, two picture slides and one full-bleed picture."""
    return deck_of(5, s1=slide('Ocean tides', 'A short guide', layout='cover') | {'image_query': 'sea'},
                   s2=slide('Tide pools', 'Rock pools.\nCrabs hide.', layout='image_split') | {'image_query': 'rock pool'},
                   s3=slide('The coast', 'Low tide at dawn.', layout='image_split') | {'image_query': 'coast'},
                   s4=slide('A harbour', 'Boats wait for the tide.', layout='image_full') | {'image_query': 'harbour'})


def removal(request, content, **edit):
    data = {**picture_data('remove'), 'content': content}
    data['revision'].update(request=request, before=content, edit={'pictures': 'remove', **edit})
    return data


def layouts(content):
    return {s['id']: (s['layout'], s['image_query']) for s in content['sections']}


def test_pictures_taken_away_are_taken_away_without_a_model_call(monkeypatch):
    content = pictured_deck()
    monkeypatch.setattr(quality, 'repair', lambda *args, **kwargs: pytest.fail('no model call is needed'))
    # The plan says which slide loses its picture: only that one does.
    plan = [{'instruction': 'Remove the picture from the tide pools slide', 'sections': ['s2']}]
    done, summary = quality.finish({}, removal('Remove the tide pool picture', content), content, job=Job(),
                                   plan=plan)
    assert layouts(done) == {'s1': ('cover', 'sea'), 's2': ('bullets', ''), 's3': ('image_split', 'coast'),
                             's4': ('image_full', 'harbour'), 's5': ('bullets', '')}
    assert summary['pictures_removed'] == 1


@pytest.mark.parametrize('request_', ['Akbarali tasvirini olib tashla', '3-slayddagi fotoni olib tashla',
                                      'Akbarali illyustratsiyasini olib tashla', '3-slayddagi chizmani o‘chir',
                                      '3-слайддаги расмни олиб ташла'])
def test_a_removal_worded_as_the_customer_words_it_takes_only_the_slide_the_plan_names(monkeypatch, request_):
    # The plan words each part in the customer's own words; every picture in
    # the deck, the cover's too, used to go when the word was not 'rasm'.
    content = pictured_deck()
    monkeypatch.setattr(quality, 'repair', lambda *args, **kwargs: pytest.fail('no model call is needed'))
    done, summary = quality.finish({}, removal(request_, content), content, job=Job(),
                                   plan=[{'instruction': request_, 'sections': ['s3']}])
    assert layouts(done) == {**layouts(content), 's3': ('bullets', '')}
    assert summary['pictures_removed'] == 1


def test_with_no_plan_a_removal_is_never_widened_to_the_whole_deck(monkeypatch):
    before = pictured_deck()
    # The model took slide 4's picture away and kept its layout; the provider
    # dropped an invalid edit_plan.
    returned = deck_of(5, **{s['id']: s for s in before['sections'] if s['id'] != 's4'},
                       s4=before['sections'][3] | {'image_query': ''})
    monkeypatch.setattr(quality, 'repair', lambda *args, **kwargs: pytest.fail('no model call is needed'))
    # The slide the request names loses its picture, and so does the one the model took it from.
    done, summary = quality.finish({}, removal('3-slayddagi fotoni olib tashla', before), returned, job=Job())
    assert layouts(done) == {**layouts(before), 's3': ('bullets', ''), 's4': ('statement', '')}
    assert summary['pictures_removed'] == 2
    # "2-slayddagi va 3-slayddagi": both numbers beside the picture; a number
    # in another part of the sentence is not a picture's slide.
    done, _ = quality.finish({}, removal("2-slayddagi va 3-slayddagi rasmlarni olib tashla va 5-slaydga matn qo'sh",
                                         before), before, job=Job())
    assert layouts(done) == {**layouts(before), 's2': ('bullets', ''), 's3': ('bullets', '')}
    # A request that names no slide leaves every picture the model kept.
    done, summary = quality.finish({}, removal('Akbaraliga qo‘yilgan rasmni olib tashlang', before), returned,
                                   job=Job())
    assert layouts(done) == {**layouts(before), 's4': ('statement', '')}
    # A plan whose parts say nothing of pictures is no plan for them.
    written = deck_of(5, **{s['id']: s for s in before['sections'] if s['id'] != 's3'},
                      s3=before['sections'][2] | {'body': 'Akbarali is brave.\nHe helps his friends.'})
    done, _ = quality.finish({}, removal('Akbarali tasvirini olib tashla', before), written, job=Job(),
                             plan=[{'instruction': 'Akbarali obrazini to‘liq yorit', 'sections': ['s3']}])
    assert layouts(done) == layouts(before)
    # Once slides move, the third slide may not be the third section any more.
    done, _ = quality.finish({}, removal('3-slayddagi fotoni olib tashla', before, structural=True), before,
                             job=Job())
    assert layouts(done) == layouts(before)


def test_unapplied_parts_and_missing_pictures_share_the_one_repair(monkeypatch):
    before = deck_of(5)
    data = {**picture_data('add', images_wanted=1), 'content': before, 'prompt': 'Five slides about tides'}
    data['revision']['before'] = before
    calls = []

    def repair(config, data, content, report, *, job, deadline=None):
        ids = [s['id'] for s in content['sections'] if s['id'] in report.to_repair()]
        calls.append((ids, quality.request_text(report, 'en', 'pptx', ids)))
        return content, 'repaired'

    monkeypatch.setattr(quality, 'repair', repair)
    plan = [{'instruction': 'Add the dates of each tide', 'sections': ['s4']}]
    _, summary = quality.finish({}, data, before, job=Job(), plan=plan)
    assert len(calls) == 1 and summary['repair'] == 'repaired'
    ids, text = calls[0]
    assert ids == ['s2', 's4']
    assert 's4: this part of the change request is not applied yet: "Add the dates of each tide".' in text
    assert 's2: the change request asks for pictures' in text
    assert summary['found']['unapplied'] == 1 and summary['found']['pictures_missing'] == 1


# ---------------------------------------------------------------- the whole job

# ORIGINAL's slides 3 and 5 rewritten, each unlike every other slide of the deck.
CHANGED = slide('Tide gauges along the coast', 'Tide gauges record the height of the sea every few minutes at fixed '
                'stations.\nTheir records go back more than a century in some old harbours.', 'Show a gauge record.')
EXAMPLE = slide('A tide in Saint-Malo', 'In Saint-Malo the sea can rise by more than twelve metres between low and '
                'high water.\nVisitors watch the causeway to the island vanish within an hour.',
                'Show photos of the causeway.')
REQUEST = 'Make slide 3 about tide gauges and add the Saint-Malo example to slide 5.'


def revise(customer, original, request, key):
    revised = create_draft(customer, {'prompt': request, 'options': {'revise_draft_id': str(original.id)}})
    quote = generation_quote(customer, revised.id, revised.version)
    job, _ = submit_job(customer, quote.id, key)
    job = execute_job(job.id)
    revised.refresh_from_db()
    return job, revised


def planned(plan, written):
    """The model's answer to a change: `plan` as its edit_plan, `written` per section id, the rest unchanged."""
    def answer(body, number):
        request = sent(body)
        if request['revision']['request'].startswith('Quality repair'):
            return reply(body, {'s5': EXAMPLE})
        unchanged = {s['id']: s for s in request['outline']['sections']}
        result = reply(body, {**unchanged, **written})
        if 'edit_plan' in body['text']['format']['schema']['properties']:
            result = {'edit_plan': plan, **result}
        return result
    return answer


@pytest.mark.django_db
def test_a_planned_part_left_undone_is_repaired_once_with_its_own_words(customer, monkeypatch):
    original = generated_deck(customer, monkeypatch, ORIGINAL)
    before = {s['id']: s for s in unpack(original.encrypted_data)['content']['sections']}
    plan = [{'instruction': 'Make slide 3 about tide gauges', 'sections': ['s3']},
            {'instruction': 'Add the Saint-Malo example to slide 5', 'sections': ['s5']}]
    requests = model(monkeypatch, planned(plan, {'s3': CHANGED}))
    job, revised = revise(customer, original, REQUEST, 'edit-plan-unapplied')
    assert job.status == 'succeeded', job.error_code
    assert stages(job) == ['revision', 'repair']
    repair = sent(requests[-1])
    assert repair['revision']['selected_section_ids'] == ['s5']
    assert ('s5: this part of the change request is not applied yet: "Add the Saint-Malo example to slide 5". '
            'Apply it fully now.') in repair['revision']['request']
    assert 's3' not in repair['revision']['request'].split('\n', 1)[1]
    after = {s['id']: s for s in unpack(revised.encrypted_data)['content']['sections']}
    assert after['s5']['body'] == EXAMPLE['body'] and after['s3']['body'] == CHANGED['body']
    for sid in ('s1', 's2', 's4', 's6'):
        assert after[sid]['body'] == before[sid]['body'], sid


@pytest.mark.django_db
def test_a_change_whose_planned_parts_were_all_applied_gets_no_repair(customer, monkeypatch):
    original = generated_deck(customer, monkeypatch, ORIGINAL)
    plan = [{'instruction': 'Make slide 3 about tide gauges', 'sections': ['s3']},
            {'instruction': 'Add the Saint-Malo example to slide 5', 'sections': ['s5']}]
    requests = model(monkeypatch, planned(plan, {'s3': CHANGED, 's5': EXAMPLE}))
    job, _ = revise(customer, original, REQUEST, 'edit-plan-applied')
    assert job.status == 'succeeded', job.error_code
    assert stages(job) == ['revision'] and len(requests) == 1


EVERY_SLIDE = ['s1', 's2', 's3', 's4', 's5', 's6']


@pytest.mark.django_db
@pytest.mark.parametrize('request_, plan', [
    # Pictures on every slide while photos are off (no Pixabay here): the
    # model was told to use none, and none could be drawn.
    ('Har bir slaydda rasm bo‘lsin', [{'instruction': 'Har bir slaydda rasm bo‘lsin', 'sections': EVERY_SLIDE[1:]}]),
    # A new look: the guidance asks for every slide's content to be kept.
    ('Varoqlarni hammasini bir hil qil', [{'instruction': 'Make all slides look the same', 'sections': EVERY_SLIDE}]),
    # Another kind of picture is chosen when the pictures are searched for.
    ('Yosh bolacha rasm qoʻyma. Realga oʻxshasin.', [{'instruction': 'Yosh bolacha rasm qoʻyma', 'sections': ['s2']},
                                                     {'instruction': 'Realga oʻxshasin', 'sections': ['s2', 's3']}]),
], ids=['pictures-off', 'redesign', 'photo-style'])
def test_a_change_no_slide_text_can_show_is_delivered_without_a_repair(customer, monkeypatch, request_, plan):
    # Each used to send the whole deck to a repair that rewrote the content the
    # customer asked to keep, or cut every slide short for a picture never drawn.
    original = generated_deck(customer, monkeypatch, ORIGINAL)
    before = unpack(original.encrypted_data)['content']['sections']
    requests = model(monkeypatch, planned(plan, {}))
    job, revised = revise(customer, original, request_, f'edit-look-{len(request_)}')
    assert job.status == 'succeeded', job.error_code
    assert stages(job) == ['revision'] and len(requests) == 1
    after = unpack(revised.encrypted_data)['content']['sections']
    assert [(s['heading'], s['body'], s['layout']) for s in after] == [(s['heading'], s['body'], s['layout'])
                                                                         for s in before]


@pytest.mark.django_db
def test_the_edit_plan_reaches_the_check_from_the_providers_reply(customer, monkeypatch):
    original = generated_deck(customer, monkeypatch, ORIGINAL)
    plan = [{'instruction': 'Make slide 3 about tide gauges', 'sections': ['s3']}]
    seen = []
    finish = quality.finish

    def spy(*args, plan=None, **kwargs):
        seen.append(plan)
        return finish(*args, plan=plan, **kwargs)

    monkeypatch.setattr(quality, 'finish', spy)
    requests = model(monkeypatch, planned(plan, {'s3': CHANGED}))
    job, _ = revise(customer, original, 'Make slide 3 about tide gauges.', 'edit-plan-passes')
    assert job.status == 'succeeded', job.error_code
    assert 'edit_plan' in requests[0]['text']['format']['schema']['properties']
    assert seen == [plan]


@pytest.mark.django_db
def test_the_photo_style_a_change_asked_for_reaches_the_photo_search(customer, monkeypatch):
    from apps.studio import photos
    asked = []

    def fetch(sections, *, cap, budget, job=None, style='photo'):
        asked.append((style, [s['id'] for s in photos.wanted(sections)]))
        return photos.PhotoResult()

    monkeypatch.setattr(photos, 'fetch_for_deck', fetch)
    pictured = by_id([slide('How ocean tides work', 'The Moon, the Sun and the sea', 'Welcome.', 'cover'),
                      slide('Tide pools', 'Rock pools appear at low tide.\nCrabs hide under the weed.',
                            'Show a pool.', 'image_split') | {'image_query': 'rock pool'}])
    model(monkeypatch, lambda body, number: reply(body, pictured))
    authored = create_draft(customer, {'feature_id': 'ai.pptx', 'prompt': 'Two slides about tide pools with drawings',
                                       'output_locale': 'en', 'options': {'pages': 2, 'photo_style': 'illustration'}})
    assert unpack(authored.encrypted_data)['options'].get('photo_style') == 'illustration'
    quote = generation_quote(customer, authored.id, authored.version)
    job, _ = submit_job(customer, quote.id, 'edit-photo-style')
    job = execute_job(job.id)
    assert job.status == 'succeeded', job.error_code
    assert asked == [('illustration', ['s2'])]


@pytest.mark.django_db
@pytest.mark.parametrize('feature,kept,other', [('ai.pdf_topic', 'elegant', 'plain'), ('ai.pptx', 'gala', 'atlas')])
def test_a_change_keeps_the_design_it_did_not_ask_to_change(customer, monkeypatch, feature, kept, other):
    """"Test tuz" on a document made "chiroyli dizaynda" came back plain (lens3/pdf_design_elegant).

    A design the model returns is remembered over the document's only for a
    change asking for a new look.
    """
    from apps.studio import provider
    from apps.studio.domain import pack
    from apps.studio.models import GenerationDraft
    from tests.test_generation_quality import run
    written = by_id([{'heading': 'Ocean tides', 'body': 'Why the sea rises and falls twice a day.', 'notes': ''},
                     {'heading': 'The Moon', 'body': 'The Moon pulls the oceans towards it as the Earth turns.',
                      'notes': ''},
                     {'heading': 'Tide tables', 'body': 'Harbours print the times of high and low water.',
                      'notes': ''}])
    model(monkeypatch, lambda body, number: reply(body, written, title='Ocean tides'))
    job, original = run(customer, 'Three pages about ocean tides', feature=feature, pages=3, key=f'design-{kept}')
    assert job.status == 'succeeded', job.error_code
    data = unpack(original.encrypted_data)
    key = 'deck_design' if feature == 'ai.pptx' else 'doc_design'
    data['options']['template_style'][key] = kept
    GenerationDraft.objects.filter(pk=original.pk).update(encrypted_data=pack(data))

    def generate(config, data, feature_id, request_id, **kwargs):
        # The model returns a design whatever it was asked.
        return {**data['content'], 'design': other, 'edit_plan': []}, {'input_tokens': 1, 'output_tokens': 1}

    monkeypatch.setattr(provider, 'generate', generate)
    job, revised = revise(customer, original, 'Use a friendlier tone throughout.', f'design-kept-{kept}')
    assert job.status == 'succeeded', job.error_code
    assert unpack(revised.encrypted_data)['options']['template_style'][key] == kept
    job, revised = revise(customer, original, 'Use a different design.', f'design-new-{kept}')
    assert job.status == 'succeeded', job.error_code
    assert unpack(revised.encrypted_data)['options']['template_style'][key] == other


def test_only_codes_and_counts_of_the_edit_plan_are_logged(monkeypatch, caplog):
    content = deck_of(3)
    data = {**picture_data(None), 'content': content}
    monkeypatch.setattr(quality, 'repair', lambda config, data, content, report, **kwargs: (content, 'repaired'))
    plan = [{'instruction': 'Write Karimova Nodira under the title', 'sections': ['s1']}]
    with caplog.at_level('INFO', logger='apps.studio.quality'):
        _, summary = quality.finish({}, data, content, job=Job(), plan=plan)
    assert summary['found'] == {'unapplied': 1} and 'quality' in caplog.text
    assert 'Karimova' not in caplog.text and 'Karimova' not in json.dumps(summary)
