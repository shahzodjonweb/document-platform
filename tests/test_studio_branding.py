"""Real branding artifacts: a logo and a name that actually reach the page."""
import io
import pytest
from PIL import Image
from pypdf import PdfReader
from pptx import Presentation
from django.core.files.uploadedfile import SimpleUploadedFile
from apps.core.errors import DomainError
from apps.core.services import upload_file,storage_path,submit_job,execute_job
from apps.studio.branding import prepare_branding,generation_inputs,render_style
from apps.studio.domain import create_draft,update_draft,draft_data,generation_quote
from apps.studio.rendering import render_pdf
from apps.studio.slides import render_pptx
from tests.test_platform import account
from tests.test_studio import paid

pytestmark=pytest.mark.django_db

@pytest.fixture(autouse=True)
def local(settings):
    settings.DEBUG=True;settings.COMMERCE_SANDBOX_ENABLED=True;settings.ENABLE_BETA_TOOLS=True


def logo(a):
    picture=Image.new('RGB',(100,40),'#285cff');raw=io.BytesIO();picture.save(raw,format='PNG')
    return upload_file(a,SimpleUploadedFile('school-logo.png',raw.getvalue(),'image/png'))


def branding(a,asset):
    return prepare_branding(a,'ai.pdf_text',{'template_style':{'accent':'#255e49'},'branding':{'name':'Школа · Maktab','accent':'#285cff','logo_asset_id':str(asset.id)}})


def test_branding_real_logo_and_unicode_name_in_pdf_pptx(settings,tmp_path):
    a=paid(settings,account());asset=logo(a);options=branding(a,asset)
    payload={'source_ids':[],'options':options};style=render_style(a,payload)
    assert generation_inputs(payload)==[str(asset.id)]
    content={'title':'Learning','sections':[{'id':'s1','heading':'Topic','body':'A branded learning document.','notes':'Speaker notes'}],'questions':[]}
    pdf=tmp_path/'brand.pdf';slides=tmp_path/'brand.pptx'
    render_pdf(content,pdf,style=style);render_pptx(content,slides,style=style)
    reader=PdfReader(pdf);assert 'Школа' in reader.pages[0].extract_text() and len(reader.pages[0].images)==1
    deck=Presentation(slides)
    assert any(s.shape_type==13 for s in deck.slides[0].shapes)
    assert any('Школа' in s.text for s in deck.slides[0].shapes if s.has_text_frame)
    assert str(deck.slides[0].shapes[0].fill.fore_color.rgb)=='285CFF'


def test_branding_owner_plan_and_bytes_are_verified(settings):
    a=paid(settings,account());other=paid(settings,account(43));asset=logo(a)
    with pytest.raises(DomainError,match='file_unavailable'):branding(other,asset)
    with pytest.raises(DomainError,match='feature_not_in_plan'):branding(account(44),asset)
    options=branding(a,asset);storage_path(asset.object_key).write_bytes(b'changed-private-content')
    with pytest.raises(DomainError,match='file_changed'):render_style(a,{'options':options})
    with pytest.raises(DomainError,match='invalid_parameters'):prepare_branding(a,'ai.pdf_text',{'template_style':{'accent':'#000000'},'branding':{'logo_url':'https://example.invalid/private'}})


def source(a,fmt='pptx'):
    d=create_draft(a,{'feature_id':'ai.pptx' if fmt=='pptx' else 'ai.pdf_text','source_text':'Original first\n\nUntouched second','output_format':fmt,'options':{'length':2}})
    return d
