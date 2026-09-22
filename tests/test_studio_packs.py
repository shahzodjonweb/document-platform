"""Concrete composite artifacts, matching private answer keys and native slides."""
import copy
import pytest
from pypdf import PdfReader
from pptx import Presentation
from apps.core.services import submit_job,execute_job,storage_path
from apps.core.models import UsageLedger
from apps.studio.domain import create_draft,update_draft,generation_quote,unpack
from apps.studio import provider
from apps.studio.packs import slots,local_materials
from apps.studio.rendering import render_pptx
from tests.test_platform import account,login_client
from tests.test_studio import paid
from tests.test_studio_adversarial import live_configuration,raw_content

pytestmark=pytest.mark.django_db

@pytest.fixture(autouse=True)
def local(settings):
    settings.DEBUG=True;settings.ENABLE_BETA_TOOLS=True;settings.COMMERCE_SANDBOX_ENABLED=True


def pack_content():
    return {'title':'Учебный набор · O‘quv to‘plami','sections':[{'id':'s1','heading':'Review','body':'Read the provided material and explain your reasoning.','notes':'PRIVATE-NOTES'}],'questions':[{'id':f'q{i}','stem':f'Explain concept number {i}.','options':[],'answer':f'KEYONLY-{i}','explanation':f'Review concept {i}.','topic':'reasoning','marks':i} for i in range(1,6)]}


@pytest.mark.parametrize(('feature','count'),[('teacher.variants',6),('teacher.differentiated',6),('teacher.lesson_pack',4),('study.exam_pack',3),('school.weekly_pack',6)])
def test_local_pack_creates_real_named_artifacts_and_separate_keys(settings,feature,count):
    a=paid(settings,account());draft=create_draft(a,{'feature_id':feature,'source_text':'Provided content','output_locale':'ru','options':{'length':1,'question_count':5}})
    draft=update_draft(a,draft.id,{'version':1,'content':pack_content()});quote=generation_quote(a,draft.id,draft.version)
    job,_=submit_job(a,quote.id,'pack-'+feature);job=execute_job(job.id)
    assert job.status=='succeeded',job.error_code
    assert job.artifacts.count()==count and not UsageLedger.objects.filter(job=job,kind='consume').exists()
    names=set()
    for artifact in job.artifacts.select_related('file'):
        names.add(artifact.file.name);path=storage_path(artifact.file.object_key)
        if path.suffix=='.pptx':
            deck=Presentation(path)
            assert any('provided material' in shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
            assert all('PRIVATE-NOTES' not in slide.notes_slide.notes_text_frame.text for slide in deck.slides)
        else:
            text=''.join(page.extract_text() for page in PdfReader(path).pages)
            if artifact.role in ('learner_material','teacher_feedback'):assert 'KEYONLY' not in text
            elif 'answers' in artifact.file.name:assert 'KEYONLY' in text
        if artifact.role=='teacher_key':
            assert login_client(a).post('/api/v1/shares',{'artifact_id':str(artifact.id)},content_type='application/json').status_code==403
    assert len(names)==count
    if feature=='teacher.variants':assert {'a.pdf','a-answers.pdf','b.pdf','b-answers.pdf','c.pdf','c-answers.pdf'}==names
    if feature=='teacher.lesson_pack':assert 'slides.pptx' in names and 'lesson-plan.pdf' in names
    draft.refresh_from_db();assert len(unpack(draft.encrypted_data)['content']['materials'])==len(slots(feature))


def test_letter_answers_stay_paired_after_variant_choice_reorder():
    data=pack_content()
    for q in data['questions']:q.update(options=['Wrong','Correct','Other'],answer='B')
    materials=local_materials('teacher.variants',data,'en')
    assert all(q['answer']=='Correct' for material in materials for q in material['questions'])
    assert len({tuple(q['id'] for q in material['questions']) for material in materials})==3
    assert len({tuple(material['questions'][0]['options']) for material in materials})==3


def test_live_pack_schema_and_actual_credit_ceiling(settings,monkeypatch):
    a=paid(settings,account());live_configuration();draft=create_draft(a,{'feature_id':'teacher.variants','source_text':'Review','options':{'length':1,'question_count':5}})
    quote=generation_quote(a,draft.id,1);raw=raw_content();raw['materials']=local_materials('teacher.variants',pack_content(),'en')
    monkeypatch.setattr(provider,'generate',lambda *args,**kwargs:(raw,{'input_tokens':300,'output_tokens':600}))
    job,_=submit_job(a,quote.id,'live-pack-mocked');job=execute_job(job.id)
    assert job.status=='succeeded',job.error_code
    assert job.artifacts.count()==6 and 0<job.settled_meters['ai_credits']<=quote.meters['ai_credits']
    assert set(provider.schema_for('teacher.variants')['properties']['materials']['items']['properties']['key']['enum'])=={'a','b','c'}


def test_missing_pack_questions_fails_without_generic_success(settings):
    a=paid(settings,account());draft=create_draft(a,{'feature_id':'teacher.variants','source_text':'Only a title','options':{'length':1}});quote=generation_quote(a,draft.id,1)
    job,_=submit_job(a,quote.id,'empty-pack');job=execute_job(job.id)
    assert job.status=='failed' and job.error_code=='pack_content_required' and not job.artifacts.exists()


def test_many_short_lines_are_paginated_as_editable_slides(tmp_path):
    content={'title':'Long source','sections':[{'heading':'Long content','body':'\n'.join(f'Unique line {i:03d}' for i in range(55)),'notes':''}],'questions':[]}
    result=render_pptx(content,tmp_path/'many-lines.pptx')
    deck=Presentation(tmp_path/'many-lines.pptx');texts='\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert result['page_count']>=5
    assert all(f'Unique line {i:03d}' in texts for i in range(55))
    assert all(len(shape.text_frame.paragraphs)<=11 for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
