"""Stock photos: the one place the server fetches a URL it did not write.

Every network call here goes through a fake `_open`, and the suite-wide guard in
conftest makes any test that forgets fail rather than reach Pixabay.
"""
import io
import json
from datetime import timedelta
from urllib.parse import parse_qs, urlsplit

import pytest
from PIL import Image
from django.utils import timezone

from apps.studio import photos
from apps.studio.models import PhotoSearch, ProviderUsage
from operations.integrations import save_config

pytestmark = pytest.mark.django_db
KEY = '12345678-0123456789abcdef0123456789abcdef'


def jpeg(width=1600, height=1000, **kwargs):
    buffer = io.BytesIO()
    Image.new('RGB', (width, height), (110, 140, 160)).save(buffer, 'JPEG', quality=85, **kwargs)
    return buffer.getvalue()


def hit(n, url=None, tags=None):
    value = {'id': n, 'type': 'photo', 'imageWidth': 1920, 'imageHeight': 1280,
             'pageURL': f'https://pixabay.com/photos/x-{n}/',
             'largeImageURL': url or f'https://pixabay.com/get/img{n}_1280.jpg',
             'webformatURL': f'https://pixabay.com/get/img{n}_640.jpg'}
    if tags is not None:
        value['tags'] = tags
    return value


class Response(io.BytesIO):
    def __enter__(self):
        return self

    def __exit__(self, *args):
        self.close()


class Network:
    """A stand-in Pixabay: searches answer with `hits`, downloads with `image`."""

    def __init__(self, hits=None, image=None, fail=None):
        self.hits = [hit(n) for n in range(1, 9)] if hits is None else hits
        self.image, self.fail = image or jpeg(), fail
        self.calls = []

    def __call__(self, request, timeout, handler):
        self.calls.append((request.full_url, timeout, handler))
        if self.fail:
            raise self.fail
        if request.full_url.startswith(photos.API):
            # Like Pixabay, a hit found for a search is tagged with its words,
            # unless the test gave it tags of its own.
            query = parse_qs(urlsplit(request.full_url).query)['q'][0]
            answer = [{'tags': ', '.join(query.split()), **value} for value in self.hits]
            return Response(json.dumps({'total': 99, 'totalHits': 99, 'hits': answer}).encode())
        return Response(self.image)

    @property
    def searches(self):
        return [url for url, _, _ in self.calls if url.startswith(photos.API)]


@pytest.fixture
def pixabay(monkeypatch):
    save_config('pixabay', {'pixabay_key': KEY, 'enabled': 'true'})
    network = Network()
    monkeypatch.setattr(photos, '_open', network)
    return network


def slides(*queries, layout='image_split'):
    first = {'id': 's1', 'heading': 'Deck', 'body': 'Sub', 'notes': '', 'layout': 'cover',
             'items': [], 'columns': [], 'image_query': ''}
    return [first] + [{'id': f's{n + 2}', 'heading': 'H', 'body': 'a\nb', 'notes': '', 'layout': layout,
                       'items': [], 'columns': [], 'image_query': query} for n, query in enumerate(queries)]


# ---------------------------------------------------------------- the pin


@pytest.mark.parametrize('url,ok', [
    ('https://pixabay.com/get/abc.jpg', True),
    ('https://cdn.pixabay.com/photo/2020/x.jpg', True),
    ('https://PIXABAY.com/get/abc.jpg', True),
    ('http://pixabay.com/get/abc.jpg', False),
    ('https://pixabay.com.evil.net/x.jpg', False),
    ('https://evilpixabay.com/x.jpg', False),
    ('https://user:pw@pixabay.com/x.jpg', False),
    ('https://pixabay.com:8443/x.jpg', False),
    ('https://pixabay.com /x.jpg', False),
    ('https://169.254.169.254/latest/meta-data', False),
    ('file:///etc/passwd', False),
    ('https://pixabay.com/' + 'x' * 2100, False),
    (None, False),
])
def test_only_pixabay_urls_are_ever_fetched(url, ok):
    assert photos._image_url_ok(url) is ok


def test_a_download_may_only_be_redirected_to_another_pixabay_host():
    import urllib.request
    request = urllib.request.Request('https://pixabay.com/get/x.jpg')
    with pytest.raises(ValueError):
        photos._PinnedRedirect().redirect_request(request, None, 302, 'Found', {}, 'https://evil.example/x.jpg')
    hops = photos._PinnedRedirect()
    for _ in range(photos.MAX_REDIRECTS):
        followed = hops.redirect_request(request, None, 302, 'Found', {}, 'https://cdn.pixabay.com/x.jpg')
        assert followed.full_url == 'https://cdn.pixabay.com/x.jpg'
    with pytest.raises(ValueError):
        hops.redirect_request(request, None, 302, 'Found', {}, 'https://cdn.pixabay.com/x.jpg')


def test_every_search_is_safe_and_refuses_redirects(pixabay):
    from apps.studio.provider import _NoRedirect
    photos.search('modern office', KEY)
    url, timeout, handler = pixabay.calls[-1]
    query = parse_qs(urlsplit(url).query)
    assert url.startswith('https://pixabay.com/api/?')
    assert query['safesearch'] == ['true'], 'Pixabay defaults safesearch to false'
    assert query['image_type'] == ['photo'] and query['lang'] == ['en']
    assert query['orientation'] == ['horizontal'] and query['q'] == ['modern office']
    assert isinstance(handler, _NoRedirect), 'the key is in the URL, so the API may not redirect'
    assert timeout == photos.SEARCH_TIMEOUT


def test_a_hit_pointing_anywhere_but_pixabay_is_dropped(pixabay):
    pixabay.hits = [hit(1, 'https://evil.example/x.jpg'), {'id': 'x'}, hit(2), {'type': 'video'}]
    found = photos.search('office', KEY)
    assert [h['id'] for h in found] == [2]


# ---------------------------------------------------------------- the bytes


def test_a_photo_is_resized_reencoded_and_stripped_of_exif():
    exif = Image.Exif()
    exif[0x0112] = 6            # rotated 90 degrees
    exif[0x010F] = 'CameraMaker'
    prepared = photos.prepare(jpeg(3000, 2000, exif=exif.tobytes()))
    with Image.open(io.BytesIO(prepared['jpeg'])) as picture:
        assert picture.format == 'JPEG'
        assert max(picture.size) <= photos.LONG_EDGE
        assert picture.size == (prepared['width'], prepared['height'])
        assert picture.height > picture.width, 'the orientation was applied'
        assert not picture.getexif(), 'nothing from the source file survives'
    assert len(prepared['jpeg']) <= photos.MAX_JPEG


def _webp():
    buffer = io.BytesIO()
    Image.new('RGB', (1600, 1000)).save(buffer, 'WEBP')
    return buffer.getvalue()


def _gif():
    buffer = io.BytesIO()
    frames = [Image.new('P', (400, 300), n) for n in range(3)]
    frames[0].save(buffer, 'GIF', save_all=True, append_images=frames[1:])
    return buffer.getvalue()


@pytest.mark.parametrize('raw', [
    pytest.param(lambda: _webp(), id='webp'),
    pytest.param(lambda: _gif(), id='animated'),
    pytest.param(lambda: jpeg(3100, 2000), id='too-many-pixels'),
    pytest.param(lambda: jpeg()[:600], id='truncated'),
    pytest.param(lambda: b'not an image at all', id='garbage'),
])
def test_bytes_that_are_not_a_plain_photo_are_refused(raw):
    with pytest.raises(Exception):
        photos.prepare(raw())


def test_an_oversized_download_is_refused_before_it_is_decoded(pixabay):
    pixabay.image = b'x' * (photos.IMAGE_BYTES + 10)
    with pytest.raises(ValueError):
        photos.download('https://pixabay.com/get/big.jpg')


# ---------------------------------------------------------------- the deck


def test_photos_are_fetched_for_the_slides_that_want_them(pixabay):
    result = photos.fetch_for_deck(slides('office', 'team meeting'), cap=6, budget=30)
    assert set(result.photos) == {'s2', 's3'} and not result.failed and not result.warnings
    assert result.photos['s2']['id'] != result.photos['s3']['id'], 'no photo twice in a deck'
    assert result.photos['s2']['page_url'].startswith('https://pixabay.com/')


def test_a_search_is_cached_for_a_day_and_only_its_hash_is_stored(pixabay):
    photos.fetch_for_deck(slides('office'), cap=6, budget=30)
    assert len(pixabay.searches) == 1
    photos.fetch_for_deck(slides('office'), cap=6, budget=30)
    assert len(pixabay.searches) == 1, 'a second deck within the day reuses the answer'
    row = PhotoSearch.objects.get()
    assert 'office' not in row.key and 'office' not in json.dumps(row.hits).replace('pixabay.com/photos', '')
    PhotoSearch.objects.update(expires_at=timezone.now() - timedelta(seconds=1))
    photos.fetch_for_deck(slides('office'), cap=6, budget=30)
    assert len(pixabay.searches) == 2, 'an expired answer is asked for again'


def test_the_plan_decides_how_many_photos_a_deck_gets(pixabay):
    result = photos.fetch_for_deck(slides('a', 'b', 'c', 'd', 'e'), cap=2, budget=30)
    assert len(result.photos) == 2 and result.limited == 3
    assert result.warnings == ['images_limited']
    assert list(result.photos) == ['s2', 's3'], 'the first slides in the deck get them'


@pytest.mark.parametrize('failure', [TimeoutError('slow'), OSError('down'), ValueError('bad')])
def test_a_provider_that_fails_costs_only_the_photos(pixabay, failure):
    pixabay.fail = failure
    result = photos.fetch_for_deck(slides('office', 'team'), cap=6, budget=30)
    assert result.photos == {} and result.failed == 2 and result.warnings == ['images_skipped']


def test_nothing_is_fetched_when_photos_are_switched_off(monkeypatch):
    network = Network()
    monkeypatch.setattr(photos, '_open', network)
    save_config('pixabay', {'pixabay_key': KEY, 'enabled': 'false'})
    result = photos.fetch_for_deck(slides('office'), cap=6, budget=30)
    assert network.calls == [] and result.photos == {} and result.warnings == []


def test_the_shared_rate_limit_holds_calls_under_pixabays(pixabay, monkeypatch):
    from apps.core.errors import DomainError

    def exhausted(*args, **kwargs):
        raise DomainError('rate_limited', 429)

    monkeypatch.setattr('apps.core.customer_auth.auth_limit', exhausted)
    result = photos.fetch_for_deck(slides('office'), cap=6, budget=30)
    assert pixabay.searches == [] and result.warnings == ['images_skipped']


# ---------------------------------------------------------------- relevance


class Model:
    """A stand-in for the model that chooses photos: answers `picks`, or fails."""

    def __init__(self, picks=None, fail=None, status='completed', text=None):
        self.picks, self.fail, self.status, self.text = picks, fail, status, text
        self.requests = []

    def __call__(self, request, timeout):
        body = json.loads(request.data)
        self.requests.append({'url': request.full_url, 'timeout': timeout, 'body': body,
                              'headers': dict(request.header_items())})
        if self.fail:
            raise self.fail
        text = self.text if self.text is not None else json.dumps({'picks': self.picks})
        return Response(json.dumps({
            'id': f'resp_{len(self.requests)}', 'model': 'test-model', 'status': self.status,
            'output': [{'type': 'message', 'content': [{'type': 'output_text', 'text': text}]}],
            'usage': {'input_tokens': 900, 'output_tokens': 40}}).encode())

    def shown(self, slide):
        """How many photos the model was shown for a slide."""
        content = self.requests[-1]['body']['input'][0]['content']
        return sum(1 for part in content if part['type'] == 'input_text'
                   and part['text'].startswith(f'Slide {slide}, photo '))


@pytest.fixture
def model(monkeypatch, pixabay):
    from apps.studio import photo_choice
    save_config('ai', {'mode': 'openai', 'model': 'test-model', 'api_key': 'sk-offline-test'})
    fake = Model()
    monkeypatch.setattr(photo_choice, '_post', fake)
    return fake


def test_hits_sharing_no_word_with_the_subject_are_not_candidates():
    hits = [photos._hit({**hit(n), 'tags': tags}) for n, tags in
            enumerate(['sunset, beach, sea', 'forklift, truck, industry', 'warehouse, forklift, shelves',
                       'warehouses, logistics'], 1)]
    ranked = photos.rank(hits, 'warehouse forklift')
    assert [h['id'] for h in ranked] == [3, 4, 2], 'both words first, the main subject next, the sunset never'
    assert photos.rank(hits, '') == []


def test_the_model_chooses_between_candidates_and_its_choice_is_used(model, pixabay):
    pixabay.hits = [hit(1, tags='office, desk'), hit(2, tags='office, meeting, people')]
    model.picks = [{'slide': 's2', 'photo': 2}]
    result = photos.fetch_for_deck(slides('office meeting'), cap=6, budget=60)
    # Ranked: hit 2 shares both words, so it is candidate 1; hit 1 is candidate 2.
    assert result.photos['s2']['id'] == 1, 'the model picked candidate 2, not the best tag match'
    assert model.shown('s2') == 2 and not result.warnings


def test_the_model_is_shown_each_slide_text_and_small_previews_never_pixabay_urls(model, pixabay):
    model.picks = [{'slide': 's2', 'photo': 1}]
    sections = slides('office')
    sections[1]['heading'], sections[1]['body'] = 'Quarterly sales review', 'Revenue grew\nTwo new clients'
    photos.fetch_for_deck(sections, cap=6, budget=60)
    request = model.requests[-1]
    body = request['body']
    assert request['url'] == 'https://api.openai.com/v1/responses' and body['store'] is False
    assert body['text']['format']['strict'] is True and body['model'] == 'test-model'
    text = json.dumps(body)
    assert 'Quarterly sales review Revenue grew Two new clients' in text
    images = [part for part in body['input'][0]['content'] if part['type'] == 'input_image']
    assert len(images) == photos.CANDIDATES and all(part['detail'] == 'low' for part in images)
    assert all(part['image_url'].startswith('data:image/jpeg;base64,') for part in images)
    assert 'pixabay.com' not in text and KEY not in text, 'no Pixabay URL or key goes to the model'
    previewed = [url for url, _, _ in pixabay.calls if url.endswith('_640.jpg')]
    assert len(previewed) == photos.CANDIDATES


def test_when_no_candidate_fits_the_slide_is_text_not_the_most_popular_photo(model, pixabay):
    model.picks = [{'slide': 's2', 'photo': 0}]
    result = photos.fetch_for_deck(slides('office'), cap=6, budget=60)
    assert result.photos == {} and result.unmatched == 1 and result.warnings == ['images_skipped']
    downloads = [url for url, _, _ in pixabay.calls if url.endswith('_1280.jpg')]
    assert downloads == [], 'nothing is downloaded for a slide the model turned down'


@pytest.mark.parametrize('answer', [
    [],                                              # left the slide out
    [{'slide': 's2', 'photo': 9}],                   # a photo it was not shown
    [{'slide': 's9', 'photo': 1}],                   # a slide that does not exist
])
def test_an_answer_that_does_not_name_a_shown_photo_means_none(model, pixabay, answer):
    model.picks = answer
    result = photos.fetch_for_deck(slides('office'), cap=6, budget=60)
    assert result.photos == {} and result.unmatched == 1


@pytest.mark.parametrize('failure', [
    {'fail': TimeoutError('slow')}, {'fail': OSError('down')}, {'status': 'incomplete'},
    {'text': 'not json'}, {'text': json.dumps({'picks': [{'slide': 's2'}]})},
])
def test_a_choice_that_fails_falls_back_to_the_best_tag_match(model, pixabay, failure):
    from apps.studio.models import ProviderAttempt
    for field, value in failure.items():
        setattr(model, field, value)
    pixabay.hits = [hit(1, tags='sunset, sea'), hit(2, tags='office, desk')]
    result = photos.fetch_for_deck(slides('office'), cap=6, budget=60)
    assert result.photos['s2']['id'] == 2, 'the tag match, never the unrelated popular photo'
    attempt = ProviderAttempt.objects.get(stage='photo_pick')
    assert attempt.status in ('failed', 'rejected')


def test_without_a_model_only_a_tag_match_is_used(pixabay):
    pixabay.hits = [hit(1, tags='sunset, sea'), hit(2, tags='mountain, lake')]
    result = photos.fetch_for_deck(slides('office'), cap=6, budget=30)
    assert result.photos == {} and result.warnings == ['images_skipped']


def test_the_same_photo_is_never_used_twice_even_if_the_model_picks_it_twice(model, pixabay):
    model.picks = [{'slide': 's2', 'photo': 1}, {'slide': 's3', 'photo': 1}]
    result = photos.fetch_for_deck(slides('office', 'office'), cap=6, budget=60)
    assert list(result.photos) == ['s2'] and result.unmatched == 1


def test_a_subject_too_narrow_to_find_is_searched_again_as_its_main_subject(pixabay):
    pixabay.hits = [hit(1, tags='nothing, alike')]
    photos.fetch_for_deck(slides('rusty harbour crane cargo'), cap=6, budget=30)
    queries = [parse_qs(urlsplit(url).query)['q'][0] for url in pixabay.searches]
    assert queries == ['rusty harbour crane cargo', 'rusty harbour']


def test_the_choice_is_accounted_like_any_other_provider_call(model, pixabay):
    from apps.studio.models import ProviderAttempt
    model.picks = [{'slide': 's2', 'photo': 1}]
    photos.fetch_for_deck(slides('office'), cap=6, budget=60)
    attempt = ProviderAttempt.objects.get(stage='photo_pick')
    assert attempt.status == 'succeeded' and attempt.input_tokens == 900 and attempt.output_tokens == 40
    assert attempt.requested_model == 'test-model'


def test_the_choice_is_skipped_when_too_little_time_is_left(model, pixabay):
    result = photos.fetch_for_deck(slides('office'), cap=6, budget=photos.PICK_MIN_SECONDS)
    assert model.requests == [] and set(result.photos) == {'s2'}, 'the tag match stands in'


def test_the_model_is_asked_for_literal_photo_subjects():
    from apps.studio.layouts import photo_guidance
    text = photo_guidance(2)
    assert 'literally shows' in text and 'main subject first' in text
    assert 'Never an abstract idea' in text and 'in English, even when the deck is in another language' in text


# ---------------------------------------------------------------- the job


@pytest.fixture
def paid_deck(settings, monkeypatch):
    from apps.studio import provider
    from tests.test_page_fill import live
    customer = live(settings)
    answer = {'title': 'Deck', 'answer_supported': True, 'citations': [], 'questions': [],
              'sections': slides('modern office', 'team meeting')}
    monkeypatch.setattr(provider, 'generate', lambda *args, **kwargs: (answer, {'input_tokens': 1, 'output_tokens': 1}))
    return customer


def _run(customer):
    from apps.core.services import execute_job, submit_job
    from apps.studio.domain import SLIDES, create_draft, generation_quote
    draft = create_draft(customer, {'feature_id': SLIDES, 'prompt': 'A 3 slide deck about our office.'})
    quote = generation_quote(customer, draft.id, draft.version)
    job, _ = submit_job(customer, quote.id, f'photos-{draft.id}')
    return quote, execute_job(job.id)


def test_a_paid_deck_carries_its_photos_and_their_licence_record(paid_deck, pixabay):
    from pptx import Presentation
    from apps.core.services import storage_path
    quote, job = _run(paid_deck)
    assert job.status == 'succeeded', job.error_code
    artifact = job.artifacts.get().file
    deck = Presentation(str(storage_path(artifact.object_key)))
    assert len(deck.slides) == 3
    assert sum(shape.shape_type == 13 for slide in deck.slides for shape in slide.shapes) == 2
    assert [entry['slide'] for entry in artifact.metadata['photos']] == [2, 3]
    assert all(entry['provider'] == 'pixabay' and entry['id'] for entry in artifact.metadata['photos'])
    assert ProviderUsage.objects.get(job=job, provider='pixabay').outcome == 'succeeded'
    assert 'images_skipped' not in job.warnings


def test_a_paid_deck_whose_photos_fail_is_still_delivered_and_billed_the_same(paid_deck, pixabay):
    from apps.core.models import UsageLedger
    pixabay.fail = OSError('down')
    quote, job = _run(paid_deck)
    assert job.status == 'succeeded', job.error_code
    assert job.artifacts.get().file.page_count == 3, 'the photo slides are drawn as text'
    assert 'images_skipped' in job.warnings
    consumed = sum(UsageLedger.objects.filter(job=job, kind='consume', meter='ai_credits')
                   .values_list('amount', flat=True))
    assert 0 < consumed <= quote.meters['ai_credits']
    assert ProviderUsage.objects.get(job=job, provider='pixabay').outcome == 'failed'


def test_the_key_never_leaves_the_encrypted_store(paid_deck, pixabay, caplog):
    pixabay.fail = OSError(f'refused https://pixabay.com/api/?key={KEY}')
    _, job = _run(paid_deck)
    record = json.dumps({'warnings': job.warnings, 'error': job.error_code,
                         'metadata': job.artifacts.get().file.metadata,
                         'usage': list(ProviderUsage.objects.values())}, default=str)
    assert KEY not in record and KEY not in caplog.text


def test_the_model_is_told_the_plan_allowance_only_when_photos_are_on(settings):
    from apps.studio.domain import SLIDES, create_draft, unpack
    from tests.test_platform import account
    settings.DEBUG = True
    customer = account()
    options = lambda: unpack(create_draft(customer, {'feature_id': SLIDES, 'prompt': 'A 4 slide deck.',
                                                     'options': {'image_cap': 99}}).encrypted_data)['options']
    assert options()['image_cap'] == 0, 'off by default, and a client cannot raise it'
    save_config('pixabay', {'pixabay_key': KEY, 'enabled': 'true'})
    from apps.core.policy import plan_limits
    assert options()['image_cap'] == plan_limits(customer)['max_deck_images']


# ---------------------------------------------------------------- the admin


def test_the_pixabay_key_is_encrypted_and_never_shown(settings):
    from operations.models import IntegrationConfig
    from tests.test_operations import staff_client
    client, _ = staff_client()
    response = client.post('/ops/integrations', {'integration': 'pixabay', 'action': 'save',
                                                 'pixabay_key': KEY, 'enabled': 'true',
                                                 'reason': 'Turn on deck photos.'})
    assert response.status_code == 302
    row = IntegrationConfig.objects.get(pk='pixabay')
    assert KEY.encode() not in bytes(row.encrypted_secrets) and KEY not in json.dumps(row.configuration)
    page = client.get('/ops/integrations').content.decode()
    assert KEY not in page and 'pixabay-settings' in page
    from operations.models import AuditLog
    assert KEY not in json.dumps(list(AuditLog.objects.values()), default=str)


@pytest.mark.parametrize('bad', ['abc', '123-XYZ', 'x' * 60, '12345678-short'])
def test_a_key_that_is_not_a_pixabay_key_is_refused(bad):
    from apps.core.errors import DomainError
    with pytest.raises(DomainError, match='invalid_pixabay_key'):
        save_config('pixabay', {'pixabay_key': bad, 'enabled': 'false'})


def test_photos_cannot_be_switched_on_without_a_key():
    from apps.core.errors import DomainError
    with pytest.raises(DomainError, match='pixabay_not_configured'):
        save_config('pixabay', {'enabled': 'true'})


def test_the_photo_allowance_is_an_editable_plan_limit():
    from operations import plans
    from apps.core.policy import limits_for_plan
    assert [limits_for_plan(plan)['max_deck_images'] for plan in ('free', 'plus', 'premium')] == [2, 6, 12]
    assert plans.FIELDS['max_deck_images'] == (0, 60)
    with pytest.raises(plans.PlanError):
        plans.save('free', {'max_deck_images': '40'})
