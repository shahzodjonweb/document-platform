"""Synthetic HTTP → queue → worker → download checks; never impersonate a customer."""
import io
import json
import mimetypes
import re
import tempfile
import time
import uuid
import zipfile
from importlib import import_module
from pathlib import Path
from urllib.error import HTTPError
from urllib.request import Request, urlopen
from http.cookies import SimpleCookie

from datetime import timedelta

from django.conf import settings
from django.utils import timezone
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen.canvas import Canvas
from apps.core.models import Account, FileAsset, Job


API_URL = 'http://127.0.0.1:8000/api/v1/'



def _as_channel_member(account):
    """Free accounts are asked to join the owner's Telegram channels before a
    task runs. A synthetic account cannot join, so it is recorded as a member
    for the hour of the audit; the rule stays on for everyone else."""
    from apps.core.channel_gate import _config as channel_config
    from apps.core.models import ChannelMembership
    now = timezone.now()
    for channel in channel_config()['channels']:
        ChannelMembership.objects.update_or_create(account=account, chat=channel['chat'], defaults={
            'is_member': True, 'checked_at': now, 'expires_at': now + timedelta(hours=1)})

class ProbeFailure(Exception):
    def __init__(self, code):
        self.code = code if re.fullmatch(r'[a-zA-Z0-9_.-]{1,100}', str(code)) else 'probe_failed'


def require(condition, code):
    if not condition:
        raise ProbeFailure(code)


class Customer:
    """Provision a test-only session server-side, then use the actual HTTP API.

    This tests authorization/CSRF and all document endpoints, not external login.
    Session cookies stay in memory and are deleted when the probe finishes.
    """
    def __init__(self, audit_id, feature):
        self.account = Account.objects.create(is_test=True, display_name=f'Service audit {audit_id} {feature}')
        _as_channel_member(self.account)
        self.session = import_module(settings.SESSION_ENGINE).SessionStore()
        self.session['customer_account_id'] = str(self.account.pk)
        self.session['customer_auth_version'] = self.account.auth_version
        self.session.set_expiry(900)
        self.session.save()
        self.cookies = {settings.SESSION_COOKIE_NAME: self.session.session_key}
        self.csrf = None
        data = self.json('GET', 'auth/session')
        require(data.get('authenticated') and data['user']['is_test'], 'test_session')
        self.csrf = data['csrf_token']

    def request(self, method, path, data=None, content_type='application/json', headers=None, csrf=True):
        headers = {'Host': 'pdfmaster-admin.orderdesk.live', 'X-Forwarded-Proto': 'https',
                   'Origin': 'https://pdfmaster-admin.orderdesk.live',
                   'Referer': 'https://pdfmaster-admin.orderdesk.live/',
                   'Cookie': '; '.join(f'{k}={v}' for k, v in self.cookies.items()),
                   'Content-Type': content_type, **(headers or {})}
        if csrf and self.csrf:
            headers['X-CSRFToken'] = self.csrf
        request = Request(API_URL + path.removeprefix('/api/v1/'),
                          data=data, method=method, headers=headers)
        try:
            response = urlopen(request, timeout=100)
        except HTTPError as exc:
            response = exc
        with response:
            body = response.read(12 * 1024 * 1024 + 1)
            require(len(body) <= 12 * 1024 * 1024, 'response_limit')
            for value in response.headers.get_all('Set-Cookie', []):
                cookie = SimpleCookie(); cookie.load(value)
                self.cookies.update({key: morsel.value for key, morsel in cookie.items()})
            return response.status, body, response.headers

    def json(self, method, path, data=None, expected=(200, 201), headers=None):
        status, body, _ = self.request(method, path, json.dumps(data).encode() if data is not None else None, headers=headers)
        try:
            result = json.loads(body)
        except ValueError:
            raise ProbeFailure('non_json_response') from None
        if status not in expected:
            raise ProbeFailure(result.get('error', {}).get('code', 'http_' + str(status)))
        return result

    def upload(self, path, password=None):
        boundary = 'Audit' + uuid.uuid4().hex
        data = b''
        if password:
            data += f'--{boundary}\r\nContent-Disposition: form-data; name="password"\r\n\r\n{password}\r\n'.encode()
        data += (f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{path.name}"\r\n'
                 f'Content-Type: {mimetypes.guess_type(path.name)[0] or "application/octet-stream"}\r\n\r\n').encode()
        data += path.read_bytes() + f'\r\n--{boundary}--\r\n'.encode()
        status, body, _ = self.request('POST', 'files/uploads', data, 'multipart/form-data; boundary=' + boundary)
        result = json.loads(body)
        require(status == 201, result.get('error', {}).get('code', 'upload_failed'))
        return result

    def retire(self):
        self.session.delete()
        if not Job.objects.filter(account=self.account, status__in=('queued', 'running', 'finalizing')).exists():
            FileAsset.objects.filter(account=self.account).exclude(state='deleted').update(expires_at=timezone.now())


def curved_gray_document(*, with_mask=False):
    """Generate generic print on bowed, shaded paper and a ribbed gray desk.

    This is independent of detector internals and contains no customer image,
    identity, address or signature. All four colored corner labels are visible.
    """
    import cv2
    import numpy as np
    width, height = 450, 680
    size = (720, 900)
    y, x = np.mgrid[:height, :width].astype(np.float32)
    x /= width - 1; y /= height - 1
    shadow = 13 * np.exp(-((x - .18) ** 2 / .10 + (y - .72) ** 2 / .18))
    light = 133 + 34 * x + 13 * y - shadow
    light += np.random.default_rng(1729).normal(0, 1.1, light.shape)
    light = np.clip(light, 130, 180).astype(np.uint8)
    paper = Image.fromarray(np.repeat(light[:, :, None], 3, axis=2))
    draw = ImageDraw.Draw(paper)
    font = lambda size: ImageFont.truetype('processors/assets/fonts/NotoSans-Regular.ttf', size)
    draw.text((80, 38), 'DOCUMENT CHECK', font=font(21), fill=(25, 25, 25))
    for top, value in [(84, 'Reference: SAMPLE A12'), (111, 'Route: Depot A to Depot B'),
                       (138, 'Date: sample entry')]:
        draw.text((30, top), value, font=font(16), fill=(38, 38, 38))
    draw.rectangle((29, 192, 419, 452), outline=(33, 33, 33), width=2)
    for top in (238, 292, 345, 399):
        draw.line((29, top, 419, top), fill=(42, 42, 42), width=2)
    for left in (239, 320):
        draw.line((left, 192, left, 452), fill=(42, 42, 42), width=2)
    for left, value in ((40, 'Item'), (251, 'Units'), (330, 'Weight')):
        draw.text((left, 205), value, font=font(15), fill=(25, 25, 25))
    for top, values in [(250, ('Parcel A', '02', '1.5')), (304, ('Parcel B', '01', '2.0')),
                        (357, ('Parcel C', '03', '0.8')), (410, ('Total', '06', '4.3'))]:
        for left, value in zip((40, 251, 334), values):
            draw.text((left, top), value, font=font(16), fill=(30, 30, 30))
    draw.text((31, 492), 'Received: sample acknowledgement', font=font(16), fill=(32, 32, 32))
    draw.line([(57, 558), (91, 531), (104, 560), (136, 538),
               (160, 556), (186, 540), (210, 553), (247, 546)], fill=(17, 42, 119), width=4)
    draw.text((30, 594), 'Inspect all items before accepting.', font=font(16), fill=(30, 30, 30))
    for position, value, color in zip(
            ((12, 10), (width - 53, 10), (12, height - 34), (width - 53, height - 34)),
            ('A1', 'B2', 'C3', 'D4'),
            ((110, 18, 21), (14, 88, 26), (17, 42, 119), (108, 21, 101))):
        draw.text(position, value, font=font(21), fill=color)
    for position, value in zip(
            ((4, height // 2 - 14), (width - 42, height // 2 - 14),
             (width // 2 - 15, 2), (width // 2 - 15, height - 28)),
            ('E5', 'F6', 'G7', 'H8')):
        draw.text(position, value, font=font(21), fill=(108, 21, 101))
    margin = 26
    yy, xx = np.mgrid[:height + 2 * margin, :width + 2 * margin].astype(np.float32)
    normalized_x = (xx - margin) / (width - 1)
    normalized_y = (yy - margin) / (height - 1)
    map_x = xx - margin - 17 * np.sin(np.pi * np.clip(normalized_y, 0, 1))
    map_y = yy - margin + 11 * np.sin(np.pi * np.clip(normalized_x, 0, 1))
    rgb = cv2.remap(np.asarray(paper), map_x, map_y, cv2.INTER_LINEAR,
                    borderMode=cv2.BORDER_CONSTANT, borderValue=(0, 0, 0))
    alpha = cv2.remap(np.full((height, width), 255, np.uint8), map_x, map_y,
                      cv2.INTER_LINEAR, borderMode=cv2.BORDER_CONSTANT, borderValue=0)
    corners = np.array(((140, 90), (601, 143), (561, 819), (88, 746)), np.float32)
    source = np.array(((margin, margin), (width - 1 + margin, margin),
                       (width - 1 + margin, height - 1 + margin), (margin, height - 1 + margin)), np.float32)
    transform = cv2.getPerspectiveTransform(source, corners)
    warped = cv2.warpPerspective(rgb, transform, size, flags=cv2.INTER_LINEAR)
    mask = cv2.warpPerspective(alpha, transform, size, flags=cv2.INTER_LINEAR).astype(np.float32) / 255
    yy, xx = np.mgrid[:size[1], :size[0]].astype(np.float32)
    desk = 112 + 6 * np.sin((yy + .16 * xx) * np.pi / 12) + 4 * xx / size[0] + 2 * yy / size[1]
    cast_shadow = cv2.GaussianBlur(mask, (0, 0), 5)
    desk = np.clip(desk - 9 * cast_shadow * (1 - mask), 0, 255)
    background = np.repeat(desk[:, :, None], 3, axis=2)
    photographed = background * (1 - mask[:, :, None]) + warped * mask[:, :, None]
    photographed = cv2.GaussianBlur(np.clip(photographed, 0, 255).astype(np.uint8), (3, 3), .55)
    image = Image.fromarray(photographed)
    return (image, np.round(mask * 255).astype(np.uint8)) if with_mask else image


def validate_fit_canvas(reader, image):
    """Verify the actual PDF image placement fills a bounded, borderless page."""
    from pypdf.generic import ContentStream
    page = reader.pages[0]
    width, height = float(page.mediabox.width), float(page.mediabox.height)
    require(abs(max(width, height) - 842) < .001
            and abs(width / height - image.width / image.height) < .00001,
            'curved_document_fitted_page')
    require(page.rotation % 360 == 0 and page.cropbox == page.mediabox,
            'curved_document_page_viewport')
    current = (1., 0., 0., 1., 0., 0.)
    stack, placements = [], []
    for operands, operator in ContentStream(page.get_contents(), reader).operations:
        if operator == b'q':
            stack.append(current)
        elif operator == b'Q':
            require(bool(stack), 'curved_document_canvas_state')
            current = stack.pop()
        elif operator == b'cm':
            require(len(operands) == 6, 'curved_document_canvas_state')
            a, b, c, d, e, f = current
            x, y, z, w, u, v = map(float, operands)
            current = (a*x + c*y, b*x + d*y, a*z + c*w, b*z + d*w,
                       a*u + c*v + e, b*u + d*v + f)
        elif operator == b'Do':
            placements.append(current)
    require(len(placements) == 1 and not stack, 'curved_document_canvas_state')
    expected = (width, 0., 0., height, 0., 0.)
    require(all(abs(actual - target) < .001 for actual, target in zip(placements[0], expected)),
            'curved_document_borderless_image_placement')
    return {'borderless_page_verified': True}


def side_label_data(image):
    """Count known purple writing and measure its distance from output edges."""
    import numpy as np
    pixels = np.asarray(image).astype(np.int16)
    ink = (pixels[:, :, 0] > pixels[:, :, 1] + 35) & (pixels[:, :, 2] > pixels[:, :, 1] + 35)
    height, width = ink.shape
    regions = (ink[height // 3:2 * height // 3, :width // 4],
               ink[height // 3:2 * height // 3, 3 * width // 4:],
               ink[:height // 4, width // 3:2 * width // 3],
               ink[3 * height // 4:, width // 3:2 * width // 3])
    coordinates = [np.where(region) for region in regions]
    counts = [len(y) for y, x in coordinates]
    require(min(counts) > 45, 'curved_document_all_side_writing')
    gaps = (int(coordinates[0][1].min()) / width,
            (width - 1 - (3 * width // 4 + int(coordinates[1][1].max()))) / width,
            int(coordinates[2][0].min()) / height,
            (height - 1 - (3 * height // 4 + int(coordinates[3][0].max()))) / height)
    return counts, gaps


def validate_paper_boundary(original, true_paper_mask, image):
    """Measure remaining desk using the generated paper alpha, not detector flags."""
    import cv2
    import numpy as np
    source = np.asarray(original)
    rgb = source.astype(np.int16)
    gray = cv2.cvtColor(source, cv2.COLOR_RGB2GRAY)
    neutral = rgb.max(axis=2) - rgb.min(axis=2) < 6
    desk_values = gray[(true_paper_mask < 8) & neutral]
    paper_values = gray[(true_paper_mask > 248) & neutral & (gray > 125)]
    require(desk_values.size > 100 and paper_values.size > 100, 'curved_document_fixture_regions')
    desk_top, paper_low = float(np.percentile(desk_values, 98)), float(np.percentile(paper_values, 2))
    require(paper_low > desk_top + 8, 'curved_document_fixture_contrast')
    threshold = (desk_top + paper_low) / 2
    rgb = np.asarray(image).astype(np.int16)
    gray = cv2.cvtColor(np.asarray(image), cv2.COLOR_RGB2GRAY)
    desk = ((rgb.max(axis=2) - rgb.min(axis=2) < 6)
            & (gray >= float(np.percentile(desk_values, 2)) - 12) & (gray < threshold))
    band = max(2, round(min(image.size) * .025))
    border = np.ones(gray.shape, dtype=bool)
    border[band:-band, band:-band] = False
    corners = (desk[:band, :band], desk[:band, -band:], desk[-band:, :band], desk[-band:, -band:])
    require(float(desk[border].mean()) < .03 and max(float(corner.mean()) for corner in corners) < .08,
            'curved_document_photographed_desk_removed')
    return {'paper_boundary_verified': True}


def validate_curved_scan(output, original, *, enabled, enhanced=True, true_paper_mask=None):
    """Check actual PDF pixels, rather than trusting an engine status flag."""
    import numpy as np
    reader = PdfReader(io.BytesIO(output))
    require(len(reader.pages) == 1 and len(reader.pages[0].images) == 1, 'curved_document_page')
    image = reader.pages[0].images[0].image.convert('RGB')
    pixels = np.asarray(image).astype(np.int16)
    flags = validate_fit_canvas(reader, image)
    if not enabled:
        require(image.size == original.size and np.array_equal(np.asarray(image), np.asarray(original)),
                'curved_document_opt_out_preserves_all_pixels')
        return {**flags, 'crop_verified': False, 'effect_verified': False, 'original_pixels_preserved': True}
    require(image.width * image.height < original.width * original.height * .80,
            'curved_document_crop')
    require(1.15 < image.height / image.width < 1.8, 'curved_document_aspect')
    regions = (pixels[:image.height // 4, :image.width // 4],
               pixels[:image.height // 4, 3 * image.width // 4:],
               pixels[3 * image.height // 4:, :image.width // 4],
               pixels[3 * image.height // 4:, 3 * image.width // 4:])
    rules = (lambda p: (p[:, :, 0] > p[:, :, 1] + 35) & (p[:, :, 0] > p[:, :, 2] + 35),
             lambda p: (p[:, :, 1] > p[:, :, 0] + 25) & (p[:, :, 1] > p[:, :, 2] + 25),
             lambda p: (p[:, :, 2] > p[:, :, 0] + 35) & (p[:, :, 2] > p[:, :, 1] + 25),
             lambda p: (p[:, :, 0] > p[:, :, 1] + 35) & (p[:, :, 2] > p[:, :, 1] + 35))
    require(all(int(rule(region).sum()) > 25 for rule, region in zip(rules, regions)),
            'curved_document_all_corner_writing')
    before_counts, _ = side_label_data(original)
    after_counts, gaps = side_label_data(image)
    require(all(after >= before * .75 for before, after in zip(before_counts, after_counts)),
            'curved_document_side_writing_preserved')
    require(all(gap < limit for gap, limit in zip(gaps, (.03, .06, .025, .022))),
            'curved_document_no_extra_image_padding')
    if enhanced:
        neutral = (pixels.max(axis=2) - pixels.min(axis=2) < 10) & (pixels.min(axis=2) > 100)
        require(float(neutral.mean()) > .5 and float(np.median(pixels[neutral])) > 210,
                'curved_document_gray_paper_brightened')
    else:
        require(true_paper_mask is not None, 'curved_document_fixture_regions')
        flags.update(validate_paper_boundary(original, true_paper_mask, image))
    # Camera blur and perspective interpolation soften print edges. These
    # marks still have over 110 levels of contrast against the required paper.
    require(int((pixels.max(axis=2) < 100).sum()) > 3000, 'curved_document_print_preserved')
    return {**flags, 'crop_verified': True, 'effect_verified': enhanced,
            'all_corner_writing_preserved': True, 'all_side_writing_preserved': True,
            'tight_page_edges_verified': True}


def tinted_document_photo(*, paper_kind='yellow', desk_kind='bright', curved=True,
                          rotation=None, with_mask=False):
    """Anonymous matte paper, edge writing and a blue stamp on a patterned desk.

    The alpha is created before the simulated camera, independently of scanner
    decisions. It is used only by the audit to recognize photographed desk.
    """
    import cv2
    import numpy as np
    width, height = 450, 680
    yy, xx = np.mgrid[:height, :width].astype(np.float32)
    x, y = xx / (width - 1), yy / (height - 1)
    base = np.array((214, 194, 151) if paper_kind == 'yellow' else (153, 155, 157), np.float32)
    light = .92 + .08 * x + .035 * y - .045 * np.exp(-((x - .2)**2 / .12 + (y - .7)**2 / .18))
    noise = np.random.default_rng(20261006).normal(0, .7, (height, width, 1))
    paper = Image.fromarray(np.clip(base * light[:, :, None] + noise, 0, 255).astype(np.uint8))
    draw = ImageDraw.Draw(paper)
    font = lambda size: ImageFont.truetype('processors/assets/fonts/NotoSans-Regular.ttf', size)
    draw.text((75, 40), 'DOCUMENT SAMPLE', font=font(22), fill=(24, 24, 24))
    for top, value in ((94, 'Reference: SAMPLE 27'), (128, 'Review all sample entries')):
        draw.text((35, top), value, font=font(17), fill=(32, 32, 32))
    draw.rectangle((32, 188, 416, 428), outline=(40, 40, 40), width=2)
    for top in (236, 284, 332, 380):
        draw.line((32, top, 416, top), fill=(43, 43, 43), width=2)
    draw.line((274, 188, 274, 428), fill=(43, 43, 43), width=2)
    for top, value in ((202, 'Item             Units'), (250, 'Sample A         02'),
                       (298, 'Sample B         04'), (346, 'Sample C         01'), (394, 'Total            07')):
        draw.text((47, top), value, font=font(17), fill=(27, 27, 27))
    draw.text((37, 475), 'Checked: generic sample', font=font(17), fill=(28, 28, 28))
    blue = (15, 44, 150)
    draw.ellipse((279, 510, 400, 610), outline=blue, width=4)
    draw.text((306, 541), 'TEST', font=font(22), fill=blue)
    draw.line(((55, 573), (81, 549), (103, 580), (125, 551),
               (148, 578), (178, 559), (217, 572)), fill=blue, width=4)
    draw.text((36, 627), 'Keep this written note.', font=font(17), fill=(30, 30, 30))
    for position, value, color in zip(
            ((12, 10), (width - 53, 10), (12, height - 34), (width - 53, height - 34)),
            ('A1', 'B2', 'C3', 'D4'),
            ((135, 13, 17), (12, 105, 23), (160, 76, 10), (130, 18, 117))):
        draw.text(position, value, font=font(21), fill=color)
    for position, value in zip(
            ((4, height // 2 - 14), (width - 42, height // 2 - 14),
             (width // 2 - 15, 2), (width // 2 - 15, height - 28)), ('E5', 'F6', 'G7', 'H8')):
        draw.text(position, value, font=font(21), fill=(12, 100, 100))
    margin = 26 if curved else 0
    if curved:
        yy, xx = np.mgrid[:height + 2 * margin, :width + 2 * margin].astype(np.float32)
        map_x = xx - margin - 17 * np.sin(np.pi * np.clip((yy - margin) / (height - 1), 0, 1))
        map_y = yy - margin + 11 * np.sin(np.pi * np.clip((xx - margin) / (width - 1), 0, 1))
        rgb = cv2.remap(np.asarray(paper), map_x, map_y, cv2.INTER_LINEAR,
                        borderMode=cv2.BORDER_CONSTANT, borderValue=(0, 0, 0))
        alpha = cv2.remap(np.full((height, width), 255, np.uint8), map_x, map_y,
                          cv2.INTER_LINEAR, borderMode=cv2.BORDER_CONSTANT, borderValue=0)
    else:
        rgb, alpha = np.asarray(paper), np.full((height, width), 255, np.uint8)
    corners = np.array(((140, 90), (601, 143), (561, 819), (88, 746)), np.float32)
    source = np.array(((margin, margin), (width - 1 + margin, margin),
                       (width - 1 + margin, height - 1 + margin), (margin, height - 1 + margin)), np.float32)
    size = (720, 900)
    transform = cv2.getPerspectiveTransform(source, corners)
    warped = cv2.warpPerspective(rgb, transform, size, flags=cv2.INTER_LINEAR)
    mask = cv2.warpPerspective(alpha, transform, size, flags=cv2.INTER_LINEAR)
    coverage = mask.astype(np.float32) / 255
    desk = Image.new('RGB', size, (235, 238, 236) if desk_kind == 'bright' else (77, 82, 79))
    pattern = ImageDraw.Draw(desk)
    petals = (221, 218, 229) if desk_kind == 'bright' else (91, 82, 99)
    leaves = (210, 228, 216) if desk_kind == 'bright' else (79, 97, 86)
    for top in range(-30, size[1], 85):
        for left in range(-25, size[0], 92):
            for dx, dy in ((-12, 0), (12, 0), (0, -12), (0, 12)):
                pattern.ellipse((left + dx - 11, top + dy - 11,
                                 left + dx + 11, top + dy + 11), fill=petals)
            pattern.ellipse((left + 22, top + 20, left + 48, top + 32), fill=leaves)
    background = np.asarray(desk).astype(np.float32)
    shade = cv2.GaussianBlur(coverage, (0, 0), 5) * (1 - coverage) * 9
    background -= shade[:, :, None]
    photographed = background * (1 - coverage[:, :, None]) + warped * coverage[:, :, None]
    photographed = cv2.GaussianBlur(np.clip(photographed, 0, 255).astype(np.uint8), (3, 3), .55)
    image, mask_image = Image.fromarray(photographed), Image.fromarray(mask)
    if rotation is not None:
        image, mask_image = image.transpose(rotation), mask_image.transpose(rotation)
    return (image, np.asarray(mask_image)) if with_mask else image


def tinted_boundary_metrics(original, true_paper_mask, image):
    """Compare real source materials; either brightness polarity is valid."""
    import cv2
    import numpy as np
    source = cv2.cvtColor(np.asarray(original), cv2.COLOR_RGB2LAB).astype(np.float32)
    require((true_paper_mask < 8).sum() > 100 and (true_paper_mask > 248).sum() > 100,
            'tinted_document_fixture_regions')
    desk = np.median(source[true_paper_mask < 8], axis=0)
    paper = np.median(source[true_paper_mask > 248], axis=0)
    require(float(np.linalg.norm(desk - paper)) > 15, 'tinted_document_fixture_materials')
    pixels = cv2.cvtColor(np.asarray(image), cv2.COLOR_RGB2LAB).astype(np.float32)
    desk_like = np.linalg.norm(pixels - desk, axis=2) < np.linalg.norm(pixels - paper, axis=2)
    band = max(2, round(min(image.size) * .025))
    border = np.ones(desk_like.shape, dtype=bool)
    border[band:-band, band:-band] = False
    corners = (desk_like[:band, :band], desk_like[:band, -band:],
               desk_like[-band:, :band], desk_like[-band:, -band:])
    return {'border_desk_fraction': float(desk_like[border].mean()),
            'corner_desk_fractions': [float(corner.mean()) for corner in corners]}


def tinted_writing_metrics(image):
    """Locate known colored marks, independently of detection and engine flags."""
    import cv2
    import numpy as np
    pixels = np.asarray(image).astype(np.int16)
    height, width = pixels.shape[:2]
    regions = (pixels[:height // 4, :width // 4], pixels[:height // 4, 3 * width // 4:],
               pixels[3 * height // 4:, :width // 4], pixels[3 * height // 4:, 3 * width // 4:])
    rules = (lambda p: (p[:, :, 0] > p[:, :, 1] + 45) & (p[:, :, 0] > p[:, :, 2] + 45),
             lambda p: (p[:, :, 1] > p[:, :, 0] + 35) & (p[:, :, 1] > p[:, :, 2] + 35),
             lambda p: (p[:, :, 0] > p[:, :, 1] + 45) & (p[:, :, 1] > p[:, :, 2] + 35),
             lambda p: (p[:, :, 0] > p[:, :, 1] + 45) & (p[:, :, 2] > p[:, :, 1] + 45))
    corners = [int(rule(region).sum()) for rule, region in zip(rules, regions)]
    # Yellow paper and pastel desk have no blue/cyan ink chroma.
    cyan_strength = np.clip(np.minimum(pixels[:, :, 1] - pixels[:, :, 0],
                                       pixels[:, :, 2] - pixels[:, :, 0]), 0, 255)
    cyan = (cyan_strength > 25) & (np.abs(pixels[:, :, 1] - pixels[:, :, 2]) < 45)
    regions = (cyan[height // 3:2 * height // 3, :width // 4],
               cyan[height // 3:2 * height // 3, 3 * width // 4:],
               cyan[:height // 4, width // 3:2 * width // 3],
               cyan[3 * height // 4:, width // 3:2 * width // 3])
    strengths = (cyan_strength[height // 3:2 * height // 3, :width // 4],
                 cyan_strength[height // 3:2 * height // 3, 3 * width // 4:],
                 cyan_strength[:height // 4, width // 3:2 * width // 3],
                 cyan_strength[3 * height // 4:, width // 3:2 * width // 3])
    coordinates, shapes, masses = [], [], []
    for region, strength in zip(regions, strengths):
        y, x = np.where(region)
        require(len(y) > 35, 'tinted_document_all_side_writing')
        coordinates.append((y, x))
        box = strength[y.min():y.max() + 1, x.min():x.max() + 1].astype(np.float32)
        shapes.append(cv2.resize(box, (64, 64), interpolation=cv2.INTER_LINEAR) > 25)
        masses.append(float(box.mean()))
    gaps = (int(coordinates[0][1].min()) / width,
            (width - 1 - (3 * width // 4 + int(coordinates[1][1].max()))) / width,
            int(coordinates[2][0].min()) / height,
            (height - 1 - (3 * height // 4 + int(coordinates[3][0].max()))) / height)
    blue = (pixels[:, :, 2] > pixels[:, :, 0] + 50) & (pixels[:, :, 2] > pixels[:, :, 1] + 45)
    stamp = blue[height // 2:15 * height // 16, width // 2:15 * width // 16]
    handwriting = blue[height // 2:15 * height // 16, :width // 2]
    _, _, stamp_stats, _ = cv2.connectedComponentsWithStats(stamp.astype(np.uint8), connectivity=8)
    return {'corner_counts': corners, 'side_shapes': shapes, 'side_mass': masses,
            'side_gaps': gaps, 'blue_stamp_pixels': int(stamp.sum()),
            'blue_stamp_components': sum(int(area) > 10 for area in stamp_stats[1:, 4]),
            'blue_handwriting_pixels': int(handwriting.sum())}


def validate_tinted_writing(original, image):
    import cv2
    import numpy as np
    before, after = tinted_writing_metrics(original), tinted_writing_metrics(image)
    require(min(after['corner_counts']) > 25, 'tinted_document_all_corner_writing')
    require(after['blue_stamp_pixels'] > 150 and after['blue_handwriting_pixels'] > 100
            and after['blue_stamp_components'] >= 5,
            'tinted_document_blue_stamp_and_handwriting')
    for old_shape, new_shape, old_mass, new_mass in zip(
            before['side_shapes'], after['side_shapes'], before['side_mass'], after['side_mass']):
        # The photographed label leans with the camera perspective, whereas
        # its rectified counterpart is upright. Allow that small bounded shear
        # and one normalized pixel of interpolation; neither restores erased
        # letters or absent strokes. Ink density is checked independently.
        coverage = 0.
        for shear in np.linspace(-.2, .2, 17):
            transform = np.array(((1, 0, 0), (shear, 1, -31.5 * shear)), np.float32)
            aligned = cv2.warpAffine(new_shape.astype(np.uint8), transform, (64, 64),
                                     flags=cv2.INTER_NEAREST)
            supported = cv2.dilate(aligned, np.ones((3, 3), np.uint8)) > 0
            coverage = max(coverage, float((old_shape & supported).sum()) / int(old_shape.sum()))
        require(coverage >= .75
                and new_mass >= old_mass * .75, 'tinted_document_side_writing_preserved')
    require(all(gap < limit for gap, limit in zip(after['side_gaps'], (.035, .065, .03, .027))),
            'tinted_document_no_extra_image_padding')
    return {'all_corner_writing_preserved': True, 'all_side_writing_preserved': True,
            'blue_stamp_and_handwriting_preserved': True, 'tight_page_edges_verified': True}


def validate_tinted_scan(output, original, true_paper_mask, *, enhanced):
    """Assert physical PDF placement, preserved ink and genuine camera cleanup."""
    import cv2
    import numpy as np
    reader = PdfReader(io.BytesIO(output))
    require(len(reader.pages) == 1 and len(reader.pages[0].images) == 1, 'tinted_document_page')
    image = reader.pages[0].images[0].image.convert('RGB')
    flags = validate_fit_canvas(reader, image)
    require(image.width * image.height < original.width * original.height * .80,
            'tinted_document_crop')
    flags.update(validate_tinted_writing(original, image))
    if enhanced:
        gray = cv2.cvtColor(np.asarray(image), cv2.COLOR_RGB2GRAY)
        before = cv2.cvtColor(np.asarray(original), cv2.COLOR_RGB2GRAY)[true_paper_mask > 248]
        print_pixels = gray[gray < 100]
        require(float(np.percentile(gray, 75)) > float(np.percentile(before, 75)) + 20
                and print_pixels.size > 3000
                and float(np.percentile(gray, 75)) - float(np.median(print_pixels)) > 110,
                'tinted_document_readability_effect')
    else:
        metrics = tinted_boundary_metrics(original, true_paper_mask, image)
        require(metrics['border_desk_fraction'] < .04 and max(metrics['corner_desk_fractions']) < .10,
                'tinted_document_photographed_desk_removed')
        flags['paper_boundary_verified'] = True
    return {**flags, 'crop_verified': True, 'effect_verified': enhanced}


def fixtures(folder):
    def pdf(name, labels, form=False):
        path = folder / name
        canvas = Canvas(str(path), pagesize=(400, 600), pageCompression=0)
        for label in labels:
            canvas.setFont('Helvetica', 18); canvas.drawString(40, 550, label)
            if form:
                canvas.acroForm.textfield(name='customer', x=40, y=430, width=260, height=45, fontSize=12)
            canvas.showPage()
        canvas.save()
        return path
    paths = {'pdf': pdf('pages.pdf', ['alpha', 'beta', 'gamma']), 'other': pdf('other.pdf', ['delta']),
             'form': pdf('form.pdf', ['Form'], form=True)}
    paths['compress'] = pdf('compress.pdf', ['Repeated readable text ' * 60])
    for color in ('red', 'blue'):
        paths[color] = folder / (color + '.png'); Image.new('RGB', (100, 80), color).save(paths[color])
    paths['scan'] = folder / 'document-photo.png'
    scan = Image.new('RGB', (700, 900), (46, 64, 54))
    draw = ImageDraw.Draw(scan)
    draw.polygon([(130, 80), (570, 125), (590, 810), (70, 750)], fill=(244, 244, 242))
    font = ImageFont.truetype('processors/assets/fonts/NotoSans-Regular.ttf', 23)
    for number, top in enumerate(range(180, 600, 48), 1):
        draw.text((155, top), f'{number}. Document scan 123', font=font, fill=(35, 35, 35))
    draw.text((150, 675), 'Blue note: keep this ink', font=font, fill=(20, 55, 145))
    scan.save(paths['scan'])
    paths['curved_scan'] = folder / 'curved-gray-document.png'
    curved, paper_mask = curved_gray_document(with_mask=True)
    curved.save(paths['curved_scan'])
    paths['curved_mask'] = folder / 'curved-gray-document-mask.png'
    Image.fromarray(paper_mask).save(paths['curved_mask'])
    paths['tinted_scan'] = folder / 'tinted-paper-patterned-desk.png'
    tinted, paper_mask = tinted_document_photo(with_mask=True)
    tinted.save(paths['tinted_scan'])
    paths['tinted_mask'] = folder / 'tinted-paper-mask.png'
    Image.fromarray(paper_mask).save(paths['tinted_mask'])
    for language, text in [('eng', 'HELLO DOCUMENT 2026'), ('rus', 'ПРИВЕТ ДОКУМЕНТ 2026'), ('uzb', 'SALOM HUJJAT 2026')]:
        paths[language] = folder / (language + '.png')
        scan = Image.new('RGB', (1500, 320), 'white')
        ImageDraw.Draw(scan).text((50, 100), text, font=ImageFont.truetype('processors/assets/fonts/NotoSans-Regular.ttf', 58), fill='black')
        scan.save(paths[language])
    paths['protected'] = folder / 'protected.pdf'
    writer = PdfWriter(clone_from=paths['other']); writer.encrypt('synthetic-audit-password'); writer.write(paths['protected'])
    paths['table'] = folder / 'table.pdf'
    from reportlab.platypus import Table, TableStyle
    from reportlab.lib import colors
    canvas = Canvas(str(paths['table']), pagesize=(400, 600))
    table = Table([['Item', 'Value'], ['Alpha', '=2+2'], ['Beta', '20']], colWidths=[150, 150], rowHeights=30)
    table.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 1, colors.black)]))
    table.wrapOn(canvas, 300, 200); table.drawOn(canvas, 40, 350); canvas.save()
    paths['docx'] = Path('processors/fixtures/office-multilingual.docx')
    paths['pptx'] = Path('processors/fixtures/office-multilingual.pptx')
    return paths


def cases():
    add = {'type': 'add_text', 'page': 1, 'x': 40, 'y': 100, 'text': 'Salom Привет', 'font_size': 18}
    highlight = {'type': 'highlight', 'page': 1, 'x': 38, 'y': 35, 'width': 180, 'height': 24}
    note = {'type': 'note', 'page': 1, 'x': 300, 'y': 60, 'text': 'Примечание'}
    return [
        ('pdf.merge', ['pdf', 'other'], {}), ('pdf.compress', ['compress'], {}),
        ('pdf.split', ['pdf'], {'ranges': ['1-2', '3']}),
        ('pdf.extract_pages', ['pdf'], {'pages': '3,1'}), ('pdf.delete_pages', ['pdf'], {'pages': '2'}),
        ('pdf.reorder', ['pdf'], {'order': [3, 2, 1]}), ('pdf.rotate', ['pdf'], {'pages': '1,3', 'angle': 90}),
        ('pdf.images_to_pdf', ['red', 'blue'], {}), ('pdf.to_images', ['pdf'], {'format': 'png', 'pages': '3,1', 'dpi': 72}),
        ('pdf.images_to_pdf', ['scan'], {'auto_crop': True, 'enhance_text': True}),
        ('pdf.images_to_pdf', ['scan'], {'auto_crop': False, 'enhance_text': False}),
        ('pdf.images_to_pdf', ['curved_scan'], {}),
        ('pdf.images_to_pdf', ['curved_scan'], {'auto_crop': True, 'enhance_text': False}),
        ('pdf.images_to_pdf', ['curved_scan'], {'auto_crop': False, 'enhance_text': False}),
        ('pdf.images_to_pdf', ['tinted_scan'], {}),
        ('pdf.images_to_pdf', ['tinted_scan'], {'auto_crop': True, 'enhance_text': False}),
        ('pdf.protect', ['other'], {}), ('pdf.unlock_known', ['protected'], {}),
        ('convert.word_to_pdf', ['docx'], {}), ('convert.pptx_to_pdf', ['pptx'], {}),
        ('ocr.extract_text', ['eng'], {'language': 'eng'}), ('ocr.searchable_pdf', ['eng'], {'language': 'eng'}),
        ('convert.pdf_to_docx', ['other'], {'mode': 'text'}), ('convert.pdf_to_xlsx', ['table'], {}),
        ('editor.visual', ['other'], {'commands': [add, highlight, note]}),
        ('editor.add_text', ['other'], {'commands': [add]}), ('editor.highlight', ['other'], {'commands': [highlight]}),
        ('editor.annotate', ['other'], {'commands': [note]}),
        ('editor.fill_forms', ['form'], {'commands': [{'type': 'fill_form', 'field': 'customer', 'value': 'Павел Salom'}]}),
        ('editor.signature_image', ['other', 'red'], {'commands': [{'type': 'signature', 'page': 1, 'x': 40, 'y': 100, 'width': 140, 'height': 45, 'input_index': 1}]}),
        ('ocr.extract_text', ['rus'], {'language': 'rus'}), ('ocr.extract_text', ['uzb'], {'language': 'uzb'}),
    ]


def validate_output(feature, outputs, parameters, inputs=None):
    if feature == 'pdf.to_images':
        with zipfile.ZipFile(io.BytesIO(outputs[0])) as archive:
            require(archive.namelist() == ['page-0003.png', 'page-0001.png'], 'image_archive_order')
            for name in archive.namelist():
                image = Image.open(io.BytesIO(archive.read(name))); image.verify()
        return
    if feature == 'ocr.extract_text':
        expected = {'eng': 'HELLO', 'rus': 'ПРИВЕТ', 'uzb': 'HUJJAT'}[parameters['language']]
        require(expected in outputs[0].decode(), 'ocr_text'); return
    if feature == 'convert.pdf_to_docx':
        from docx import Document
        require('delta' in '\n'.join(p.text for p in Document(io.BytesIO(outputs[0])).paragraphs), 'editable_word_text'); return
    if feature == 'convert.pdf_to_xlsx':
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(outputs[0]), data_only=False); sheet = book.worksheets[0]
        require(sheet['A2'].value == 'Alpha' and sheet['B2'].value == '=2+2' and sheet['B2'].data_type == 's', 'table_cells_safe'); return
    readers = [PdfReader(io.BytesIO(output)) for output in outputs]
    if feature == 'pdf.protect':
        require(readers[0].is_encrypted and readers[0].decrypt('synthetic-audit-password'), 'pdf_encryption')
    pages = readers[0].pages
    texts = [page.extract_text().strip() for page in pages]
    expected = {'pdf.merge': ['alpha', 'beta', 'gamma', 'delta'], 'pdf.extract_pages': ['gamma', 'alpha'],
                'pdf.delete_pages': ['alpha', 'gamma'], 'pdf.reorder': ['gamma', 'beta', 'alpha'],
                'pdf.unlock_known': ['delta'], 'pdf.protect': ['delta']}
    if feature in expected: require(texts == expected[feature], 'page_text_order')
    if feature == 'pdf.unlock_known': require(not readers[0].is_encrypted, 'unlocked')
    if feature == 'pdf.split': require([len(r.pages) for r in readers] == [2, 1] and readers[1].pages[0].extract_text().strip() == 'gamma', 'split_pages')
    if feature == 'pdf.rotate': require([page.rotation for page in pages] == [90, 0, 90], 'page_rotation')
    if feature == 'pdf.compress': require('Repeated readable text' in texts[0], 'compressed_text')
    if feature == 'pdf.images_to_pdf':
        scan_case = inputs in (['scan'], ['curved_scan'], ['tinted_scan']) or 'auto_crop' in parameters
        require(len(pages) == (1 if scan_case else 2) and all(len(p.images) > 0 for p in pages), 'image_pages')
        if 'auto_crop' in parameters and inputs not in (['curved_scan'], ['tinted_scan']):
            picture = pages[0].images[0].image.convert('RGB')
            if parameters['auto_crop']:
                require(picture.width < 650 and picture.height < 850, 'document_crop')
                require(sum(b > r + 35 and b > g + 25 for r, g, b in picture.get_flattened_data()) > 100, 'document_colored_ink')
            else:
                require(picture.size == (700, 900) and picture.getpixel((20, 20)) == (46, 64, 54), 'document_opt_out')
    if feature in ('convert.word_to_pdf', 'convert.pptx_to_pdf'):
        require(len(pages) == 3 and all(phrase in text for phrase, text in zip(
            ['Document pages remain readable', 'Hujjat sahifalari aniq ko‘rinadi', 'Страницы документа остаются читаемыми'], texts)), 'office_multilingual_pages')
    if feature == 'ocr.searchable_pdf': require('HELLO' in texts[0] and len(pages[0].images), 'searchable_text_layer')
    if feature in ('editor.visual', 'editor.add_text'): require('Salom Привет' in texts[0], 'unicode_overlay')
    if feature in ('editor.visual', 'editor.highlight', 'editor.annotate'):
        types = {a.get_object()['/Subtype'] for a in pages[0].get('/Annots', [])}
        expected = {'editor.visual': {'/Highlight', '/Text'}, 'editor.highlight': {'/Highlight'}, 'editor.annotate': {'/Text'}}[feature]
        require(expected.issubset(types), 'serialized_annotations')
    if feature == 'editor.fill_forms': require(readers[0].get_fields()['customer']['/V'] == 'Павел Salom', 'unicode_form')
    if feature == 'editor.signature_image': require(len(pages[0].images) == 1, 'signature_image')


def check_documents(audit_id, emit):
    results = []
    deadline = time.monotonic() + 410
    with tempfile.TemporaryDirectory(prefix='pdfmaster-service-audit-') as directory:
        paths = fixtures(Path(directory))
        for feature, inputs, parameters in cases():
            customer = None
            start = time.monotonic()
            row = {'feature': feature, 'variant': parameters.get('language', 'default'), 'scope': 'http_queue_worker_download'}
            if feature == 'pdf.images_to_pdf' and 'auto_crop' in parameters:
                row['variant'] = 'scanning_on' if parameters['auto_crop'] else 'scanning_off'
            if inputs == ['curved_scan']:
                row['variant'] = ('curved_gray_defaults' if not parameters else
                                  'curved_gray_crop_only' if parameters['auto_crop'] else 'curved_gray_off')
            if inputs == ['tinted_scan']:
                row['variant'] = 'tinted_paper_defaults' if not parameters else 'tinted_paper_crop_only'
            try:
                require(time.monotonic() < deadline, 'audit_time_budget')
                customer = Customer(audit_id, feature)
                assets = [customer.upload(paths[key], 'synthetic-audit-password' if key == 'protected' else None) for key in inputs]
                payload = {'feature_id': feature, 'input_ids': [a['id'] for a in assets], 'parameters': parameters}
                if feature == 'pdf.protect':
                    payload['secret_id'] = customer.json('POST', 'secrets', {'password': 'synthetic-audit-password'})['id']
                quote = customer.json('POST', 'quotes', payload)
                key = 'service-audit-' + uuid.uuid4().hex
                job = customer.json('POST', 'jobs', {'quote_id': quote['id']}, headers={'Idempotency-Key': key})
                stop = min(deadline, time.monotonic() + 60)
                while job['status'] in ('queued', 'running', 'finalizing') and time.monotonic() < stop:
                    time.sleep(0.5); job = customer.json('GET', 'jobs/' + job['id'])
                require(job['status'] in ('succeeded', 'no_op'), (job.get('error') or {}).get('code', 'job_' + job['status']))
                require(Job.objects.get(pk=job['id'], account=customer.account).attempt_count >= 1, 'worker_attempt')
                outputs = []
                for artifact in job['artifacts']:
                    status, body, headers = customer.request('GET', artifact['download_url'])
                    require(status == 200 and headers.get('Content-Disposition', '').startswith('attachment'), 'download')
                    outputs.append(body)
                require(bool(outputs), 'missing_output')
                validate_output(feature, outputs, parameters, inputs)
                if inputs == ['curved_scan']:
                    enabled = parameters.get('auto_crop', True)
                    enhanced = parameters.get('enhance_text', True)
                    require(job.get('parameters', {}).get('auto_crop') is enabled
                            and job.get('parameters', {}).get('enhance_text') is enhanced,
                            'curved_document_quoted_switches')
                    require(all(job.get('parameters', {}).get(key) == value for key, value in
                                {'paper_size': 'fit', 'orientation': 'auto', 'margin': 0}.items()),
                            'curved_document_quoted_layout')
                    with Image.open(paths['curved_scan']) as original, Image.open(paths['curved_mask']) as mask:
                        import numpy as np
                        row.update(validate_curved_scan(outputs[0], original.convert('RGB'), enabled=enabled,
                                                        enhanced=enhanced, true_paper_mask=np.asarray(mask)))
                if inputs == ['tinted_scan']:
                    enhanced = parameters.get('enhance_text', True)
                    require(job.get('parameters', {}).get('auto_crop') is True
                            and job.get('parameters', {}).get('enhance_text') is enhanced,
                            'tinted_document_quoted_switches')
                    require(all(job.get('parameters', {}).get(key) == value for key, value in
                                {'paper_size': 'fit', 'orientation': 'auto', 'margin': 0}.items()),
                            'tinted_document_quoted_layout')
                    with Image.open(paths['tinted_scan']) as original, Image.open(paths['tinted_mask']) as mask:
                        import numpy as np
                        row.update(validate_tinted_scan(outputs[0], original.convert('RGB'), np.asarray(mask),
                                                        enhanced=enhanced))
                if feature == 'pdf.compress': require(len(outputs[0]) < paths['compress'].stat().st_size, 'compression_size')
                if feature == 'pdf.merge':
                    preview = job['artifacts'][0]['preview_url']
                    status, body, _ = customer.request('GET', preview)
                    require(status == 200 and body.startswith(b'\x89PNG'), 'preview_png')
                    retry = customer.json('POST', 'jobs', {'quote_id': quote['id']}, headers={'Idempotency-Key': key})
                    require(retry['id'] == job['id'], 'idempotency')
                    status, _, _ = customer.request('POST', 'quotes', json.dumps(payload).encode(), csrf=False)
                    require(status == 403, 'csrf_boundary')
                    other = Customer(audit_id, 'ownership')
                    try:
                        for path in ('jobs/' + job['id'], job['artifacts'][0]['download_url']):
                            status, _, _ = other.request('GET', path); require(status == 404, 'owner_boundary')
                        for fid in ('editor.existing_text', 'editor.insert_images', 'editor.replace_images', 'editor.redact'):
                            response = customer.json('POST', 'quotes', {**payload, 'feature_id': fid}, expected=(403,))
                            require(response.get('error', {}).get('code') == 'feature_not_in_plan', 'paid_access_boundary')
                    finally: other.retire()
                    row['access_idempotency_preview'] = 'passed'
                require(job['settled_meters'] == job['meters'], 'usage_settlement')
                row.update(status='passed', outputs=len(outputs))
            except Exception as exc:
                code = getattr(exc, 'code', 'probe_failed')
                row.update(status='failed', error=type(exc).__name__, code=code if re.fullmatch(r'[a-zA-Z0-9_.-]{1,100}', str(code)) else 'probe_failed')
            finally:
                if customer: customer.retire()
            row['seconds'] = round(time.monotonic() - start, 2)
            results.append(row); emit(row)
    return results
