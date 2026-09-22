"""Real OCR and editable conversion adapters with explicit quality limitations."""
from __future__ import annotations
import io
import math
import os
import shutil
import subprocess
import tempfile
from contextlib import closing
from pathlib import Path

ASSETS = Path(__file__).resolve().parent / 'assets'
LANGUAGES = ('eng', 'rus', 'uzb', 'eng+rus', 'eng+uzb', 'eng+rus+uzb')
LANGUAGE = {'type': 'string', 'enum': list(LANGUAGES)}
PAGES = {'type': 'string', 'maxLength': 6000}
SCHEMAS = {
 'ocr.extract_text': {'type':'object','additionalProperties':False,'properties':{'language':LANGUAGE,'pages':PAGES,'dpi':{'type':'integer','minimum':150,'maximum':300}}},
 'ocr.searchable_pdf': {'type':'object','additionalProperties':False,'properties':{'language':LANGUAGE,'pages':PAGES,'dpi':{'type':'integer','minimum':150,'maximum':300}}},
 'convert.pdf_to_docx': {'type':'object','additionalProperties':False,'properties':{'mode':{'type':'string','enum':['text','ocr']},'language':LANGUAGE,'pages':PAGES}},
 'convert.pdf_to_xlsx': {'type':'object','additionalProperties':False,'properties':{'strategy':{'type':'string','enum':['lines','text']},'pages':PAGES}},
}


def tesseract_runtime():
    configured = os.environ.get('PDFMASTER_TESSERACT_BIN')
    executable = configured or shutil.which('tesseract')
    available = bool(executable and Path(executable).is_file() and all((ASSETS/'tessdata'/f'{lang}.traineddata').is_file() for lang in ('eng','rus','uzb')))
    return {'available': available, 'executable': executable, 'languages': list(LANGUAGES), 'engine': 'tesseract'}


def normalize(feature, parameters, metadata=None):
    from .engine import ProcessorError, parse_pages, MAX_PAGES
    if not isinstance(parameters, dict) or set(parameters)-set(SCHEMAS[feature]['properties']):
        raise ProcessorError('invalid_parameters')
    value=dict(parameters)
    value.setdefault('pages','all')
    count=metadata[0].get('page_count') if metadata else MAX_PAGES
    parse_pages(value['pages'],count or MAX_PAGES)
    if feature.startswith('ocr.') or feature=='convert.pdf_to_docx':
        value.setdefault('language','eng')
        if value['language'] not in LANGUAGES:raise ProcessorError('invalid_parameters')
    if feature.startswith('ocr.'):
        value.setdefault('dpi',200)
        if type(value['dpi']) is not int or not 150<=value['dpi']<=300:raise ProcessorError('invalid_parameters')
    if feature=='convert.pdf_to_docx':
        value.setdefault('mode','text')
        if value['mode'] not in ('text','ocr'):raise ProcessorError('invalid_parameters')
    if feature=='convert.pdf_to_xlsx':
        value.setdefault('strategy','lines')
        if value['strategy'] not in ('lines','text'):raise ProcessorError('invalid_parameters')
    return value


def rendered_pages(path, metadata, parameters):
    """Yield one owned PIL image at a time; all decoded/rendered pixels bounded."""
    from PIL import Image,ImageOps
    from .engine import ProcessorError,parse_pages,MAX_PIXELS,MAX_RENDER_PIXELS
    if metadata['kind']=='image':
        with Image.open(path,formats=['PNG','JPEG','WEBP']) as raw:
            picture=ImageOps.exif_transpose(raw).convert('RGB')
            yield 0,picture
        return
    import pypdfium2 as pdfium
    total=0
    scale=parameters.get('dpi',200)/72
    with pdfium.PdfDocument(str(path)) as document:
        document.init_forms()
        for index in parse_pages(parameters['pages'],len(document)):
            with closing(document[index]) as page:
                width,height=page.get_size()
                pixels=math.ceil(width*scale)*math.ceil(height*scale)
                total+=pixels
                if pixels>MAX_PIXELS or total>MAX_RENDER_PIXELS:raise ProcessorError('render_pixel_limit')
                bitmap=page.render(scale=scale)
                try: picture=bitmap.to_pil().convert('RGB')
                finally: bitmap.close()
                try: yield index,picture
                finally: picture.close()


def run_ocr(path, metadata, parameters, scratch, searchable=False):
    from .engine import ProcessorError
    runtime=tesseract_runtime()
    if not runtime['available']:raise ProcessorError('engine_unavailable')
    texts=[];pdfs=[]
    for index,picture in rendered_pages(path,metadata,parameters):
        source=scratch/f'scan-{index:04d}.png';target=scratch/f'ocr-{index:04d}'
        picture.save(source,dpi=(parameters.get('dpi',200),)*2)
        args=[runtime['executable'],str(source),str(target),'--tessdata-dir',str(ASSETS/'tessdata'),'-l',parameters['language'],'--oem','1','--psm','3']
        if searchable:
            # Explicit configs do not depend on a system tessdata/configs directory.
            args+=['-c','tessedit_create_pdf=1','-c','tessedit_create_txt=1']
        try:
            proc=subprocess.run(args,stdin=subprocess.DEVNULL,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,timeout=35,check=False)
        except subprocess.TimeoutExpired:raise ProcessorError('processor_timeout') from None
        text_path=target.with_suffix('.txt')
        if proc.returncode or not text_path.is_file():raise ProcessorError('ocr_failed')
        text=text_path.read_text(encoding='utf-8').strip()
        if len(text)>2_000_000:raise ProcessorError('output_size_limit')
        texts.append((index,text))
        if searchable:
            pdf=target.with_suffix('.pdf')
            if not pdf.is_file():raise ProcessorError('ocr_failed')
            pdfs.append(pdf)
        source.unlink(missing_ok=True)
    if not any(text.strip() for _,text in texts):raise ProcessorError('no_text_found')
    return texts,pdfs


def artifact(path,page_count,mime):
    from .engine import ProcessorError,MAX_OUTPUT_BYTES
    if not path.is_file() or not 0<path.stat().st_size<=MAX_OUTPUT_BYTES:raise ProcessorError('output_invalid')
    os.chmod(path,0o600)
    return {'path':str(path.resolve()),'name':path.name,'mime_type':mime,'page_count':page_count,'size_bytes':path.stat().st_size}


def execute(feature, paths, parameters, out, metadata):
    from .engine import ProcessorError,_pdf_reader,_write_pdf,parse_pages,inspect_file,office_runtime
    from pypdf import PdfWriter
    source=paths[0];info=metadata[0]
    warnings=[]
    with tempfile.TemporaryDirectory(prefix='advanced-',dir=out) as temp:
        scratch=Path(temp)
        if feature.startswith('ocr.'):
            texts,pdfs=run_ocr(source,info,parameters,scratch,feature=='ocr.searchable_pdf')
            warnings=['printed_text_ocr_review_required']
            if feature=='ocr.extract_text':
                destination=out/'recognized-text.txt'
                destination.write_text('\n\n\f\n\n'.join(text for _,text in texts),encoding='utf-8')
                artifacts=[artifact(destination,len(texts),'text/plain; charset=utf-8')]
            else:
                writer=PdfWriter()
                for path in pdfs:
                    for page in _pdf_reader(path).pages:writer.add_page(page)
                artifacts=[_write_pdf(writer,out/'searchable.pdf')]
            engine='tesseract'
        elif feature=='convert.pdf_to_docx':
            from docx import Document
            from docx.shared import Pt
            if parameters['mode']=='ocr':
                texts,_=run_ocr(source,info,{**parameters,'dpi':200},scratch)
                warnings.append('printed_text_ocr_review_required')
            else:
                reader=_pdf_reader(source)
                texts=[(i,reader.pages[i].extract_text(extraction_mode='layout')) for i in parse_pages(parameters['pages'],len(reader.pages))]
            if not any(text.strip() for _,text in texts):raise ProcessorError('scan_requires_ocr')
            if sum(len(text) for _,text in texts)>2_000_000:raise ProcessorError('output_size_limit')
            document=Document();normal=document.styles['Normal'];normal.font.name='Noto Sans';normal.font.size=Pt(11)
            for index,(_,text) in enumerate(texts):
                if index:document.add_page_break()
                for line in text.splitlines():
                    if line.strip():document.add_paragraph(line.rstrip())
            destination=out/'editable-document.docx';document.save(destination)
            check=Document(destination)
            if not any(p.text for p in check.paragraphs):raise ProcessorError('output_invalid')
            count=inspect_file(destination)['page_count'] if office_runtime()['available'] else len(texts)
            artifacts=[artifact(destination,count,'application/vnd.openxmlformats-officedocument.wordprocessingml.document')]
            warnings+=['editable_text_reflowed_layout_not_preserved']
            engine='pypdf/python-docx'
        else:
            import pdfplumber
            from openpyxl import Workbook,load_workbook
            from openpyxl.styles import Font,PatternFill
            book=Workbook();book.remove(book.active)
            row_count=0
            with pdfplumber.open(source) as document:
                for page_index in parse_pages(parameters['pages'],len(document.pages)):
                    settings={'vertical_strategy':parameters['strategy'],'horizontal_strategy':parameters['strategy']}
                    for table_index,table in enumerate(document.pages[page_index].extract_tables(settings),1):
                        if not table or max(map(len,table),default=0)<2:continue
                        sheet=book.create_sheet(f'Page {page_index+1} table {table_index}')
                        for row in table:
                            row_count+=1
                            if row_count>20_000 or len(row)>100:raise ProcessorError('table_limit')
                            for value in row:
                                if value and len(value)>32_000:raise ProcessorError('table_limit')
                            sheet.append([value or '' for value in row])
                            for cell in sheet[sheet.max_row]:
                                # Formula-looking source text stays literal, never executable.
                                cell.data_type='s'
                        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='235CC5')
                        sheet.freeze_panes='A2'
                        for column in sheet.columns:sheet.column_dimensions[column[0].column_letter].width=24
            if not book.worksheets:raise ProcessorError('no_tables_found')
            destination=out/'extracted-tables.xlsx';book.save(destination)
            reloaded=load_workbook(destination,data_only=False)
            if any(cell.data_type=='f' for sheet in reloaded for row in sheet for cell in row):raise ProcessorError('output_invalid')
            artifacts=[artifact(destination,info['page_count'],'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')]
            warnings=['native_tables_only_review_cells','extracted_values_are_literal_text']
            engine='pdfplumber/openpyxl'
    return artifacts,{'engine':engine,'warnings':warnings}
