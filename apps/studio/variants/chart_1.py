"""Chart variants: every quantity measured against the slide's largest, as columns, stems, bars, dots and tracks."""
from pptx.util import Inches

from . import variant
from ..slide_layouts import _text_lines, _units
from ..slides import WIDTH_SAFETY, contrast_ratio
from ..slides import _from_hls as from_hls
from ..slides import _hls as hls
from ..slides import _luminance as luminance
from .kit import (FOOT, FRAME, LEFT, RIGHT, RULE, WIDTH, columns_of, disc, entries, eyebrow, figure, header_band,
                  header_compact, header_number, headline, hex_of, hues, inks, line, numbers, on, oval, readable,
                  rect, ring, rounded, rule, strong, text)

TOP_EDGE = FRAME[1]
# The classic drawer's word limits: a bar's category, its figure, and the slide's unit.
LABEL, VALUE, UNIT = 6, 3, 6
HEADING = 14


# ---------------------------------------------------------------- reading


def _plot(section):
    """(items, values, peak): the bars as the classic drawer takes them, values never below 0, peak never 0."""
    items = entries('chart', section)
    values = [max(0.0, value) for value in numbers(items)]
    return items, values, (max(values) if values else 0) or 1.0


def _unit(section):
    unit = columns_of(section)[:1]
    return unit[0] if unit else ''


def _plain(value, words):
    parts = str(value or '').split()
    return ' '.join(parts[:words]) + ('…' if len(parts) > words else '')


# ---------------------------------------------------------------- measuring


def _bold(ctx, bold, heading):
    return bool(bold) or bool(heading and ctx.heading_bold)


def _lines(ctx, value, size, width, *, bold=False, heading=False):
    return _text_lines(value, size, Inches(width), bold=_bold(ctx, bold, heading))


def _wide(ctx, value, size, *, bold=False, heading=False):
    """Inches `value` takes on one line at `size` points, as the fitter measures it."""
    return 0.001 * size * _units(value) * WIDTH_SAFETY * (1.06 if _bold(ctx, bold, heading) else 1) / 72 + 0.02


def _settles(ctx, value, size, width, height, lines, *, bold, heading, spacing=1.15):
    if not value:
        return True
    rows = _lines(ctx, value, size, width, bold=bold, heading=heading)
    if rows > lines or rows * size * spacing > height * 72 - 3:
        return False
    return all(_lines(ctx, word, size, width, bold=bold, heading=heading) == 1 for word in value.split())


def _label_size(ctx, items, width, height, start, floor=11, *, lines=2, bold=False, heading=False):
    """The one size every label is set at: the largest from `start` at which each keeps whole words within
    `lines` lines of `width` and `height` inches.

    A label as long as a sentence keeps to 12 pt, as running text does. Labels
    that hold at no size, or only by breaking a word, mark the slide
    shortened: the variant is then passed over for it (variants.choose).
    """
    labels = [_plain(item.get('label'), LABEL) for item in items]
    floor = max(floor, 12) if any(len(value) > 40 for value in labels) else floor

    def held(size):
        return all(_settles(ctx, value, size, width, height, lines, bold=bold, heading=heading) for value in labels)
    size = start
    while size > floor and not held(size):
        size -= 1
    if not held(size):
        ctx.shortened = True
    return size


def _label_tall(ctx, items, width, size, *, bold=False, heading=False):
    """Inches the tallest label takes at `size` in a box `width` wide."""
    rows = max([_lines(ctx, _plain(item.get('label'), LABEL), size, width, bold=bold, heading=heading)
                for item in items if item.get('label')] or [1])
    return rows * size * 1.15 / 72 + 0.06


def _figure_size(ctx, items, width, start, floor=14):
    """The one size every figure is set at: the largest from `start` at which each keeps to one line `width` wide."""
    values = [_plain(item.get('value'), VALUE) for item in items]
    size = start
    while size > floor and any(_lines(ctx, value, size, width, heading=True) > 1 for value in values if value):
        size -= 2
    return size


def _figure_tall(size):
    """The height of a one-line figure's box at `size` points."""
    return (size * 1.1 + 4) / 72


def _figure_wide(ctx, items, size):
    """Inches the widest figure takes at `size` in the heading face."""
    return max([_wide(ctx, _plain(item.get('value'), VALUE), size, heading=True)
                for item in items if item.get('value')] or [0.3])


def _heading_fit(ctx, section, width, size, lines):
    """(size, inches) the slide's heading takes in `width` at most `lines` lines, as kit.headline sets it."""
    value = _plain(section.get('heading'), HEADING)
    count = _lines(ctx, value, size, width, heading=True)
    while size > 18 and count > lines:
        size -= 2
        count = _lines(ctx, value, size, width, heading=True)
    return size, max(1, count) * size * 1.08 / 72 + 0.1


# ---------------------------------------------------------------- colour


def _dark(ctx, fill):
    """Whether the fill (role or hex) is a dark ground, as slides._reads_on judges one."""
    return luminance(hex_of(ctx, fill)) <= 0.18


def _paint(ctx, fill='surface'):
    """How a bar is painted on the fill role `fill`: {'fill', 'line', 'edge'}.

    The accent itself where it shows there. A pale brand colour on a light
    ground keeps its fill and gains an outline in the accent moved until it
    reads (`line`), which is also the colour of anything drawn as a line or
    ring (`edge`). A dark accent on a dark ground (navy on near-black) all but
    vanishes, and a filled disc could not be told from a ring: there the
    marks are filled in that lifted accent too.
    """
    if contrast_ratio(hex_of(ctx, 'accent'), hex_of(ctx, fill)) >= 2.2:
        return {'fill': 'accent', 'line': None, 'edge': 'accent'}
    edge = readable(ctx, 'accent', fill, 3.0)
    if _dark(ctx, fill):
        return {'fill': edge, 'line': None, 'edge': edge}
    return {'fill': 'accent', 'line': edge, 'edge': edge}


def _series(ctx, n):
    """kit.hues(ctx, n) for long bars: the design's accent and second colour as they are, the colours turned
    round the wheel from them toned down.

    kit.hues sets the turned colours at the accent's own strength, so a vivid
    accent brings neon lime, mint and magenta, garish over a bar's whole
    length. Each turned colour keeps its hue at a calmer saturation and a
    mid lightness; a grey accent's tones are kept as they are.
    """
    paints = hues(ctx, n)
    if hls(hex_of(ctx, 'accent'))[2] < 0.12:
        return paints
    brand = {hex_of(ctx, 'accent'), hex_of(ctx, 'secondary')}
    tempered = []
    for paint in paints:
        fill = hex_of(ctx, paint['fill'])
        if fill not in brand:
            hue, lightness, saturation = hls(fill)
            fill = from_hls(hue, min(0.48, max(0.38, lightness)), min(saturation, 0.42))
            paint = {'fill': fill, 'ink': on(fill), 'muted': on(fill), 'text': readable(ctx, fill, 'surface', 4.5)}
        tempered.append(paint)
    return tempered


def _hue(ctx, paint):
    """A bar's fill from `_series` made to show on the slide: (paint, outline options).

    A pale fill on a light ground (a brand yellow) keeps its colour and gains
    an outline in that colour moved until it reads; a dark one on a dark
    ground (a navy accent) is lifted, with its inks worked out again for it.
    """
    fill = hex_of(ctx, paint['fill'])
    if contrast_ratio(fill, hex_of(ctx, 'surface')) >= 2.2:
        return paint, {}
    edge = readable(ctx, fill, 'surface', 3.0)
    if _dark(ctx, 'surface'):
        return {**paint, 'fill': edge, 'ink': on(edge), 'muted': on(edge),
                'text': readable(ctx, edge, 'surface', 4.5)}, {}
    return paint, {'line': edge, 'line_width': 0.75}


def _outline(paint):
    return {'line': paint['line'], 'line_width': 0.75} if paint['line'] else {}


def _blend(ctx, top, bottom, share):
    """`share` of the colour `top` over `bottom` (roles or hex), as one solid hex colour."""
    one, other = hex_of(ctx, top), hex_of(ctx, bottom)
    return ''.join(f'{round(int(one[i:i + 2], 16) * share + int(other[i:i + 2], 16) * (1 - share)):02X}'
                   for i in (0, 2, 4))


def _show_unit(ctx, section, box, *, colour='muted', align='right'):
    """The slide's unit (columns[0]) as a quiet caption in `box`; whether there was one."""
    unit = _unit(section)
    if unit:
        text(ctx, box, unit, size=12, floor=11, colour=colour, align=align, anchor='middle', words=UNIT)
    return bool(unit)


def _slots(n, x, width, widest, gap):
    """`n` equal slots of at most `widest`, `gap` apart and centred in `x`..`x + width`: [(x, width), ...]."""
    size = min(widest, (width - gap * (n - 1)) / n)
    start = x + (width - n * size - (n - 1) * gap) / 2
    return [(start + k * (size + gap), size) for k in range(n)]


# ---------------------------------------------------------------- columns


@variant('chart', 'columns')
def columns(ctx, section, lines):
    """Accent columns rising from a thin baseline under a compact heading, each figure set above its column and
    each category centred beneath it."""
    top = header_compact(ctx, section)
    items, values, peak = _plot(section)
    n = len(items)
    gap = 0.5 if n <= 4 else 0.32 if n <= 6 else 0.2
    slots = _slots(n, LEFT, WIDTH, 2.5, gap)
    slot = slots[0][1]
    reach = 0.25
    start, end = max(LEFT, slots[0][0] - reach), min(RIGHT, slots[-1][0] + slot + reach)
    label = _label_size(ctx, items, slot, 1.0, 16 if n <= 4 else 15 if n <= 6 else 14, lines=4)
    label_tall = _label_tall(ctx, items, slot, label)
    size = _figure_size(ctx, items, slot, 30 if n <= 4 else 26 if n <= 6 else 22)
    value_tall = _figure_tall(size)
    chart_top = top + (0.45 if _show_unit(ctx, section, (end - 5.5, top, 5.5, 0.3)) else 0.15)
    base = FOOT - label_tall - 0.16
    lift = 0.1
    tallest = base - chart_top - value_tall - lift
    width = min(1.3, slot * 0.62)
    paint = _paint(ctx)
    for item, value, (x, w) in zip(items, values, slots):
        tall = max(0.06, tallest * value / peak)
        cx = x + w / 2
        rect(ctx, (cx - width / 2, base - tall, width, tall), paint['fill'], **_outline(paint))
        figure(ctx, (x, base - tall - lift - value_tall, w, value_tall), item.get('value'), size=size,
               colour='accent_text', align='center', anchor='bottom', words=VALUE)
        text(ctx, (x, base + 0.16, w, label_tall), item.get('label'), size=label, floor=label,
             colour='ink', align='center', words=LABEL)
    line(ctx, (start, base), (end, base), 'muted', 1.0)


# ---------------------------------------------------------------- lollipop


@variant('chart', 'lollipop')
def lollipop(ctx, section, lines):
    """Beside the slide's number and heading, each category sits at the left of a thin stem that runs from a
    shared baseline to a haloed accent dot at its value, the figure beside the dot."""
    top = header_number(ctx, section) + 0.05
    items, values, peak = _plot(section)
    n = len(items)
    room = FOOT - top
    row = min(1.0 if n <= 3 else 0.8, room / n)
    label_w = 3.3
    label = _label_size(ctx, items, label_w, row - 0.02, 18 if row >= 0.7 else 16, lines=3)
    size = _figure_size(ctx, items, 1.6, 22 if row >= 0.6 else 18, floor=14)
    value_w = _figure_wide(ctx, items, size) + 0.1
    value_tall = _figure_tall(size)
    dot, halo = (0.26, 0.46) if row >= 0.55 else (0.22, 0.38)
    base = LEFT + label_w + 0.3
    reach = RIGHT - value_w - 0.12 - halo / 2 - base
    y0 = top + (room - n * row) / 2
    # The unit heads the figures' side, just above the first row.
    _show_unit(ctx, section, (RIGHT - 5.5, max(top - 0.38, y0 - 0.3), 5.5, 0.3))
    line(ctx, (base, y0 + 0.08), (base, y0 + n * row - 0.08), 'muted', 1.25)
    paint = _paint(ctx)
    for k, (item, value) in enumerate(zip(items, values)):
        y = y0 + k * row
        cy = y + row / 2
        cx = base + max(halo / 2 + 0.06, reach * value / peak)
        line(ctx, (base, cy), (cx - dot / 2, cy), paint['edge'], 2.0)
        oval(ctx, disc(cx, cy, halo), 'accent_soft')
        oval(ctx, disc(cx, cy, dot), paint['fill'], **_outline(paint))
        text(ctx, (LEFT, y + 0.01, label_w, row - 0.02), item.get('label'), size=label, floor=label,
             colour='ink', align='right', anchor='middle', words=LABEL)
        figure(ctx, (cx + halo / 2 + 0.1, cy - value_tall / 2 - 0.02, value_w, value_tall), item.get('value'),
               size=size, colour='ink', anchor='middle', words=VALUE)


# ---------------------------------------------------------------- hue bars


@variant('chart', 'hue_bars')
def hue_bars(ctx, section, lines):
    """The heading in a column at the left, centred on thick bars beside it, each bar in a colour of its own with
    its figure set inside its end, or just past a bar too short to hold it."""
    side = 3.4
    x0 = LEFT + side + 0.55
    items, values, peak = _plot(section)
    n = len(items)
    top = TOP_EDGE + (0.45 if _unit(section) else 0.1)
    room = FOOT - top
    row = min(0.95, room / n)
    bar = min(0.6, row * 0.64)
    label_w = 2.3
    label = _label_size(ctx, items, label_w, row - 0.04, 16 if row >= 0.7 else 15, lines=3)
    size = _figure_size(ctx, items, 1.6, min(20, int((bar * 72 - 4) / 1.1)), floor=12)
    value_w = _figure_wide(ctx, items, size) + 0.04
    value_tall = _figure_tall(size)
    start = x0 + label_w + 0.25
    reach = RIGHT - start
    paints = _series(ctx, n)
    y0 = top + (room - n * row) / 2
    _show_unit(ctx, section, (x0, max(TOP_EDGE, y0 - 0.4), RIGHT - x0, 0.3))

    # The heading's column: the slide's number at the top, the heading and its rule centred on the bars.
    eyebrow(ctx, (LEFT, TOP_EDGE, side, 0.3))
    heading_size, heading_tall = _heading_fit(ctx, section, side, 30, 6)
    block = 0.22 + heading_tall
    y = min(FOOT - block, max(TOP_EDGE + 0.55, y0 + n * row / 2 - block / 2))
    rule(ctx, (LEFT, y, 0.9, RULE))
    headline(ctx, section, (LEFT, y + 0.22, side, heading_tall), size=heading_size, lines=6, anchor='top')

    for k, (item, value, paint) in enumerate(zip(items, values, paints)):
        paint, outline = _hue(ctx, paint)
        y = y0 + k * row
        by = y + (row - bar) / 2
        length = max(0.1, reach * value / peak)
        rounded(ctx, (start, by, length, bar), paint['fill'], radius=0.07, **outline)
        text(ctx, (x0, y + 0.02, label_w, row - 0.04), item.get('label'), size=label, floor=label,
             colour='ink', align='right', anchor='middle', words=LABEL)
        middle = by + (bar - value_tall) / 2 - 0.02
        if length >= value_w + 0.4:
            figure(ctx, (start + length - 0.16 - value_w, middle, value_w, value_tall), item.get('value'),
                   size=size, colour=paint['ink'], align='right', anchor='middle', words=VALUE)
        else:
            figure(ctx, (start + length + 0.12, middle, value_w, value_tall), item.get('value'),
                   size=size, colour=paint['text'], anchor='middle', words=VALUE)


# ---------------------------------------------------------------- dot rows


@variant('chart', 'dot_rows')
def dot_rows(ctx, section, lines):
    """Under the heading on a strong band, a pictogram row per category: ten small discs, as many filled as its
    share of the largest value and the rest left as rings, the category at the left and the figure at the right."""
    top = header_band(ctx, section, fill=strong(ctx), height=1.5)
    items, values, peak = _plot(section)
    n = len(items)
    unit_top = top - 0.06
    top += 0.32 if _unit(section) else 0.05
    room = FOOT - top
    row = min(0.78, room / n)
    space = 0.4
    size = _figure_size(ctx, items, 1.6, 24 if row >= 0.6 else 20, floor=14)
    value_w = max(1.0, _figure_wide(ctx, items, size) + 0.05)
    dot = min(0.44, row * 0.6)
    # Ten discs a little more than a disc apart, beside the figures; the categories take the rest of the row.
    pitch = min(dot * 1.5, (WIDTH - value_w - 2 * space - 2.6) / 10)
    dot = min(dot, pitch * 0.74)
    x_dots = RIGHT - value_w - space - 10 * pitch
    label_w = x_dots - space - LEFT
    label = _label_size(ctx, items, label_w, row - 0.08, 17 if row >= 0.65 else 15, lines=3, bold=True,
                        heading=True)
    paint = _paint(ctx)
    y0 = top + (room - n * row) / 2
    _show_unit(ctx, section, (RIGHT - 5.5, max(unit_top, y0 - 0.36), 5.5, 0.3))
    for k, (item, value) in enumerate(zip(items, values)):
        y = y0 + k * row
        cy = y + row / 2
        if k:
            line(ctx, (LEFT, y), (RIGHT, y), 'accent_soft', 0.75)
        filled = round(10 * value / peak)
        filled = max(1, filled) if value > 0 else 0
        for d in range(10):
            cx = x_dots + pitch * (d + 0.5)
            if d < filled:
                oval(ctx, disc(cx, cy, dot), paint['fill'], **_outline(paint))
            else:
                ring(ctx, disc(cx, cy, dot - 0.03), paint['edge'], width=1.25)
        text(ctx, (LEFT, y + 0.04, label_w, row - 0.08), item.get('label'), size=label, floor=label,
             colour='ink', bold=True, heading=True, anchor='middle', words=LABEL)
        figure(ctx, (RIGHT - value_w, cy - _figure_tall(size) / 2 - 0.02, value_w, _figure_tall(size)),
               item.get('value'), size=size, colour='accent_text', align='right', anchor='middle', words=VALUE)


# ---------------------------------------------------------------- progress tracks


@variant('chart', 'progress_tracks')
def progress_tracks(ctx, section, lines):
    """Under a centred heading, a rounded pale track per category, filled in the accent to its share of the
    largest value, with the category above the track and the figure at its right end — in two columns for five
    or more."""
    eyebrow(ctx, (LEFT, TOP_EDGE, WIDTH, 0.3), align='center')
    headline(ctx, section, (LEFT + 0.8, 0.95, WIDTH - 1.6, 1.08), size=30, lines=2, align='center')
    rule(ctx, (LEFT + WIDTH / 2 - 0.45, 2.16, 0.9, RULE))
    unit = _show_unit(ctx, section, (LEFT + WIDTH / 2 - 3.0, 2.3, 6.0, 0.28), align='center')
    top = 2.78 if unit else 2.5
    items, values, peak = _plot(section)
    n = len(items)
    per = 2 if n >= 5 else 1
    rows = -(-n // per)
    gap = 0.8
    width = min(9.4, (WIDTH - gap * (per - 1)) / per)
    starts = [LEFT + (WIDTH - per * width - (per - 1) * gap) / 2 + c * (width + gap) for c in range(per)]
    room = FOOT - top
    track = 0.24 if rows <= 3 else 0.2
    size = _figure_size(ctx, items, 1.8, 22 if rows <= 3 else 20, floor=14)
    value_w = _figure_wide(ctx, items, size) + 0.06
    label_w = width - value_w - 0.25
    label = _label_size(ctx, items, label_w, 0.62, 17 if rows <= 3 else 16, lines=2, bold=True, heading=True)
    head = max(_label_tall(ctx, items, label_w, label, bold=True, heading=True), _figure_tall(size))
    item_tall = head + 0.1 + track
    pitch = min(item_tall + 0.55, room / rows)
    y0 = top + (room - rows * pitch + (pitch - item_tall)) / 2
    paint = _paint(ctx, 'card')
    for k, (item, value) in enumerate(zip(items, values)):
        row, column = (k % rows, k // rows) if per > 1 else (k, 0)
        x, y = starts[column], y0 + row * pitch
        ty = y + head + 0.1
        rounded(ctx, (x, ty, width, track), 'card', radius=track / 2)
        rounded(ctx, (x, ty, max(track, width * value / peak), track), paint['fill'], radius=track / 2,
                **_outline(paint))
        text(ctx, (x, y, label_w, head), item.get('label'), size=label, floor=label, colour='ink',
             bold=True, heading=True, anchor='bottom', words=LABEL)
        figure(ctx, (x + width - value_w, y, value_w, head), item.get('value'), size=size, colour='accent_text',
               align='right', anchor='bottom', words=VALUE)


# ---------------------------------------------------------------- panel columns


@variant('chart', 'panel_columns')
def panel_columns(ctx, section, lines):
    """A strong panel fills the frame with the heading at its top; below it a rounded capsule per category,
    filled from the foot to its share of the largest value, the figures in a row above the capsules and the
    categories beneath."""
    panel = inks(strong(ctx))
    rect(ctx, FRAME, panel['fill'])
    pad = 0.5
    inner = WIDTH - 2 * pad
    eyebrow(ctx, (LEFT + pad, TOP_EDGE + 0.32, 3.0, 0.3), colour=panel['muted'])
    _show_unit(ctx, section, (RIGHT - pad - 4.5, TOP_EDGE + 0.32, 4.5, 0.3), colour=panel['muted'])
    rule(ctx, (LEFT + pad, TOP_EDGE + 0.74, 0.9, RULE), colour=panel['accent'])
    heading_size, heading_tall = _heading_fit(ctx, section, inner * 0.75, 28, 2)
    headline(ctx, section, (LEFT + pad, TOP_EDGE + 0.9, inner * 0.75, heading_tall), size=heading_size, lines=2,
             anchor='top', colour=panel['ink'])
    items, values, peak = _plot(section)
    n = len(items)
    gap = 0.45 if n <= 4 else 0.3 if n <= 6 else 0.16
    slots = _slots(n, LEFT + pad, inner, 2.2, gap)
    slot = slots[0][1]
    label = _label_size(ctx, items, slot, 0.95, 15 if n <= 4 else 14 if n <= 6 else 13, lines=4)
    label_tall = _label_tall(ctx, items, slot, label)
    size = _figure_size(ctx, items, slot, 30 if n <= 4 else 26 if n <= 6 else 22)
    value_tall = _figure_tall(size)
    chart_top = TOP_EDGE + 0.9 + heading_tall + 0.3
    base = FOOT - 0.4 - label_tall - 0.16
    track_top = chart_top + value_tall + 0.14
    width = min(0.56, slot * 0.4)
    tall = base - track_top
    # The accent where it shows on the panel; else the panel's own ink.
    fill = 'accent' if contrast_ratio(hex_of(ctx, 'accent'), hex_of(ctx, panel['fill'])) >= 3.0 else panel['accent']
    track = _blend(ctx, panel['ink'], panel['fill'], 0.16)
    for item, value, (x, w) in zip(items, values, slots):
        cx = x + w / 2
        rounded(ctx, (cx - width / 2, track_top, width, tall), track, radius=width / 2)
        filled = max(width, tall * value / peak)
        rounded(ctx, (cx - width / 2, base - filled, width, filled), fill, radius=width / 2)
        figure(ctx, (x, chart_top, w, value_tall), item.get('value'), size=size, colour=panel['ink'],
               align='center', anchor='bottom', words=VALUE)
        text(ctx, (x, base + 0.16, w, label_tall), item.get('label'), size=label, floor=label,
             colour=panel['muted'], align='center', words=LABEL)
