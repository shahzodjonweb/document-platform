"""Slide decks that look designed rather than generated.

A slide is not a short page. It is a headline and a few bullets, with what the
presenter would say kept in the notes — so this module builds from that shape
rather than flowing prose into a box.

Three things here are load-bearing and easy to break by accident:

* **The first shape added to every slide is the accent element**, filled with
  the accent exactly as the customer's branding supplied it. Derived text
  colours are contrast-corrected; the brand colour itself never is, because
  repainting someone's brand to win a contrast check is worse than the check.
* **One slide per section, always.** The slide count is what was quoted and
  billed. Nothing here adds or removes a slide: it sets the deck tighter, moves
  a slide to two columns, and only then drops bullets.
* **No binary parts may survive the save.** `processors/engine.py` refuses any
  `.bin` member on upload, so a deck carrying python-pptx's inherited printer
  settings would be rejected by the platform that produced it.

Text is written as real paragraphs with wrapping left to the viewer. The old
renderer measured line breaks itself and wrote one paragraph per wrapped line
with `word_wrap=False`, which meant the "editable deck" fell apart the moment
anyone edited it.
"""
import colorsys
import io
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from .rendering import FIT_STEPS, SENTENCE, font_path

# python-pptx cannot embed fonts: `font.name` is a name the viewer resolves for
# itself. So the deck names faces that ship with Microsoft Office on Windows and
# macOS and cover Cyrillic, rather than the Noto Sans the server happens to have.
# If these substitute badly in Google Slides or Keynote, Georgia/Verdana are the
# duller but even safer fallback — two constants, nothing else changes.
HEADING_FONT = 'Cambria'
BODY_FONT = 'Corbel'
BULLET_FONT = 'Arial'

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)
MARGIN = Inches(0.85)

# A list slide carries real explanation: up to six bullets, each a sentence.
# They used to be five fragments of fourteen words, and decks read as empty.
MAX_BULLETS = 6
BULLET_WORDS = 20
MIN_BODY_SIZE = 14
MAX_BODY_SIZE = 20
# The smallest a list may go on a slide that still does not fit the deck size.
FLOOR_BODY_SIZE = 13
# Measured metrics come from reportlab's Noto; the viewer renders Corbel. The
# estimate only has to be close enough to choose a size, so it is padded rather
# than made exact.
WIDTH_SAFETY = 1.10
LINE_MULTIPLE = 1.22

MARKER = re.compile(r'^\s*(?:[•·▪◦‣*\-–—]+\s*|\(?[0-9]{1,2}[.)]\s+|\(?[a-zA-Z][.)]\s+)')


def _register_font():
    if 'PDFMaster' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('PDFMaster', str(font_path())))


# ---------------------------------------------------------------- colour


def _hex(value):
    value = (value or '').lstrip('#')
    return value if re.fullmatch(r'[0-9a-fA-F]{6}', value) else '255E49'


def _rgb(value):
    return RGBColor.from_string(_hex(value).upper())


def _hls(value):
    red, green, blue = (int(_hex(value)[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return colorsys.rgb_to_hls(red, green, blue)


def _from_hls(hue, lightness, saturation):
    red, green, blue = colorsys.hls_to_rgb(hue % 1.0, min(1.0, max(0.0, lightness)),
                                           min(1.0, max(0.0, saturation)))
    return '%02X%02X%02X' % (round(red * 255), round(green * 255), round(blue * 255))


def _luminance(value):
    channels = []
    for index in (0, 2, 4):
        channel = int(_hex(value)[index:index + 2], 16) / 255
        channels.append(channel / 12.92 if channel <= 0.03928 else ((channel + 0.055) / 1.055) ** 2.4)
    return 0.2126 * channels[0] + 0.7152 * channels[1] + 0.0722 * channels[2]


def contrast_ratio(one, other):
    """WCAG contrast between two hex colours, lighter over darker."""
    first, second = _luminance(one), _luminance(other)
    high, low = max(first, second), min(first, second)
    return (high + 0.05) / (low + 0.05)


def _reads_on(value, background, target):
    """Move a colour away from its background until it reads, or run out of room.

    Which way is away depends on the background, so this serves a light deck and
    a dark one with the same call.
    """
    hue, lightness, saturation = _hls(value)
    step = -0.03 if _luminance(background) > 0.18 else 0.03
    for _ in range(40):
        if contrast_ratio(value, background) >= target:
            break
        lightness += step
        if not 0.0 <= lightness <= 1.0:
            break
        value = _from_hls(hue, lightness, saturation)
    return value


# Surface, secondary surface and soft-tint lightnesses per theme. Everything
# else is derived and contrast-corrected, so a theme is three numbers and a
# decision about what the cover and the dividers are filled with.
THEME_GROUND = {
    'light': (0.975, 0.94, 0.86),
    'dark': (0.11, 0.165, 0.26),
    'bold': (0.975, 0.94, 0.86),
}


def _soft(accent, secondary, lightness):
    """The quiet tint behind marks and cells: the accent's, or a design's second colour."""
    hue, _, saturation = _hls(secondary or accent)
    return _from_hls(hue, lightness, min(saturation, 0.45 if secondary else 0.35))


def palette(accent, theme='light', secondary=None, paper=None, accent_headings=False):
    """A small set of roles derived from the one colour a customer can choose.

    `accent` is returned untouched — it is the brand, and it is what fills the
    accent shape. Everything that carries text is contrast-corrected against the
    surface it sits on, which is most of what separates a designed deck from a
    generated one, and is what lets an arbitrary accent be safe.
    """
    accent = _hex(accent).upper()
    theme = theme if theme in THEME_GROUND else 'light'
    hue, _, saturation = _hls(accent)
    ground, alt, soft = THEME_GROUND[theme]
    dark = ground < 0.5
    surface = _from_hls(hue, ground, min(saturation, 0.12 if dark else 0.10))
    surface_alt = _from_hls(hue, alt, min(saturation, 0.14))
    if paper and _hls(paper)[1] >= 0.5 and not dark:
        # A design's own paper: cream, aqua, warm stone. Panels sit a step darker.
        surface = _hex(paper).upper()
        paper_hue, paper_light, paper_sat = _hls(surface)
        surface_alt = _from_hls(paper_hue, paper_light - 0.045, min(paper_sat + 0.05, 0.4))
    elif paper and _hls(paper)[1] < 0.5 and dark:
        # A dark design's own ground: a green board, deep navy. Panels a step lighter.
        surface = _hex(paper).upper()
        paper_hue, paper_light, paper_sat = _hls(surface)
        surface_alt = _from_hls(paper_hue, paper_light + 0.055, min(paper_sat, 0.4))
    roles = {
        'accent': accent,
        'theme': theme,
        'surface': surface,
        'surface_alt': surface_alt,
        'accent_soft': _soft(accent, secondary, soft),
        'ink': _from_hls(hue, 0.95 if dark else 0.13, min(saturation, 0.08 if dark else 0.25)),
        'muted': _from_hls(hue, 0.66 if dark else 0.45, min(saturation, 0.15)),
        'accent_text': accent,
    }
    roles['ink'] = _reads_on(roles['ink'], surface, 7.0)
    roles['muted'] = _reads_on(roles['muted'], surface, 4.5)
    roles['accent_text'] = _reads_on(roles['accent_text'], surface, 4.5)
    # Text on the accent itself: start from whichever end already reads better,
    # then push it until it does. Picking the better of black and white is not
    # enough — a mid-tone accent is too far from both.
    on_dark, on_light = _from_hls(hue, 0.13, min(saturation, 0.2)), 'FFFFFF'
    best = on_light if contrast_ratio(on_light, accent) >= contrast_ratio(on_dark, accent) else on_dark
    roles['on_accent'] = _reads_on(best, accent, 4.5)
    # A bold deck puts the accent itself behind the cover and the dividers; the
    # others keep them quiet, so the accent stays an accent.
    if theme == 'bold':
        roles['cover_fill'], roles['cover_ink'] = accent, roles['on_accent']
        roles['band'], roles['band_ink'] = accent, roles['on_accent']
    else:
        roles['cover_fill'] = 'FFFFFF' if theme == 'light' and not paper else surface
        roles['cover_ink'] = roles['ink'] if theme == 'dark' else _reads_on(roles['ink'], 'FFFFFF', 7.0)
        roles['band'], roles['band_ink'] = roles['surface_alt'], roles['ink']
    roles['cover_muted'] = (roles['cover_ink'] if theme == 'bold'
                            else _reads_on(roles['muted'], roles['cover_fill'], 4.5))
    roles['band_muted'] = (roles['band_ink'] if theme == 'bold'
                           else _reads_on(roles['muted'], roles['band'], 4.5))
    # Text that sits on a filled panel — a card, a tinted cell — is corrected
    # against that panel rather than against the slide behind it.
    roles['card'] = roles['surface_alt']
    roles['card_ink'] = _reads_on(roles['ink'], roles['card'], 7.0)
    roles['card_muted'] = _reads_on(roles['muted'], roles['card'], 4.5)
    roles['card_accent'] = _reads_on(accent, roles['card'], 4.5)
    # Headlines in the accent are a design's choice; large text needs 3:1, and
    # accent_text already reads at 4.5:1 on the surface.
    roles['heading'] = roles['accent_text'] if accent_headings else roles['ink']
    roles['soft_ink'] = _reads_on(roles['ink'], roles['accent_soft'], 7.0)
    roles['soft_muted'] = _reads_on(roles['muted'], roles['accent_soft'], 4.5)
    # More fills a composition (apps/studio/compositions) can set text on, each
    # with the ink that reads there. The design's second colour as a fill:
    roles['secondary'] = _hex(secondary).upper() if secondary else roles['accent_soft']
    roles['on_secondary'] = ink_on(roles['secondary'])
    # A deep, near-black version of the accent for night panels and big grounds.
    roles['deep'] = _from_hls(hue, 0.16 if not dark else 0.07, min(saturation, 0.6))
    roles['on_deep'] = ink_on(roles['deep'], 7.0)
    roles['deep_muted'] = _reads_on(_from_hls(hue, 0.72, min(saturation, 0.25)), roles['deep'], 4.5)
    # A dark gradient from the accent to the second colour, darkened until white
    # reads on both ends; and a light one from the surface to the soft tint,
    # with ink that reads on its darker end.
    roles['gradient'] = (_darken_for('FFFFFF', accent), _darken_for('FFFFFF', roles['secondary']))
    roles['on_gradient'] = 'FFFFFF'
    roles['gradient_light'] = (surface, roles['accent_soft'])
    roles['on_gradient_light'] = roles['soft_ink']
    return roles


def ink_on(fill, target=4.5):
    """Text that reads on `fill`: from whichever end already reads better, pushed until it does."""
    hue, _, saturation = _hls(fill)
    light, dark = 'FFFFFF', _from_hls(hue, 0.12, min(saturation, 0.2))
    best = light if contrast_ratio(light, fill) >= contrast_ratio(dark, fill) else dark
    return _reads_on(best, fill, target)


def _darken_for(text, fill, target=4.5):
    """`fill` made darker, keeping its hue, until `text` reads on it."""
    hue, lightness, saturation = _hls(fill)
    while contrast_ratio(text, _from_hls(hue, lightness, saturation)) < target and lightness > 0.05:
        lightness -= 0.02
    return _from_hls(hue, lightness, saturation)


# ---------------------------------------------------------------- content


def bullets_of(body, limit=MAX_BULLETS, words=BULLET_WORDS):
    """The lines of a slide, as bullets, with the model's own markers removed.

    Returns `(bullets, shortened)`. A model told "one bullet per line" will still
    sometimes write "• " at the start of each, and rendering that under a real
    bullet glyph gives "• • Revenue rose" — stripping it is the single most
    visible line in this function.

    A body that arrives as one long paragraph is split into sentences instead.
    That is what makes a legacy draft, a locally authored one, or a model that
    ignored the brief render as a deck rather than as one enormous bullet.
    """
    lines = [MARKER.sub('', line).strip() for line in (body or '').split('\n')]
    lines = [line for line in lines if line]
    if len(lines) == 1 and len(lines[0]) > words * 8:
        lines = [part.strip() for part in SENTENCE.split(lines[0]) if part.strip()]
    shortened = False
    kept = []
    for line in lines[:limit]:
        pieces = line.split(' ')
        if len(pieces) > words * 2:
            line = ' '.join(pieces[:words * 2]).rstrip(' ,;:') + '…'
            shortened = True
        kept.append(line)
    if len(lines) > limit:
        shortened = True
        if kept:
            kept[-1] = kept[-1].rstrip('.') + '…'
    return kept, shortened


# ---------------------------------------------------------------- theme


def _theme_part(deck):
    return deck.slide_masters[0].part.part_related_by(RT.THEME)


def apply_theme(deck, roles, faces=(HEADING_FONT, BODY_FONT)):
    """Repaint the inherited Office theme so the deck is not stock Calibri blue.

    python-pptx has no theme API and loads `theme1.xml` as an opaque part, so
    this rewrites its bytes. Only the colour values and the two font names are
    touched: the format scheme, the object defaults and the eleven stock layouts
    stay exactly as Office wrote them, so "New Slide" in PowerPoint still works
    and there is no way to produce a file that needs repairing.

    Pinned to python-pptx==1.0.2 by the writable `blob` on a generic part; the
    guard test in tests/test_slides_deck.py fails loudly if that moves.
    """
    from lxml import etree
    part = _theme_part(deck)
    theme = etree.fromstring(part.blob)
    scheme = theme.find(qn('a:themeElements'))
    if scheme is None:
        return
    colours = scheme.find(qn('a:clrScheme'))
    fonts = scheme.find(qn('a:fontScheme'))
    wanted = {'dk1': roles['ink'], 'dk2': roles['ink'], 'lt1': 'FFFFFF', 'lt2': roles['surface'],
              'accent1': roles['accent'], 'accent2': roles['accent_soft'], 'accent3': roles['muted'],
              'accent4': roles['ink'], 'accent5': roles['surface_alt'], 'accent6': roles['accent_text'],
              'hlink': roles['accent_text'], 'folHlink': roles['muted']}
    if colours is not None:
        for name, value in wanted.items():
            slot = colours.find(qn('a:' + name))
            if slot is None:
                continue
            for child in list(slot):
                slot.remove(child)
            srgb = slot.makeelement(qn('a:srgbClr'), {'val': value})
            slot.append(srgb)
    if fonts is not None:
        for tag, face in (('a:majorFont', faces[0]), ('a:minorFont', faces[1])):
            group = fonts.find(qn(tag))
            if group is None:
                continue
            latin = group.find(qn('a:latin'))
            if latin is not None:
                latin.set('typeface', face)
    part.blob = etree.tostring(theme, xml_declaration=True, encoding='UTF-8', standalone=True)


def strip_binary_parts(deck):
    """Leave no binary part behind.

    `processors/engine.py` refuses any `.bin` member on upload, so a deck still
    carrying python-pptx's inherited printer settings would be rejected by the
    platform that made it. The thumbnail goes for a different reason: it is a
    blank 4:3 Office preview that becomes the file's icon for a 16:9 deck.
    """
    for relationship in list(deck.part.rels.values()):
        if relationship.reltype.endswith('/printerSettings'):
            deck.part.drop_rel(relationship.rId)
    package = deck.part.package
    for relationship in list(package._rels.values()):
        if relationship.reltype.endswith('/thumbnail'):
            package.drop_rel(relationship.rId)


def apply_fonts(deck, look):
    """Put a design's fonts on every run the layouts wrote.

    The layouts write the two house faces; a design swaps them here in one
    pass, so no drawing function needs to know which design it is drawing.
    """
    faces = {HEADING_FONT: look['heading_font'], BODY_FONT: look['body_font']}
    if faces == {HEADING_FONT: HEADING_FONT, BODY_FONT: BODY_FONT} and not look['heading_bold']:
        return
    for slide in deck.slides:
        for latin in slide._element.iter(qn('a:latin')):
            face = latin.get('typeface')
            if face not in faces:
                continue
            latin.set('typeface', faces[face])
            if face == HEADING_FONT and look['heading_bold']:
                latin.getparent().set('b', '1')


# ---------------------------------------------------------------- text


def _frame(shape, anchor=MSO_ANCHOR.TOP):
    """A text frame that reflows when it is edited.

    python-pptx writes `wrap="none"` with `spAutoFit` by default, which is why
    the old decks grew their boxes off the slide instead of rewrapping.
    """
    frame = shape.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    return frame


def _paragraph(frame, index):
    return frame.paragraphs[0] if index == 0 else frame.add_paragraph()


def _write(paragraph, text, *, font, size, colour, bold=False, spacing=1.15,
           before=0, after=0, align=PP_ALIGN.LEFT):
    paragraph.text = text
    paragraph.alignment = align
    paragraph.line_spacing = spacing
    paragraph.space_before = Pt(before)
    paragraph.space_after = Pt(after)
    for run in paragraph.runs:
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = _rgb(colour)
    return paragraph


def _bullet(paragraph, colour, size, char='•'):
    """A real hanging bullet: the glyph sits outside the text, not above it.

    python-pptx has no bullet API. `marL` with a matching negative `indent` is
    what makes a wrapped bullet line up under its own text instead of under the
    dot — the difference between a list and someone having typed a dot.
    """
    properties = paragraph._pPr if paragraph._pPr is not None else paragraph._p.get_or_add_pPr()
    hang = Pt(size * 0.95)
    properties.set('marL', str(int(hang)))
    properties.set('indent', str(-int(hang)))
    for tag, attributes in ((qn('a:buClr'), None),
                            (qn('a:buFont'), {'typeface': BULLET_FONT, 'pitchFamily': '34', 'charset': '0'}),
                            (qn('a:buChar'), {'char': char})):
        for existing in properties.findall(tag):
            properties.remove(existing)
        element = properties.makeelement(tag, attributes or {})
        if tag == qn('a:buClr'):
            element.append(element.makeelement(qn('a:srgbClr'), {'val': _hex(colour).upper()}))
        properties.append(element)


def _no_bullet(paragraph):
    properties = paragraph._pPr if paragraph._pPr is not None else paragraph._p.get_or_add_pPr()
    for tag in ('a:buClr', 'a:buFont', 'a:buChar', 'a:buAutoNum', 'a:buNone'):
        for existing in properties.findall(qn(tag)):
            properties.remove(existing)
    properties.append(properties.makeelement(qn('a:buNone'), {}))


def _text_width(text, size):
    _register_font()
    return pdfmetrics.stringWidth(text, 'PDFMaster', size) * WIDTH_SAFETY


def _wrapped_lines(text, size, width_emu):
    width = Emu(int(width_emu)).pt
    return max(1, -(-_text_width(text, size) // max(width, 1)))


# ---------------------------------------------------------------- layout


class Zone:
    """Where something sits on a slide, in EMU.

    Geometry is named rather than written inline at the call site, because that
    is the seam a later image variant needs: a half-width picture layout becomes
    another zone table, not another renderer.
    """

    __slots__ = ('left', 'top', 'width', 'height')

    def __init__(self, left, top, width, height):
        self.left, self.top, self.width, self.height = int(left), int(top), int(width), int(height)

    def box(self):
        return self.left, self.top, self.width, self.height


CONTENT_WIDTH = int(SLIDE_WIDTH - 2 * MARGIN)
ZONES = {
    'eyebrow': Zone(MARGIN, Inches(0.62), CONTENT_WIDTH, Inches(0.3)),
    'headline': Zone(MARGIN, Inches(1.0), CONTENT_WIDTH, Inches(1.25)),
    'rule': Zone(MARGIN, Inches(2.38), Inches(1.1), Pt(3)),
    'body': Zone(MARGIN, Inches(2.75), CONTENT_WIDTH, Inches(3.65)),
    'body_left': Zone(MARGIN, Inches(2.75), int(CONTENT_WIDTH * 0.47), Inches(3.65)),
    'body_right': Zone(int(MARGIN + CONTENT_WIDTH * 0.53), Inches(2.75),
                       int(CONTENT_WIDTH * 0.47), Inches(3.65)),
    'column_rule': Zone(int(MARGIN + CONTENT_WIDTH * 0.5), Inches(2.85), Pt(2), Inches(3.4)),
    'cover_title': Zone(MARGIN, Inches(2.35), int(CONTENT_WIDTH * 0.78), Inches(2.1)),
    'cover_rule': Zone(MARGIN, Inches(4.65), Inches(1.6), Pt(4)),
    'cover_subtitle': Zone(MARGIN, Inches(4.95), int(CONTENT_WIDTH * 0.7), Inches(0.9)),
    'statement': Zone(MARGIN, Inches(2.5), int(CONTENT_WIDTH * 0.86), Inches(2.6)),
    'divider_band': Zone(0, Inches(2.55), SLIDE_WIDTH, Inches(2.4)),
    'divider_title': Zone(MARGIN, Inches(3.05), int(CONTENT_WIDTH * 0.8), Inches(1.4)),
    'footer': Zone(MARGIN, Inches(6.82), int(CONTENT_WIDTH * 0.7), Inches(0.32)),
    'number': Zone(int(SLIDE_WIDTH - MARGIN - Inches(1.0)), Inches(6.82), Inches(1.0), Inches(0.32)),
    'accent_bar': Zone(0, 0, Inches(0.18), SLIDE_HEIGHT),
    # Deliberately bleeds off the bottom-right corner: a disc that runs past the
    # edge reads as a design device, where one tucked inside reads as a bubble.
    # It is added before any text, so it always sits behind it.
    'cover_mark': Zone(int(SLIDE_WIDTH - Inches(3.4)), int(SLIDE_HEIGHT - Inches(3.4)),
                       Inches(4.6), Inches(4.6)),
}


def layout_for(index, bullets, total):
    """How a slide is laid out when it did not choose. See `layouts.legacy`."""
    from .layouts import legacy
    return legacy(index, bullets, total)


def _shape(slide, zone, colour, shape=MSO_SHAPE.RECTANGLE):
    element = slide.shapes.add_shape(shape, *zone.box())
    element.fill.solid()
    element.fill.fore_color.rgb = _rgb(colour)
    element.line.fill.background()
    element.shadow.inherit = False
    return element


def _slide_number(slide, zone, roles, colour=None, number=1):
    """A real slide-number field, so deleting a slide renumbers the rest.

    PowerPoint recomputes the field, but a previewer that shows the stored text
    instead (Quick Look, a chat thumbnail) would print "1" on every slide, so
    the stored text is the right number too.
    """
    colour = colour or roles['muted']
    frame = _frame(slide.shapes.add_textbox(*zone.box()))
    paragraph = frame.paragraphs[0]
    _write(paragraph, '', font=BODY_FONT, size=10, colour=colour, align=PP_ALIGN.RIGHT)
    field = paragraph._p.makeelement(qn('a:fld'), {'id': '{B7A4E0C1-4C1E-4D3E-9F2A-0A1B2C3D4E5F}',
                                                   'type': 'slidenum'})
    properties = field.makeelement(qn('a:rPr'), {'lang': 'en-US'})
    fill = properties.makeelement(qn('a:solidFill'), {})
    fill.append(fill.makeelement(qn('a:srgbClr'), {'val': _hex(colour).upper()}))
    properties.append(fill)
    latin = properties.makeelement(qn('a:latin'), {'typeface': BODY_FONT})
    properties.append(latin)
    field.append(properties)
    text = field.makeelement(qn('a:t'), {})
    text.text = str(number)
    field.append(text)
    paragraph._p.append(field)


# ---------------------------------------------------------------- slides


FILLED = ('cover', 'closing')


def _ground(slide, roles, kind, has_photo=False, mark=True):
    """The background, then the accent element — always `shapes[0]`.

    The accent bar is filled with the brand colour exactly as given, on every
    slide; tests/test_slides_deck.py pins it. A photo, if the layout has one, is
    added straight after this and so always sits beneath the chrome and text.
    A design with art draws that instead of the cover's disc (`mark=False`).
    """
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(roles['cover_fill'] if kind in FILLED else roles['surface'])
    _shape(slide, ZONES['accent_bar'], roles['accent'])
    if kind == 'cover' and not has_photo and mark:
        _shape(slide, ZONES['cover_mark'], roles['accent_soft'], MSO_SHAPE.OVAL)


# ---------------------------------------------------------------- art

# Where a design's pattern goes, in inches, and which way it fades out: always
# towards the text, so a pattern never reads as a box. Nothing here overlaps a
# text zone, the footer or the logo, so text is never set over a picture.
ART_PPI = 120
ART_ZONES = {
    ('cover', 'side'): ((10.15, 0, 3.183, 6.6), 'left'),
    ('cover', 'top'): ((0.18, 0, 13.153, 2.05), 'down'),
    ('cover', 'corner'): ((8.6, 0, 4.733, 2.2), 'corner'),
    # A divider's band ends in the pattern, past the end of its title.
    ('section', 'band'): ((10.45, 2.55, 2.883, 2.4), 'left'),
    # A strip under the footer of every other slide, for designs with `edge`.
    ('content', 'edge'): ((0.18, 7.26, 13.153, 0.24), 'ends'),
}
# A full-bleed cover photo, stopping above the footer, and the solid panel the
# title is set on. The cover drawer's photo zones sit inside the panel.
COVER_PHOTO_ZONE = (0.18, 0, 13.153, 6.55)
COVER_PANEL_ZONE = (0.18, 1.75, 7.0, 4.8)


def _art_mask(pattern):
    from pathlib import Path
    from PIL import Image
    if pattern not in _ART_MASKS:
        path = Path(__file__).resolve().parent / 'assets' / 'deck' / f'{pattern}.png'
        with Image.open(path) as mask:
            _ART_MASKS[pattern] = mask.convert('L')
    return _ART_MASKS[pattern]


_ART_MASKS = {}
_ART_PNGS = {}


def art_png(pattern, width, height, colour, strength, fade):
    """A pattern cut to a zone, tinted, faded and with its strength as opacity.

    The crop is taken from the mask's top-right corner at a fixed scale, so a
    pattern is the same size on every zone and every deck. Kept per process:
    identical pictures are also stored once in the file, which python-pptx
    deduplicates by content.
    """
    key = (pattern, width, height, colour, strength, fade)
    if key in _ART_PNGS:
        return _ART_PNGS[key]
    from PIL import Image, ImageChops
    mask = _art_mask(pattern)
    w, h = round(width * ART_PPI), round(height * ART_PPI)
    alpha = mask.crop((mask.width - w, 0, mask.width, h)).point(lambda value: round(value * strength))
    ramp = None
    if fade == 'left':
        ramp = Image.linear_gradient('L').rotate(90, expand=True)
        ramp = ramp.resize((w, h)).point(lambda value: min(255, value * 2))
    elif fade == 'down':
        ramp = Image.linear_gradient('L').transpose(Image.FLIP_TOP_BOTTOM).resize((w, h))
        ramp = ramp.point(lambda value: min(255, value * 2))
    elif fade == 'corner':
        ramp = Image.radial_gradient('L').point(lambda value: 255 - value)
        ramp = ramp.crop((0, 128, 128, 256)).resize((w, h))
        ramp = ramp.point(lambda value: min(255, value * 2))
    elif fade == 'ends':
        half = Image.linear_gradient('L').rotate(90, expand=True)
        ramp = Image.new('L', (w, h))
        ramp.paste(half.resize((w // 2, h)), (0, 0))
        ramp.paste(half.transpose(Image.FLIP_LEFT_RIGHT).resize((w - w // 2, h)), (w // 2, 0))
        ramp = ramp.point(lambda value: min(255, value * 4))
    if ramp is not None:
        alpha = ImageChops.multiply(alpha, ramp)
    picture = Image.new('RGBA', (w, h), tuple(int(_hex(colour)[i:i + 2], 16) for i in (0, 2, 4)) + (0,))
    picture.putalpha(alpha)
    buffer = io.BytesIO()
    picture.save(buffer, 'PNG', optimize=True)
    _ART_PNGS[key] = buffer.getvalue()
    return _ART_PNGS[key]


def art_colour(look, roles, filled):
    """The pattern's colour: the design's, or the ink that reads on a filled accent."""
    if filled:
        return roles['on_accent']
    if look['art'].get('colour') == 'secondary' and look.get('secondary'):
        return _hex(look['secondary']).upper()
    return roles['accent']


def _art(slide, roles, look, kind):
    """A design's pattern on a slide that has room for one, or nothing."""
    art = look.get('art')
    if not art:
        return None
    if kind in FILLED:
        place = ('cover', art['place'])
    elif kind == 'section':
        place = ('section', 'band')
    elif art.get('edge'):
        place = ('content', 'edge')
    else:
        return None
    (left, top, width, height), fade = ART_ZONES[place]
    # On a bold deck the cover and the band are the accent itself.
    filled = roles['theme'] == 'bold' and kind in FILLED + ('section',)
    # A thin pattern on a pale ground reads fainter than the same on a fill.
    strength = art['strength'] * (1.4 if roles['theme'] == 'light' and not filled else 1.0)
    strength = min(0.85, strength * (0.6 if place[0] == 'content' else 1.0))
    data = art_png(art['pattern'], width, height, art_colour(look, roles, filled), strength, fade)
    return slide.shapes.add_picture(io.BytesIO(data), Inches(left), Inches(top), Inches(width), Inches(height))


def _chrome(slide, roles, style, kind, number=1, quiet=None):
    """Brand name, slide number and logo — the same on every slide."""
    quiet = roles.get(quiet, quiet) or (roles['cover_muted'] if kind in FILLED else roles['muted'])
    brand = (style or {}).get('brand_name', '')
    if brand:
        frame = _frame(slide.shapes.add_textbox(*ZONES['footer'].box()))
        _write(frame.paragraphs[0], brand, font=BODY_FONT, size=10, colour=quiet)
        _no_bullet(frame.paragraphs[0])
    if kind != 'cover':
        _slide_number(slide, ZONES['number'], roles, quiet, number)
    logo = (style or {}).get('logo_png')
    if logo:
        from PIL import Image
        with Image.open(io.BytesIO(logo)) as picture:
            scale = min(1.5 / picture.width, 0.48 / picture.height)
            width, height = picture.width * scale, picture.height * scale
        slide.shapes.add_picture(io.BytesIO(logo), Inches(12.65 - width), Inches(7.15 - height),
                                 width=Inches(width), height=Inches(height))


def _headline(slide, zone, text, roles, *, size, colour=None, anchor=MSO_ANCHOR.BOTTOM, lines=2,
              align=PP_ALIGN.LEFT, bold=False):
    """A headline that shrinks to fit rather than being renamed "Section"."""
    while size > 18 and _wrapped_lines(text, size, zone.width) > lines:
        size -= 2
    frame = _frame(slide.shapes.add_textbox(*zone.box()), anchor)
    paragraph = _write(frame.paragraphs[0], text, font=HEADING_FONT, size=size, bold=bold,
                       colour=colour or roles.get('heading', roles['ink']), spacing=1.08, align=align)
    _no_bullet(paragraph)
    return paragraph


def _bullet_block(slide, zone, bullets, roles, size):
    frame = _frame(slide.shapes.add_textbox(*zone.box()))
    for index, bullet in enumerate(bullets):
        paragraph = _write(_paragraph(frame, index), bullet, font=BODY_FONT, size=size,
                           colour=roles['ink'], spacing=LINE_MULTIPLE,
                           before=0 if index == 0 else size * 0.55)
        # Exact leading, as `_fits` assumes. A relative 122% is 122% of the
        # font's own line height — about a fifth taller than the estimate —
        # which a full list of sentence-long bullets turns into an overflow.
        paragraph.line_spacing = Pt(size * LINE_MULTIPLE)
        _bullet(paragraph, roles['accent_text'], size)
    return frame


def _eyebrow(slide, text, roles):
    frame = _frame(slide.shapes.add_textbox(*ZONES['eyebrow'].box()))
    _no_bullet(_write(frame.paragraphs[0], text, font=BODY_FONT, size=10,
                      colour=roles['muted'], bold=True))


def _fits(bullets, size, zone):
    """Roughly how tall a bullet block runs, without wrapping anything for real."""
    height = 0
    for index, bullet in enumerate(bullets):
        lines = _wrapped_lines(bullet, size, zone.width - Pt(size))
        height += lines * size * LINE_MULTIPLE + (0 if index == 0 else size * 0.55)
    return height <= Emu(zone.height).pt


# The layouts whose body is a list at the deck-wide size, and the zone it fills.
# Every other layout sets its own type, so it does not pull the deck's size down.
LIST_ZONES = {'bullets': 'body', 'two_column': 'body_left'}


def _deck_size(prepared):
    """One body size for the whole deck, so slides do not change size partway."""
    from .slide_layouts import Z, LEFT, TOP, BOTTOM
    zones = {**{kind: ZONES[name] for kind, name in LIST_ZONES.items()},
             'image_split': Z(LEFT, TOP, 13.333 * 0.52 - LEFT - 0.45, BOTTOM - TOP)}
    lists = [(zones[kind], lines) for kind, lines, section in prepared
             if kind in zones and not _titled_columns(kind, section)]
    for step in range(FIT_STEPS + 1):
        size = MAX_BODY_SIZE - (MAX_BODY_SIZE - MIN_BODY_SIZE) * step / FIT_STEPS
        if all(_fits(lines, size, zone) for zone, lines in lists):
            return size
    return MIN_BODY_SIZE


def _titled_columns(kind, section):
    titled = [item for item in section.get('items') or [] if item.get('label') and item.get('text')]
    return kind == 'two_column' and section.get('layout') == 'two_column' and len(titled) == 2


# Layouts that draw the section's body as lines, so a cut there is a real cut.
LINE_LAYOUTS = {'cover', 'section', 'bullets', 'two_column', 'statement', 'image_split',
                'image_full', 'closing'}


QUESTIONS_LABEL = {'en': 'Questions', 'uz': 'Savollar', 'ru': 'Вопросы', 'uz_cyrl': 'Саволлар'}
ANSWERS_LABEL = {'en': 'Answers', 'uz': 'Javoblar', 'ru': 'Ответы', 'uz_cyrl': 'Жавоблар'}
LETTERED = re.compile(r'^\s*\(?[A-Za-zА-Яа-я][).]\s')


def question_sections(content, locale='en', script='latn'):
    """The deck's questions as slides of their own, five to a slide.

    Each question is numbered, with its options on one line beneath it. The
    answers go in the speaker notes, where the presenter sees them and the
    audience does not. One slide per question made decks of mostly empty
    slides, and pushed "10 slides, 5 questions at the end" to fifteen.
    """
    from .pages import QUESTIONS_PER_SLIDE
    questions = content.get('questions') or []
    groups = [questions[start:start + QUESTIONS_PER_SLIDE]
              for start in range(0, len(questions), QUESTIONS_PER_SLIDE)]
    if locale == 'uz' and script == 'cyrl':
        locale = 'uz_cyrl'
    label = QUESTIONS_LABEL.get(locale, QUESTIONS_LABEL['en'])
    sections = []
    for group_index, group in enumerate(groups):
        entries, answers = [], []
        for offset, question in enumerate(group):
            number = group_index * QUESTIONS_PER_SLIDE + offset + 1
            options = [option if LETTERED.match(option) else f'{chr(65 + position)}) {option}'
                       for position, option in enumerate(question.get('options') or [])]
            entries.append({'number': number, 'stem': question['stem'], 'options': options})
            if question.get('answer'):
                answers.append(f'{number}. {question["answer"]}')
        heading = label if len(groups) == 1 else f'{label} ({group_index + 1}/{len(groups)})'
        notes = (ANSWERS_LABEL.get(locale, ANSWERS_LABEL['en']) + ': ' + '; '.join(answers)) if answers else ''
        sections.append({'id': f'q{group_index + 1}', 'heading': heading, 'body': '', 'notes': notes,
                         '_questions': entries})
    return sections


def _within(zone, box):
    """`zone` (EMU) cut down to `box` (inches): a photo kept inside a framing card."""
    left, top = max(zone.left, Inches(box[0])), max(zone.top, Inches(box[1]))
    right = min(zone.left + zone.width, Inches(box[0] + box[2]))
    bottom = min(zone.top + zone.height, Inches(box[1] + box[3]))
    return Zone(left, top, right - left, bottom - top)


def render_pptx(content, path, locale='en', role='user_document', style=None, photos=None):
    """One section, one slide — with the first one as the deck's title slide.

    The cover is `sections[0]` rather than an extra slide, so the count the
    customer asked for is the count delivered and the price is unchanged. With a
    single section there is nothing to introduce, so it stays a content slide.

    Each slide is drawn with the layout it chose, if its content can fill it,
    and otherwise with that layout's fallback. `photos` maps a section id to an
    already-fetched, already-validated picture; this function never touches the
    network, so a deck renders the same with or without one.
    """
    from . import layouts
    from .slide_layouts import DRAW, PHOTO_ZONES, Ctx, place_photo
    style = style or {}
    photos = photos or {}
    from .deck_designs import look as design_look
    look = design_look(style)
    roles = palette(look['accent'], look['theme'], look['secondary'], look.get('paper'),
                    look.get('accent_headings', False))
    sections = [dict(section) for section in content['sections']] + question_sections(
        content, locale, style.get('uz_script', 'latn'))

    total = len(sections)
    shortened = False
    prepared = []
    for index, section in enumerate(sections):
        if section.get('_questions'):
            prepared.append(('questions', [], section))
            continue
        lines, cut = bullets_of(section['body'])
        if not lines and section.get('items'):
            lines, cut = bullets_of('\n'.join(layouts.as_lines(section)))
        has_photo = section.get('id') in photos
        kind = layouts.resolve(section, lines, index, total, has_photo)
        # Titled columns draw their items, not these lines: a cut here is no cut.
        if kind in LINE_LAYOUTS and not _titled_columns(kind, section):
            shortened = shortened or cut
        if kind == 'cover':
            section = {**section, '_title': content['title']}
            if not section['heading'].strip():
                section['heading'] = content['title']
            lines = lines[:1]
        prepared.append((kind, lines, section))

    deck = Presentation()
    deck.slide_width, deck.slide_height = SLIDE_WIDTH, SLIDE_HEIGHT
    apply_theme(deck, roles, (look['heading_font'], look['body_font']))
    size = _deck_size(prepared)
    drawn, pictured = [], []

    from .compositions import Canvas, compose, zone as inch_zone
    arrangement = compose(look.get('composition', 'classic'))
    classic = arrangement.name == 'classic'
    for index, (kind, lines, section) in enumerate(prepared):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        photo = photos.get(section.get('id')) if kind in PHOTO_ZONES else None
        # The composition draws everything under the text and, for a cover,
        # closing slide or divider, says where the text goes.
        role_of = kind if kind in ('cover', 'closing', 'section') else 'content'
        plan = arrangement.ground(Canvas(slide, roles, look, role_of), role_of, photo is not None)
        if photo is not None:
            try:
                if plan is not None and plan.photo:
                    place_photo(slide, inch_zone(plan.photo), photo)
                elif not classic:
                    place_photo(slide, _within(PHOTO_ZONES[kind], arrangement.content_box), photo)
                elif kind == 'cover' and look['cover_photo']:
                    # A photo design's cover: the picture across the slide and
                    # the title on a solid panel, never on the picture itself.
                    place_photo(slide, Zone(*(Inches(value) for value in COVER_PHOTO_ZONE)), photo)
                    _shape(slide, Zone(*(Inches(value) for value in COVER_PANEL_ZONE)), roles['cover_fill'])
                else:
                    place_photo(slide, PHOTO_ZONES[kind], photo)
                pictured.append({'slide': index + 1, **{key: photo[key] for key in photo if key != 'jpeg'}})
            except Exception:
                # A picture python-pptx cannot place leaves the slide to its text
                # layout rather than costing the customer the deck.
                photo = None
                kind = layouts.resolve(section, lines, index, total, False)
        _chrome(slide, roles, style, kind, index + 1, quiet=plan.quiet if plan is not None else None)
        context = Ctx(slide, roles, style, index, total, size, locale, photo is not None, plan)
        DRAW[kind](context, section, lines)
        # A classic design's pattern goes in its fixed places; a composition
        # puts its own pattern where it has room.
        if classic and not (kind == 'cover' and photo is not None):
            _art(slide, roles, look, kind)
        shortened = shortened or context.shortened
        drawn.append(kind)
        # The cover introduces the deck; there is nothing for a presenter to say
        # over it. Learner and public copies never carry notes at all.
        notes = '' if kind == 'cover' or role in ('learner_material', 'public_preview') else section.get('notes', '')
        slide.notes_slide.notes_text_frame.text = notes

    apply_fonts(deck, look)
    deck.core_properties.title = content['title']
    deck.core_properties.author = 'PDF Master'
    strip_binary_parts(deck)
    deck.save(str(path))
    return {'path': str(path), 'name': path.name,
            'mime_type': 'application/vnd.openxmlformats-officedocument.presentationml.presentation',
            'page_count': len(deck.slides), 'role': role, 'shortened': shortened,
            'metadata': {'layouts': drawn, 'photos': pictured, 'design': look['id']}}
