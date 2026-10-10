"""Draw a rendered deck's slides as images, from the shapes in the file.

No PowerPoint or LibreOffice is needed: this reads the same shapes check.py
reads — solid, transparent and gradient fills, rectangles, rounded panels,
ovals, polygons, lines, pictures and wrapped text — and paints them with
Pillow. Fonts are approximated with Noto Sans, so a preview shows the layout,
colours and proportions faithfully and the typefaces only roughly. It is what
the theme picker's thumbnails are made from, and how a composition is reviewed.
"""
import io
import math
from functools import lru_cache
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.oxml.ns import qn

from .check import EMU, SLIDE_H, SLIDE_W, _background, _describe

FONT = Path(__file__).resolve().parents[3] / 'processors' / 'assets' / 'fonts' / 'NotoSans-Regular.ttf'


@lru_cache(maxsize=64)
def _font(size):
    return ImageFont.truetype(str(FONT), max(6, int(round(size))))


def _hex(value):
    return tuple(int(value[i:i + 2], 16) for i in (0, 2, 4))


def slide_images(deck, width=960):
    """Every slide of `deck` (path or Presentation) as an RGB image `width` pixels wide."""
    presentation = deck if hasattr(deck, 'slides') else Presentation(str(deck))
    return [render(slide, width) for slide in presentation.slides]


def render(slide, width=960):
    scale = width / SLIDE_W
    height = round(SLIDE_H * scale)
    image = Image.new('RGBA', (width, height), _hex(_background(slide)[1]) + (255,))
    for shape, described in zip(slide.shapes, (_describe(shape) for shape in slide.shapes)):
        layer = Image.new('RGBA', image.size, (0, 0, 0, 0))
        draw = ImageDraw.Draw(layer)
        box = [round(value * scale) for value in described['box']]
        x0, y0, x1, y1 = box[0], box[1], box[0] + box[2], box[1] + box[3]
        if described['kind'] == 'picture':
            _picture(layer, shape, (x0, y0, x1, y1))
        elif described['kind'] == 'line':
            (ax, ay), (bx, by) = described['segment']
            outline = _outline(shape)
            if outline:
                draw.line([(ax * scale, ay * scale), (bx * scale, by * scale)], fill=_hex(outline[0]) + (255,),
                          width=max(1, round(outline[1] / 72 * scale)))
        else:
            mask = Image.new('L', image.size, 0)
            _geometry_mask(ImageDraw.Draw(mask), described, scale, (x0, y0, x1, y1), shape)
            fill = described['fill']
            if fill:
                kind, value, alpha = fill
                if kind == 'gradient':
                    paint = _gradient(image.size, (x0, y0, x1, y1), value, _angle(shape))
                else:
                    paint = Image.new('RGBA', image.size, _hex(value) + (255,))
                layer.paste(paint, (0, 0), mask.point(lambda v: round(v * alpha)))
            outline = _outline(shape)
            if outline:
                _geometry_mask(draw, described, scale, (x0, y0, x1, y1), shape,
                               outline=_hex(outline[0]) + (255,), stroke=max(1, round(outline[1] / 72 * scale)))
            if described['text']:
                _text(draw, shape, (x0, y0, x1, y1), scale)
        image = Image.alpha_composite(image, layer)
    return image.convert('RGB')


def _geometry_mask(draw, described, scale, box, shape, outline=None, stroke=0):
    x0, y0, x1, y1 = box
    fill = None if outline else 255
    if described['geometry'] == 'oval':
        draw.ellipse(box, fill=fill, outline=outline, width=stroke)
    elif described['geometry'] == 'polygon' and described.get('points'):
        points = [(x * scale, y * scale) for x, y in described['points']]
        if outline:
            draw.line(points + points[:1], fill=outline, width=stroke)
        else:
            draw.polygon(points, fill=fill)
    else:
        radius = _radius(shape, box)
        if radius:
            draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=stroke)
        else:
            draw.rectangle(box, fill=fill, outline=outline, width=stroke)


def _radius(shape, box):
    properties = shape._element.find(qn('p:spPr'))
    preset = properties.find(qn('a:prstGeom')) if properties is not None else None
    if preset is None or preset.get('prst') != 'roundRect':
        return 0
    guide = preset.find('.//' + qn('a:gd'))
    share = int(guide.get('fmla').split()[-1]) / 100000 if guide is not None else 0.16667
    return round(share * min(box[2] - box[0], box[3] - box[1]))


def _outline(shape):
    properties = shape._element.find(qn('p:spPr'))
    line = properties.find(qn('a:ln')) if properties is not None else None
    if line is None or line.find(qn('a:noFill')) is not None:
        return None
    colour = line.find('.//' + qn('a:srgbClr'))
    if colour is None:
        return None
    return colour.get('val').upper(), int(line.get('w', 12700)) / 12700


def _angle(shape):
    properties = shape._element.find(qn('p:spPr'))
    linear = properties.find('.//' + qn('a:lin')) if properties is not None else None
    return int(linear.get('ang', 0)) / 60000 if linear is not None else 0


def _gradient(size, box, stops, angle):
    """A linear gradient over `box`, DrawingML angle: clockwise from left-to-right."""
    import numpy
    x0, y0, x1, y1 = box
    width, height = max(1, x1 - x0), max(1, y1 - y0)
    start, end = numpy.array(_hex(stops[0]), float), numpy.array(_hex(stops[-1]), float)
    radians = math.radians(angle)
    dx, dy = math.cos(radians), math.sin(radians)
    span = abs(width * dx) + abs(height * dy) or 1
    xs, ys = numpy.meshgrid(numpy.arange(width) - width / 2, numpy.arange(height) - height / 2)
    t = numpy.clip((xs * dx + ys * dy) / span + 0.5, 0, 1)[..., None]
    pixels = (start + (end - start) * t).round().astype(numpy.uint8)
    alpha = numpy.full((height, width, 1), 255, numpy.uint8)
    tile = Image.fromarray(numpy.concatenate([pixels, alpha], axis=2), 'RGBA')
    canvas = Image.new('RGBA', size, (0, 0, 0, 0))
    canvas.paste(tile, (x0, y0))
    return canvas


def _picture(layer, shape, box):
    x0, y0, x1, y1 = box
    try:
        picture = Image.open(io.BytesIO(shape.image.blob)).convert('RGBA')
    except Exception:
        return
    width, height = picture.size
    left, top = shape.crop_left, shape.crop_top
    right, bottom = shape.crop_right, shape.crop_bottom
    picture = picture.crop((round(width * left), round(height * top),
                            round(width * (1 - right)), round(height * (1 - bottom))))
    picture = picture.resize((max(1, x1 - x0), max(1, y1 - y0)))
    layer.alpha_composite(picture, (max(0, x0), max(0, y0)))


def _text(draw, shape, box, scale):
    x0, y0, x1, y1 = box
    frame = shape.text_frame
    anchor = frame._txBody.find(qn('a:bodyPr')).get('anchor', 't')
    blocks = []
    for paragraph in frame.paragraphs:
        text = ''.join(node.text or '' for node in paragraph._p.iter(qn('a:t')))
        if not text.strip():
            continue
        size, colour, bold = 18.0, '000000', False
        for run in paragraph._p.iter(qn('a:rPr')):
            if run.get('sz'):
                size = int(run.get('sz')) / 100
            bold = bold or run.get('b') == '1'
            solid = run.find(qn('a:solidFill'))
            if solid is not None and solid.find(qn('a:srgbClr')) is not None:
                colour = solid.find(qn('a:srgbClr')).get('val')
        align = (paragraph._p.pPr.get('algn') if paragraph._p.pPr is not None else None) or 'l'
        pixels = size / 72 * scale
        font = _font(pixels)
        lines = _wrap(text, font, (x1 - x0) - 0.1 * scale)
        blocks.append((lines, font, pixels * 1.15, _hex(colour), align, bold))
    total = sum(len(lines) * leading for lines, _, leading, *_ in blocks)
    y = y0 + 0.05 * scale
    if anchor == 'b':
        y = y1 - total - 0.05 * scale
    elif anchor == 'ctr':
        y = y0 + ((y1 - y0) - total) / 2
    for lines, font, leading, colour, align, bold in blocks:
        for line in lines:
            length = font.getlength(line)
            x = x0 + 0.05 * scale
            if align == 'ctr':
                x = x0 + ((x1 - x0) - length) / 2
            elif align == 'r':
                x = x1 - length - 0.05 * scale
            draw.text((x, y), line, font=font, fill=colour + (255,), stroke_width=1 if bold else 0,
                      stroke_fill=colour + (255,))
            y += leading


def _wrap(text, font, width):
    lines = []
    for raw in text.split('\n'):
        words, line = raw.split(), ''
        for word in words:
            candidate = f'{line} {word}'.strip()
            if font.getlength(candidate) <= width or not line:
                line = candidate
            else:
                lines.append(line)
                line = word
        lines.append(line)
    return lines


def contact_sheet(images, columns, gap=12, background=(228, 228, 228), labels=None):
    """Images in a grid, optionally labelled underneath, as one image."""
    width, height = images[0].size
    label_height = 22 if labels else 0
    rows = math.ceil(len(images) / columns)
    sheet = Image.new('RGB', (columns * width + (columns + 1) * gap,
                              rows * (height + label_height) + (rows + 1) * gap), background)
    draw = ImageDraw.Draw(sheet)
    for index, image in enumerate(images):
        x = gap + (index % columns) * (width + gap)
        y = gap + (index // columns) * (height + label_height + gap)
        sheet.paste(image, (x, y))
        if labels:
            draw.text((x, y + height + 3), labels[index], fill=(30, 30, 30), font=_font(15))
    return sheet
