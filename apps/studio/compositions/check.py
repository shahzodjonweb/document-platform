"""Read a rendered deck back and say where text is hard to read or misplaced.

Used by tests/test_compositions.py on every composition, and by anyone drawing
a new one. For each piece of text it samples the box the text lives in (all of
it: someone editing the deck can type into the whole box) against everything
drawn beneath it, finds the colour actually under each point — blending any
transparent fill with what is under that, and taking both ends of a gradient —
and reports:

* text that does not reach its contrast (4.5:1, or 3:1 at 24 pt and above);
* text over a picture or across a line;
* two text boxes overlapping, or a box running off the slide;
* on covers, closing slides and dividers, a single-paragraph text that would
  not fit its box at its size (measured as slides.py measures).
"""
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from ..slides import _wrapped_lines, contrast_ratio

EMU = 914400
SLIDE_W, SLIDE_H = 13.333, 7.5


def problems(deck, kinds=None, accent=None):
    """Every problem in `deck` (a path or a Presentation), as 'slide N: …' lines.

    `kinds` is the layout of each slide (render_pptx's metadata['layouts']);
    the fit check runs on covers, closing slides and dividers. `accent`, when
    given, is checked as the fill of every slide's first shape.
    """
    presentation = deck if hasattr(deck, 'slides') else Presentation(str(deck))
    found = []
    for number, slide in enumerate(presentation.slides, 1):
        kind = kinds[number - 1] if kinds else ''
        found += [f'slide {number} ({kind}): {problem}' for problem in _slide(slide, kind, accent)]
    return found


def _slide(slide, kind, accent):
    shapes = [_describe(shape) for shape in slide.shapes]
    background = _background(slide)
    if accent and (not shapes or shapes[0]['fill'] != ('solid', accent.lstrip('#').upper(), 1.0)):
        yield f'the first shape is not the accent element filled {accent}'
    texts = [(index, shape) for index, shape in enumerate(shapes) if shape['text']]
    for index, text in texts:
        label = repr(text['text'][:40])
        left, top, width, height = text['box']
        if left < -0.01 or top < -0.01 or left + width > SLIDE_W + 0.01 or top + height > SLIDE_H + 0.01:
            yield f'{label} runs off the slide'
        target = 3.0 if text['size'] >= 24 else 4.5
        worst, reason = None, ''
        for point in _samples(text['box']):
            for colour in _colours_at(point, shapes[:index], background):
                if colour == 'picture':
                    reason = 'sits on a picture'
                    break
                ratio = contrast_ratio(text['colour'], colour)
                if worst is None or ratio < worst[0]:
                    worst = (ratio, colour)
            if reason:
                break
        if reason:
            yield f'{label} {reason}'
        elif worst and worst[0] < target - 0.05:
            yield f'{label} in #{text["colour"]} reads {worst[0]:.2f}:1 on #{worst[1]} (needs {target}:1)'
        for below in shapes[:index]:
            if below['kind'] == 'line' and _segment_hits(below['segment'], _inset(text['box'], 0.02)):
                yield f'{label} has a line across it'
        if kind in ('cover', 'closing', 'section') and text['paragraphs'] == 1 and text['size']:
            lines = _wrapped_lines(text['text'], text['size'], int(width * EMU))
            needed = lines * text['size'] * text['spacing'] / 72
            if needed > height + 0.04:
                yield f'{label} needs {needed:.2f} in at {text["size"]} pt; its box is {height:.2f} in'
    for position, (_, one) in enumerate(texts):
        for _, other in texts[position + 1:]:
            if _overlap(_inset(one['box'], 0.02), _inset(other['box'], 0.02)):
                yield f'{one["text"][:30]!r} overlaps {other["text"][:30]!r}'


def _describe(shape):
    box = tuple(value / EMU for value in (shape.left or 0, shape.top or 0, shape.width or 0, shape.height or 0))
    found = {'box': box, 'kind': 'shape', 'fill': None, 'text': '', 'geometry': 'rect'}
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        found['kind'] = 'picture'
        return found
    if shape.shape_type == MSO_SHAPE_TYPE.LINE or shape._element.tag == qn('p:cxnSp'):
        found['kind'] = 'line'
        flip = shape._element.spPr.find(qn('a:xfrm'))
        x1, y1, x2, y2 = box[0], box[1], box[0] + box[2], box[1] + box[3]
        if flip is not None and flip.get('flipV') == '1':
            y1, y2 = y2, y1
        if flip is not None and flip.get('flipH') == '1':
            x1, x2 = x2, x1
        found['segment'] = ((x1, y1), (x2, y2))
        return found
    found['fill'] = _fill(shape)
    found['geometry'], found['points'] = _geometry(shape, box)
    if getattr(shape, 'has_text_frame', False) and shape.has_text_frame and shape.text_frame.text.strip():
        colours, sizes = [], []
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph._p.iter(qn('a:rPr')):
                fill = run.find(qn('a:solidFill'))
                colour = fill.find(qn('a:srgbClr')) if fill is not None else None
                if colour is not None:
                    colours.append(colour.get('val').upper())
                if run.get('sz'):
                    sizes.append(int(run.get('sz')) / 100)
        found['text'] = shape.text_frame.text.strip()
        found['colour'] = colours[0] if colours else '000000'
        found['size'] = max(sizes) if sizes else 0
        found['paragraphs'] = len([p for p in shape.text_frame.paragraphs if p.text.strip()])
        spacing = shape.text_frame.paragraphs[0].line_spacing
        found['spacing'] = (spacing.pt / found['size'] if hasattr(spacing, 'pt') and found['size']
                            else float(spacing) if isinstance(spacing, float) else 1.15)
    return found


def _fill(shape):
    properties = shape._element.find(qn('p:spPr'))
    if properties is None:
        return None
    solid = properties.find(qn('a:solidFill'))
    if solid is not None and solid.find(qn('a:srgbClr')) is not None:
        colour = solid.find(qn('a:srgbClr'))
        alpha = colour.find(qn('a:alpha'))
        return ('solid', colour.get('val').upper(), int(alpha.get('val')) / 100000 if alpha is not None else 1.0)
    gradient = properties.find(qn('a:gradFill'))
    if gradient is not None:
        stops = [stop.find(qn('a:srgbClr')).get('val').upper() for stop in gradient.iter(qn('a:gs'))
                 if stop.find(qn('a:srgbClr')) is not None]
        return ('gradient', tuple(stops), 1.0)
    return None


def _geometry(shape, box):
    """'rect', 'oval' or 'polygon' (with its points in inches) for hit tests."""
    properties = shape._element.find(qn('p:spPr'))
    preset = properties.find(qn('a:prstGeom')) if properties is not None else None
    if preset is not None:
        name = preset.get('prst')
        if name == 'ellipse':
            return 'oval', None
        if name == 'rtTriangle':
            left, top, width, height = box
            return 'polygon', [(left, top), (left, top + height), (left + width, top + height)]
        return 'rect', None
    custom = properties.find(qn('a:custGeom')) if properties is not None else None
    if custom is not None:
        path = custom.find(qn('a:pathLst')).find(qn('a:path'))
        scale_x = box[2] / max(1, int(path.get('w', 1))) * 914400 / EMU
        scale_y = box[3] / max(1, int(path.get('h', 1))) * 914400 / EMU
        points = []
        for step in path:
            point = step.find(qn('a:pt'))
            if point is not None:
                points.append((box[0] + int(point.get('x')) * scale_x / 914400 * EMU,
                               box[1] + int(point.get('y')) * scale_y / 914400 * EMU))
        return 'polygon', points
    return 'rect', None


def _background(slide):
    fill = slide.background._cSld.bg
    if fill is None:
        return ('solid', 'FFFFFF', 1.0)
    solid = fill.find('.//' + qn('a:solidFill'))
    if solid is not None and solid.find(qn('a:srgbClr')) is not None:
        return ('solid', solid.find(qn('a:srgbClr')).get('val').upper(), 1.0)
    return ('solid', 'FFFFFF', 1.0)


def _samples(box, columns=9, rows=5):
    left, top, width, height = _inset(box, 0.03)
    return [(left + width * column / (columns - 1), top + height * row / (rows - 1))
            for column in range(columns) for row in range(rows)]


def _colours_at(point, shapes, background):
    """Every colour that can be under `point`: one, or both ends of a gradient; 'picture' for a picture."""
    for depth in range(len(shapes) - 1, -1, -1):
        shape = shapes[depth]
        if shape['kind'] == 'line' or shape['text'] or not _contains(shape, point):
            continue
        if shape['kind'] == 'picture':
            return ['picture']
        fill = shape['fill']
        if fill is None:
            continue
        kind, value, alpha = fill
        colours = list(value) if kind == 'gradient' else [value]
        if alpha >= 0.999:
            return colours
        below = _colours_at(point, shapes[:depth], background)
        if 'picture' in below:
            return ['picture']
        return [_blend(colour, under, alpha) for colour in colours for under in below]
    kind, value, _ = background
    return [value]


def _contains(shape, point):
    x, y = point
    left, top, width, height = shape['box']
    if not (left <= x <= left + width and top <= y <= top + height):
        return False
    if shape['geometry'] == 'oval':
        rx, ry = width / 2, height / 2
        return ((x - left - rx) / max(rx, 1e-6)) ** 2 + ((y - top - ry) / max(ry, 1e-6)) ** 2 <= 1
    if shape['geometry'] == 'polygon' and shape.get('points'):
        return _inside_polygon(point, shape['points'])
    return True


def _inside_polygon(point, points):
    x, y = point
    inside = False
    for (x1, y1), (x2, y2) in zip(points, points[1:] + points[:1]):
        if (y1 > y) != (y2 > y) and x < (x2 - x1) * (y - y1) / ((y2 - y1) or 1e-9) + x1:
            inside = not inside
    return inside


def _blend(top, bottom, alpha):
    channels = [round(int(top[i:i + 2], 16) * alpha + int(bottom[i:i + 2], 16) * (1 - alpha)) for i in (0, 2, 4)]
    return '%02X%02X%02X' % tuple(channels)


def _inset(box, margin):
    left, top, width, height = box
    return (left + margin, top + margin, max(0.0, width - 2 * margin), max(0.0, height - 2 * margin))


def _overlap(one, other):
    return (one[0] < other[0] + other[2] and other[0] < one[0] + one[2]
            and one[1] < other[1] + other[3] and other[1] < one[1] + one[3])


def _segment_hits(segment, box):
    """Whether a line segment passes through a box (sampled along its length)."""
    (x1, y1), (x2, y2) = segment
    for step in range(41):
        t = step / 40
        x, y = x1 + (x2 - x1) * t, y1 + (y2 - y1) * t
        if box[0] <= x <= box[0] + box[2] and box[1] <= y <= box[1] + box[3]:
            return True
    return False
