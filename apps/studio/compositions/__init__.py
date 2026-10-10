"""Compositions: how a deck design arranges its slides, not only how it colours them.

A design used to be a palette and a pair of fonts on one fixed arrangement —
an accent bar down the left, the title low on the left of the cover, a band
across every divider — so every deck read as the same template in another
colour. A composition is the arrangement: what the cover, the dividers and the
closing slide look like and where their text goes, and how every content slide
is framed. A design (apps/studio/deck_designs.py) is a palette, fonts and a
composition.

A composition draws with a `Canvas` and returns a `Plan` of where each piece of
text goes, at what size and in which colour. Two rules keep any composition
safe, and tests/test_compositions.py checks them on every slide of every
composition in every theme:

* **Text sits on a known fill.** A composition sets each text colour from the
  role that reads on the fill under it (`ink` on `surface`, `on_accent` on
  `accent`, `card_ink` on `card` ...): the palette has corrected each pair.
  Nothing but that fill is under the text — no picture, no line, no edge.
* **Content slides keep their text zones.** The eyebrow, headline, body and
  footer of the 20-odd content layouts are fitted and tested where they are
  (apps/studio/slide_layouts.py). A composition frames them — panels behind
  the whole content area, bars, rules, marks and pattern outside it — and
  never moves them. Covers, dividers and closing slides are its to arrange.

The first shape on every slide is the accent element, filled with the accent
exactly as given (tests pin it): `Canvas.accent` must be the first thing drawn.
"""
from dataclasses import dataclass, field

from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SLIDE_W, SLIDE_H = 13.333, 7.5
# Where every content layout writes, in inches: eyebrow at 0.62, headline from
# 1.0, body to 6.4, footer and slide number at 6.82-7.14, between x 0.85 and
# 12.48. A content frame keeps this box on the surface colour.
CONTENT_TEXT = (0.85, 0.62, 11.633, 6.52)
# The footer (brand name) and slide-number boxes every slide may carry.
FOOTER = (0.85, 6.82, 8.143, 0.32)
NUMBER = (11.483, 6.82, 1.0, 0.32)

COMPOSITIONS = {}


def composition(cls):
    """Register a composition class under its `name`."""
    COMPOSITIONS[cls.name] = cls()
    return cls


@dataclass
class Text:
    """One piece of text a plan places: a box in inches, a colour, and how it sits."""
    box: tuple
    colour: str
    size: int
    align: str = 'left'        # left, center, right
    anchor: str = 'bottom'     # top, middle, bottom
    lines: int = 3
    upper: bool = False
    bold: bool = False
    numeral: bool = False      # a divider eyebrow written as "02" rather than "02 / 07"


@dataclass
class Plan:
    """Where a cover, closing slide or divider writes; None leaves a piece out."""
    title: Text = None
    subtitle: Text = None
    rule: tuple = None              # (box, colour)
    eyebrow: Text = None            # a divider's "02 / 07"
    kicker: Text = None             # a divider's one-line kicker
    title_with_kicker: Text = None  # the divider title when it has a kicker
    quiet: str = None               # footer and slide-number colour here
    photo: tuple = None             # where a photo goes on this slide, if it has one


@dataclass
class Canvas:
    """Drawing on one slide, in inches, with colours named by role or given as hex."""
    slide: object
    roles: dict
    look: dict
    kind: str
    drawn: list = field(default_factory=list)

    def colour(self, value):
        value = self.roles.get(value, value)
        return value.lstrip('#').upper()

    def background(self, value):
        self.slide.background.fill.solid()
        self.slide.background.fill.fore_color.rgb = _rgb(self.colour(value))

    def accent(self, box, shape='rect', **options):
        """The accent element: the first shape on the slide, in the accent itself."""
        if len(self.slide.shapes):
            raise RuntimeError('The accent element must be the first shape on the slide.')
        return self.shape(shape, box, 'accent', **options)

    def rect(self, box, fill, **options):
        return self.shape('rect', box, fill, **options)

    def rounded(self, box, fill, radius=0.12, **options):
        return self.shape('rounded', box, fill, radius=radius, **options)

    def oval(self, box, fill, **options):
        return self.shape('oval', box, fill, **options)

    def shape(self, kind, box, fill, *, radius=0.12, line=None, line_width=1.0, alpha=None,
              gradient=None, angle=90):
        """A filled shape. `fill` None leaves it unfilled (an outline needs `line`)."""
        geometry = {'rect': MSO_SHAPE.RECTANGLE, 'rounded': MSO_SHAPE.ROUNDED_RECTANGLE,
                    'oval': MSO_SHAPE.OVAL, 'triangle': MSO_SHAPE.RIGHT_TRIANGLE}[kind]
        element = self.slide.shapes.add_shape(geometry, *_emu(box))
        if kind == 'rounded':
            # The radius in inches, as a share of the shorter side.
            element.adjustments[0] = min(0.5, radius / max(0.01, min(box[2], box[3])))
        self._paint(element, fill, line, line_width, alpha, gradient, angle)
        self.drawn.append(element)
        return element

    def polygon(self, points, fill, *, line=None, line_width=1.0, alpha=None, gradient=None, angle=90):
        """A filled polygon through `points` in inches."""
        scale = 914400
        builder = self.slide.shapes.build_freeform(round(points[0][0] * scale), round(points[0][1] * scale), scale=1)
        builder.add_line_segments([(round(x * scale), round(y * scale)) for x, y in points[1:]], close=True)
        element = builder.convert_to_shape()
        self._paint(element, fill, line, line_width, alpha, gradient, angle)
        self.drawn.append(element)
        return element

    def line(self, start, end, colour, width=1.0):
        """A straight line between two points in inches, `width` in points."""
        element = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(start[0]), Inches(start[1]),
                                                  Inches(end[0]), Inches(end[1]))
        element.line.color.rgb = _rgb(self.colour(colour))
        element.line.width = Pt(width)
        self.drawn.append(element)
        return element

    def art(self, box, fade='left', strength=None, colour=None):
        """The design's pattern in `box`, tinted, if it has one; nothing otherwise."""
        art = self.look.get('art')
        if not art:
            return None
        import io
        from ..slides import art_png
        data = art_png(art['pattern'], box[2], box[3], self.colour(colour or 'accent'),
                       min(0.85, strength if strength is not None else art['strength']), fade)
        element = self.slide.shapes.add_picture(io.BytesIO(data), *_emu(box))
        self.drawn.append(element)
        return element

    def _paint(self, element, fill, line, line_width, alpha, gradient, angle):
        if gradient:
            element.fill.gradient()
            element.fill.gradient_angle = angle
            stops = element.fill.gradient_stops
            stops[0].color.rgb = _rgb(self.colour(gradient[0]))
            stops[0].position = 0.0
            stops[1].color.rgb = _rgb(self.colour(gradient[1]))
            stops[1].position = 1.0
        elif fill is None:
            element.fill.background()
        else:
            element.fill.solid()
            element.fill.fore_color.rgb = _rgb(self.colour(fill))
            if alpha is not None:
                _set_alpha(element.fill._xPr.find(qn('a:solidFill')), alpha)
        if line:
            element.line.color.rgb = _rgb(self.colour(line))
            element.line.width = Pt(line_width)
        else:
            element.line.fill.background()
        element.shadow.inherit = False


def _rgb(value):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(value)


def _emu(box):
    return tuple(Inches(value) for value in box)


def zone(box):
    """A slides.Zone (EMU) for a box in inches."""
    from ..slides import Zone
    return Zone(*_emu(box))


def _set_alpha(solid_fill, alpha):
    colour = solid_fill[0]
    for old in colour.findall(qn('a:alpha')):
        colour.remove(old)
    element = colour.makeelement(qn('a:alpha'), {'val': str(int(round(alpha * 100000)))})
    colour.append(element)


class Composition:
    """The arrangement a design wears. Subclasses draw each kind of slide.

    `ground(canvas, kind, photo)` draws everything under the text of one slide
    — background, accent element, panels, shapes, pattern — and returns a
    `Plan` for a cover, closing slide or divider (`kind` 'cover', 'closing',
    'section') or None for a content slide. `photo` is True when the slide has
    a picture, which goes where `Plan.photo` says (covers) or where the layout
    puts it (content slides, inside `content_box`).
    """
    name = ''
    # Content-slide pictures stay inside this box (inches), e.g. a framing card.
    content_box = (0.0, 0.0, SLIDE_W, SLIDE_H)

    def ground(self, canvas, kind, photo=False):
        raise NotImplementedError


def compose(name):
    """The composition registered as `name`, or the classic one."""
    from . import catalog  # noqa: F401  (registers every composition)
    return COMPOSITIONS.get(name) or COMPOSITIONS['classic']
