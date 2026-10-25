"""How each slide layout is drawn.

One drawer per layout id in `layouts.CATALOGUE`, registered with `@draws`. A
drawer receives a slide on which the ground (background and accent bar), any
photo and the chrome (brand, slide number, logo) are already in place, and adds
only the layout's own text and shapes. That keeps the shape order the same on
every slide, and keeps any one layout readable on its own.

Two rules hold for every drawer. Text is only ever written on the slide's own
surface or on a panel whose colour role was contrast-corrected for it — never
over a photo. And every field is clipped to a word limit, marking the deck as
shortened, rather than being allowed to run off its box.

These are the classic drawers. A designed deck draws its content slides in
variants of them (apps/studio/variants) — other arrangements of the same
fields — and falls back to these whenever a variant would not hold a slide.
"""
import dataclasses
import functools
import re

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt
from reportlab.pdfbase import pdfmetrics

from . import slides as kit
from .layouts import PRO, parse_number

DRAW = {}

LEFT, WIDTH = 0.85, 13.333 - 2 * 0.85
TOP, BOTTOM = 2.75, 6.4
LABELS = {
    'pros': {'en': 'For', 'uz': 'Afzalliklari', 'ru': 'За'},
    'cons': {'en': 'Against', 'uz': 'Kamchiliklari', 'ru': 'Против'},
    'vs': {'en': 'vs', 'uz': 'va', 'ru': 'vs'},
}


def draws(*kinds):
    def register(function):
        for kind in kinds:
            DRAW[kind] = function
        return function
    return register


class Ctx:
    """What a drawer needs to know about the slide it is drawing.

    `heading_bold` is the design's: slides.apply_fonts sets the heading face
    bold after drawing, so text in that face is measured bold from the start.
    """

    __slots__ = ('slide', 'roles', 'style', 'index', 'total', 'size', 'locale', 'photo', 'shortened', 'plan',
                 'heading_bold', 'squeezed')

    def __init__(self, slide, roles, style, index, total, size, locale, photo=False, plan=None,
                 heading_bold=False):
        self.slide, self.roles, self.style = slide, roles, style
        self.index, self.total, self.size, self.locale = index, total, size, locale
        self.photo, self.shortened, self.plan = photo, False, plan
        self.heading_bold = heading_bold
        # Set by a variant whose list had to go well below the deck's size
        # (apps/studio/variants/kit.bullets): it is then not used for the slide.
        self.squeezed = False


def Z(left, top, width, height):
    return kit.Zone(Inches(left), Inches(top), Inches(width), Inches(height))


def clip(ctx, text, words):
    """A field held to a word limit, with the deck marked shortened if it was cut."""
    parts = str(text or '').split()
    if len(parts) <= words:
        return ' '.join(parts)
    ctx.shortened = True
    return ' '.join(parts[:words]).rstrip(',;:.') + '…'


def cut(text):
    """`text` one step shorter, visibly: its last word gone, or a lone word's last letter."""
    text = text.rstrip('…')
    words = text.split()
    return (' '.join(words[:-1]) if len(words) > 1 else text[:-1]).rstrip(',;:.') + '…'


def fit(text, size, zone, lines, floor):
    """The largest size up to `size` at which the text keeps to `lines` lines."""
    while size > floor and kit._wrapped_lines(text, size, zone.width) > lines:
        size -= 2
    return size


@functools.lru_cache(maxsize=50000)
def _units(text):
    """`text`'s width in the font's own units: whole numbers, so a line's is its words' sum."""
    kit._register_font()
    return pdfmetrics.stringWidth(text, 'PDFMaster', 1000)


def _text_lines(text, size, width, *, bold=False):
    """Count wrapped lines, including unused space at the end of each line.

    Measuring the entire string and dividing by the width misses word breaks,
    especially in a narrow card. The small extra width for bold text also
    leaves room for Office's substituted heading font. A line's width is added
    up from its words' rather than measured again for every word, which made
    fitting one long question take seconds.
    """
    available = max(1, Emu(int(width)).pt / (1.06 if bold else 1))

    def over(units):
        # slides._text_width, in the order reportlab scales units to points.
        return 0.001 * size * units * kit.WIDTH_SAFETY > available

    space, total = _units(' '), 0
    for line in str(text).split('\n'):
        current, count = None, 1  # the width of the line so far; None while it is empty
        for word in line.split():
            units = _units(word)
            if current is not None and over(current + space + units):
                count += 1
                current = None
            if over(units):
                # Viewers break unspaced identifiers too. Count their actual
                # character widths rather than treating a long token as one line.
                for character in word:
                    if current is not None and over(current + _units(character)):
                        count += 1
                        current = None
                    current = (current or 0) + _units(character)
            else:
                current = units if current is None else current + space + units
        total += count
    return total


def _paragraph_height(paragraph, zone):
    width = zone.width - (Pt(paragraph['size'] * 0.95) if paragraph.get('bullet') else 0)
    width -= paragraph.get('inset_right', 0)
    width /= paragraph.get('width_safety', 1)
    lines = _text_lines(paragraph['text'], paragraph['size'], width,
                        bold=paragraph.get('bold', False))
    return (lines * paragraph['size'] * paragraph.get('spacing', 1.15)
            + paragraph.get('before', 0) + paragraph.get('after', 0))


def _fit_paragraphs(ctx, zone, paragraphs, *, shorten=True):
    """Fit a complete block with readable floors and explicit line leading.

    Fonts shrink together before any content is removed. If even the floors
    cannot hold unusually long fields, keep each paragraph, visibly elide the
    longest one and report the existing shortened warning to the caller.
    """
    fitted = [dict(paragraph) for paragraph in paragraphs]
    for paragraph in fitted:
        # A design with bold headings sets the heading face bold after drawing
        # (slides.apply_fonts). Measured regular, its text ran out of the box.
        if ctx.heading_bold and paragraph.get('font') == kit.HEADING_FONT:
            paragraph['bold'] = True
    available = Emu(zone.height).pt - 3  # allow for glyph ascenders/descenders

    def height():
        return sum(_paragraph_height(paragraph, zone) for paragraph in fitted)

    while height() > available:
        reducible = [paragraph for paragraph in fitted
                     if paragraph['size'] > paragraph.get('floor', paragraph['size'])]
        if not reducible:
            break
        for paragraph in reducible:
            paragraph['size'] = max(paragraph['floor'], paragraph['size'] - 1)
    while shorten:
        heights = [_paragraph_height(paragraph, zone) for paragraph in fitted]
        if sum(heights) <= available:
            break
        candidates = [index for index, paragraph in enumerate(fitted) if len(paragraph['text'].rstrip('…')) > 1]
        if not candidates:
            break
        # The tallest paragraph loses words until the block fits or another
        # one is the tallest (the first of equals), which is then cut instead.
        tallest = max(candidates, key=heights.__getitem__)
        earlier = max([heights[index] for index in candidates if index < tallest], default=-1)
        later = max([heights[index] for index in candidates if index > tallest], default=-1)
        rest = sum(heights) - heights[tallest]

        def done(text):
            if len(text.rstrip('…')) <= 1:
                return True
            own = _paragraph_height({**fitted[tallest], 'text': text}, zone)
            return (sum(own if index == tallest else value for index, value in enumerate(heights)) <= available
                    or own <= earlier or own < later)

        keep = max(available - rest, earlier, later) / heights[tallest]
        fitted[tallest]['text'] = _cut_until(fitted[tallest]['text'], done, keep)
        ctx.shortened = True
    return fitted


def _cut_until(text, done, keep):
    """The longest visible cut of `text` that is `done`.

    The cuts are a word at a time from the end, then a lone word's letters,
    each ending in '…'. Every cut is shorter than the one before, so the right
    one is found by halves from the share `keep` expected to fit. Measuring
    after every word, as this used to, took seconds for one long question.
    """
    text = ' '.join(text.rstrip('…').split())
    first = text.find(' ') if ' ' in text else len(text)
    points = [index for index in range(len(text) - 1, 0, -1) if text[index] == ' ']
    points = points + list(range(first - 1, -1, -1)) or [0]

    def version(step):
        return text[:points[step]].rstrip(',;:.') + '…'

    low, high = 0, len(points) - 1  # the last cut leaves one letter at most, so it is done
    step = next((index for index, point in enumerate(points) if point <= keep * len(text)), high)
    while low < high:
        if done(version(step)):
            high = step
        else:
            low = step + 1
        step = (low + high) // 2
    return version(low)


def fit_text(ctx, text, zone, *, size, floor=12, spacing=1.15, bold=False):
    """Fit one cell; draw it with exact point leading of ``size * spacing``."""
    paragraph = _fit_paragraphs(ctx, zone, [dict(text=text, size=size, floor=floor,
                                                spacing=spacing, bold=bold)])[0]
    return paragraph['text'], paragraph['size']


def _put_paragraphs(ctx, zone, paragraphs, anchor=MSO_ANCHOR.TOP):
    frame = kit._frame(ctx.slide.shapes.add_textbox(*zone.box()), anchor)
    for index, values in enumerate(_fit_paragraphs(ctx, zone, paragraphs)):
        paragraph = kit._write(kit._paragraph(frame, index), values['text'],
                               font=values.get('font', kit.BODY_FONT), size=values['size'],
                               colour=values['colour'], bold=values.get('bold', False),
                               before=values.get('before', 0), after=values.get('after', 0),
                               align=values.get('align', PP_ALIGN.LEFT))
        # Relative line spacing is based on a viewer's font metrics, which can
        # be considerably taller than the point size used by the estimator.
        paragraph.line_spacing = Pt(values['size'] * values.get('spacing', 1.15))
        if values.get('inset_right'):
            paragraph._p.get_or_add_pPr().set('marR', str(int(values['inset_right'])))
        if values.get('bullet'):
            kit._bullet(paragraph, values['bullet_colour'], values['size'], values['bullet'])
        else:
            kit._no_bullet(paragraph)
    return frame


def put(ctx, zone, value, *, size, colour, font=None, bold=False, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, spacing=1.15, exact_spacing=False):
    """One paragraph of text in a zone, with wrapping left to the viewer."""
    frame = kit._frame(ctx.slide.shapes.add_textbox(*zone.box()), anchor)
    paragraph = kit._write(frame.paragraphs[0], value, font=font or kit.BODY_FONT, size=size,
                           colour=colour, bold=bold, align=align, spacing=spacing)
    if exact_spacing:
        paragraph.line_spacing = Pt(size * spacing)
    kit._no_bullet(paragraph)
    return frame


def fitted(ctx, zone, text, *, size, colour, floor=12, bold=False, font=None, align=PP_ALIGN.LEFT,
           spacing=1.15, anchor=MSO_ANCHOR.TOP):
    """One field that shrinks to its box, and only then is shortened.

    Every box a model's words go into is measured: Russian and Uzbek run a
    third longer than English, and a label that fits in English overflowed
    its box in both until it was.
    """
    return _put_paragraphs(ctx, zone, [dict(text=text, size=size, floor=floor, colour=colour, bold=bold,
                                            font=font or kit.BODY_FONT, align=align, spacing=spacing)], anchor)


def fit_headline(ctx, zone, text, *, size, bold=False, **options):
    """`slides._headline` measured at the weight the design draws it in, with any cut reported."""
    return kit._headline(ctx.slide, zone, text, ctx.roles, size=size, bold=bold or ctx.heading_bold, ctx=ctx,
                         **options)


def fit_bullets(ctx, lines, zone):
    """A list at the deck's size if it fits, else a little smaller, else shortened.

    Sentence-long bullets make a full slide the normal case, so a slide whose
    list still runs long steps down to the floor on its own, and only then
    loses words from its longest bullet — the deck is marked shortened.
    """
    lines, size = list(lines), ctx.size
    while size > kit.FLOOR_BODY_SIZE and not kit._fits(lines, size, zone):
        size -= 0.5
    while lines and not kit._fits(lines, size, zone):
        longest = max(range(len(lines)), key=lambda index: len(lines[index]))
        words = lines[longest].rstrip('…').split()
        if len(words) > 6:
            lines[longest] = ' '.join(words[:-2]).rstrip(',;:.') + '…'
        else:
            lines.pop(longest)
        ctx.shortened = True
    return lines, size


def header(ctx, section, *, width=None, size=32):
    """Eyebrow, headline and accent rule — the top of most content slides."""
    width = width or WIDTH
    # Every box, not just the text in it, stays out of a photo's zone: someone
    # editing the deck can type into the whole box.
    put(ctx, Z(LEFT, 0.62, width, 0.3), f'{ctx.index + 1:02d} / {ctx.total:02d}', size=10,
        colour=ctx.roles['muted'], bold=True)
    fit_headline(ctx, Z(LEFT, 1.0, width, 1.25), clip(ctx, section['heading'], 14), size=size)
    kit._shape(ctx.slide, kit.ZONES['rule'], ctx.roles['accent'])


def items_of(section):
    return section.get('items') or []


def columns_of(section):
    return [column for column in section.get('columns') or [] if column]


def body_text(section):
    return ' '.join(str(section.get('body') or '').split())


# ---------------------------------------------------------------- existing


ALIGN = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}
ANCHOR = {'top': MSO_ANCHOR.TOP, 'middle': MSO_ANCHOR.MIDDLE, 'bottom': MSO_ANCHOR.BOTTOM}


def _planned(ctx, spec, text, words, *, headline):
    """One piece of text where a composition's plan puts it (apps/studio/compositions)."""
    value = clip(ctx, text, words)
    if spec.upper:
        value = value.upper()
    colour = ctx.roles.get(spec.colour, spec.colour)
    if headline:
        fit_headline(ctx, Z(*spec.box), value, size=spec.size, colour=colour, anchor=ANCHOR[spec.anchor],
                     lines=spec.lines, align=ALIGN[spec.align], bold=spec.bold)
    else:
        frame = fitted(ctx, Z(*spec.box), value, size=spec.size, colour=colour, floor=min(12, spec.size),
                       bold=spec.bold, align=ALIGN[spec.align])
        frame.vertical_anchor = ANCHOR[spec.anchor]


def _planned_cover(ctx, title, lines):
    """A cover or closing slide as its composition's plan lays it out."""
    plan = ctx.plan
    several = len(lines) > 1 and bool(plan.details or plan.subtitle)
    # A title page's lines may need the room under the title: the title's box
    # ends `lift` higher (it is set from its foot) and the rule goes up with it,
    # the plan's or one the composition drew itself (`rising`).
    lift = plan.lift if several and plan.details else 0
    spec = plan.title
    if lift:
        spec = dataclasses.replace(spec, box=(*spec.box[:3], spec.box[3] - lift))
        for shape in plan.rising:
            shape.top = shape.top - Inches(lift)
    _planned(ctx, spec, title, 14, headline=True)
    if plan.rule:
        box, colour = plan.rule
        kit._shape(ctx.slide, Z(box[0], box[1] - lift, *box[2:]), ctx.roles.get(colour, colour))
    if several:
        spec = plan.details or plan.subtitle
        byline(ctx, Z(*spec.box), lines, size=spec.size, colour=ctx.roles.get(spec.colour, spec.colour),
               align=ALIGN[spec.align], anchor=ANCHOR[spec.anchor], upper=spec.upper, bold=spec.bold)
    elif lines and plan.subtitle:
        _planned(ctx, plan.subtitle, lines[0], 24, headline=False)


# A title page carries what its writer dictated under the title — the school,
# the kind of work, who prepared it, who accepts it, the group, the city and
# year — one line each. Every line the box holds is drawn, up to eight: a cover
# that kept only the first lost the student's name, and one that kept five lost
# the city and year.
TITLE_LINES = 8
GLYPH = re.compile(r'^\s*[•·▪◦‣*\-–—]+\s*')


def title_lines(section, lines):
    """The lines under a cover's or closing slide's title.

    A body written as several lines is drawn line by line, as written, up to
    `TITLE_LINES` (only a list marker is taken off: "A. Karimov" keeps its
    initial). A one-line subtitle is the first of `lines`, as it always was.
    """
    written = [GLYPH.sub('', line).strip() for line in str(section.get('body') or '').split('\n')]
    written = [line for line in written if line]
    return written if len(written) > 1 else lines[:1]


def byline(ctx, zone, lines, *, size, colour, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, upper=False, bold=False):
    """Several title-page lines in one box, a paragraph each: smaller together first, then fewer.

    As many whole lines as the box holds at its floor are drawn, up to
    `TITLE_LINES`; the rest are left out and the deck is marked shortened.
    """
    floor = min(12, size)
    paragraphs = [dict(text=clip(ctx, line.upper() if upper else line, 24), size=size, floor=floor,
                       colour=colour, bold=bold, align=align, spacing=1.15)
                  for line in lines[:TITLE_LINES]]
    # Whole lines past what the box holds at its floor are left out: cutting
    # each to fit would leave a column of initials ("M…", "T…") instead. A
    # long line that wraps takes the room of two.
    available, used, kept = Emu(zone.height).pt - 3, 0, 0
    for paragraph in paragraphs:
        used += _paragraph_height({**paragraph, 'size': floor}, zone)
        if kept and used > available:
            break
        kept += 1
    if kept < len(lines):
        ctx.shortened = True
    return _put_paragraphs(ctx, zone, paragraphs[:kept], anchor)


@draws('cover')
def cover(ctx, section, lines):
    roles, photo = ctx.roles, ctx.photo
    lines = title_lines(section, lines)
    if ctx.plan:
        return _planned_cover(ctx, section['heading'].strip() or section.get('_title', ''), lines)
    title = section['heading'].strip() or section.get('_title', '')
    zone = Z(LEFT, 2.35, WIDTH * 0.5, 2.1) if photo else kit.ZONES['cover_title']
    fit_headline(ctx, zone, clip(ctx, title, 14), size=44, colour=roles['cover_ink'], lines=3)
    # On a bold deck the cover is the accent, so the rule has to be the ink.
    kit._shape(ctx.slide, kit.ZONES['cover_rule'],
               roles['cover_ink'] if roles['theme'] == 'bold' else roles['accent'])
    if len(lines) > 1:
        # Down to just above the footer; a photo design's title panel ends at 6.55.
        byline(ctx, Z(LEFT, 4.95, WIDTH * 0.5, 1.55) if photo else Z(LEFT, 4.95, WIDTH * 0.7, 1.7), lines,
               size=18, colour=roles['cover_muted'])
    elif lines:
        subtitle_zone = Z(LEFT, 4.95, WIDTH * 0.5, 0.9) if photo else kit.ZONES['cover_subtitle']
        fitted(ctx, subtitle_zone, clip(ctx, lines[0], 24), size=18, floor=14, colour=roles['cover_muted'],
               spacing=1.25)


@draws('section')
def section_(ctx, section, lines):
    roles = ctx.roles
    kicker = lines[0] if lines and section.get('layout') not in (None, '', 'auto') else ''
    if ctx.plan:
        plan = ctx.plan
        if plan.eyebrow:
            number = f'{ctx.index + 1:02d}' if plan.eyebrow.numeral else f'{ctx.index + 1:02d} / {ctx.total:02d}'
            _planned(ctx, plan.eyebrow, number, 3, headline=False)
        title = plan.title_with_kicker if kicker and plan.title_with_kicker else plan.title
        _planned(ctx, title, section['heading'], 14, headline=True)
        if kicker and plan.kicker:
            _planned(ctx, plan.kicker, kicker, 20, headline=False)
        return
    kit._shape(ctx.slide, kit.ZONES['divider_band'], roles['band'])
    kit._eyebrow(ctx.slide, f'{ctx.index + 1:02d} / {ctx.total:02d}', roles)
    zone = Z(LEFT, 2.95, WIDTH * 0.8, 1.2) if kicker else kit.ZONES['divider_title']
    fit_headline(ctx, zone, clip(ctx, section['heading'], 14), size=34, colour=roles['band_ink'],
                 anchor=MSO_ANCHOR.TOP, lines=2)
    if kicker:
        fitted(ctx, Z(LEFT, 4.2, WIDTH * 0.8, 0.6), clip(ctx, kicker, 20), size=16, floor=12,
               colour=roles['band_muted'])


@draws('statement')
def statement(ctx, section, lines):
    roles = ctx.roles
    kit._eyebrow(ctx.slide, f'{ctx.index + 1:02d} / {ctx.total:02d}', roles)
    # The headline steps back so the one idea can carry the slide. Setting the
    # statement smaller than its own heading made it read as a subtitle.
    fit_headline(ctx, kit.ZONES['headline'], section['heading'], size=20, colour=roles['muted'],
                 anchor=MSO_ANCHOR.BOTTOM, lines=2)
    kit._shape(ctx.slide, kit.ZONES['rule'], roles['accent'])
    line = lines[0] if lines else body_text(section)
    fitted(ctx, kit.ZONES['statement'], line, font=kit.HEADING_FONT, size=40 if len(line) <= 52 else 32,
           floor=24, colour=roles['ink'], spacing=1.14)


@draws('bullets')
def bullets(ctx, section, lines):
    header(ctx, section)
    lines, size = fit_bullets(ctx, lines, kit.ZONES['body'])
    kit._bullet_block(ctx.slide, kit.ZONES['body'], lines, ctx.roles, size)


@draws('two_column')
def two_column(ctx, section, lines):
    roles = ctx.roles
    header(ctx, section)
    titled = [item for item in items_of(section) if item.get('label') and item.get('text')]
    if len(titled) == 2 and section.get('layout') == 'two_column':
        for zone, item in zip((kit.ZONES['body_left'], kit.ZONES['body_right']), titled):
            _put_paragraphs(ctx, zone, [
                dict(text=clip(ctx, item['label'], 8), font=kit.HEADING_FONT, size=22, floor=18,
                     colour=roles['accent_text'], bold=True, after=10),
                dict(text=clip(ctx, item['text'], 60), size=17, floor=13, colour=roles['ink'],
                     spacing=kit.LINE_MULTIPLE)])
    else:
        half = -(-len(lines) // 2)
        for zone, part in ((kit.ZONES['body_left'], lines[:half]), (kit.ZONES['body_right'], lines[half:])):
            part, size = fit_bullets(ctx, part, zone)
            kit._bullet_block(ctx.slide, zone, part, roles, size)
    kit._shape(ctx.slide, kit.ZONES['column_rule'], roles['accent_soft'])


# ---------------------------------------------------------------- text


@draws('quote')
def quote(ctx, section, lines):
    roles = ctx.roles
    kit._eyebrow(ctx.slide, f'{ctx.index + 1:02d} / {ctx.total:02d}', roles)
    if section['heading'].strip():
        fit_headline(ctx, kit.ZONES['headline'], section['heading'], size=20, colour=roles['muted'], lines=2)
    put(ctx, Z(LEFT, 2.25, 1.0, 1.4), '“', size=110, colour=roles['accent_text'],
        font=kit.HEADING_FONT, spacing=1.0)
    words = clip(ctx, body_text(section), 45)
    fitted(ctx, Z(LEFT + 1.0, 2.7, WIDTH * 0.8, 2.6), words, size=30 if len(words) <= 140 else 25, floor=18,
           colour=roles['ink'], font=kit.HEADING_FONT, spacing=1.2)
    who = items_of(section)[0]
    attribution = ', '.join(part for part in (clip(ctx, who.get('label'), 8), clip(ctx, who.get('text'), 12)) if part)
    fitted(ctx, Z(LEFT + 1.0, 5.45, WIDTH * 0.8, 0.6), f'— {attribution}', size=16, floor=12, colour=roles['muted'])


@draws('closing')
def closing(ctx, section, lines):
    roles = ctx.roles
    lines = title_lines(section, lines)
    if ctx.plan:
        return _planned_cover(ctx, section['heading'], lines)
    fit_headline(ctx, kit.ZONES['cover_title'], clip(ctx, section['heading'], 14), size=44,
                 colour=roles['cover_ink'], lines=3)
    kit._shape(ctx.slide, kit.ZONES['cover_rule'],
               roles['cover_ink'] if roles['theme'] == 'bold' else roles['accent'])
    if len(lines) > 1:
        byline(ctx, Z(LEFT, 4.95, WIDTH * 0.7, 1.7), lines, size=18, colour=roles['cover_muted'])
    elif lines:
        fitted(ctx, kit.ZONES['cover_subtitle'], clip(ctx, lines[0], 24), size=18, floor=14,
               colour=roles['cover_muted'], spacing=1.25)


# ---------------------------------------------------------------- data


@draws('big_number')
def big_number(ctx, section, lines):
    roles = ctx.roles
    header(ctx, section, size=28)
    item = items_of(section)[0]
    value = clip(ctx, item.get('value'), 3)
    zone = Z(LEFT, 2.75, 5.6, 2.6)
    put(ctx, zone, value, size=fit(value, 110, zone, 1, 60), colour=roles['accent_text'],
        font=kit.HEADING_FONT, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0)
    paragraphs = [dict(text=clip(ctx, item.get('label'), 10), font=kit.HEADING_FONT, size=28, floor=18,
                       colour=roles['ink'], after=8)]
    if item.get('text'):
        paragraphs.append(dict(text=clip(ctx, item['text'], 40), size=16, floor=12, colour=roles['muted'],
                               spacing=kit.LINE_MULTIPLE))
    _put_paragraphs(ctx, Z(LEFT + 6.0, 3.0, WIDTH - 6.0, 2.8), paragraphs, MSO_ANCHOR.MIDDLE)


@draws('stats')
def stats(ctx, section, lines):
    roles = ctx.roles
    header(ctx, section)
    tiles = items_of(section)
    gap = 0.35
    width = (WIDTH - gap * (len(tiles) - 1)) / len(tiles)
    for position, item in enumerate(tiles):
        left = LEFT + position * (width + gap)
        kit._shape(ctx.slide, Z(left, 3.0, width, 0.06), roles['accent'])
        value = clip(ctx, item.get('value'), 3)
        zone = Z(left, 3.2, width, 1.1)
        # One line if a smaller size gets it there, and held to the box's height
        # too: a value with words after its figure wrapped into the label below.
        # Leading of 1.2 sets the figure where single spacing in this face did.
        fitted(ctx, zone, value, size=fit(value, 52, zone, 1, 28), floor=20, colour=roles['accent_text'],
               font=kit.HEADING_FONT, spacing=1.2)
        fitted(ctx, Z(left, 4.4, width, 0.7), clip(ctx, item.get('label'), 8), size=18, floor=12,
               colour=roles['ink'], bold=True)
        if item.get('text'):
            fitted(ctx, Z(left, 5.1, width, 1.3), clip(ctx, item['text'], 20), size=14, floor=12,
                   colour=roles['muted'])


@draws('chart')
def chart(ctx, section, lines):
    """Horizontal bars drawn from shapes.

    A python-pptx chart embeds an Excel workbook under /ppt/embeddings/, which
    our own upload check refuses — a customer could not re-upload their deck.
    Shapes render the same in every viewer and stay editable as shapes.
    """
    roles = ctx.roles
    header(ctx, section)
    bars = items_of(section)
    numbers = [parse_number(item.get('value')) or 0 for item in bars]
    top, span = 2.85, BOTTOM - 2.85
    unit = columns_of(section)[:1]
    if unit:
        put(ctx, Z(LEFT, 2.55, WIDTH, 0.3), clip(ctx, unit[0], 6), size=12, colour=roles['muted'],
            align=PP_ALIGN.RIGHT)
    # Few bars get thicker rows rather than leaving the lower half of the slide empty.
    row = min(0.85, span / len(bars))
    label_width, value_width = WIDTH * 0.27, WIDTH * 0.13
    area = WIDTH - label_width - value_width - 0.3
    peak = max(numbers) or 1
    for position, (item, number) in enumerate(zip(bars, numbers)):
        y = top + position * row
        fitted(ctx, Z(LEFT, y, label_width, row), clip(ctx, item.get('label'), 6), size=14, floor=10,
               colour=roles['ink'], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        length = max(0.04, area * number / peak)
        bar = kit._shape(ctx.slide, Z(LEFT + label_width + 0.15, y + row * 0.2, length, row * 0.6),
                         roles['accent'])
        # An outline in the corrected accent keeps a pale brand colour visible.
        bar.line.fill.solid()
        bar.line.fill.fore_color.rgb = kit._rgb(roles['accent_text'])
        bar.line.width = Pt(0.75)
        # Measured like its label: with eight bars, words after the figure
        # wrapped in the value column and ran past the row.
        fitted(ctx, Z(LEFT + label_width + 0.3 + length, y, value_width + (area - length), row),
               clip(ctx, item.get('value'), 3), size=14, floor=10, colour=roles['ink'], bold=True,
               anchor=MSO_ANCHOR.MIDDLE)


@draws('table')
def table(ctx, section, lines):
    roles = ctx.roles
    header(ctx, section)
    columns = columns_of(section)[:3]
    rows = items_of(section)
    height = min(BOTTOM - TOP, 0.62 * (len(rows) + 1))
    frame = ctx.slide.shapes.add_table(len(rows) + 1, len(columns), *Z(LEFT, TOP, WIDTH, height).box())
    grid = frame.table
    grid.horz_banding = False
    grid.first_row = True

    def cell(row, column, value, fill, colour, bold=False):
        target = grid.cell(row, column)
        target.fill.solid()
        target.fill.fore_color.rgb = kit._rgb(fill)
        target.margin_left = target.margin_right = Inches(0.12)
        target.margin_top = target.margin_bottom = Inches(0.06)
        target.vertical_anchor = MSO_ANCHOR.MIDDLE
        zone = kit.Zone(0, 0, grid.columns[column].width - target.margin_left - target.margin_right,
                        grid.rows[row].height - target.margin_top - target.margin_bottom)
        value, size = fit_text(ctx, value, zone, size=16 if row == 0 else 15, bold=bold)
        paragraph = target.text_frame.paragraphs[0]
        kit._write(paragraph, value, font=kit.BODY_FONT, size=size,
                   colour=colour, bold=bold)
        paragraph.line_spacing = Pt(size * 1.15)
        kit._no_bullet(paragraph)

    for column, name in enumerate(columns):
        cell(0, column, clip(ctx, name, 5), roles['accent'], roles['on_accent'], bold=True)
    for row, item in enumerate(rows, 1):
        fill, ink = (roles['surface'], roles['ink']) if row % 2 else (roles['card'], roles['card_ink'])
        values = [item.get('label', ''), item.get('text', ''), item.get('value', '')]
        if section.get('layout') == 'chart' and len(columns) == 2:
            # A chart's figure is its value, not its optional text. Keep any
            # context alongside it when a chart falls back to two columns.
            values = [values[0], ' — '.join(part for part in (values[2], values[1]) if part)]
        for column in range(len(columns)):
            cell(row, column, clip(ctx, values[column], 12), fill, ink, bold=column == 0)


@draws('comparison')
def comparison(ctx, section, lines):
    roles = ctx.roles
    header(ctx, section)
    first, second = columns_of(section)[:2]
    label_width = WIDTH * 0.26
    column_width = (WIDTH - label_width - 0.7) / 2
    a_left = LEFT + label_width
    b_left = a_left + column_width + 0.7
    for left, name in ((a_left, first), (b_left, second)):
        kit._shape(ctx.slide, Z(left, 2.7, column_width, 0.6), roles['accent'])
        zone = Z(left, 2.7, column_width, 0.6)
        value, size = fit_text(ctx, clip(ctx, name, 5), zone, size=18, bold=True)
        put(ctx, zone, value, size=size, colour=roles['on_accent'], bold=True,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, exact_spacing=True)
    kit._shape(ctx.slide, Z(a_left + column_width + 0.1, 2.73, 0.5, 0.5), roles['card'], MSO_SHAPE.OVAL)
    put(ctx, Z(a_left + column_width + 0.1, 2.73, 0.5, 0.5), LABELS['vs'][ctx.locale], size=12,
        colour=roles['card_accent'], bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    rows = items_of(section)
    row = min(0.75, (BOTTOM - 3.45) / len(rows))
    for position, item in enumerate(rows):
        y = 3.45 + position * row
        if position:
            kit._shape(ctx.slide, Z(LEFT, y, WIDTH, 0.015), roles['accent_soft'])
        zone = Z(LEFT, y, label_width - 0.2, row)
        value, size = fit_text(ctx, clip(ctx, item.get('label'), 6), zone, size=14, bold=True)
        put(ctx, zone, value, size=size, colour=roles['muted'], bold=True,
            anchor=MSO_ANCHOR.MIDDLE, exact_spacing=True)
        for left, key in ((a_left, 'text'), (b_left, 'value')):
            zone = Z(left, y, column_width, row)
            value, size = fit_text(ctx, clip(ctx, item.get(key), 14), zone, size=15)
            put(ctx, zone, value, size=size, colour=roles['ink'], align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE, exact_spacing=True)


# ---------------------------------------------------------------- structure


@draws('agenda')
def agenda(ctx, section, lines):
    roles = ctx.roles
    header(ctx, section)
    entries = [item for item in items_of(section) if item.get('label')][:6]
    columns = 2 if len(entries) > 4 else 1
    per_column = -(-len(entries) // columns)
    width = (WIDTH - 0.5 * (columns - 1)) / columns
    row = (BOTTOM - TOP) / per_column
    for position, item in enumerate(entries):
        column, line = divmod(position, per_column)
        left, y = LEFT + column * (width + 0.5), TOP + line * row
        put(ctx, Z(left, y, 0.95, row), f'{position + 1:02d}', size=26, colour=roles['accent_text'],
            font=kit.HEADING_FONT)
        paragraphs = [dict(text=clip(ctx, item['label'], 8), size=20, floor=16, colour=roles['ink'],
                           bold=True, after=2)]
        if item.get('text'):
            paragraphs.append(dict(text=clip(ctx, item['text'], 22), size=14, floor=12,
                                   colour=roles['muted']))
        _put_paragraphs(ctx, Z(left + 1.0, y + 0.02, width - 1.0, row - 0.04), paragraphs)


@draws('questions')
def questions(ctx, section, lines):
    """Up to five questions: the number, the question, and its options on one line."""
    roles = ctx.roles
    header(ctx, section)
    entries = section.get('_questions') or []
    if not entries:
        return
    row = (BOTTOM - TOP + 0.3) / len(entries)
    for position, entry in enumerate(entries):
        y = TOP - 0.15 + position * row
        put(ctx, Z(LEFT, y, 0.7, min(row, 0.5)), f'{entry["number"]}.', size=20,
            colour=roles['accent_text'], font=kit.HEADING_FONT)
        paragraphs = [dict(text=entry['stem'], size=18, floor=12, colour=roles['ink'], bold=True, after=2)]
        if entry.get('options'):
            paragraphs.append(dict(text='  ·  '.join(entry['options']), size=14, floor=11, colour=roles['muted']))
        _put_paragraphs(ctx, Z(LEFT + 0.75, y + 0.02, WIDTH - 0.75, row - 0.06), paragraphs)


@draws('timeline')
def timeline(ctx, section, lines):
    roles = ctx.roles
    header(ctx, section)
    events = items_of(section)[:6]
    kit._shape(ctx.slide, Z(LEFT, 4.0, WIDTH, 0.04), roles['accent_soft'])
    slot = WIDTH / len(events)
    for position, item in enumerate(events):
        centre = LEFT + slot * (position + 0.5)
        left, width = centre - slot / 2 + 0.06, slot - 0.12
        kit._shape(ctx.slide, Z(centre - 0.13, 3.89, 0.26, 0.26), roles['accent'], MSO_SHAPE.OVAL)
        fitted(ctx, Z(left, 3.05, width, 0.7), clip(ctx, item.get('value'), 4), size=16, floor=11,
               colour=roles['accent_text'], bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
        fitted(ctx, Z(left, 4.35, width, 0.7), clip(ctx, item.get('label'), 7), size=16, floor=11,
               colour=roles['ink'], bold=True, align=PP_ALIGN.CENTER)
        if item.get('text'):
            fitted(ctx, Z(left, 5.05, width, 1.5), clip(ctx, item['text'], 16), size=13, floor=12,
                   colour=roles['muted'], align=PP_ALIGN.CENTER)


@draws('process')
def process(ctx, section, lines):
    roles = ctx.roles
    header(ctx, section)
    steps = [item for item in items_of(section) if item.get('label')][:5]
    slot = WIDTH / len(steps)
    first = LEFT + slot / 2
    kit._shape(ctx.slide, Z(first, 3.33, slot * (len(steps) - 1), 0.04), roles['accent_soft'])
    for position, item in enumerate(steps):
        centre = LEFT + slot * (position + 0.5)
        left, width = centre - slot / 2 + 0.08, slot - 0.16
        kit._shape(ctx.slide, Z(centre - 0.35, 3.0, 0.7, 0.7), roles['accent'], MSO_SHAPE.OVAL)
        put(ctx, Z(centre - 0.35, 3.0, 0.7, 0.7), str(position + 1), size=20, colour=roles['on_accent'],
            bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        fitted(ctx, Z(left, 3.9, width, 0.7), clip(ctx, item['label'], 7), size=17, floor=11, colour=roles['ink'],
               bold=True, align=PP_ALIGN.CENTER)
        if item.get('text'):
            # Down to just above the footer band: a step's sentence needs the height.
            fitted(ctx, Z(left, 4.6, width, 1.95), clip(ctx, item['text'], 22), size=14, floor=12,
                   colour=roles['muted'], align=PP_ALIGN.CENTER)


@draws('pros_cons')
def pros_cons(ctx, section, lines):
    roles = ctx.roles
    header(ctx, section)
    names = columns_of(section)
    groups = (
        ([item for item in items_of(section) if item.get('value', '').strip().lower() in PRO], '+',
         names[0] if len(names) > 0 else LABELS['pros'][ctx.locale]),
        ([item for item in items_of(section) if item.get('value', '').strip().lower() not in PRO], '–',
         names[1] if len(names) > 1 else LABELS['cons'][ctx.locale]),
    )
    gap = 0.3
    # A five-to-one split should give the longer argument more room. Keep both
    # panels substantial while accounting for all accepted points on each side.
    share = min(0.75, max(0.25, (len(groups[0][0]) + 0.5) / (sum(len(g[0]) for g in groups) + 1)))
    widths = [(WIDTH - gap) * share, (WIDTH - gap) * (1 - share)]
    blocks = []
    for position, (points, glyph, name) in enumerate(groups):
        width = widths[position]
        left = LEFT if position == 0 else LEFT + widths[0] + gap
        kit._shape(ctx.slide, Z(left, TOP, width, BOTTOM - TOP), roles['card'])
        kit._shape(ctx.slide, Z(left, TOP, width, 0.07), roles['accent'])
        paragraphs = [dict(text=clip(ctx, name, 5), font=kit.HEADING_FONT, size=24, floor=20,
                           colour=roles['card_accent'], bold=True, after=8)]
        paragraphs.extend(dict(text=clip(ctx, item.get('label') or item.get('text'), 20),
                               size=19, floor=13, colour=roles['card_ink'], spacing=1.2,
                               before=5 if index else 0, bullet=glyph,
                               bullet_colour=roles['card_accent'])
                          for index, item in enumerate(points))
        zone = Z(left + 0.25, TOP + 0.25, width - 0.5, BOTTOM - TOP - 0.45)
        blocks.append((zone, _fit_paragraphs(ctx, zone, paragraphs)))
    # A short opposing argument still uses the same type scale as the long one.
    heading_size = min(paragraphs[0]['size'] for _, paragraphs in blocks)
    body_size = min(paragraph['size'] for _, paragraphs in blocks for paragraph in paragraphs[1:])
    for zone, paragraphs in blocks:
        for index, paragraph in enumerate(paragraphs):
            paragraph['size'] = heading_size if index == 0 else body_size
        _put_paragraphs(ctx, zone, paragraphs)


@draws('cards')
def cards(ctx, section, lines):
    roles = ctx.roles
    header(ctx, section)
    entries = [item for item in items_of(section) if item.get('label')][:6]
    per_row = 2 if len(entries) == 4 else 3 if len(entries) > 3 else len(entries)
    rows = -(-len(entries) // per_row)
    gap = 0.2 if rows > 1 else 0.3
    width = (WIDTH - gap * (per_row - 1)) / per_row
    # Dense grids can use the space just below the rule and above the footer.
    # This keeps six complete cards at readable sizes without adding a slide.
    start, end = (2.6, 6.6) if rows > 1 else (TOP, BOTTOM)
    height = (end - start - gap * (rows - 1)) / rows
    blocks = []
    for position, item in enumerate(entries):
        row, column = divmod(position, per_row)
        left, top = LEFT + column * (width + gap), start + row * (height + gap)
        paragraphs = []
        if item.get('value'):
            paragraphs.append(dict(text=clip(ctx, item['value'], 3).upper(), size=11, floor=10,
                                   colour=roles['card_accent'], bold=True, after=2))
        paragraphs.append(dict(text=clip(ctx, item['label'], 7), font=kit.HEADING_FONT,
                               size=19, floor=14, colour=roles['card_ink'], bold=True, after=3))
        if item.get('text'):
            paragraphs.append(dict(text=clip(ctx, item['text'], 28 if rows == 1 else 25),
                                   size=14, floor=12, colour=roles['card_muted'], spacing=1.15))
        padding = 0.1 if rows > 1 else 0.2
        zone = Z(left + 0.2, top + padding, width - 0.4, height - 2 * padding)
        blocks.append((Z(left, top, width, height), zone, paragraphs))
    # Long translated fields may not fit a six-card grid even at its readable
    # floor. Wide rows make better use of the same slide before any clipping.
    if any(sum(_paragraph_height(p, zone) for p in _fit_paragraphs(ctx, zone, paragraphs, shorten=False))
           > Emu(zone.height).pt - 3 for _, zone, paragraphs in blocks):
        _card_rows(ctx, entries)
        return
    for panel, zone, paragraphs in blocks:
        kit._shape(ctx.slide, panel, roles['card'])
        _put_paragraphs(ctx, zone, paragraphs)


def _card_rows(ctx, entries):
    """A compact card arrangement for complete, unusually verbose fields."""
    gap = 2 / 72
    height = (6.7 - 2.5 - gap * (len(entries) - 1)) / len(entries)
    for index, item in enumerate(entries):
        top = 2.5 + index * (height + gap)
        kit._shape(ctx.slide, Z(LEFT, top, WIDTH, height), ctx.roles['card'])
        zone = Z(LEFT + 0.2, top + 0.02, WIDTH - 0.4, height - 0.04)
        tag = clip(ctx, item.get('value'), 3).upper()
        tag_width = min(zone.width * 0.35, Pt(max(90, kit._text_width(tag, 10) + 8))) if tag else 0
        paragraphs = [dict(text=clip(ctx, item['label'], 7), size=16, floor=14,
                           font=kit.HEADING_FONT, colour=ctx.roles['card_ink'], bold=True,
                           spacing=1.1, after=1, inset_right=tag_width, width_safety=1.1)]
        if item.get('text'):
            # A compact row has no spare line. Corbel can substitute to the
            # wider DejaVu Sans on Linux, so a line that only just fits Noto's
            # measured width must be budgeted as wrapping before choosing size.
            paragraphs.append(dict(text=clip(ctx, item['text'], 28 if len(entries) <= 3 else 25), size=14, floor=12,
                                   colour=ctx.roles['card_muted'], spacing=1.1, width_safety=1.1))
        _put_paragraphs(ctx, zone, paragraphs)
        if tag:
            tag_zone = kit.Zone(zone.left + zone.width - tag_width, zone.top, tag_width, Pt(18))
            tag, size = fit_text(ctx, tag, tag_zone, size=10, floor=10, bold=True)
            put(ctx, tag_zone, tag, size=size, colour=ctx.roles['card_accent'], bold=True,
                align=PP_ALIGN.RIGHT, exact_spacing=True)


@draws('matrix')
def matrix(ctx, section, lines):
    roles = ctx.roles
    header(ctx, section)
    cells = [item for item in items_of(section) if item.get('label')][:4]
    axes = columns_of(section)[:2]
    bottom = BOTTOM - (0.35 if axes else 0)
    gap = 0.12
    width, height = (WIDTH - gap) / 2, (bottom - TOP - gap) / 2
    for position, item in enumerate(cells):
        row, column = divmod(position, 2)
        tinted = (row + column) % 2 == 1
        fill = roles['accent_soft'] if tinted else roles['card']
        ink = roles['soft_ink'] if tinted else roles['card_ink']
        muted = roles['soft_muted'] if tinted else roles['card_muted']
        left, top = LEFT + column * (width + gap), TOP + row * (height + gap)
        kit._shape(ctx.slide, Z(left, top, width, height), fill)
        paragraphs = [dict(text=clip(ctx, item['label'], 6), font=kit.HEADING_FONT, size=20, floor=16,
                           colour=ink, bold=True, after=6)]
        if item.get('text'):
            paragraphs.append(dict(text=clip(ctx, item['text'], 28), size=15, floor=12, colour=muted,
                                   spacing=kit.LINE_MULTIPLE))
        _put_paragraphs(ctx, Z(left + 0.3, top + 0.2, width - 0.6, height - 0.4), paragraphs)
    if axes:
        fitted(ctx, Z(LEFT, bottom + 0.05, WIDTH, 0.3), '  ·  '.join(clip(ctx, axis, 5) for axis in axes),
               size=12, floor=9, colour=roles['muted'], align=PP_ALIGN.CENTER)


# ---------------------------------------------------------------- photo

# Where each photo layout puts its picture. Nothing here overlaps a text zone
# or the footer band, so text never has to be read against a photograph.
PHOTO_ZONES = {
    'cover': Z(13.333 * 0.56, 0, 13.333 * 0.44, 6.55),
    'image_split': Z(13.333 * 0.52, 0, 13.333 * 0.48, 6.55),
    'image_full': Z(0.18, 0, 13.333 - 0.18, 5.3),
}


@draws('image_split')
def image_split(ctx, section, lines):
    width = 13.333 * 0.52 - LEFT - 0.45
    header(ctx, section, width=width, size=30)
    zone = Z(LEFT, TOP, width, BOTTOM - TOP)
    lines, size = fit_bullets(ctx, lines, zone)
    kit._bullet_block(ctx.slide, zone, lines, ctx.roles, size)


@draws('image_full')
def image_full(ctx, section, lines):
    roles = ctx.roles
    fit_headline(ctx, Z(LEFT, 5.4, WIDTH * 0.7, 0.7), clip(ctx, section['heading'], 12), size=26,
                 anchor=MSO_ANCHOR.TOP, lines=1)
    caption = lines[0] if lines else ''
    if caption:
        fitted(ctx, Z(LEFT, 6.1, WIDTH * 0.7, 0.6), clip(ctx, caption, 24), size=14, floor=12,
               colour=roles['muted'])


def place_photo(slide, zone, photo):
    """A photo filling its zone, cropped rather than stretched.

    The crop is set on the picture rather than baked into the image, so the
    customer can reframe it in PowerPoint.
    """
    import io
    picture = slide.shapes.add_picture(io.BytesIO(photo['jpeg']), *zone.box())
    target = zone.width / zone.height
    source = photo['width'] / photo['height']
    if source > target:
        spare = 1 - target / source
        picture.crop_left = picture.crop_right = spare / 2
    elif source < target:
        spare = 1 - source / target
        picture.crop_top, picture.crop_bottom = spare * 0.4, spare * 0.6
    return picture
