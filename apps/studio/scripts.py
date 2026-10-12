"""Every text a reader sees, in the one alphabet its language is written in.

Uzbek and English are Latin only; Russian is Cyrillic, with Latin allowed for
names, URLs and identifiers. The model is told so in every call (see
apps/studio/language.py); this checks what came back, field by field, and puts
right what it can without losing meaning:

- a Cyrillic Uzbek word in an Uzbek document is the same word in the other
  alphabet, so it is transliterated (apps/studio/uzbek.py);
- a Russian name or term in an English document is written in Latin letters
  the usual English way (Khiva, Tashkent's "Toshkent" stays recognisable);
- a word mixing the two alphabets is first put into the one most of its
  letters use (uzbek.settle), so it keeps its sound, then into the document's
  own: "Мoсква" is Moskva, not Mockva; "Pythоn" stays Python in Russian.

A whole passage written in the wrong language is not transliterated into
gibberish: apps/studio/quality.py sends it back to the model. Only verbatim
source quotes (citations) keep their own alphabet.
"""
import re
import unicodedata

from . import uzbek
from .language import CYRILLIC

# Russian letters in Latin, the way English text writes Russian names (BGN/PCGN, simplified).
RUSSIAN = {
    'а': 'a', 'б': 'b', 'в': 'v', 'г': 'g', 'д': 'd', 'е': 'e', 'ё': 'yo', 'ж': 'zh', 'з': 'z', 'и': 'i',
    'й': 'y', 'к': 'k', 'л': 'l', 'м': 'm', 'н': 'n', 'о': 'o', 'п': 'p', 'р': 'r', 'с': 's', 'т': 't',
    'у': 'u', 'ф': 'f', 'х': 'kh', 'ц': 'ts', 'ч': 'ch', 'ш': 'sh', 'щ': 'shch', 'ъ': '', 'ы': 'y', 'ь': '',
    'э': 'e', 'ю': 'yu', 'я': 'ya',
}
# The other Cyrillic letters a text can carry: Uzbek, Ukrainian, Kazakh, Tajik, Serbian, historical.
OTHER = {
    'ў': 'o‘', 'қ': 'q', 'ғ': 'g‘', 'ҳ': 'h', 'і': 'i', 'ї': 'yi', 'є': 'ye', 'ґ': 'g', 'ә': 'a', 'ң': 'ng',
    'ө': 'o', 'ү': 'u', 'ұ': 'u', 'һ': 'h', 'ј': 'j', 'љ': 'lj', 'њ': 'nj', 'ћ': 'c', 'ђ': 'dj', 'џ': 'dz',
    'ѕ': 'dz', 'ѓ': 'gj', 'ќ': 'kj', 'ӣ': 'i', 'ӯ': 'u', 'ҷ': 'j', 'ҙ': 'z', 'ҡ': 'q', 'ҫ': 's', 'ӑ': 'a',
    'ӗ': 'e', 'ӳ': 'u', 'ѣ': 'e', 'ѳ': 'f', 'ѵ': 'i', 'ӏ': 'i',
}
VOWELS = set('аеёиоуыэюяў')
WORD = re.compile(r"[\wЀ-ԯ'‘’ʻʼ`]+")


def _letter(letter):
    """Any Cyrillic letter in Latin, keeping its case; marks that only decorate are dropped."""
    lower = letter.lower()
    value = RUSSIAN.get(lower, OTHER.get(lower))
    if value is None:
        if unicodedata.combining(letter):
            return ''
        # A letter no table knows: its decomposed base, or the last word of its name.
        base = ''.join(c for c in unicodedata.normalize('NFKD', lower) if not unicodedata.combining(c))
        if base != lower and not CYRILLIC.search(base):
            value = base
        else:
            name = unicodedata.name(letter, '').split(' ')[-1].lower()
            value = re.sub('[^a-z]', '', name)[:2]
    # A capital starts its digraph and no more: Хива is Khiva; russian_to_latin capitalises all-caps words.
    return value[:1].upper() + value[1:] if letter.isupper() else value


def russian_to_latin(word):
    out = []
    for position, letter in enumerate(word):
        lower = letter.lower()
        if not CYRILLIC.match(letter):
            out.append(letter)
            continue
        value = _letter(letter)
        # е is "ye" at the start of a word and after a vowel: Ереван → Yerevan.
        if lower == 'е' and (position == 0 or word[position - 1].lower() in VOWELS):
            value = 'Ye' if letter.isupper() else 'ye'
        if letter.isupper() and word.isupper() and len(word) > 1:
            value = value.upper()
        out.append(value)
    return ''.join(out)


def _english(text):
    def fix(match):
        word = uzbek.settle(match.group(0))
        if not CYRILLIC.search(word):
            return word
        return uzbek.to_latin(word) if uzbek.UZBEK_LETTERS.search(word) else russian_to_latin(word)
    return WORD.sub(fix, text)


def _russian(text):
    """Russian keeps Latin names; a word mixing the alphabets keeps the one most of it is in (Pythоn, мaркет)."""
    return WORD.sub(lambda match: uzbek.settle(match.group(0)), text)


def to_script(text, locale):
    """`text` in the alphabet `locale` is written in."""
    if not isinstance(text, str) or not text:
        return text
    if locale == 'ru':
        return _russian(text)
    # A letter typed as base and mark (и + ˘) would leave the mark on its Latin letter (ĭ): compose it first.
    text = unicodedata.normalize('NFC', text) if CYRILLIC.search(text) else text
    if locale == 'uz':
        text = uzbek.clean_text(text, 'latn')
    elif locale == 'en':
        text = _english(text)
    if not CYRILLIC.search(text):
        return text
    # Whatever a table missed is still never delivered in Cyrillic: word by word, so an all-caps word
    # stays all caps, then any mark standing alone.
    text = WORD.sub(lambda match: russian_to_latin(match.group(0)), text)
    return CYRILLIC.sub(lambda match: _letter(match.group(0)), text)


def fields(content):
    """(path, text) for every field of `content` a reader sees, citations excepted."""
    yield ('title',), content.get('title', '')
    for i, section in enumerate(content.get('sections', [])):
        for key in ('heading', 'body', 'notes'):
            if isinstance(section.get(key), str):
                yield ('sections', i, key), section[key]
        for j, item in enumerate(section.get('items') or []):
            for key in ('label', 'text', 'value'):
                if isinstance(item.get(key), str):
                    yield ('sections', i, 'items', j, key), item[key]
        for j, column in enumerate(section.get('columns') or []):
            if isinstance(column, str):
                yield ('sections', i, 'columns', j), column
    for i, question in enumerate(content.get('questions', [])):
        for key in ('stem', 'answer', 'explanation', 'topic'):
            if isinstance(question.get(key), str):
                yield ('questions', i, key), question[key]
        for j, option in enumerate(question.get('options') or []):
            if isinstance(option, str):
                yield ('questions', i, 'options', j), option


def violations(content, locale):
    """Paths of the fields written in the wrong alphabet (never their text)."""
    if locale not in ('uz', 'en'):
        return []
    return [path for path, text in fields(content) if isinstance(text, str) and CYRILLIC.search(text)]


def _set(content, path, value):
    target = content
    for key in path[:-1]:
        target = target[key]
    target[path[-1]] = value


def enforce(content, locale):
    """(content in its alphabet, paths that were corrected). Raises if any is still wrong."""
    import copy
    result = copy.deepcopy(content)
    fixed = []
    for path, text in list(fields(result)):
        value = to_script(text, locale)
        if value != text:
            _set(result, path, value)
            if CYRILLIC.search(text or '') or locale == 'ru':
                fixed.append(path)
    left = violations(result, locale)
    if left:
        raise ScriptPolicyError(left)
    return result, fixed


class ScriptPolicyError(ValueError):
    def __init__(self, paths):
        super().__init__(f'{len(paths)} field(s) still in the wrong alphabet')
        self.paths = paths


# --- The delivered file --------------------------------------------------------

def artifact_text(path):
    """All the text a reader can see in a rendered .pptx or .pdf, notes included."""
    path = str(path)
    if path.endswith('.pptx'):
        from pptx import Presentation
        deck = Presentation(path)
        parts = [deck.core_properties.title or '', deck.core_properties.subject or '']

        def walk(shapes):
            for shape in shapes:
                if shape.shape_type == 6:  # group
                    walk(shape.shapes)
                if getattr(shape, 'has_text_frame', False) and shape.has_text_frame:
                    parts.append(shape.text_frame.text)
                if getattr(shape, 'has_table', False) and shape.has_table:
                    parts.extend(cell.text for row in shape.table.rows for cell in row.cells)
                if getattr(shape, 'has_chart', False) and shape.has_chart:
                    chart = shape.chart
                    for plot in chart.plots:
                        parts.extend(str(c) for c in plot.categories)
                        parts.extend(series.name or '' for series in plot.series)
        for slide in deck.slides:
            walk(slide.shapes)
            if slide.has_notes_slide:
                parts.append(slide.notes_slide.notes_text_frame.text)
        return '\n'.join(parts)
    from pypdf import PdfReader
    reader = PdfReader(path)
    meta = reader.metadata or {}
    return '\n'.join([str(meta.get('/Title') or ''), str(meta.get('/Subject') or '')]
                     + [page.extract_text() or '' for page in reader.pages])


def artifact_violations(path, locale, allowed=''):
    """Cyrillic words in a rendered Uzbek or English file that are not part of a source quote."""
    if locale not in ('uz', 'en'):
        return []
    quoted = set(WORD.findall(allowed or ''))
    return sorted({word for word in WORD.findall(artifact_text(path))
                   if CYRILLIC.search(word) and word not in quoted})
